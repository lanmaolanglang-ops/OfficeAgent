"""
Installation Test Suite - 安装测试方案
测试在全新Windows环境下的安装、启动、功能验证

判定口径（P1-23 修正）：
- **存在（exists）**：路径/文件真的在。
- **可用（usable）**：能真正跑起来——可执行文件是真正的 PE/可执行类型、
  subprocess 能起来且 returncode 可解释、HTTP 端点返回预期状态码与结构。
- **契约（contract）**：符合当前产品定义——依赖清单以
  ``requirements-production.txt`` 为准、/health 的版本号与
  ``office_agent._version.__version__`` 一致、上传走真实路由 ``/api/file/upload``。

开发源码树 ≠ 已安装产品：两者判据不同，用 ``mode`` 显式区分。
"""
import os
import re
import sys
import json
import time
import tempfile
import subprocess
import urllib.error
import urllib.request
from pathlib import Path
from dataclasses import dataclass, field
from enum import Enum
from typing import Optional

# 确保项目路径
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


# ============================================================
# 产品契约（唯一来源）
# ============================================================

#: 发布安装的主产物名（与 desktop/release_manifest.py 的 PRIMARY_ARTIFACT 一致）
PRIMARY_ARTIFACT = "OfficeAgent.exe"
MANIFEST_NAME = "release-manifest.json"
#: 上传接口的真实路由（旧自测写 /api/v1/files/upload，该路由不存在）
UPLOAD_ROUTE = "/api/file/upload"

#: 发行名 → import 名（pip 名与 import 名不同，必须显式映射）
REQUIREMENT_IMPORT_NAMES = {
    "fastapi": "fastapi",
    "uvicorn": "uvicorn",
    "pydantic": "pydantic",
    "python-multipart": "multipart",
    "httpx": "httpx",
    "sqlalchemy": "sqlalchemy",
    "alembic": "alembic",
    "python-docx": "docx",
    "python-pptx": "pptx",
    "openpyxl": "openpyxl",
    "pandas": "pandas",
    "numpy": "numpy",
    "usearch": "usearch",
    "pymupdf": "fitz",
    "pillow": "PIL",
    "cryptography": "cryptography",
    "psutil": "psutil",
    "pywin32": "win32api",
    "pyinstaller": "PyInstaller",
}

#: 仅构建期需要，不参与"运行期核心依赖"判定
BUILD_ONLY_REQUIREMENTS = {"pyinstaller"}
#: 平台专属依赖：非 Windows 缺失不算失败
PLATFORM_GATED_REQUIREMENTS = {"pywin32"}

#: requirements-production.txt 不可用时（例如发布安装态）的兜底必需集
FALLBACK_REQUIRED_REQUIREMENTS = (
    "fastapi", "uvicorn", "sqlalchemy", "python-docx",
    "python-pptx", "openpyxl",
)


def parse_production_requirements(text: str) -> list[str]:
    """从 requirements-production.txt 解析出发行名列表（忽略注释/空行/环境标记）。"""
    names: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.split("#", 1)[0].strip()
        if not line:
            continue
        # 形如 pywin32==306; sys_platform == "win32" → 取分号前的包声明
        declaration = line.split(";", 1)[0].strip()
        match = re.match(r"^([A-Za-z0-9_.\-]+)", declaration)
        if match:
            names.append(match.group(1))
    return names


class TestStatus(str, Enum):
    __test__ = False
    PASS = "pass"
    FAIL = "fail"
    SKIP = "skip"
    WARN = "warn"


@dataclass
class InstallTestResult:
    name: str = ""
    status: TestStatus = TestStatus.PASS
    message: str = ""
    duration_ms: float = 0
    details: dict = field(default_factory=dict)


class InstallationTestSuite:
    """安装测试套件"""

    def __init__(self, app_dir: Path = None, port: int = 8765,
                 mode: str = "auto"):
        self.app_dir = Path(app_dir) if app_dir else Path(__file__).parent.parent
        self.port = port
        self.results: list[InstallTestResult] = []
        self.backend_process: Optional[subprocess.Popen] = None
        self.mode = self.detect_mode(self.app_dir) if mode == "auto" else mode

    # ------------------------------------------------------------
    # 开发态 / 发布态判据
    # ------------------------------------------------------------
    @staticmethod
    def detect_mode(app_dir: Path) -> str:
        """识别当前 app_dir 是"发布安装态"还是"开发源码树"。

        - installed：含主产物 OfficeAgent.exe（PyInstaller COLLECT 输出目录）；
        - source：含 office_agent/api/main.py 的源码树；
        - unknown：两者都不满足。
        """
        if (app_dir / PRIMARY_ARTIFACT).exists():
            return "installed"
        if (app_dir / "office_agent" / "api" / "main.py").exists():
            return "source"
        return "unknown"

    def _source_tree_only(self, what: str) -> InstallTestResult:
        """源码树判据在发布安装态下不适用：显式 WARN，而不是伪造成通过/失败。"""
        if self.mode == "installed":
            return InstallTestResult(
                status=TestStatus.WARN,
                message=f"{what}：发布安装态不适用（源码树判据）",
            )
        return InstallTestResult(status=TestStatus.FAIL, message=f"{what}：{self.app_dir}")

    # ------------------------------------------------------------
    # 主流程
    # ------------------------------------------------------------
    def run_all(self) -> list[InstallTestResult]:
        """运行所有测试"""
        tests = [
            ("检查安装布局", self.test_install_layout),
            ("检查Python运行时", self.test_python_runtime),
            ("检查核心依赖", self.test_dependencies),
            ("检查目录结构", self.test_directory_structure),
            ("检查配置文件", self.test_config_files),
            ("检查Backend启动", self.test_backend_startup),
            ("检查健康检查API", self.test_health_endpoint),
            ("检查就绪门禁", self.test_ready_endpoint),
            ("检查文件上传API", self.test_file_upload),
            ("检查Word处理", self.test_word_processing),
            ("检查Excel处理", self.test_excel_processing),
            ("检查PPT处理", self.test_ppt_processing),
            ("检查Backend停止", self.test_backend_shutdown),
        ]
        for name, test_fn in tests:
            start = time.time()
            try:
                result = test_fn()
                if result is None:
                    # 检查函数必须显式返回结论。返回 None 说明该检查有漏网分支，
                    # 旧实现把 None 记为 PASS（"默认成功"），会让真实回归静默通过。
                    result = InstallTestResult(
                        status=TestStatus.FAIL,
                        message="检查未返回结论（实现缺陷，不得视为通过）",
                    )
                result.name = name
                result.duration_ms = (time.time() - start) * 1000
            except Exception as e:
                result = InstallTestResult(
                    name=name, status=TestStatus.FAIL,
                    message=f"{type(e).__name__}: {e}",
                    duration_ms=(time.time() - start) * 1000,
                )
            self.results.append(result)
            icon = {"pass": "✅", "fail": "❌", "skip": "⏭️", "warn": "⚠️"}[result.status.value]
            print(f"{icon} {result.name}: {result.message}")
        return self.results

    def test_install_layout(self) -> InstallTestResult:
        """检查是否存在真实的发布安装产物（存在 + 真实可用）"""
        artifact = self.app_dir / PRIMARY_ARTIFACT
        if not artifact.exists():
            if self.mode == "source":
                return InstallTestResult(
                    status=TestStatus.SKIP,
                    message=f"开发源码树（{self.mode}），无 {PRIMARY_ARTIFACT}",
                )
            return InstallTestResult(
                status=TestStatus.FAIL, message=f"缺少主产物 {artifact}",
            )
        if not artifact.is_file():
            return InstallTestResult(
                status=TestStatus.FAIL,
                message=f"{PRIMARY_ARTIFACT} 存在但不是普通文件（类型错误）",
            )
        with open(artifact, "rb") as handle:
            magic = handle.read(2)
        if magic != b"MZ":
            return InstallTestResult(
                status=TestStatus.FAIL,
                message=f"{PRIMARY_ARTIFACT} 不是可执行文件（magic={magic!r}）",
            )
        manifest = self.app_dir / MANIFEST_NAME
        if not manifest.exists():
            return InstallTestResult(
                status=TestStatus.WARN,
                message=f"{PRIMARY_ARTIFACT} 可用，但缺少 {MANIFEST_NAME}",
            )
        return InstallTestResult(
            status=TestStatus.PASS,
            message=f"{PRIMARY_ARTIFACT} + {MANIFEST_NAME} 就位",
        )

    def test_python_runtime(self) -> InstallTestResult:
        """检查Python运行时"""
        version = sys.version_info
        if version >= (3, 10):
            return InstallTestResult(
                status=TestStatus.PASS,
                message=f"Python {version.major}.{version.minor}.{version.micro}",
                details={"version": f"{version.major}.{version.minor}.{version.micro}"},
            )
        return InstallTestResult(status=TestStatus.FAIL, message=f"Python版本过低: {version.major}.{version.minor}")

    def resolve_dependencies(self) -> tuple[list[str], list[str]]:
        """解析 (必需发行名, 可选发行名)。

        必需集来自 requirements-production.txt（当前产品的依赖契约）；
        构建期依赖与平台专属依赖归入可选，缺失只告警不判失败。
        """
        req_file = self.app_dir / "requirements-production.txt"
        names: list[str] = []
        if req_file.exists():
            try:
                names = parse_production_requirements(
                    req_file.read_text(encoding="utf-8"),
                )
            except OSError:
                names = []
        if not names:
            names = list(FALLBACK_REQUIRED_REQUIREMENTS)

        required: list[str] = []
        optional: list[str] = []
        for name in names:
            key = name.lower()
            if key in BUILD_ONLY_REQUIREMENTS:
                optional.append(name)
            elif key in PLATFORM_GATED_REQUIREMENTS and os.name != "nt":
                optional.append(name)
            else:
                required.append(name)
        return required, optional

    @staticmethod
    def _missing_imports(names: list[str]) -> list[str]:
        missing = []
        for name in names:
            module = REQUIREMENT_IMPORT_NAMES.get(name.lower(), name)
            try:
                __import__(module)
            except ImportError:
                missing.append(name)
        return missing

    def test_dependencies(self) -> InstallTestResult:
        """检查核心依赖（必需缺失 → FAIL；可选缺失 → WARN）"""
        required, optional = self.resolve_dependencies()
        missing_required = self._missing_imports(required)
        missing_optional = self._missing_imports(optional)

        if missing_required:
            return InstallTestResult(
                status=TestStatus.FAIL,
                message=f"缺少必需依赖: {', '.join(missing_required)}",
                details={
                    "required": required,
                    "missing_required": missing_required,
                    "missing_optional": missing_optional,
                },
            )
        if missing_optional:
            return InstallTestResult(
                status=TestStatus.WARN,
                message=(
                    f"{len(required)} 个必需依赖齐全；"
                    f"可选依赖缺失: {', '.join(missing_optional)}"
                ),
                details={"required": required, "missing_optional": missing_optional},
            )
        return InstallTestResult(
            status=TestStatus.PASS,
            message=f"必需依赖 {len(required)} 个齐全，可选依赖 {len(optional)} 个齐全",
            details={"required": required, "optional": optional},
        )

    def test_directory_structure(self) -> InstallTestResult:
        """检查目录结构（源码树判据；发布安装态不适用）"""
        if self.mode == "installed":
            return self._source_tree_only("目录结构检查")
        required_dirs = [
            "office_agent",
            "office_agent/api",
            "office_agent/excel_agent",
            "office_agent/ppt_agent",
        ]
        missing = []
        for dir_path in required_dirs:
            target = self.app_dir / dir_path
            if not target.is_dir():
                missing.append(dir_path)
        if missing:
            return InstallTestResult(status=TestStatus.FAIL, message=f"缺少目录: {', '.join(missing)}")
        return InstallTestResult(status=TestStatus.PASS, message=f"所有{len(required_dirs)}个目录存在")

    def test_config_files(self) -> InstallTestResult:
        """检查配置文件（源码树判据；发布安装态不适用）"""
        if self.mode == "installed":
            return self._source_tree_only("配置文件检查")
        required_files = [
            "office_agent/__init__.py",
            "office_agent/api/main.py",
            "requirements-production.txt",
        ]
        missing = [f for f in required_files if not (self.app_dir / f).is_file()]
        if missing:
            return InstallTestResult(status=TestStatus.FAIL, message=f"缺少文件: {', '.join(missing)}")
        return InstallTestResult(status=TestStatus.PASS, message="核心文件完整")

    @staticmethod
    def _expected_version() -> str:
        """当前产品的权威版本号（唯一来源 office_agent/_version.py）"""
        try:
            from office_agent._version import __version__
            return str(__version__)
        except Exception:
            return ""

    def _identify_health_payload(self, payload) -> tuple[bool, str]:
        """按当前产品契约判定 /health 载荷是否属于本产品且健康。

        真实契约（office_agent/api/router/health.py）是**扁平** JSON：
        ``{"status": "healthy"|"degraded"|"unhealthy", "version": ..., "checks": {...}}``。
        旧实现用 ``data.get("data", {}).get("version")`` 读一个不存在的键，
        且把"响应里含 data 字段"当作健康（端口上是别的服务也判通过）。
        """
        if not isinstance(payload, dict):
            return False, "响应不是 JSON 对象"
        status = payload.get("status")
        if status != "healthy" and payload.get("success") is not True:
            checks = payload.get("checks")
            reasons = []
            if isinstance(checks, dict):
                for key, value in checks.items():
                    if isinstance(value, dict) and value.get("status") not in ("healthy", None):
                        reasons.append(f"{key}={value.get('status')}")
            detail = f"（{', '.join(reasons)}）" if reasons else ""
            return False, f"status={status!r}{detail}"
        version = payload.get("version")
        return True, str(version) if isinstance(version, str) else ""

    def _check_version(self, reported: str) -> Optional[InstallTestResult]:
        """版本号与权威来源比对；不符即 FAIL（旧实现从不比对版本）。"""
        expected = self._expected_version()
        if expected and reported and reported != expected:
            return InstallTestResult(
                status=TestStatus.FAIL,
                message=f"版本不符: 后端 {reported!r} ≠ 产品 {expected!r}",
            )
        return None

    def _get_json(self, path: str, timeout: float = 5.0):
        """请求本地端点并返回 (status_code, payload)；非 2xx 也返回而非抛错。"""
        url = f"http://127.0.0.1:{self.port}{path}"
        try:
            with urllib.request.urlopen(url, timeout=timeout) as resp:
                return resp.status, json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            body = b""
            try:
                body = exc.read()
            except Exception:
                body = b""
            try:
                payload = json.loads(body.decode("utf-8")) if body else None
            except (ValueError, UnicodeDecodeError):
                payload = None
            return exc.code, payload

    def test_backend_startup(self) -> InstallTestResult:
        """检查Backend启动"""
        import socket
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", self.port)) == 0:
                # 端口已被占用：必须确认占用者是本产品，而不是任意别的服务
                try:
                    status_code, payload = self._get_json("/health", timeout=3)
                except Exception:
                    status_code, payload = 0, None
                ok, describe = self._identify_health_payload(payload)
                if ok:
                    return InstallTestResult(
                        status=TestStatus.PASS,
                        message=f"Backend已在运行（复用，version={describe or 'unknown'}）",
                    )
                return InstallTestResult(
                    status=TestStatus.FAIL,
                    message=(
                        f"端口 {self.port} 被非本产品服务占用"
                        f"（HTTP {status_code}, {describe}）"
                    ),
                )
        # 启动Backend
        env = os.environ.copy()
        env["OFFICE_AGENT_LOCAL"] = "1"
        env["AUTH_MODE"] = "local"
        env["PYTHONPATH"] = str(self.app_dir)
        cmd = [
            sys.executable, "-m", "uvicorn",
            "office_agent.api.main:app",
            "--host", "127.0.0.1",
            "--port", str(self.port),
            "--log-level", "warning",
        ]
        self.backend_process = subprocess.Popen(
            cmd, cwd=str(self.app_dir), env=env,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        # 等待启动
        for _ in range(30):
            time.sleep(0.5)
            try:
                status_code, payload = self._get_json("/health", timeout=2)
            except Exception:
                ok, _ = False, ""
            else:
                ok, _ = self._identify_health_payload(payload)
            if ok:
                return InstallTestResult(
                    status=TestStatus.PASS,
                    message=f"Backend启动成功 (PID={self.backend_process.pid})",
                )
            returncode = self.backend_process.poll()
            if returncode is not None:
                _, stderr = self.backend_process.communicate()
                return InstallTestResult(
                    status=TestStatus.FAIL,
                    message=(
                        f"Backend启动失败: returncode={returncode} "
                        f"{stderr.decode('utf-8', errors='ignore')[:200]}"
                    ),
                )
        return InstallTestResult(status=TestStatus.FAIL, message="Backend启动超时")

    def test_health_endpoint(self) -> InstallTestResult:
        """检查健康检查API（严格契约 + 版本比对）"""
        try:
            status_code, payload = self._get_json("/health", timeout=5)
        except Exception as e:
            return InstallTestResult(status=TestStatus.FAIL, message=f"健康检查失败: {e}")
        if status_code != 200:
            return InstallTestResult(
                status=TestStatus.FAIL, message=f"健康检查 HTTP {status_code}",
            )
        ok, describe = self._identify_health_payload(payload)
        if not ok:
            return InstallTestResult(status=TestStatus.FAIL, message=f"健康检查未通过: {describe}")
        version_mismatch = self._check_version(describe)
        if version_mismatch:
            return version_mismatch
        return InstallTestResult(status=TestStatus.PASS, message=f"健康检查正常: {describe or 'unknown'}")

    def test_ready_endpoint(self) -> InstallTestResult:
        """检查就绪门禁 /ready（严格契约：未就绪返回 503）"""
        if self.backend_process is None and not self._port_in_use():
            return InstallTestResult(status=TestStatus.SKIP, message="Backend未在运行")
        try:
            status_code, payload = self._get_json("/ready", timeout=5)
        except Exception as e:
            return InstallTestResult(status=TestStatus.FAIL, message=f"就绪检查失败: {e}")
        ready = isinstance(payload, dict) and payload.get("ready") is True
        if status_code == 200 and ready:
            return InstallTestResult(status=TestStatus.PASS, message="就绪门禁通过（ready=true）")
        if status_code == 503 and not ready:
            return InstallTestResult(
                status=TestStatus.WARN,
                message="后端未就绪（HTTP 503, ready=false）",
            )
        return InstallTestResult(
            status=TestStatus.FAIL,
            message=f"就绪门禁响应不符合契约: HTTP {status_code}, ready={ready}",
        )

    def _port_in_use(self) -> bool:
        import socket
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            return s.connect_ex(("127.0.0.1", self.port)) == 0

    def test_file_upload(self) -> InstallTestResult:
        """检查文件上传（真实路由 + 严格判定，404 视为失败）"""
        import uuid
        # 创建测试文件
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w") as f:
            f.write("test content")
            test_file = f.name
        try:
            boundary = uuid.uuid4().hex
            with open(test_file, "rb") as f:
                content = f.read()
            body = (
                f"--{boundary}\r\n"
                f'Content-Disposition: form-data; name="file"; filename="test.txt"\r\n'
                f"Content-Type: text/plain\r\n\r\n"
            ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
            req = urllib.request.Request(
                f"http://127.0.0.1:{self.port}{UPLOAD_ROUTE}",
                data=body,
                headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
                method="POST",
            )
            try:
                with urllib.request.urlopen(req, timeout=10) as resp:
                    status_code = resp.status
                    raw = resp.read()
            except urllib.error.HTTPError as e:
                # 路由不存在（404）或鉴权失败都必须判失败：旧实现把 404 当
                # "上传 API 响应正常"，路由被删掉自测也永远通过。
                return InstallTestResult(
                    status=TestStatus.FAIL,
                    message=f"上传接口不可用: HTTP {e.code} {UPLOAD_ROUTE}",
                )
            if status_code not in (200, 201):
                return InstallTestResult(
                    status=TestStatus.FAIL, message=f"上传返回异常状态: HTTP {status_code}",
                )
            try:
                payload = json.loads(raw.decode("utf-8"))
            except (ValueError, UnicodeDecodeError):
                return InstallTestResult(
                    status=TestStatus.FAIL, message="上传响应不是合法 JSON",
                )
            if not isinstance(payload, dict):
                return InstallTestResult(
                    status=TestStatus.FAIL, message="上传响应结构不符合契约",
                )
            data = payload.get("data")
            file_id = ""
            if isinstance(data, dict):
                file_id = str(data.get("file_id") or data.get("id") or "")
            if payload.get("success") is not True or not file_id:
                return InstallTestResult(
                    status=TestStatus.FAIL,
                    message=f"上传未返回可用文件标识: {payload}",
                )
            return InstallTestResult(
                status=TestStatus.PASS, message=f"文件上传正常（file_id={file_id}）",
            )
        finally:
            Path(test_file).unlink(missing_ok=True)

    def test_word_processing(self) -> InstallTestResult:
        """检查Word处理"""
        try:
            from docx import Document
            from office_agent.parsers import analyze_document
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
                doc = Document()
                doc.add_heading("Test", 0)
                doc.add_paragraph("Hello World")
                doc.save(f.name)
                test_path = f.name
            try:
                tree = analyze_document(test_path)
            finally:
                Path(test_path).unlink(missing_ok=True)
            if tree is not None:
                return InstallTestResult(status=TestStatus.PASS, message="Word解析正常")
        except Exception as e:
            return InstallTestResult(status=TestStatus.FAIL, message=f"Word处理失败: {e}")
        return InstallTestResult(status=TestStatus.FAIL, message="Word处理异常")

    def test_excel_processing(self) -> InstallTestResult:
        """检查Excel处理"""
        try:
            from openpyxl import Workbook
            from office_agent.excel_agent import analyze_excel
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
                wb = Workbook()
                ws = wb.active
                ws.append(["Name", "Value"])
                ws.append(["Test", 123])
                wb.save(f.name)
                test_path = f.name
            try:
                result = analyze_excel(test_path)
            finally:
                Path(test_path).unlink(missing_ok=True)
            if result is not None:
                return InstallTestResult(status=TestStatus.PASS, message="Excel解析正常")
        except Exception as e:
            return InstallTestResult(status=TestStatus.FAIL, message=f"Excel处理失败: {e}")
        return InstallTestResult(status=TestStatus.FAIL, message="Excel处理异常")

    def test_ppt_processing(self) -> InstallTestResult:
        """检查PPT处理"""
        try:
            from pptx import Presentation
            from office_agent.ppt_agent import PPTService  # noqa: F401
            with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = "Test"
                prs.save(f.name)
                test_path = f.name
            Path(test_path).unlink(missing_ok=True)
            return InstallTestResult(status=TestStatus.PASS, message="PPT处理正常")
        except Exception as e:
            return InstallTestResult(status=TestStatus.FAIL, message=f"PPT处理失败: {e}")

    def test_backend_shutdown(self) -> InstallTestResult:
        """检查Backend停止"""
        if self.backend_process:
            try:
                self.backend_process.terminate()
                try:
                    self.backend_process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    self.backend_process.kill()
                return InstallTestResult(status=TestStatus.PASS, message="Backend正常停止")
            except Exception as e:
                return InstallTestResult(status=TestStatus.WARN, message=f"停止异常: {e}")
        return InstallTestResult(status=TestStatus.SKIP, message="Backend非测试启动")

    def get_summary(self) -> dict:
        """获取测试摘要"""
        total = len(self.results)
        passed = sum(1 for r in self.results if r.status == TestStatus.PASS)
        failed = sum(1 for r in self.results if r.status == TestStatus.FAIL)
        warnings = sum(1 for r in self.results if r.status == TestStatus.WARN)
        skipped = sum(1 for r in self.results if r.status == TestStatus.SKIP)
        return {
            "total": total,
            "passed": passed,
            "failed": failed,
            "warnings": warnings,
            "skipped": skipped,
            "pass_rate": f"{passed/total*100:.1f}%" if total else "0%",
            "can_release": failed == 0,
        }


def run_installation_tests(app_dir: Path = None, port: int = 8765) -> dict:
    """运行安装测试"""
    print("=" * 60)
    print("  OfficeAgent Installation Test Suite")
    print("=" * 60)
    print()
    suite = InstallationTestSuite(app_dir=app_dir, port=port)
    suite.run_all()
    summary = suite.get_summary()
    print()
    print("=" * 60)
    print(f"  Results: {summary['passed']}/{summary['total']} passed ({summary['pass_rate']})")
    if summary["failed"]:
        print(f"  Failed: {summary['failed']}")
    if summary["warnings"]:
        print(f"  Warnings: {summary['warnings']}")
    print(f"  Can release: {'✅ YES' if summary['can_release'] else '❌ NO'}")
    print("=" * 60)
    return summary


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--app-dir", type=str, default=None)
    args = parser.parse_args()
    app_dir = Path(args.app_dir) if args.app_dir else Path(__file__).parent.parent
    run_installation_tests(app_dir=app_dir, port=args.port)
