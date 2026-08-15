"""
Word Agent 评估器
"""
from pathlib import Path
from .evaluator import BaseEvaluator, EvalResult, EvalType, EvalStatus, EvalMetric


class WordEvaluator(BaseEvaluator):
    """Word Agent 输出质量评估"""

    def evaluate_format(self, input_path: str | Path, output_path: str | Path,
                        expected_font: str = None) -> EvalResult:
        """评估Word格式输出质量"""
        start = time.time()
        metrics = []
        try:
            from docx import Document
            doc = Document(str(output_path))

            # 1. 文件可打开
            metrics.append(EvalMetric("文件完整性", 100, weight=2.0))

            # 2. 段落数量
            para_count = len(doc.paragraphs)
            para_score = min(100, para_count * 10) if para_count > 0 else 0
            metrics.append(EvalMetric("段落完整性", para_score, weight=1.0,
                                      details=f"{para_count}个段落"))

            # 3. 标题识别
            headings = [p for p in doc.paragraphs if p.style and "Heading" in str(p.style)]
            heading_score = min(100, len(headings) * 25) if headings else 50
            metrics.append(EvalMetric("标题识别", heading_score, weight=1.5,
                                      details=f"{len(headings)}个标题"))

            # 4. 字体一致性
            if expected_font:
                fonts_ok = 0
                fonts_total = 0
                for p in doc.paragraphs[:20]:
                    for run in p.runs:
                        fonts_total += 1
                        if run.font.name == expected_font:
                            fonts_ok += 1
                font_score = (fonts_ok / fonts_total * 100) if fonts_total > 0 else 50
                metrics.append(EvalMetric("字体一致性", font_score, weight=1.5,
                                          details=f"{fonts_ok}/{fonts_total}"))

            # 5. 表格
            tables = doc.tables
            table_score = 100 if tables else 80
            metrics.append(EvalMetric("表格完整性", table_score, weight=1.0,
                                      details=f"{len(tables)}个表格"))

            status = EvalStatus.PASS if all(m.score >= 60 for m in metrics) else EvalStatus.PARTIAL

        except Exception as e:
            metrics.append(EvalMetric("错误", 0, details=str(e)))
            status = EvalStatus.ERROR

        return EvalResult(
            eval_type=EvalType.WORD,
            test_name="word_format_eval",
            status=status,
            metrics=metrics,
            duration=time.time() - start,
        )

    def evaluate_quality(self, doc_path: str | Path) -> EvalResult:
        """评估Word文档质量"""
        start = time.time()
        metrics = []
        try:
            from docx import Document
            doc = Document(str(doc_path))

            # 内容完整性
            total_text = sum(len(p.text) for p in doc.paragraphs)
            content_score = min(100, total_text / 10) if total_text > 0 else 0
            metrics.append(EvalMetric("内容完整性", content_score, weight=2.0,
                                      details=f"{total_text}字符"))

            # 结构合理性
            has_title = any(p.style and "Heading 1" in str(p.style) for p in doc.paragraphs)
            structure_score = 100 if has_title else 60
            metrics.append(EvalMetric("结构合理性", structure_score, weight=1.5))

            # 格式规范性
            empty_paras = sum(1 for p in doc.paragraphs if not p.text.strip())
            total_paras = len(doc.paragraphs)
            empty_ratio = empty_paras / total_paras if total_paras > 0 else 1
            format_score = max(0, 100 - empty_ratio * 50)
            metrics.append(EvalMetric("格式规范性", format_score, weight=1.0,
                                      details=f"空段落率{empty_ratio:.1%}"))

            status = EvalStatus.PASS if all(m.score >= 60 for m in metrics) else EvalStatus.PARTIAL

        except Exception as e:
            metrics.append(EvalMetric("错误", 0, details=str(e)))
            status = EvalStatus.ERROR

        return EvalResult(
            eval_type=EvalType.WORD,
            test_name="word_quality_eval",
            status=status,
            metrics=metrics,
            duration=time.time() - start,
        )


import time
