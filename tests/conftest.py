"""
pytest 共享 fixtures
"""
import sys
import tempfile
import shutil
from pathlib import Path

import pytest

# 确保项目根目录在 path 中
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


@pytest.fixture
def temp_dir():
    """临时目录"""
    d = tempfile.mkdtemp(prefix="office_agent_test_")
    yield Path(d)
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def sample_docx(temp_dir):
    """创建测试用 Word 文档"""
    from docx import Document
    doc = Document()
    doc.add_heading("测试文档", level=1)
    doc.add_paragraph("这是第一段正文内容。")
    doc.add_heading("第一章", level=2)
    doc.add_paragraph("这是第一章的内容。")
    doc.add_paragraph("这是第二段，包含一些格式测试。")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for i in range(3):
        for j in range(3):
            table.cell(i, j).text = f"单元格{i}{j}"
    path = temp_dir / "sample.docx"
    doc.save(str(path))
    yield path


@pytest.fixture
def sample_xlsx(temp_dir):
    """创建测试用 Excel 文件"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "销售数据"
    ws.append(["月份", "销售额", "成本", "利润"])
    data = [
        [1, 10000, 6000, 4000],
        [2, 12000, 7000, 5000],
        [3, 15000, 8000, 7000],
        [4, 11000, 6500, 4500],
        [5, 13000, 7500, 5500],
        [6, 18000, 9000, 9000],
    ]
    for row in data:
        ws.append(row)
    ws2 = wb.create_sheet("库存")
    ws2.append(["产品", "数量", "单价"])
    ws2.append(["A", 100, 50])
    ws2.append(["B", 200, 30])
    path = temp_dir / "sample.xlsx"
    wb.save(str(path))
    yield path


@pytest.fixture
def sample_pptx(temp_dir):
    """创建测试用 PPT 文件"""
    from pptx import Presentation
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "测试演示文稿"
    slide.placeholders[1].text = "副标题"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "内容页"
    path = temp_dir / "sample.pptx"
    prs.save(str(path))
    yield path


@pytest.fixture
def security_config():
    """安全配置"""
    from office_agent.security import SecurityConfig
    return SecurityConfig(
        jwt_secret_key="test-secret-key-for-testing-only",
        access_token_expire=3600,
        max_file_size=10 * 1024 * 1024,
        sandbox_timeout=10,
    )


class _FakeSemanticEmbedder:
    """Deterministic semantic backend used by RAG regression tests."""

    model_id = "semantic-test-v1"
    version = "test"
    semantic = True
    dimension = 4

    def embed(self, texts):
        return [[0.1, 0.2, 0.3, 0.4] for _ in texts]

    def embed_query(self, text):
        return [0.1, 0.2, 0.3, 0.4]


@pytest.fixture
def fake_semantic_embedder(monkeypatch):
    """Replace the RAG semantic factory with a deterministic backend."""
    from office_agent.knowledge_base import embeddings as embeddings_module

    embedder = _FakeSemanticEmbedder()
    monkeypatch.setattr(
        embeddings_module, "create_semantic_embedder", lambda *a, **kw: embedder,
    )
    return embedder
