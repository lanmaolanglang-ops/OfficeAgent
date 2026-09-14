# -*- coding: utf-8 -*-
"""P3 batch C regression tests (PPT)."""
from office_agent.ppt_agent.models import (
    PPTOutline, SlideContent, items_per_page, continuation_title,
)


# ---------- P3-47: single-source capacity + continuation title ----------
def test_p3_47_shared_capacity_and_title():
    assert items_per_page("toc") == 6
    assert items_per_page("content") == 8
    assert items_per_page("two_column") == 6
    assert items_per_page("timeline") == 5
    assert items_per_page("data_cards") == 4
    assert continuation_title("标题", 0) == "标题"
    assert continuation_title("标题", 1) == "标题（续 2）"


def test_p3_47_fix_outline_splits_all_overflow_with_shared_title():
    from office_agent.ppt_agent.quality_checker import PPTQualityChecker
    outline = PPTOutline(title="t")
    outline.add_slide(SlideContent(layout="cover", title="c"))
    outline.add_slide(SlideContent(layout="content", title="要点页",
                                   bullets=[f"b{i}" for i in range(20)]))
    qc = PPTQualityChecker()
    fixed = qc.fix_outline(outline, qc.check_outline(outline))
    pages = [s for s in fixed.slides if s.layout == "content"]
    assert [len(s.bullets) for s in pages] == [8, 8, 4]
    assert pages[0].title == "要点页"
    assert pages[1].title == "要点页（续 2）"
    assert pages[2].title == "要点页（续 3）"


# ---------- P3-49: AI parser maps chart / quote / body_text ----------
def test_p3_49_ai_parser_maps_chart_quote_body():
    from office_agent.ppt_agent.content_planner import ContentPlanner
    planner = ContentPlanner.__new__(ContentPlanner)
    data = {"slides": [
        {"layout": "chart", "title": "趋势", "chart_type": "pie",
         "chart_categories": ["A", "B"], "chart_series": [("s", [1, 2])]},
        {"layout": "quote", "title": "金句", "quote_text": "专注",
         "quote_source": "书"},
        {"layout": "content", "title": "正文", "body_text": "整段"},
    ]}
    out = planner._parse_ai_outline(data, "d", "professional")
    chart = next(s for s in out.slides if s.layout == "chart")
    assert chart.chart_type == "pie"
    assert chart.chart_categories == ["A", "B"]
    assert chart.chart_series == [("s", [1, 2])]
    quote = next(s for s in out.slides if s.layout == "quote")
    assert quote.quote_text == "专注" and quote.quote_source == "书"
    assert next(s for s in out.slides if s.title == "正文").body_text == "整段"


# ---------- P3-50: minimal preset uses a Windows-bundled CJK font ----------
def test_p3_50_minimal_font_available_on_windows():
    from office_agent.ppt_agent.slide_designer import StylePresets
    _, fonts = StylePresets.minimal()
    assert fonts.title_cn == "微软雅黑"
    assert fonts.body_cn == "微软雅黑"


# ---------- P3-44/45: only valid hex; readable foreground by luminance ----------
def test_p3_44_45_color_helpers():
    from office_agent.ppt_agent.template_analyzer import (
        ThemeColorInfo, _readable_on_dark, _is_hex6,
    )
    assert _readable_on_dark("#1F4E79") == "#FFFFFF"
    assert _readable_on_dark("#F2F2F2") == "#000000"
    assert ThemeColorInfo(dk2="#F2F2F2").to_color_scheme().text_light == "#000000"
    assert not _is_hex6("lt1")
    assert _is_hex6("FFC000")


# ---------- P3-37/38/39/42 + real PPTX ----------
def test_p3_37_38_39_42_real_pptx(tmp_path):
    from office_agent.ppt_agent.ppt_service import PPTService
    outline = PPTOutline(title="真实", author="作者张三")
    outline.add_slide(SlideContent(layout="cover", title="封面",
                                   notes="notes不应覆盖作者"))
    outline.add_slide(SlideContent(layout="section", title="第一章"))
    outline.add_slide(SlideContent(layout="section", title="第二章"))
    outline.add_slide(SlideContent(
        layout="data_cards", title="指标",
        data=[("指标", i, "个") for i in range(1, 10)]))  # 9 -> 4+4+1
    out = tmp_path / "c.pptx"
    res = PPTService().generate(outline, str(out))
    assert res.success, res.message
    assert res.slide_count == 6  # cover + 2 section + 3 card pages

    from pptx import Presentation
    prs = Presentation(str(out))
    texts = []
    for sl in prs.slides:
        for shp in sl.shapes:
            if shp.has_text_frame:
                texts.append(shp.text_frame.text)
    joined = "\n".join(texts)
    assert "作者张三" in joined
    assert "notes不应覆盖作者" not in joined
    assert "01" in joined and "02" in joined
    assert "单击" not in joined and "Click to" not in joined


# ---------- P3-52: pie chart colors each slice ----------
def test_p3_52_pie_colors_each_point(tmp_path):
    from office_agent.ppt_agent.ppt_service import PPTService
    from pptx import Presentation
    from pptx.oxml.ns import qn
    outline = PPTOutline(title="pie")
    outline.add_slide(SlideContent(layout="cover", title="c"))
    outline.add_slide(SlideContent(
        layout="chart", title="占比", chart_type="pie",
        chart_categories=["A", "B", "C", "D"],
        chart_series=[("销量", [10, 20, 30, 40])]))
    out = tmp_path / "pie.pptx"
    assert PPTService().generate(outline, str(out)).success
    prs = Presentation(str(out))
    checked = False
    for sl in prs.slides:
        for shp in sl.shapes:
            if shp.has_chart:
                series = shp.chart.series[0]
                per_point = []
                for pt in series.points:
                    per_point.append(len(pt._element.findall(".//" + qn("a:srgbClr"))))
                assert len(per_point) == 4
                assert all(c >= 1 for c in per_point)
                checked = True
    assert checked
