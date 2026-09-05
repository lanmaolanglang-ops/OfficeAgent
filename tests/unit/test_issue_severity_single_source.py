"""
IssueSeverity 枚举字面量剩余债务收尾专项

钉住五件事：
1. 合法值表与 values()/is_valid()/normalize() 权威助手语义；
2. 三个质量问题数据类（Word QualityIssue / Excel ExcelQualityIssue /
   PPT PPTQualityIssue）在构造边界规范化合法值、拒绝未知值（不静默漂移）；
3. JSON/API 序列化保持原字符串值兼容；
4. 完整检查流程产出的 severity 全部合法（历史字符串数据兼容路径）；
5. AST 源码守卫：Excel/PPT 质量域不得重新引入裸 severity 字面量。
"""
import ast
import inspect

import pytest

from office_agent.quality.checker import IssueSeverity, QualityIssue
from office_agent.excel_agent.models import ExcelQualityIssue
from office_agent.ppt_agent.quality_checker import PPTQualityIssue


# ------------------------------------------------------------------ enum helpers

class TestIssueSeverityHelpers:
    def test_values_are_stable_legacy_strings(self):
        # 序列化兼容锚点：值表与历史字符串完全一致，顺序即枚举声明顺序
        assert IssueSeverity.values() == ["error", "warning", "info"]

    @pytest.mark.parametrize("value", ["error", "warning", "info"])
    def test_is_valid_accepts_all_legal_values(self, value):
        assert IssueSeverity.is_valid(value) is True

    @pytest.mark.parametrize("value", [
        "critical", "high", "medium", "low", "moderate", "", "errors",
        None, 1, ["error"],
    ])
    def test_is_valid_rejects_unknown_values(self, value):
        assert IssueSeverity.is_valid(value) is False

    @pytest.mark.parametrize("raw,expected", [
        ("ERROR", "error"),
        (" Warning ", "warning"),
        ("info", "info"),
    ])
    def test_normalize_accepts_case_and_whitespace(self, raw, expected):
        assert IssueSeverity.normalize(raw) == expected

    def test_normalize_rejects_unknown_without_fallback(self):
        with pytest.raises(ValueError):
            IssueSeverity.normalize("critical")

    def test_normalize_unknown_with_explicit_fallback(self):
        assert IssueSeverity.normalize("critical",
                                       fallback=IssueSeverity.WARNING.value) == "warning"


# ------------------------------------------------------------------ construction boundary

class TestConstructionBoundary:
    @pytest.mark.parametrize("cls,kwargs", [
        (QualityIssue, {"type": "font", "message": "m"}),
        (ExcelQualityIssue, {"issue_type": "formula", "message": "m"}),
        (PPTQualityIssue, {"slide_index": -1, "issue_type": "template", "message": "m"}),
    ])
    def test_legal_values_normalized_at_construction(self, cls, kwargs):
        issue = cls(severity=" Warning ", **kwargs)
        assert issue.severity == "warning"

    @pytest.mark.parametrize("cls,kwargs", [
        (QualityIssue, {"type": "font", "message": "m"}),
        (ExcelQualityIssue, {"issue_type": "formula", "message": "m"}),
        (PPTQualityIssue, {"slide_index": -1, "issue_type": "template", "message": "m"}),
    ])
    def test_unknown_severity_rejected_not_silent(self, cls, kwargs):
        with pytest.raises(ValueError):
            cls(severity="critical", **kwargs)

    def test_excel_default_severity_is_enum_value(self):
        assert ExcelQualityIssue().severity == IssueSeverity.WARNING.value


# ------------------------------------------------------------------ serialization compat

class TestSerializationCompat:
    def test_word_issue_to_dict_stays_plain_strings(self):
        issue = QualityIssue(type="font", severity="error", message="m")
        data = issue.to_dict()
        assert data["severity"] == "error"
        assert isinstance(data["severity"], str)

    def test_excel_issue_to_dict_stays_plain_strings(self):
        issue = ExcelQualityIssue(severity="info", message="m")
        assert issue.to_dict()["severity"] == "info"

    def test_ppt_issue_to_dict_stays_plain_strings(self):
        issue = PPTQualityIssue(slide_index=0, issue_type="content",
                                severity="warning", message="m")
        assert issue.to_dict()["severity"] == "warning"

    def test_historical_string_values_still_accepted(self):
        # 历史数据兼容：以旧裸字符串构造的对象与枚举值构造完全等价
        legacy = ExcelQualityIssue(severity="error", message="m")
        modern = ExcelQualityIssue(severity=IssueSeverity.ERROR.value, message="m")
        assert legacy.to_dict() == modern.to_dict()


# ------------------------------------------------------------------ full-pipeline legality

class TestPipelineSeverityLegality:
    def test_excel_check_report_severities_all_legal(self, tmp_path):
        from openpyxl import Workbook
        from office_agent.excel_agent.quality_checker import ExcelQualityChecker

        wb = Workbook()
        ws = wb.active
        ws.append(["col"])
        ws.append(["=1/0"])
        wb.create_sheet("empty_sheet")
        path = tmp_path / "q.xlsx"
        wb.save(path)

        report = ExcelQualityChecker().check(str(path))
        assert report.issues, "构造的坏文件应产生质量 issue"
        assert all(IssueSeverity.is_valid(i.severity) for i in report.issues)

    def test_ppt_outline_check_severities_all_legal(self):
        from office_agent.ppt_agent.models import PPTOutline, SlideContent
        from office_agent.ppt_agent.quality_checker import PPTQualityChecker

        outline = PPTOutline(title="")
        outline.slides = [
            SlideContent(layout="content", title="", bullets=[]),
            SlideContent(layout="content", title="ok", bullets=["a"] * 12),
        ]
        report = PPTQualityChecker().check_outline(outline)
        assert report.issues, "空标题/过多要点应产生质量 issue"
        assert all(IssueSeverity.is_valid(i.severity) for i in report.issues)


# ------------------------------------------------------------------ AST source guards

class TestNoBareSeverityLiterals:
    QUALITY_MODULES = [
        "office_agent.excel_agent.quality_checker",
        "office_agent.ppt_agent.quality_checker",
    ]
    _SEVERITIES = {"error", "warning", "info"}

    @pytest.mark.parametrize("module_name", QUALITY_MODULES)
    def test_no_bare_severity_keyword_literal(self, module_name):
        module = __import__(module_name, fromlist=["x"])
        tree = ast.parse(inspect.getsource(module))
        offenders = [
            kw.value.lineno
            for node in ast.walk(tree) if isinstance(node, ast.Call)
            for kw in node.keywords
            if kw.arg == "severity"
            and isinstance(kw.value, ast.Constant)
            and kw.value.value in self._SEVERITIES
        ]
        assert not offenders, f"{module_name} severity= 裸字面量行: {offenders}"

    @pytest.mark.parametrize("module_name", QUALITY_MODULES)
    def test_no_bare_severity_comparison_literal(self, module_name):
        module = __import__(module_name, fromlist=["x"])
        tree = ast.parse(inspect.getsource(module))
        offenders = []
        for node in ast.walk(tree):
            if not isinstance(node, ast.Compare):
                continue
            if not (isinstance(node.left, ast.Attribute) and node.left.attr == "severity"):
                continue
            for comparator in node.comparators:
                if (isinstance(comparator, ast.Constant)
                        and comparator.value in self._SEVERITIES):
                    offenders.append(node.lineno)
        assert not offenders, f"{module_name} severity 比较裸字面量行: {offenders}"
