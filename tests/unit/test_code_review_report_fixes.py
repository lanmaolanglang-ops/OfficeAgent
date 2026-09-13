"""CODE_REVIEW_REPORT.md 中已确认缺陷的回归测试。"""
import asyncio
import json
import logging
import threading

import openpyxl


def test_security_config_from_env_uses_default_without_attribute_error(monkeypatch):
    from office_agent.security.config import DEFAULT_JWT_SECRET, SecurityConfig

    monkeypatch.delenv("OFFICE_AGENT_JWT_SECRET", raising=False)
    assert SecurityConfig.from_env().jwt_secret_key is DEFAULT_JWT_SECRET


def test_missing_development_storage_path_is_an_info(tmp_path):
    from office_agent.config_system.schemas import GlobalConfig, StorageConfig
    from office_agent.config_system.validators import ConfigValidator, Severity

    config = GlobalConfig(storage=StorageConfig(local_path=str(tmp_path / "missing")))
    issues = ConfigValidator().validate_global(config)
    assert any(issue.severity is Severity.INFO for issue in issues)


def _make_merged_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "模板"
    ws.append(["名称", "金额"])
    ws.append(["旧数据", 1])
    ws.merge_cells("A3:B3")
    ws["A3"] = "合并说明"
    wb.save(path)


def test_template_fill_public_entrypoints_handle_merged_cells(tmp_path):
    from office_agent.excel_agent.template_analyzer import ExcelTemplateAnalyzer

    template = tmp_path / "template.xlsx"
    list_output = tmp_path / "list.xlsx"
    dict_output = tmp_path / "dict.xlsx"
    _make_merged_template(template)

    analyzer = ExcelTemplateAnalyzer()
    analyzer.fill_template(str(template), [["甲", 10]], str(list_output))
    analyzer.fill_with_dict(
        str(template), [{"名称": "乙", "金额": 20}], str(dict_output)
    )

    assert list_output.exists()
    assert dict_output.exists()


def _formula_sheet():
    from office_agent.excel_agent.models import ColumnInfo, DataProfile, SheetInfo

    sheet = SheetInfo(
        name="数据",
        row_count=3,
        col_count=3,
        columns=[
            ColumnInfo("分类", 0, data_type="text", semantic_type="category"),
            ColumnInfo("销售额", 1, data_type="number", semantic_type="amount"),
            ColumnInfo("备注", 2, data_type="text"),
        ],
    )
    return sheet, DataProfile(sheets=[sheet], total_rows=3, total_sheets=1)


def test_aggregate_formulas_do_not_overwrite_each_other():
    from office_agent.excel_agent.formula_generator import FormulaGenerator

    _sheet, profile = _formula_sheet()
    formulas = FormulaGenerator(profile).generate_from_text("计算销售额合计和平均", "数据")
    aggregate = [f for f in formulas if f.category in {"sum", "average"}]
    assert {f.category for f in aggregate} == {"sum", "average"}
    assert len({f.target_cell for f in aggregate}) == 2


def test_logical_growth_and_rank_results_use_first_empty_column():
    from office_agent.excel_agent.formula_generator import FormulaGenerator

    sheet, profile = _formula_sheet()
    logical = FormulaGenerator(profile).generate_from_text(
        "如果销售额大于100则达标否则未达标", "数据"
    )
    assert logical and all(f.target_cell.startswith("D") for f in logical)

    generator = FormulaGenerator(profile)
    growth = generator.generate_from_text("计算销售额环比", "数据")
    assert growth and all(f.target_cell.startswith("D") for f in growth)
    assert all(f.target_cell.startswith("D") for f in generator.generate_rank_formulas(1, sheet))


def test_conditional_formulas_use_more_than_five_unique_categories(tmp_path):
    from office_agent.excel_agent.data_analyzer import DataAnalyzer
    from office_agent.excel_agent.formula_generator import FormulaGenerator

    source = tmp_path / "categories.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["分类", "销售额"])
    for index in range(7):
        ws.append([f"类别{index}", index + 1])
    wb.save(source)

    profile = DataAnalyzer().analyze(str(source))
    assert len(profile.get_sheet("数据").columns[0].unique_values) == 7
    formulas = FormulaGenerator(profile).generate_from_text("按分类统计销售额", "数据")
    labels = [f for f in formulas if f.category == "label"]
    assert len(labels) == 7


def test_spaced_sheet_scatter_uses_the_two_numeric_columns(tmp_path):
    from office_agent.excel_agent.chart_generator import ChartGenerator
    from office_agent.excel_agent.excel_service import ExcelService
    from office_agent.excel_agent.models import ChartSpec

    output = tmp_path / "scatter.xlsx"
    service = ExcelService()
    service.create(str(output))
    ws = service.wb.active
    ws.title = "Sheet 1"
    ws.append(["名称", "X值", "Y值"])
    ws.append(["a", 1, 10])
    ws.append(["b", 2, 20])

    spec = ChartSpec(
        chart_type="scatter", data_range="B1:C3", x_title="X值", y_title="Y值"
    )
    ChartGenerator()._render_chart(service, spec, "Sheet 1")
    chart = ws._charts[0]
    assert "'Sheet 1'!$B$2:$B$3" in str(chart.series[0].xVal.numRef.f)
    assert "'Sheet 1'!$C$2:$C$3" in str(chart.series[0].yVal.numRef.f)


def test_scatter_keeps_non_adjacent_x_and_y_columns(tmp_path):
    from office_agent.excel_agent.chart_generator import ChartGenerator
    from office_agent.excel_agent.excel_service import ExcelService
    from office_agent.excel_agent.models import ColumnInfo, DataProfile, SheetInfo

    sheet = SheetInfo(
        name="数据", row_count=2, col_count=5,
        columns=[
            ColumnInfo("名称", 0, data_type="text"),
            ColumnInfo("X值", 1, data_type="number"),
            ColumnInfo("备注1", 2, data_type="text"),
            ColumnInfo("备注2", 3, data_type="text"),
            ColumnInfo("Y值", 4, data_type="number"),
        ],
    )
    generator = ChartGenerator(DataProfile(sheets=[sheet]))
    spec = generator.generate_from_text("用X值和Y值画散点图", "数据")[0]
    assert spec.x_values_range == "B2:B3"
    assert spec.y_values_range == "E2:E3"

    service = ExcelService().create(str(tmp_path / "scatter-spaced.xlsx"), "数据")
    ws = service.get_sheet("数据")
    ws.append(["名称", "X值", "备注1", "备注2", "Y值"])
    ws.append(["a", 1, "", "", 10])
    ws.append(["b", 2, "", "", 20])
    generator._render_chart(service, spec, "数据")
    chart = ws._charts[0]
    assert "'数据'!$B$2:$B$3" in str(chart.series[0].xVal.numRef.f)
    assert "'数据'!$E$2:$E$3" in str(chart.series[0].yVal.numRef.f)


def test_empty_sheet_does_not_generate_or_render_charts(tmp_path):
    from office_agent.excel_agent.chart_generator import ChartGenerator
    from office_agent.excel_agent.excel_service import ExcelService
    from office_agent.excel_agent.models import ChartSpec, ColumnInfo, DataProfile, SheetInfo

    sheet = SheetInfo(
        name="数据", row_count=0, col_count=2,
        columns=[ColumnInfo("分类", 0, data_type="text"), ColumnInfo("值", 1, data_type="number")],
    )
    generator = ChartGenerator(DataProfile(sheets=[sheet]))
    assert generator.generate_from_text("画柱状图", "数据") == []
    assert generator.auto_charts("数据") == []

    service = ExcelService().create(str(tmp_path / "empty.xlsx"), "数据")
    service.get_sheet("数据").append(["分类", "值"])
    generator._render_chart(
        service, ChartSpec(chart_type="column", data_range="B1:B1", categories_range="A2:A1"), "数据"
    )
    assert service.get_sheet("数据")._charts == []
    assert not any(change.startswith("添加图表") for change in service.changes)


def test_dedup_count_and_ifs_templates_generate_formulas():
    from office_agent.excel_agent.formula_generator import FormulaGenerator

    _sheet, profile = _formula_sheet()
    generator = FormulaGenerator(profile)
    dedup = generator.generate_from_text("销售额去重计数", "数据")
    assert len(dedup) == 1
    assert dedup[0].category == "dedup_count"
    assert dedup[0].formula.startswith("=SUMPRODUCT(")

    logical = generator.generate_from_text(
        "销售额分级，大于等于90为优秀，大于等于60为及格，否则不及格", "数据"
    )
    assert logical and all(item.category == "ifs" for item in logical)
    assert all(item.formula.startswith("=IFS(") for item in logical)
    assert logical[0].formula == '=IFS(B2>=90,"优秀",B2>=60,"及格",TRUE,"不及格")'


def test_ppt_corrupt_base_template_falls_back(tmp_path):
    from office_agent.ppt_agent.models import PPTOutline, SlideContent
    from office_agent.ppt_agent.ppt_service import PPTService

    corrupt = tmp_path / "corrupt.pptx"
    corrupt.write_bytes(b"not a pptx")
    output = tmp_path / "output.pptx"
    outline = PPTOutline(title="回退测试")
    outline.add_slide(SlideContent(layout="cover", title="回退测试"))
    outline._base_template_path = str(corrupt)

    result = PPTService().generate(outline, str(output))
    assert result.success, result.message
    assert output.exists()


def test_knowledge_base_refits_all_documents_and_survives_restart(tmp_path):
    from office_agent.knowledge_base.knowledge_base import OfficeKnowledgeBase

    kb = OfficeKnowledgeBase(str(tmp_path))
    kb.add_text("alpha orchard policy", title="alpha")
    kb.add_text("zebraword compliance rule", title="zebra")
    assert "zebraword" in kb.embedder.vocabulary
    assert all(len(item.embedding) == kb.embedder.dimension for item in kb.store._chunks)

    vocabulary_before_restart = set(kb.embedder.vocabulary)
    restarted = OfficeKnowledgeBase(str(tmp_path))
    assert restarted.embedder._fitted is True
    assert restarted.search("zebraword", min_score=0).best_score > 0

    restarted.add_text("quasar retention standard", title="quasar")
    assert vocabulary_before_restart <= set(restarted.embedder.vocabulary)
    assert "quasar" in restarted.embedder.vocabulary
    assert all(
        len(item.embedding) == restarted.embedder.dimension
        for item in restarted.store._chunks
    )


def test_keyword_score_is_normalized():
    from office_agent.knowledge_base.models import KnowledgeChunk
    from office_agent.knowledge_base.vector_store import VectorStore

    store = VectorStore()
    chunk = KnowledgeChunk(content="甲乙丙丁", keywords=["甲乙", "乙丙", "丙丁"])
    assert store._keyword_match_score({"甲乙", "乙丙", "丙丁"}, chunk) == 1.0


def test_chunk_keywords_are_cached_across_searches(monkeypatch):
    from office_agent.knowledge_base.models import KnowledgeChunk
    from office_agent.knowledge_base.vector_store import VectorStore

    store = VectorStore()
    content = "alpha orchard policy"
    original = store._extract_keywords
    content_extractions = 0

    def counted(text):
        nonlocal content_extractions
        if text == content:
            content_extractions += 1
        return original(text)

    monkeypatch.setattr(store, "_extract_keywords", counted)
    store.add_chunks([KnowledgeChunk(content=content)])
    store.search("alpha", min_score=0)
    store.search("orchard", min_score=0)
    assert content_extractions == 1


def test_model_log_handler_matches_prefixed_logger(monkeypatch):
    from office_agent.logging_system.handlers import ModelCallLogHandler

    called = threading.Event()
    handler = ModelCallLogHandler()
    monkeypatch.setattr(handler, "_write_to_db", lambda _record: called.set())
    record = logging.LogRecord(
        "office_agent.model.openai", logging.INFO, __file__, 1, "call", (), None
    )
    handler.emit(record)
    assert called.wait(1)


def test_trace_context_is_async_isolated_and_released():
    from office_agent.logging_system.tracer import (
        _trace_context_var, get_trace_context, trace_span,
    )

    async def worker(name):
        with trace_span(name) as span:
            await asyncio.sleep(0)
            assert get_trace_context().root is span
            return id(get_trace_context())

    async def run_workers():
        return await asyncio.gather(worker("a"), worker("b"))

    context_ids = asyncio.run(run_workers())
    assert context_ids[0] != context_ids[1]
    assert _trace_context_var.get() is None


def test_ipv6_host_parser_accepts_bracketed_loopback():
    from office_agent.api.middleware.local_guard import _hostname

    assert _hostname("[::1]:8765") == "::1"
    assert _hostname("localhost:8765") == "localhost"


def test_task_repository_records_start_and_duration(tmp_path, request):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from office_agent.database.base import Base
    from office_agent.database.repository import TaskRepository
    import office_agent.database.models  # noqa: F401 - register all mappings

    engine = create_engine(f"sqlite:///{tmp_path / 'tasks.db'}")
    request.addfinalizer(engine.dispose)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        repo = TaskRepository(session)
        task = repo.create_task("excel", "test")
        session.commit()
        repo.start_task(task.id)
        repo.complete_task(task.id, duration_ms=321)
        session.commit()
        session.refresh(task)
        assert task.started_at is not None
        assert task.finished_at is not None
        assert task.duration_ms == 321
    finally:
        session.close()


def test_new_skill_and_workflow_are_persisted(tmp_path, request):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from office_agent.config_system.config_manager import ConfigManager
    from office_agent.database.base import Base
    from office_agent.database.repository import SkillConfigRepository, WorkflowConfigRepository
    import office_agent.database.models  # noqa: F401

    engine = create_engine(f"sqlite:///{tmp_path / 'config.db'}")
    request.addfinalizer(engine.dispose)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    manager = ConfigManager.__new__(ConfigManager)
    manager._session_factory = factory
    manager._persist_skill("new-skill", {"description": "d", "tools": ["x"]})
    manager._persist_workflow("new-flow", {"description": "d", "steps": [{"id": "a"}]})

    session = factory()
    try:
        assert SkillConfigRepository(session).get_by_name("new-skill") is not None
        assert WorkflowConfigRepository(session).get_by_name("new-flow") is not None
    finally:
        session.close()


def test_image_input_detects_file_mime(tmp_path):
    from office_agent.vision_gateway.vision_models import ImageInput

    image = tmp_path / "photo.jpg"
    image.write_bytes(b"jpeg")
    assert ImageInput.from_file(str(image)).get_mime() == "image/jpeg"


def test_failed_visual_page_has_empty_score_details():
    from office_agent.quality.visual_checker import PageVisualResult

    assert PageVisualResult().score_details == {}


def test_fixed_point_line_spacing_keeps_exact_rule():
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from office_agent.parsers.format_parser import parse_format_rule
    from office_agent.services.word_service import WordService

    parsed = parse_format_rule("正文固定值20磅")
    assert parsed["line_spacing"] == 20
    assert parsed["line_spacing_rule"] == "exactly"
    config = WordService().config_from_dict(parsed)
    paragraph = Document().add_paragraph("正文")
    WordService()._apply_paragraph_format(
        paragraph, config.body_font, config.body_paragraph
    )
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert abs(paragraph.paragraph_format.line_spacing.pt - 20) < 0.01


def test_short_plain_text_is_kept_as_ppt_bullet():
    from office_agent.ppt_agent.content_planner import ContentPlanner

    outline = ContentPlanner().plan_from_text("### 页面\n短句")
    page = next(slide for slide in outline.slides if slide.title == "页面")
    assert page.bullets == ["短句"]


def test_presentation_filename_is_sanitized():
    from office_agent.ppt_agent.ppt_orchestrator import _safe_presentation_name

    name = _safe_presentation_name(r"..\evil:*?<>|")
    assert name != ""
    assert ".." not in name
    assert not any(char in name for char in '\\/:*?"<>|')


def test_chart_data_dicts_are_normalized(tmp_path):
    from office_agent.ppt_agent.models import PPTOutline, SlideContent
    from office_agent.ppt_agent.ppt_service import PPTService

    outline = PPTOutline(title="图表")
    outline.add_slide(SlideContent(layout="cover", title="图表"))
    outline.add_slide(SlideContent(
        layout="chart", title="数据",
        data=[{"label": "甲", "value": 1}, {"name": "乙", "amount": "2"}],
    ))
    output = tmp_path / "chart.pptx"
    assert PPTService().generate(outline, str(output)).success


def test_excel_write_and_formula_skip_merged_followers(tmp_path):
    from office_agent.excel_agent.excel_service import ExcelService
    from office_agent.excel_agent.models import FormulaSpec

    service = ExcelService()
    service.create(str(tmp_path / "merged.xlsx"), "数据")
    ws = service.get_sheet("数据")
    ws.merge_cells("A1:B1")
    service.write_data("数据", [["标题", "不能写入"]])
    service.add_formula(FormulaSpec("=1", "B1"), "数据")
    assert ws["A1"].value == "标题"


def test_mixed_type_sort_keeps_blanks_last():
    from office_agent.excel_agent.excel_service import ExcelService

    service = ExcelService().create("unused.xlsx", "数据")
    service.write_data("数据", [["值"], [2], [None], ["10"], [1]])
    service.sort_data("数据", 0, ascending=True)
    values = [service.get_sheet("数据").cell(row=i, column=1).value for i in range(2, 6)]
    assert values[-1] is None
    assert set(values[:-1]) == {1, 2, "10"}


def test_openpyxl_fallback_row_count_excludes_header(tmp_path):
    from office_agent.excel_agent.data_analyzer import DataAnalyzer

    source = tmp_path / "rows.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["列"])
    wb.active.append([1])
    wb.active.append([2])
    wb.save(source)
    profile = DataAnalyzer()._analyze_profile_with_openpyxl(str(source))
    assert profile.sheets[0].row_count == 2


def test_trend_labels_follow_source_rows_with_missing_values():
    from office_agent.excel_agent.analysis_engine import AnalysisEngine

    rows = [["一月", 10], ["二月"], ["三月", 20], ["四月", 5]]
    trend = AnalysisEngine()._analyze_trend(rows, ["月份", "销售额"], 1, 0)
    assert [period["period"] for period in trend.periods] == ["三月", "四月"]
    assert trend.max_growth_period == "三月"


def test_pie_chart_has_per_point_colors(tmp_path):
    from office_agent.excel_agent.excel_service import ExcelService
    from office_agent.excel_agent.models import ChartSpec

    service = ExcelService().create(str(tmp_path / "pie.xlsx"), "数据")
    service.write_data("数据", [["分类", "值"], ["甲", 1], ["乙", 2], ["丙", 3]])
    service.add_chart(ChartSpec(
        chart_type="pie", data_range="B1:B4", categories_range="A2:A4"
    ), "数据")
    points = service.get_sheet("数据")._charts[0].series[0].data_points
    assert len(points) == 3


def test_deepseek_uses_official_thinking_parameter(monkeypatch):
    from office_agent.model_gateway.clients import openai_client as module
    from office_agent.model_gateway.clients.openai_client import OpenAIClient
    from office_agent.models.model_schemas import ModelConfig, ModelProvider

    captured = {}

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self):
            return json.dumps({
                "choices": [{"message": {"content": "ok"}, "finish_reason": "stop"}],
                "usage": {},
            }).encode()

    def fake_urlopen(request, timeout):
        captured.update(json.loads(request.data.decode()))
        return Response()

    monkeypatch.setattr(module, "urlopen", fake_urlopen)
    client = OpenAIClient(ModelConfig(
        id="d", provider=ModelProvider.DEEPSEEK, display_name="d",
        api_key="x", base_url="https://api.deepseek.com/v1", model="deepseek-v4-pro",
    ))
    assert client.chat([{"role": "user", "content": "hi"}]).success
    assert captured["thinking"] == {"type": "disabled"}
    assert "reasoning" not in captured


def test_expired_active_api_key_is_cleaned(monkeypatch):
    from office_agent.security.auth.token import TokenManager

    manager = TokenManager()
    key = manager.create_api_key("user", expires_days=1)
    manager._api_keys[key].expires_at = 0
    manager.cleanup_expired()
    assert key not in manager._api_keys


def test_word_title_style_is_document_title():
    from office_agent.quality_scoring.word_scorer import WordQualityScorer

    paragraph = {
        "style_name": "Title", "text": "文档标题", "main_bold": True,
        "main_size": 20, "main_font": "黑体",
    }
    result = WordQualityScorer()._analyze_headings(None, [paragraph])
    assert result["title_found"] is True
    assert result["headings"][0]["level"] == 0
