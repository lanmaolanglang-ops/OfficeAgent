"""硬编码默认值集中化回归测试（清单 665/667/668/671/672）。

 pin 五项治理的行为不变量：
 - 665：Excel/PPT/Word 评分权重为类级命名常量，合计 1.0，
   真实文件评分总分与旧版字面量加权结果一致；
 - 667：配图默认尺寸唯一来源 DEFAULT_IMAGE_SIZE = "1024x768"，
   generate() 默认参数引用该常量，源码中字面量只出现一次；
 - 668：类型嗅探采样 50 / 日期嗅探采样 20 / 唯一值预览 20 分别命名，
   默认行数与边界行为不变；
 - 671：缓存值扫描上限 _CACHED_SCAN_MAX_ROWS = 5000 已存在且被复用
   （清单滞后项，只补守卫）；
 - 672：自动质量修订最大层数 MAX_QUALITY_REVISION_DEPTH = 3，
   深度边界与取消行为不变。
"""
import ast
import inspect
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]


# ---------------------------------------------------------------------------
# 665 评分权重
# ---------------------------------------------------------------------------

class TestScorerWeights:
    def test_excel_weights_named_and_normalized(self):
        from office_agent.quality_scoring.excel_scorer import ExcelQualityScorer
        total = (ExcelQualityScorer.WEIGHT_FORMULA_ACCURACY
                 + ExcelQualityScorer.WEIGHT_ANALYSIS_ACCURACY
                 + ExcelQualityScorer.WEIGHT_CHART_APPROPRIATENESS)
        assert total == pytest.approx(1.0)

    def test_ppt_weights_named_and_normalized(self):
        from office_agent.quality_scoring.ppt_scorer import PPTQualityScorer
        total = (PPTQualityScorer.WEIGHT_VISUAL
                 + PPTQualityScorer.WEIGHT_CONTENT_COMPLETENESS
                 + PPTQualityScorer.WEIGHT_TEMPLATE_ADHERENCE)
        assert total == pytest.approx(1.0)

    def test_word_weights_named_and_normalized(self):
        from office_agent.quality_scoring.word_scorer import WordQualityScorer
        total = (WordQualityScorer.WEIGHT_FORMAT_ACCURACY
                 + WordQualityScorer.WEIGHT_HEADING_RECOGNITION
                 + WordQualityScorer.WEIGHT_LAYOUT_CONSISTENCY)
        assert total == pytest.approx(1.0)

    def test_excel_total_matches_legacy_literal_weighting(self, tmp_path):
        """旧代表性输入：总分仍等于 0.4/0.35/0.25 字面量加权。"""
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook
        from office_agent.quality_scoring.excel_scorer import ExcelQualityScorer

        wb = Workbook()
        ws = wb.active
        ws.title = "销售"
        ws["A1"] = "地区"
        ws["B1"] = "金额"
        ws["A2"] = "华东"
        ws["B2"] = 100
        ws["A3"] = "华北"
        ws["B3"] = 200
        ws["B4"] = "=SUM(B2:B3)"
        path = tmp_path / "sample.xlsx"
        wb.save(path)

        result = ExcelQualityScorer().score(str(path), expected_formulas=["SUM"])
        legacy = (result.formula_accuracy * 0.4
                  + result.analysis_accuracy * 0.35
                  + result.chart_appropriateness * 0.25)
        assert result.total_score == pytest.approx(legacy)

    def test_ppt_total_matches_legacy_literal_weighting(self, tmp_path):
        pytest.importorskip("pptx")
        from pptx import Presentation
        from office_agent.quality_scoring.ppt_scorer import PPTQualityScorer

        prs = Presentation()
        for title, body in (("封面", "年度报告"), ("内容", "营收增长")):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        path = tmp_path / "sample.pptx"
        prs.save(path)

        result = PPTQualityScorer().score(str(path), expected_slides=2)
        legacy = (result.visual_score * 0.35
                  + result.content_completeness * 0.35
                  + result.template_adherence * 0.30)
        assert result.total_score == pytest.approx(legacy)

    def test_word_total_matches_legacy_literal_weighting(self, tmp_path):
        pytest.importorskip("docx")
        from docx import Document
        from office_agent.quality_scoring.word_scorer import WordQualityScorer

        doc = Document()
        doc.add_heading("报告标题", level=1)
        doc.add_paragraph("正文第一段内容。")
        path = tmp_path / "sample.docx"
        doc.save(path)

        result = WordQualityScorer().score(str(path))
        legacy = (result.format_accuracy * 0.4
                  + result.heading_recognition * 0.35
                  + result.layout_consistency * 0.25)
        assert result.total_score == pytest.approx(legacy)

    @pytest.mark.parametrize("module_name", [
        "office_agent.quality_scoring.excel_scorer",
        "office_agent.quality_scoring.ppt_scorer",
        "office_agent.quality_scoring.word_scorer",
    ])
    def test_no_duplicate_weight_definitions(self, module_name):
        """每个 scorer 类内 WEIGHT_* 赋值恰好 3 处，无第二套定义。"""
        import importlib
        module = importlib.import_module(module_name)
        tree = ast.parse(inspect.getsource(module))
        assignments = [
            t.id for node in ast.walk(tree)
            if isinstance(node, ast.ClassDef)
            for t in node.body
            if isinstance(t, ast.Assign)
            for t in t.targets
            if isinstance(t, ast.Name) and t.id.startswith("WEIGHT_")
        ]
        assert len(assignments) == 3


# ---------------------------------------------------------------------------
# 667 配图默认尺寸
# ---------------------------------------------------------------------------

class TestDefaultImageSize:
    @staticmethod
    def _import_gateway():
        # gateway 经 image_generation.config → model_gateway → api 存在
        # 包级循环引用；先初始化 api 主包与既有套件顺序一致，仅影响导入顺序。
        import office_agent.api.main  # noqa: F401
        from office_agent.image_generation import gateway
        return gateway

    def test_default_constant_value(self):
        gateway = self._import_gateway()
        assert gateway.DEFAULT_IMAGE_SIZE == "1024x768"

    def test_generate_signature_default_uses_constant(self):
        gateway = self._import_gateway()
        default = inspect.signature(
            gateway.ImageGenerationGateway.generate).parameters["size"].default
        assert default is gateway.DEFAULT_IMAGE_SIZE

    def test_literal_single_source_in_gateway(self):
        source = (REPO_ROOT / "office_agent/image_generation/gateway.py").read_text(
            encoding="utf-8")
        assert source.count("1024x768") == 1

    def test_orchestrator_has_no_size_literal(self):
        source = (REPO_ROOT / "office_agent/ppt_agent/ppt_orchestrator.py").read_text(
            encoding="utf-8")
        assert "1024x768" not in source


# ---------------------------------------------------------------------------
# 668 数据采样
# ---------------------------------------------------------------------------

class TestSamplingDefaults:
    def test_constant_values(self):
        from office_agent.excel_agent.analysis_engine import TYPE_DETECTION_SAMPLE_SIZE
        from office_agent.excel_agent.data_analyzer import (
            DATE_DETECTION_SAMPLE_SIZE, UNIQUE_VALUES_PREVIEW_LIMIT)
        assert TYPE_DETECTION_SAMPLE_SIZE == 50
        assert DATE_DETECTION_SAMPLE_SIZE == 20
        assert UNIQUE_VALUES_PREVIEW_LIMIT == 20

    def test_type_detection_samples_first_50_only(self):
        """前 50 行定类型：第 51 行起的多数派不影响判定（默认行为不变）。"""
        from office_agent.excel_agent.analysis_engine import (
            TYPE_DETECTION_SAMPLE_SIZE, AnalysisEngine)
        engine = AnalysisEngine()
        rows = [[i + 1] for i in range(TYPE_DETECTION_SAMPLE_SIZE)]
        rows += [["文本值"]] * (TYPE_DETECTION_SAMPLE_SIZE * 2)
        assert engine._detect_column_types(rows, ["col"]) == ["number"]

    def test_type_detection_text_boundary(self):
        from office_agent.excel_agent.analysis_engine import AnalysisEngine
        engine = AnalysisEngine()
        rows = [["文本值"]] * 10
        assert engine._detect_column_types(rows, ["col"]) == ["text"]

    def test_date_detection_samples_first_20_only(self):
        """前 20 行做日期嗅探：20 行之后的非日期值不影响判定。"""
        pd = pytest.importorskip("pandas")
        from office_agent.excel_agent.data_analyzer import (
            DATE_DETECTION_SAMPLE_SIZE, DataAnalyzer)
        analyzer = DataAnalyzer()
        values = (["2026-01-%02d" % (i + 1) for i in range(DATE_DETECTION_SAMPLE_SIZE)]
                  + ["不是日期"] * 50)
        series = pd.Series(values, dtype=object)
        assert analyzer._infer_data_type(series, "日期列") == "date"


# ---------------------------------------------------------------------------
# 671 缓存扫描上限（清单滞后项：常量已存在，补守卫）
# ---------------------------------------------------------------------------

class TestCachedScanLimit:
    def test_constant_value(self):
        from office_agent.excel_agent.quality_checker import ExcelQualityChecker
        assert ExcelQualityChecker._CACHED_SCAN_MAX_ROWS == 5000

    def test_scan_uses_named_constant(self):
        from office_agent.excel_agent import quality_checker
        source = inspect.getsource(quality_checker.ExcelQualityChecker._check_cached_errors)
        assert "_CACHED_SCAN_MAX_ROWS" in source
        assert "5000" not in source


# ---------------------------------------------------------------------------
# 672 质量修订层数
# ---------------------------------------------------------------------------

def _make_worker(monkeypatch):
    from office_agent.task_queue.config import config
    from office_agent.task_queue.worker import LocalWorker
    queues = {name: dict(cfg) for name, cfg in config.TASK_QUEUES.items()}
    monkeypatch.setattr(config, "TASK_QUEUES", queues)
    return LocalWorker()


def _fake_task(revision_number):
    return SimpleNamespace(
        id="task-1", revision_number=revision_number, options_json="{}",
        instruction="生成报告", task_type="ppt", agent_name="ppt_agent",
        priority=5)


class TestQualityRevisionDepth:
    def test_constant_value(self):
        from office_agent.task_queue.worker import MAX_QUALITY_REVISION_DEPTH
        assert MAX_QUALITY_REVISION_DEPTH == 3

    def test_max_depth_stops_revision(self, monkeypatch):
        """revision_number 达到上限(3)后不再派生修订子任务。"""
        from office_agent.task_queue.worker import MAX_QUALITY_REVISION_DEPTH
        worker = _make_worker(monkeypatch)
        try:
            result = {"quality_check": {"reports": [{"issues": ["问题A"]}]}}
            child = worker._schedule_quality_revision(
                "task-1", _fake_task(MAX_QUALITY_REVISION_DEPTH), result, ["out-1"])
            assert child is None
        finally:
            worker.shutdown()

    def test_cancelled_task_stops_revision(self, monkeypatch):
        """取消事件置位时，即使深度未到上限也立即停止。"""
        worker = _make_worker(monkeypatch)
        try:
            worker._cancel_events["task-1"] = threading.Event()
            worker._cancel_events["task-1"].set()
            result = {"quality_check": {"reports": [{"issues": ["问题A"]}]}}
            child = worker._schedule_quality_revision(
                "task-1", _fake_task(1), result, ["out-1"])
            assert child is None
        finally:
            worker.shutdown()

    def test_no_issues_no_revision(self, monkeypatch):
        worker = _make_worker(monkeypatch)
        try:
            child = worker._schedule_quality_revision(
                "task-1", _fake_task(1), {"quality_check": {"reports": []}}, ["out-1"])
            assert child is None
        finally:
            worker.shutdown()

    def test_below_max_depth_creates_child_with_incremented_revision(
            self, monkeypatch, tmp_path):
        """第 1 层任务 QA 未过：派生 revision_number=2 的子任务（默认链行为不变）。"""
        worker = _make_worker(monkeypatch)
        created = {}
        input_file = tmp_path / "out.pptx"
        input_file.write_bytes(b"placeholder")

        class _FakeStorage:
            def get_file_path(self, _fid):
                return str(input_file)

        class _FakeGateway:
            def __init__(self, cancel_event=None):
                pass

            def chat(self, **_kwargs):
                raise RuntimeError("no model in test")

        class _FakeRepo:
            def __init__(self, _session):
                pass

            def create_task(self, **kwargs):
                created.update(kwargs)
                return SimpleNamespace(id="child-1")

        class _FakeSessionScope:
            def __enter__(self):
                return object()

            def __exit__(self, *exc):
                return False

        import office_agent.task_queue as tq
        from office_agent.storage import storage_service
        from office_agent.database import session as db_session
        from office_agent.database import repository as db_repo
        monkeypatch.setattr(storage_service, "get_storage_service", lambda: _FakeStorage())
        monkeypatch.setattr("office_agent.model_gateway.ModelGateway", _FakeGateway)
        monkeypatch.setattr(db_session, "session_scope", lambda: _FakeSessionScope())
        monkeypatch.setattr(db_repo, "TaskRepository", _FakeRepo)
        monkeypatch.setattr(tq, "submit_task", lambda *a, **k: "child-1", raising=False)
        try:
            result = {"quality_check": {"reports": [{"issues": ["问题A"]}]}}
            child = worker._schedule_quality_revision(
                "task-1", _fake_task(1), result, ["out-1"])
            assert child == "child-1"
            assert created["revision_number"] == 2
            assert created["parent_task_id"] == "task-1"
        finally:
            worker.shutdown()
