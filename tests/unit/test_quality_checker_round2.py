"""P3-9：quality checker 剩余子问题回归（不与 P1-20 重叠）。

本轮确认并修复的子问题：
1. 零宽字符（U+200B/200C/200D/2060）与私有区字符被一律判 ERROR——
   它们渲染不可见、多为复制粘贴/符号字体残留，降为 WARNING；
   控制字符与 U+FFFD 替换符仍为 ERROR；
2. 字体/字号检查在**首个非空 run 后早退**——首 run 正常会掩盖同段
   后续 run 的字体/字号问题；
3. ``_cn_to_int`` 只支持到九十九（"一百二十三"返回 0），调用侧正则
   的字符类也不含百/千/零。
"""
import pytest
from docx import Document as DocxDocument
from docx.shared import Pt
from docx.oxml.ns import qn

from office_agent.models.schemas import FormatConfig
from office_agent.quality.checker import QualityChecker, QualityReport


def _set_cn_font(run, font_name: str):
    """python-docx 的 font.name 不写 eastAsia，检查器读取的是 eastAsia。"""
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)


class TestGarbledSeverity:
    def test_zero_width_downgraded_and_control_char_stays_error(self):
        """零宽字符是复制粘贴残留（提示清理），控制字符才是真乱码。

        python-docx 拒绝把控制字符写进 XML，因此直接用假文档对象
        驱动分类器（_check_garbled_text 只依赖 paragraphs[].text）。
        """
        from types import SimpleNamespace
        class _FakeDoc:
            paragraphs = [
                SimpleNamespace(text="正常开头\u200b隐藏字符"),
                SimpleNamespace(text="损坏\u0007控制符"),
            ]

        report = QualityReport(file_path="fake.docx")
        QualityChecker()._check_garbled_text(_FakeDoc(), report)
        zero_width = [i for i in report.issues
                      if i.type == "garbled" and "零宽" in i.message]
        control = [i for i in report.issues
                   if i.type == "garbled" and "乱码字符" in i.message]
        assert len(zero_width) == 1 and zero_width[0].severity == "warning"
        assert len(control) == 1 and control[0].severity == "error"

    def test_replacement_char_stays_error(self, tmp_path):
        doc = DocxDocument()
        doc.add_paragraph("编码损坏\ufffd替换符")
        path = tmp_path / "replacement.docx"
        doc.save(str(path))

        report = QualityChecker().check(str(path))
        replacement = [i for i in report.issues
                       if i.type == "garbled" and "乱码字符" in i.message]
        assert len(replacement) == 1
        assert replacement[0].severity == "error"


class TestFontRunScan:
    def test_bad_font_in_later_run_is_detected(self, tmp_path):
        """回归核心：首 run 正常不能掩盖同段后续 run 的字体问题。"""
        doc = DocxDocument()
        para = doc.add_paragraph()
        _set_cn_font(para.add_run("正常前缀"), "宋体")
        _set_cn_font(para.add_run("异常字体后缀"), "黑体")
        path = tmp_path / "fonts.docx"
        doc.save(str(path))

        report = QualityChecker().check(str(path), format_config=FormatConfig())
        font_issues = [i for i in report.issues if i.type == "font"]
        assert font_issues, "同段后续 run 的字体问题不得被首 run 早退掩盖"
        assert font_issues[0].paragraph_index == 0

    def test_clean_first_run_no_false_positive(self, tmp_path):
        doc = DocxDocument()
        para = doc.add_paragraph()
        _set_cn_font(para.add_run("全部正常"), "宋体")
        path = tmp_path / "clean.docx"
        doc.save(str(path))

        report = QualityChecker().check(str(path), format_config=FormatConfig())
        assert not [i for i in report.issues if i.type == "font"]

    def test_bad_size_in_later_run_is_detected(self, tmp_path):
        doc = DocxDocument()
        para = doc.add_paragraph()
        clean = para.add_run("正常字号")
        clean.font.size = Pt(12)
        bad = para.add_run("异常字号")
        bad.font.size = Pt(20)
        path = tmp_path / "sizes.docx"
        doc.save(str(path))

        report = QualityChecker().check(str(path), format_config=FormatConfig())
        size_issues = [i for i in report.issues if i.type == "font_size"]
        assert size_issues, "同段后续 run 的字号问题不得被首 run 早退掩盖"


class TestChineseNumerals:
    @pytest.mark.parametrize("cn,value", [
        ("一", 1), ("九", 9), ("十", 10), ("十二", 12), ("二十三", 23),
        ("一百零五", 105), ("一百二十三", 123), ("三百", 300),
        ("一千零一", 1001), ("两千", 2000),
    ])
    def test_cn_to_int_supported(self, cn, value):
        assert QualityChecker._cn_to_int(cn) == value

    @pytest.mark.parametrize("cn", ["abc", "", "哈喽"])
    def test_cn_to_int_invalid_returns_zero(self, cn):
        assert QualityChecker._cn_to_int(cn) == 0

    def test_hundred_level_heading_number_resolves(self):
        """“一百零二”这类编号可解析（历史版本返回 0 → 编号链断裂误报）。"""
        assert QualityChecker._cn_to_int("一百零二") == 102

    def test_heading_with_hundreds_parses_via_regex(self):
        """调用侧正则的字符类必须包含百/千/零，否则编号根本进不了解析。"""
        import re
        text = "一百零二、附录小节"
        m = re.match(r"^([一二三四五六七八九十百千零]+)[、.]", text)
        assert m is not None
        assert QualityChecker._cn_to_int(m.group(1)) == 102
