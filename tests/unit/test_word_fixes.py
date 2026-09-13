"""
Word 处理链修复的回归测试

覆盖：
- 三线表 tblBorders 去重与 OOXML schema 顺序（避免 Word/WPS 判定文档损坏）
- 排版后正文/表格内容保留
- LLM 非法配置值（"12pt"/None/bool）的宽容解析
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

PROJECT_ROOT = Path(__file__).parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from office_agent.services.word_service import WordService  # noqa: E402
from office_agent.models.schemas import FontConfig, ParagraphConfig  # noqa: E402
from office_agent.parsers.document_structure import DocumentStructureAnalyzer  # noqa: E402
from office_agent.parsers.format_parser import parse_format_rule  # noqa: E402
from office_agent.quality.checker import QualityReport  # noqa: E402
from office_agent.quality_scoring.word_scorer import WordQualityScorer  # noqa: E402


def _make_doc_with_inline_borders(path: Path):
    """构造带内联 tblBorders 的文档（Word 生成的表格常见形态）"""
    doc = Document()
    doc.add_heading("第一章 测试", level=1)
    doc.add_paragraph("正文内容，用于测试。")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    tbl_pr = table._tbl.find(qn("w:tblPr"))
    inline = parse_xml(
        '<w:tblBorders %s><w:top w:val="single" w:sz="4"/>'
        '<w:left w:val="single" w:sz="4"/><w:bottom w:val="single" w:sz="4"/>'
        '<w:right w:val="single" w:sz="4"/></w:tblBorders>' % nsdecls("w")
    )
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(inline)
    else:
        tbl_pr.append(inline)
    doc.save(str(path))


class TestThreeLineTableBorders:
    def test_tbl_borders_dedup_and_order(self, temp_dir):
        """排版后 tblBorders 唯一且位于 tblLook 之前（schema 顺序）"""
        src = temp_dir / "in.docx"
        out = temp_dir / "out.docx"
        _make_doc_with_inline_borders(src)

        result = WordService().process(
            str(src), output_path=str(out),
            config_dict={"font": "宋体", "size": "小四"},
        )
        assert result.success, result.message

        doc2 = Document(str(out))
        tbl_pr = doc2.tables[0]._tbl.find(qn("w:tblPr"))
        borders = tbl_pr.findall(qn("w:tblBorders"))
        assert len(borders) == 1, "tblBorders 应去重为 1 个"

        look = tbl_pr.find(qn("w:tblLook"))
        if look is not None:
            children = list(tbl_pr)
            assert children.index(borders[0]) < children.index(look)

    def test_cell_bottom_border_dedup(self):
        """P5-9：_set_cell_bottom_border 应先移除已有 w:bottom 再写入新值，
        避免在已有 tcBorders 上追加第二个 w:bottom。"""
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        # 预先写入一个旧的 bottom
        tcBorders = parse_xml(
            '<w:tcBorders %s><w:bottom w:val="single" w:sz="4" '
            'w:space="0" w:color="FF0000"/></w:tcBorders>' % nsdecls("w")
        )
        tcPr.append(tcBorders)

        WordService()._set_cell_bottom_border(cell, 0.75)

        bottoms = tcBorders.findall(qn("w:bottom"))
        assert len(bottoms) == 1, f"应只剩 1 个 w:bottom，实际有 {len(bottoms)} 个"
        assert bottoms[0].get(qn("w:sz")) == "6", "0.75pt × 8 = 6"
        assert bottoms[0].get(qn("w:color")) == "000000"


class TestWordContentPreservation:
    def test_heading_numbering_preserves_run_formatting(self):
        doc = Document()
        para = doc.add_paragraph()
        prefix = para.add_run("第一章 ")
        prefix.bold = True
        english = para.add_run("Mixed")
        english.italic = True
        chinese = para.add_run("标题")
        chinese.bold = True

        WordService()._prepend_number(para, "1")
        assert para.text == "1 Mixed标题"
        assert english.italic is True
        assert chinese.bold is True
        assert english.text == "1 Mixed"

    def test_partial_heading_config_still_formats_all_levels(self):
        config = WordService().config_from_dict({
            "headings": {"1": {"font": "微软雅黑", "size": 16}},
        })
        assert set(config.headings) == {1, 2, 3, 4}
        assert config.headings[1].font.cn_font == "微软雅黑"
        assert config.headings[2].font.cn_font

    def test_create_document_supports_markdown_tables(self, temp_dir):
        output = temp_dir / "table.docx"
        result = WordService().create_document(
            "# 报告\n| 项目 | 金额 |\n| --- | ---: |\n| 甲 | 10 |",
            output_path=str(output),
        )
        assert result.success, result.message
        document = Document(output)
        assert len(document.tables) == 1
        assert document.tables[0].cell(1, 1).text == "10"


class TestRemainingWordCorrectness:
    def test_character_indent_tracks_font_size(self):
        para = Document().add_paragraph("正文")
        WordService()._apply_paragraph_format(
            para, FontConfig(size=15), ParagraphConfig(first_line_indent_chars=2),
        )
        assert para.paragraph_format.first_line_indent.pt == 30
        parsed = parse_format_rule("正文五号宋体，首行缩进2字符")
        assert parsed["first_line_indent_chars"] == 2

    def test_dominant_run_size_ignores_single_large_symbol(self):
        para = Document().add_paragraph()
        symbol = para.add_run("※")
        symbol.font.size = Pt(30)
        body = para.add_run("这是承担主要内容的正文文本")
        body.font.size = Pt(12)
        size, _ = DocumentStructureAnalyzer()._get_font_info(para)
        assert size == 12

    def test_llm_context_uses_filtered_node_position(self):
        doc = Document()
        doc.add_paragraph("")
        doc.add_paragraph("第一段普通正文内容")
        doc.add_paragraph("")
        doc.add_paragraph("第二段普通正文内容")
        contexts = []

        def callback(_text, context):
            contexts.append(context)
            return None

        DocumentStructureAnalyzer(callback).analyze(doc)
        assert contexts[-1]["prev_types"]

    def test_unnumbered_heading_does_not_advance_number_counter(self):
        doc = Document()
        doc.add_heading("说明", level=1)
        doc.add_heading("1. 正常标题", level=1)
        checker = __import__("office_agent.quality.checker", fromlist=["QualityChecker"]).QualityChecker()
        tree = checker.structure_analyzer.analyze(doc)
        report = QualityReport()
        checker._check_numbering(tree, report)
        assert not [issue for issue in report.issues if issue.type == "numbering"]

    def test_heading_style_without_space_keeps_level_and_expected_format_scores(self, temp_dir):
        doc = Document()
        style = doc.styles["Heading 2"]
        style.name = "标题2"
        doc.add_paragraph("二级标题", style=style)
        body = doc.add_paragraph()
        body_run = body.add_run("正文内容足够长，用于质量评分。")
        body_run.font.name = "Arial"
        body_run.font.size = Pt(12)
        path = temp_dir / "score.docx"
        doc.save(path)
        scorer = WordQualityScorer()
        info = scorer._analyze_headings(doc, scorer._analyze_paragraphs(doc))
        assert info["headings"][0]["level"] == 2
        normal = scorer.score(str(path)).format_accuracy
        constrained = scorer.score(str(path), {"font": "宋体"}).format_accuracy
        assert constrained < normal
    def test_content_preserved_after_formatting(self, temp_dir):
        """排版不得删除原有文本与表格内容"""
        src = temp_dir / "in.docx"
        out = temp_dir / "out.docx"
        _make_doc_with_inline_borders(src)

        result = WordService().process(
            str(src), output_path=str(out),
            config_dict={"font": "宋体", "size": 12},
        )
        assert result.success

        doc2 = Document(str(out))
        texts = " ".join(p.text for p in doc2.paragraphs)
        # 标题正文保留（"第一章"前缀按设计被规范化为多级编号"1"）
        assert "测试" in texts
        assert "正文内容，用于测试" in texts
        assert doc2.tables[0].cell(1, 1).text == "2"

    def test_malformed_config_tolerated(self, temp_dir):
        """LLM 返回 "12pt"/None 等非法值不应让整个任务失败"""
        src = temp_dir / "in.docx"
        out = temp_dir / "out.docx"
        _make_doc_with_inline_borders(src)

        result = WordService().process(
            str(src), output_path=str(out),
            config_dict={"size": "12pt", "line_spacing": None,
                         "headings": {"1": {"size": "三号"}}},
        )
        assert result.success, result.message
