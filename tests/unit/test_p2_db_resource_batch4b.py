"""Batch 4B：DB session 契约 / File cascade / ExcelService 句柄。"""
from __future__ import annotations

import pytest
from openpyxl import Workbook


class TestDbSessionContractP2_25:
    def test_get_db_rolls_back_on_exception(self):
        from office_agent.database.session import get_db

        class FakeSession:
            def __init__(self):
                self.rolled = False
                self.closed = False

            def rollback(self):
                self.rolled = True

            def close(self):
                self.closed = True

        gen = get_db()
        # 替换内部 Session
        import office_agent.database.session as sm

        fake = FakeSession()
        original = sm.SessionLocal
        sm.SessionLocal = lambda: fake
        try:
            gen = get_db()
            next(gen)  # 进入 yield
            with pytest.raises(RuntimeError):
                gen.throw(RuntimeError("boom"))
        finally:
            sm.SessionLocal = original
        assert fake.rolled is True
        assert fake.closed is True

    def test_session_scope_commits(self, tmp_path, monkeypatch):
        from office_agent.database import session as sm

        class FakeSession:
            def __init__(self):
                self.committed = False
                self.closed = False

            def commit(self):
                self.committed = True

            def rollback(self):
                pass

            def close(self):
                self.closed = True

        fake = FakeSession()
        monkeypatch.setattr(sm, "SessionLocal", lambda: fake)
        with sm.session_scope() as s:
            assert s is fake
        assert fake.committed is True
        assert fake.closed is True


class TestFileVersionsCascadeP2_26:
    def test_relationship_has_passive_deletes(self):
        from office_agent.database.models.file import File

        rel = File.__mapper__.relationships["versions"]
        assert rel.passive_deletes is True


class TestExcelServiceLifecycleP2_33:
    def test_close_is_idempotent(self, tmp_path):
        from office_agent.excel_agent.excel_service import ExcelService

        path = tmp_path / "a.xlsx"
        svc = ExcelService().create(str(path), "S")
        svc.save()
        svc.close()
        svc.close()
        assert svc.wb is None

    def test_context_manager_closes(self, tmp_path):
        from office_agent.excel_agent.excel_service import ExcelService

        src = tmp_path / "src.xlsx"
        Workbook().save(str(src))
        with ExcelService() as svc:
            svc.open(str(src))
            assert svc.wb is not None
        assert svc.wb is None

    def test_reopen_closes_previous(self, tmp_path):
        from office_agent.excel_agent.excel_service import ExcelService

        a = tmp_path / "a.xlsx"
        b = tmp_path / "b.xlsx"
        Workbook().save(str(a))
        Workbook().save(str(b))
        svc = ExcelService()
        svc.open(str(a))
        first = svc.wb
        svc.open(str(b))
        assert svc.wb is not first
        svc.close()

    def test_windows_can_delete_after_close(self, tmp_path):
        """Windows 句柄占用时 unlink 会失败；close 后必须可删。"""
        from office_agent.excel_agent.excel_service import ExcelService

        path = tmp_path / "del.xlsx"
        ExcelService().create(str(path), "S").save()
        svc = ExcelService().open(str(path))
        svc.close()
        path.unlink()
        assert not path.exists()
