# -*- coding: utf-8 -*-
"""P4/P5 cleanup 回归测试（Batch H/I/J，17 个独立根因）。

覆盖：P4-1/2 颜色归一化、P4-4 滑动窗口身份有界、P4-5 stop(timeout=0)、
P4-6 get_logs 失败可观测、P4-10 key 掩码、P4-11 存储统计前缀、
P4-31 折线图 marker、P4-35 大数值浮点匹配、P4-43 文本框 auto_size、
P4-50 失败历史窗口、P4-51 分隔线误报、P4-52 评分条钳制、
P5-4 编号句误判标题、P5-13 文件名硬化、P5-28 指标锁、P5-34 sc 删除返回码、
P5-37 关闭后读取 sheetnames。
"""
import importlib.util
import subprocess
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]


# ---------------- P4-1 / P4-2 颜色归一化 ----------------
def test_p4_2_double_hash_is_rejected():
    from office_agent.colors import parse_hex_argb, parse_hex_color
    # 旧实现 lstrip("#") 会把 ##FF0000 静默当成合法颜色
    with pytest.raises(ValueError):
        parse_hex_color("##FF0000")
    with pytest.raises(ValueError):
        parse_hex_argb("###00ff00")


def test_p4_1_single_hash_and_short_forms_still_valid():
    from office_agent.colors import parse_hex_argb, parse_hex_color
    assert parse_hex_color("#FF0000") == (255, 0, 0, 255)
    assert parse_hex_color("F00") == (255, 0, 0, 255)
    assert parse_hex_argb("#FF0000") == "FFFF0000"
    assert parse_hex_argb("80FF0000") == "80FF0000"


# ---------------- P4-4 SlidingWindow 身份表有界 ----------------
def test_p4_4_sliding_window_evicts_when_cap_below_ten():
    from office_agent.rate_limiter import SlidingWindowLimiter
    limiter = SlidingWindowLimiter(max_requests=10_000, window_seconds=60)
    limiter.MAX_IDENTITIES = 5  # cap<10：旧实现 len//10==0，一个都不淘汰
    for i in range(6):
        assert limiter.is_allowed(f"identity-{i}").allowed
    assert len(limiter._requests) == 5


# ---------------- P4-5 stop(timeout=0) 不被 `or` 吞掉 ----------------
def test_p4_5_stop_zero_timeout_preserved():
    from office_agent.runtime_manager import AppStatus, ApplicationRuntimeManager
    mgr = ApplicationRuntimeManager.__new__(ApplicationRuntimeManager)
    mgr.config = SimpleNamespace(shutdown_timeout=30)
    mgr._state_lock = threading.Lock()
    mgr._stop_event = threading.Event()
    mgr.state = SimpleNamespace(status=AppStatus.RUNNING)
    captured = {}

    class _FakeProc:
        def terminate(self):
            pass

        def send_signal(self, _sig):
            pass

        def wait(self, timeout=None):
            captured["timeout"] = timeout
            raise subprocess.TimeoutExpired(cmd="backend", timeout=timeout)

    mgr._process = _FakeProc()
    mgr._set_status = lambda _s: None
    mgr._kill_process_handle = lambda _p: None
    mgr.stop(timeout=0)
    assert captured["timeout"] == 0  # 旧实现会变成 30


# ---------------- P4-6 get_logs 读取失败不再静默 ----------------
def test_p4_6_get_logs_failure_is_logged(tmp_path, monkeypatch):
    # 直接 patch 模块 logger.warning，与全局日志级别/handler/disable 完全解耦，
    # 避免同进程内其它测试改动日志配置导致的不确定（P4-6）
    from office_agent import runtime_manager as rm
    from office_agent.runtime_manager import ApplicationRuntimeManager
    mgr = ApplicationRuntimeManager.__new__(ApplicationRuntimeManager)
    mgr.config = SimpleNamespace(log_dir=tmp_path)
    (tmp_path / "backend.log").mkdir()  # 对目录 open 会抛异常
    calls = []
    monkeypatch.setattr(rm.logger, "warning", lambda *a, **k: calls.append(a))
    out = mgr.get_logs()
    assert out == ""
    assert calls and "backend 日志失败" in calls[0][0]


# ---------------- P4-10 API key 掩码只留后 4 位 ----------------
def test_p4_10_mask_key_does_not_leak_prefix():
    from office_agent.api.router.settings import _mask_key
    masked = _mask_key("sk-test-99887766")
    assert masked == "****7766"
    assert "9988" not in masked and "sk-" not in masked
    assert _mask_key("ab") == "****"


# ---------------- P4-11 / P5-36 存储统计按前缀分类 ----------------
def test_p4_11_storage_stats_prefix_not_substring(tmp_path):
    from office_agent.api.core.file_manager import FileInfo, FileManager
    fm = FileManager(upload_dir=str(tmp_path / "up"), output_dir=str(tmp_path / "out"))
    fm.files = {
        "1": FileInfo("file_a", "a.xlsx", "p", "excel", ".xlsx", 1),
        "2": FileInfo("out_b", "b.xlsx", "p", "excel", ".xlsx", 1),
        # 含子串但前缀不匹配，不能被误计
        "3": FileInfo("draft_file_c", "c.xlsx", "p", "excel", ".xlsx", 1),
        "4": FileInfo("tmp_out_d", "d.xlsx", "p", "excel", ".xlsx", 1),
    }
    stats = fm.get_storage_stats()
    assert stats["upload_count"] == 1
    assert stats["output_count"] == 1
    assert stats["total_files"] == 4


# ---------------- P4-31 折线图 marker 真正序列化 ----------------
def test_p4_31_line_chart_marker_rendered(tmp_path):
    from openpyxl import Workbook, load_workbook
    from office_agent.excel_agent.chart_generator import ChartGenerator
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["月份", "销量"])
    for m, v in [("1月", 10), ("2月", 20), ("3月", 30)]:
        ws.append([m, v])

    class _Service:
        def __init__(self, worksheet):
            self._ws = worksheet
            self.changes = []

        def get_sheet(self, _name=None):
            return self._ws

    gen = ChartGenerator()
    spec = gen.create_chart("line", data_range="B1:B4", categories_range="A1:A4")
    gen._render_chart(_Service(ws), spec)
    assert ws._charts, "折线图应被加入工作表"
    for series in ws._charts[0].series:
        assert series.marker is not None and series.marker.symbol == "circle"

    out = tmp_path / "line_marker.xlsx"
    wb.save(out)
    reopened = load_workbook(out)
    rws = reopened["Sheet1"]
    assert rws._charts, "重新打开后图表仍在"
    for series in rws._charts[0].series:
        assert series.marker.symbol == "circle"  # 旧 chart.marker=True 不会落盘


# ---------------- P4-35 大数值浮点匹配 ----------------
def test_p4_35_large_value_label_match():
    from office_agent.excel_agent.analysis_engine import AnalysisEngine
    eng = AnalysisEngine()
    big = 1_000_000.1
    # 固定 0.001 绝对容差下，浮点噪声会让等值匹配失败
    assert eng._find_label_for_value(
        [["最大值", big]], value_col=1, target_value=big, label_col=0
    ) == "最大值"
    # 小数值仍按 0.001 下限，不放松：差距 0.002 > 0.001 视为不匹配
    assert eng._find_label_for_value(
        [["近", 0.5]], value_col=1, target_value=0.502, label_col=0
    ) == ""


# ---------------- P4-43 文本框 auto_size 只赋值一次且可用 ----------------
def test_p4_43_textbox_autosize_single_assignment():
    from pptx import Presentation
    from office_agent.ppt_agent.ppt_service import PPTService
    svc = PPTService.__new__(PPTService)
    svc._sx = svc._sy = 1.0
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = svc._add_text_box(slide, 1, 1, 4, 1, "hello",
                            color="000000", font_name="Calibri")
    assert box.text_frame.auto_size is not None


# ---------------- P4-50 / P5-48 失败历史按窗口统计 ----------------
def test_p4_50_recent_failures_window_filtered():
    from office_agent.model_gateway.failover import FailoverManager
    fm = FailoverManager.__new__(FailoverManager)
    now = time.time()
    fm.model_manager = SimpleNamespace(list_models=lambda: [
        SimpleNamespace(id="m1", display_name="M1", enabled=True, api_key="x")
    ])
    fm._failure_history = {"m1": [now - 7200, now - 5400]}  # 都是 1 小时前
    fm._is_in_cooldown = lambda _mid: False
    status = fm.get_model_status()
    assert status["m1"]["recent_failures"] == 0


# ---------------- P4-51 / P5-41 分隔线不算乱码 ----------------
def _garbled_report_for(text):
    from office_agent.quality.checker import QualityChecker, QualityReport
    qc = QualityChecker()
    report = QualityReport()
    doc = SimpleNamespace(paragraphs=[SimpleNamespace(text=text)])
    qc._check_garbled_text(doc, report)
    return report


def test_p4_51_separator_line_not_garbled():
    report = _garbled_report_for("-" * 20)
    assert report.issues == [], "纯连字符分隔线不应误报乱码"
    report2 = _garbled_report_for("_ " * 10)
    # 含空格的下划线序列同样不应命中“连续重复字符”
    assert not any("连续重复" in i.message for i in report2.issues)


def test_p4_51_real_repeat_still_flagged():
    report = _garbled_report_for("aaaaaaaa")  # 8 连字符级别的真重复
    assert any("连续重复" in i.message for i in report.issues)


# ---------------- P4-52 / P5-42 评分条钳制 ----------------
def test_p4_52_score_bar_clamped():
    from office_agent.quality_scoring.scoring_engine import ScoreResult
    over = ScoreResult._score_bar(150, width=20)
    under = ScoreResult._score_bar(-10, width=20)
    normal = ScoreResult._score_bar(50, width=20)
    assert len(over.strip("[]")) == 20 and over.count("░") == 0
    assert under.count("█") == 0
    assert normal.count("█") == 10


# ---------------- P5-4 短编号句不被误判为标题 ----------------
def test_p5_4_numbered_sentence_not_heading(tmp_path):
    from docx import Document
    from office_agent.knowledge_base.document_parser import DocumentParser
    doc = Document()
    doc.add_paragraph("1. 项目背景")                                   # 真标题
    doc.add_paragraph("2. 我们需要在本季度内完成交付，请大家配合。")   # 完整句子
    path = tmp_path / "p5_4.docx"
    doc.save(path)
    parsed = DocumentParser().parse(str(path))
    titles = [s.title for s in parsed.sections]
    assert "1. 项目背景" in titles
    assert not any("完成交付" in t for t in titles)


# ---------------- P5-13 文件名硬化 ----------------
def test_p5_13_filename_hardening():
    from office_agent.security.file_security.file_scanner import FileScanner
    fs = FileScanner()
    assert not fs.is_filename_safe("report.txt:secret")          # NTFS ADS
    assert not fs.is_filename_safe("invoice\u202eexe.docx")      # bidi 反转
    assert not fs.is_filename_safe("zero\u200bwidth.txt")        # 零宽
    assert not fs.is_filename_safe("a" * 252 + ".txt")           # 超长
    # 合法文件名（含可见 en-dash / 中文 / 下划线连字符）不受影响
    assert fs.is_filename_safe("Q3 报告—最终版_v1.docx")


# ---------------- P5-28 指标 get 与 inc 并发安全（定次、无 sleep） ----------------
def test_p5_28_metrics_get_concurrent_safe():
    from office_agent.logging_system.metrics import Counter, Gauge
    counter, gauge = Counter("c"), Gauge("g")
    errors = []

    def c_writer():
        for _ in range(20_000):
            counter.inc()

    def c_reader():
        for _ in range(20_000):
            try:
                counter.get()
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)

    def g_mix():
        for i in range(20_000):
            gauge.set(float(i))
            gauge.get()

    threads = [threading.Thread(target=c_writer) for _ in range(2)]
    threads += [threading.Thread(target=c_reader) for _ in range(2)]
    threads += [threading.Thread(target=g_mix) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors
    assert counter.get() == 40_000


# ---------------- P5-34 sc delete 返回码校验 ----------------
def _load_service_manager():
    spec = importlib.util.spec_from_file_location(
        "sm_under_test", REPO_ROOT / "desktop" / "service_manager.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_p5_34_sc_delete_must_succeed(monkeypatch, tmp_path):
    sm = _load_service_manager()
    monkeypatch.setattr(sm, "is_admin", lambda: True)

    def fail_delete(cmd, **kwargs):
        rc = 1 if "delete" in cmd else 0
        return subprocess.CompletedProcess(cmd, rc, stderr=b"access denied")

    monkeypatch.setattr(sm.subprocess, "run", fail_delete)
    assert sm.uninstall_service(tmp_path) is False

    def ok(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(sm.subprocess, "run", ok)
    assert sm.uninstall_service(tmp_path) is True


# ---------------- P5-37 read_only workbook 关闭前固化 sheetnames ----------------
def test_p5_37_excel_sheet_count_after_close(tmp_path):
    from openpyxl import Workbook
    from office_agent.knowledge_base.document_parser import DocumentParser
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "第一"
    ws1.append(["名称", "数值"])
    ws1.append(["甲", 1])
    ws2 = wb.create_sheet("第二")
    ws2.append(["名称", "数值"])
    ws2.append(["乙", 2])
    path = tmp_path / "p5_37.xlsx"
    wb.save(path)
    parsed = DocumentParser().parse(str(path))  # 内部 read_only=True
    assert parsed.metadata["sheet_count"] == 2
    assert "甲" in parsed.full_text and "乙" in parsed.full_text
