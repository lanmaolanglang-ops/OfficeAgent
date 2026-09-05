"""ExcelService ``self.wb`` 生命周期不变式回归（Mypy Phase 4 Cluster 4）。

根因：``self.wb: Optional[Workbook]`` 在 ``create()`` / ``open()`` 之前为 ``None``，
但 ``get_sheet`` / ``create_sheet`` / ``rename_sheet`` / ``delete_sheet`` /
``list_sheets`` 等方法直接访问 ``self.wb.xxx``，未对 ``None`` 做任何收窄。

修复：新增 ``workbook`` 访问器，把生命周期不变式显式化——打开前访问抛
明确 ``RuntimeError``，而非让调用方收到晦涩的
``AttributeError: 'NoneType' object has no attribute ...``。

本文件验证**运行时契约**：未打开时快速失败、打开后各方法正常。
"""
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook

from office_agent.excel_agent.excel_service import ExcelService


def test_workbook_access_before_open_raises_clear_error():
    svc = ExcelService()
    with pytest.raises(RuntimeError) as exc:
        _ = svc.workbook
    assert "尚未创建或打开" in str(exc.value)


def test_sheet_ops_before_open_raise_clear_error():
    svc = ExcelService()
    for op in (svc.get_sheet, svc.list_sheets):
        with pytest.raises(RuntimeError) as exc:
            op()
        assert "尚未创建或打开" in str(exc.value)


def test_workbook_after_create_is_usable(tmp_path):
    svc = ExcelService()
    out = tmp_path / "out.xlsx"
    svc.create(str(out))

    # 访问器返回真实 Workbook
    assert isinstance(svc.workbook, OpenpyxlWorkbook)
    assert svc.list_sheets() == ["Sheet1"]

    # 工作表操作全链路
    svc.create_sheet("数据")
    assert set(svc.list_sheets()) == {"Sheet1", "数据"}
    svc.rename_sheet("数据", "明细")
    assert "明细" in svc.list_sheets()
    svc.delete_sheet("明细")
    assert svc.list_sheets() == ["Sheet1"]
    assert svc.get_sheet("Sheet1") is not None


def test_workbook_after_open_reads_first_sheet(tmp_path):
    # 构造真实 xlsx 供 open() 读取
    src = tmp_path / "src.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws.title = "销售"
    ws["A1"] = "金额"
    wb.save(str(src))

    svc = ExcelService()
    svc.open(str(src))
    assert isinstance(svc.workbook, OpenpyxlWorkbook)
    assert svc.workbook.sheetnames[0] == "销售"
    assert svc.list_sheets() == ["销售"]
