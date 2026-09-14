"""
PPT 生成链修复的回归测试

覆盖：
- LLM 畸形输出（null 标题 / 非字符串要点 / 空表行 / 超量卡片 / 字符串图表值）不崩溃
- 4:3 等非 16:9 页面所有元素不越界
- 图表 numCache 仅含数值（避免 PowerPoint/WPS 报修复）
- 页数预算强制（LLM 超量返回被裁剪）
- 模板基底生成（主题继承 + 旧页清除）
"""
import sys
import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches
from pptx.shapes.graphfrm import GraphicFrame

PROJECT_ROOT = Path(__file__).parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from office_agent.ppt_agent.ppt_service import PPTService, hex_to_rgb  # noqa: E402
from office_agent.ppt_agent.content_planner import ContentPlanner  # noqa: E402
from office_agent.ppt_agent.slide_designer import SlideDesigner  # noqa: E402
from office_agent.ppt_agent.slide_planner import SlidePlanner  # noqa: E402
from office_agent.ppt_agent.models import PPTOutline, SlideContent  # noqa: E402
from office_agent.ppt_agent.quality_checker import PPTQualityChecker  # noqa: E402


def _hostile_outline() -> PPTOutline:
    outline = PPTOutline(title="测试")
    outline.add_slide(SlideContent(layout="cover", title=None, subtitle=None))
    outline.add_slide(SlideContent(layout="toc", title="目录",
                                   bullets=[1, 2, {"text": "三"}, None]))
    outline.add_slide(SlideContent(layout="table", title="表格页",
                                   table_data=[[], ["列A", "列B"], [None, 42], "标量行"]))
    outline.add_slide(SlideContent(layout="data_cards", title="卡片页",
                                   data=[(f"项{i}", str(i * 10), "%") for i in range(8)]))
    outline.add_slide(SlideContent(layout="chart", title="图表页", chart_type="column",
                                   chart_categories=["一月", "二月", "三月"],
                                   chart_series=[("销售额", ["100", "abc", None]),
                                                 ("成本", [1, 2, 3])]))
    outline.add_slide(SlideContent(layout="content_list", title="列表页",
                                   bullets=["项目" + "长" * 80, 3.14]))
    outline.add_slide(SlideContent(layout="two_column", title="两栏",
                                   left_content=[None, 5], right_content=[{"x": 1}]))
    outline.add_slide(SlideContent(layout="timeline", title="时间线",
                                   timeline_items=[("Q1", 42, None), ("Q2", "b", "c")]))
    outline.add_slide(SlideContent(layout="summary", title="总结", bullets=["谢谢"]))
    return outline


def _assert_shapes_within(prs: Presentation):
    w, h = prs.slide_width, prs.slide_height
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.left is None:
                continue
            assert (shape.left + (shape.width or 0)) <= w, \
                f"元素越出右边界: {shape.shape_type}"
            assert (shape.top + (shape.height or 0)) <= h, \
                f"元素越出下边界: {shape.shape_type}"


class TestPPTHostileContent:
    def test_hostile_outline_generates(self, temp_dir):
        """畸形 LLM 输出不应让生成崩溃"""
        out = temp_dir / "hostile.pptx"
        result = PPTService().generate(_hostile_outline(), str(out))
        assert result.success, result.message
        prs = Presentation(str(out))
        _assert_shapes_within(prs)

    def test_chart_numcache_numeric_only(self, temp_dir):
        """字符串/None 数值必须被转为数字（否则 numCache 非法，PowerPoint 报修复）"""
        out = temp_dir / "hostile.pptx"
        assert PPTService().generate(_hostile_outline(), str(out)).success

        prs = Presentation(str(out))
        charts = [s.chart for slide in prs.slides for s in slide.shapes
                  if isinstance(s, GraphicFrame) and s.has_chart]
        assert charts, "应生成图表"
        xml = charts[0]._chartSpace.xml
        num_caches = re.findall(r"<c:numCache>.*?</c:numCache>", xml, re.S)
        assert num_caches, "无 numCache"
        for cache in num_caches:
            vals = re.findall(r"<c:v>([^<]*)</c:v>", cache)
            assert all(re.match(r"^-?\d+(\.\d+)?$", v) for v in vals), vals

    def test_chart_rejects_series_category_length_mismatch(self):
        service = PPTService()
        service.prs = Presentation()
        slide = service.prs.slides.add_slide(service.prs.slide_layouts[6])
        with pytest.raises(ValueError, match="长度|类别"):
            service._add_chart(slide, 1, 1, 5, 3, "column",
                               ["一月", "二月"], [("销售", [1])])

    def test_chart_skips_only_bad_series_when_good_one_remains(self, temp_dir):
        """局部降级：坏系列跳过，好系列仍出图。"""
        from office_agent.ppt_agent.models import PPTOutline, SlideContent

        outline = PPTOutline(title="T")
        outline.slides = [
            SlideContent(layout="cover", title="封面"),
            SlideContent(
                layout="chart", title="图", chart_type="column",
                chart_categories=["一", "二"],
                chart_series=[("好", [1, 2]), ("坏", [9])],
            ),
        ]
        out = temp_dir / "partial_chart.pptx"
        result = PPTService().generate(outline, str(out))
        assert result.success, result.message
        prs = Presentation(str(out))
        charts = [
            s.chart for slide in prs.slides for s in slide.shapes
            if isinstance(s, GraphicFrame) and s.has_chart
        ]
        assert charts, "好系列应仍生成图表"

    def test_hex_color_validation_and_short_form(self):
        assert tuple(hex_to_rgb("#abc")) == (170, 187, 204)
        with pytest.raises(ValueError, match="无效"):
            hex_to_rgb("#12GG00")


class TestPPTScaling:
    def test_4x3_no_overflow(self, temp_dir):
        """4:3 模板尺寸下所有版式的元素都应在页面内"""
        outline = PPTOutline(title="43", slide_width=10.0, slide_height=7.5)
        for slide in [
            SlideContent(layout="cover", title="封面"),
            SlideContent(layout="content", title="内容", bullets=["a", "b"]),
            SlideContent(layout="data_cards", title="卡片",
                         data=[("x", "1", ""), ("y", "2", "")]),
            SlideContent(layout="timeline", title="时间线",
                         timeline_items=[("Q1", "启动", "立项"), ("Q2", "开发", "核心")]),
            SlideContent(layout="table", title="表格",
                         table_data=[["头1", "头2"], ["a", "b"]]),
            SlideContent(layout="summary", title="总结"),
        ]:
            outline.add_slide(slide)

        out = temp_dir / "43.pptx"
        result = PPTService().generate(outline, str(out))
        assert result.success, result.message
        _assert_shapes_within(Presentation(str(out)))


class TestSlideBudget:
    def test_budget_trims_middle_keeps_summary(self):
        planner = ContentPlanner()
        outline = PPTOutline(title="测试")
        for i in range(30):
            outline.add_slide(SlideContent(layout="content", title=f"页{i}"))
        outline.add_slide(SlideContent(layout="summary", title="总结"))

        trimmed = planner._enforce_slide_budget(outline, 10)
        assert len(trimmed.slides) == 10
        assert trimmed.slides[-1].layout == "summary"

    def test_dense_content_is_not_truncated(self):
        bullets = [f"要点{i}-" + "长" * 80 for i in range(8)]
        slide = SlideContent(layout="content", title="密集", bullets=bullets.copy())
        designed = SlideDesigner().design(PPTOutline(title="x", slides=[slide]))
        assert designed.slides[0].bullets == bullets
        assert "自动缩字" in designed.slides[0].notes

    def test_outline_data_preserves_chart_and_table_fields(self):
        outline = ContentPlanner().plan_from_outline_data("x", [{
            "layout": "chart", "title": "图表",
            "chart_type": "line", "chart_title": "趋势",
            "chart_categories": ["Q1", "Q2"],
            "chart_series": [["收入", [1, 2]]],
            "table_data": [["A"], ["B"]], "table_header": False,
        }])
        slide = outline.slides[0]
        assert slide.chart_categories == ["Q1", "Q2"]
        assert slide.chart_series == [["收入", [1, 2]]]
        assert slide.table_data == [["A"], ["B"]]
        assert slide.table_header is False

    def test_slide_reduction_preserves_order_summary_and_section_followers(self):
        outline = PPTOutline(title="x", slides=[
            SlideContent(layout="cover", title="封面"),
            SlideContent(layout="toc", title="目录"),
            SlideContent(layout="section", title="章节一"),
            SlideContent(layout="content", title="内容一"),
            SlideContent(layout="content", title="内容一扩展"),
            SlideContent(layout="section", title="章节二"),
            SlideContent(layout="content", title="内容二"),
            SlideContent(layout="summary", title="总结"),
        ])
        adjusted = ContentPlanner()._adjust_slide_count(outline, 6)
        titles = [slide.title for slide in adjusted.slides]
        assert titles[-1] == "总结"
        assert titles.index("章节一") < titles.index("内容一")
        assert titles.index("章节二") < titles.index("内容二")

    def test_missing_document_is_an_error_not_template_fallback(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            SlidePlanner().plan_from_document("x", str(tmp_path / "missing.docx"))

    def test_budget_noop_when_within(self):
        planner = ContentPlanner()
        outline = PPTOutline(title="测试")
        for i in range(5):
            outline.add_slide(SlideContent(layout="content", title=f"页{i}"))
        assert len(planner._enforce_slide_budget(outline, 10).slides) == 5

    def test_designer_paginates_without_losing_bullets_cards_or_rows(self):
        outline = PPTOutline(slides=[
            SlideContent(layout="content_list", title="列表", bullets=[str(i) for i in range(17)]),
            SlideContent(layout="data_cards", title="卡片", data=list(range(9))),
            SlideContent(layout="table", title="表格",
                         table_data=[["表头"]] + [[str(i)] for i in range(85)]),
        ])
        designed = SlideDesigner().design(outline)
        bullets = [item for slide in designed.slides for item in slide.bullets]
        cards = [item for slide in designed.slides for item in slide.data]
        table_rows = []
        for slide in designed.slides:
            if slide.layout == "table":
                table_rows.extend(slide.table_data[1:])
        assert bullets == [str(i) for i in range(17)]
        assert cards == list(range(9))
        assert table_rows == [[str(i)] for i in range(85)]

    def test_fix_outline_is_copy_on_write_and_reduces_font(self):
        outline = PPTOutline(slides=[
            SlideContent(layout="content", title="密集", bullets=["长" * 500]),
        ])
        report = PPTQualityChecker().check_outline(outline)
        fixed = PPTQualityChecker().fix_outline(outline, report)
        assert fixed is not outline
        assert outline.slides[0].body_font_size is None
        dense = next(slide for slide in fixed.slides if slide.title == "密集")
        assert dense.body_font_size is not None
        assert dense.bullets == outline.slides[0].bullets

    def test_fill_content_clamps_level(self):
        """P5-11：_fill_content 对 level 做 [0,8] 钳制，防负值/超大值。"""
        from pptx import Presentation
        from office_agent.ppt_agent.template_analyzer import TemplateAnalyzer

        prs = Presentation()
        blank = prs.slide_layouts[6]  # blank
        slide = prs.slides.add_slide(blank)
        placeholder = slide.shapes.add_textbox(
            0, 0, prs.slide_width, prs.slide_height
        )
        ta = TemplateAnalyzer()

        # 负 level → 0
        ta._fill_content(placeholder, {"bullets": [
            {"text": "a", "level": -3},
            {"text": "b", "level": 5},
            {"text": "c", "level": 99},
            {"text": "d", "level": "x"},
        ]})
        levels = [p.level for p in placeholder.text_frame.paragraphs]
        assert levels == [0, 5, 8, 0]

    def test_two_column_and_timeline_paginate_without_losing_items(self, temp_dir):
        """P5-11：固定几何版式超容量时分页，不能把条目挤出页面。"""
        output = temp_dir / "overflow.pptx"
        left = [f"左栏-{i}" for i in range(13)]
        right = [f"右栏-{i}" for i in range(11)]
        timeline = [(f"T{i}", f"节点-{i}", f"描述-{i}") for i in range(11)]
        outline = PPTOutline(slides=[
            SlideContent(
                layout="two_column", title="两栏", left_content=left,
                right_content=right, page_number=1,
            ),
            SlideContent(
                layout="timeline", title="时间线", timeline_items=timeline,
                page_number=2,
            ),
        ])

        result = PPTService().generate(outline, str(output))
        assert result.success, result.message
        prs = Presentation(str(output))
        # 两栏 3 页（13/6），时间线 3 页（11/5）。
        assert len(prs.slides) == 6
        assert result.slide_count == 6
        text = "\n".join(
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        lines = [line.strip() for line in text.splitlines()]
        for item in left + right:
            assert lines.count(f"• {item}") == 1, f"条目丢失或重复: {item}"
        for item in [part for row in timeline for part in row]:
            assert lines.count(item) == 1, f"条目丢失或重复: {item}"
        assert "两栏（续 2）" in text
        assert "时间线（续 2）" in text


class TestTemplateBaseDeck:
    def test_template_layout_indices_reject_negative_and_invalid_values(self, temp_dir):
        """P5-11：负/非法 layout_index 不得触发 Python 的负索引或类型错误。"""
        from office_agent.ppt_agent.template_analyzer import (
            LayoutInfo, TemplateAnalyzer, TemplateConfig,
        )

        config = TemplateConfig(layouts=[
            LayoutInfo(
                index=0, name="only", has_title=True,
                title_left=1.0, title_top=2.0,
                title_width=3.0, title_height=0.5,
            ),
        ])
        assert config.get_title_position(-1) == (0.5, 0.3, 12.333, 1.0)
        assert config.get_content_position(-1) == (0.5, 1.5, 12.333, 5.5)

        template = temp_dir / "layout-index-template.pptx"
        source = Presentation()
        source.save(str(template))
        expected_layout_name = source.slide_layouts[1].name
        output = temp_dir / "layout-index-output.pptx"
        TemplateAnalyzer().create_from_template(
            str(template),
            [
                {"layout_index": -1, "title": "negative"},
                {"layout_index": "bad", "title": "invalid"},
            ],
            str(output),
        )
        generated = Presentation(str(output))
        assert len(generated.slides) == 2
        assert all(
            slide.slide_layout.name == expected_layout_name
            for slide in generated.slides
        )

    def test_theme_inherited_and_slides_cleared(self, temp_dir):
        """以模板为基底：主题部件继承、模板内容页清除、输出尺寸继承"""
        from office_agent.ppt_agent.ppt_orchestrator import PPTOrchestrator

        template = temp_dir / "tpl.pptx"
        tpl = Presentation()
        tpl.slide_width = Inches(10)
        tpl.slide_height = Inches(7.5)
        s = tpl.slides.add_slide(tpl.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = (
            "OLD CONTENT TO BE CLEARED"
        )
        tpl.save(str(template))
        # 在主题 XML 植入标记色，验证继承
        marker_file = temp_dir / "tpl_marked.pptx"
        with zipfile.ZipFile(str(template)) as zin, \
                zipfile.ZipFile(str(marker_file), "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item.startswith("ppt/theme/theme"):
                    data = data.replace(b'<a:srgbClr val="4F81BD"/>',
                                        b'<a:srgbClr val="CA00CA"/>')
                zout.writestr(item, data)

        out = temp_dir / "out.pptx"
        orch = PPTOrchestrator(model_gateway=None, image_gateway=None)
        result = orch.generate_with_template(template_path=str(marker_file),
                                             theme="季度经营汇报", slide_count=5,
                                             output_path=str(out))
        assert result.success, result.message

        # 主题标记继承
        with zipfile.ZipFile(str(out)) as z:
            themes = [n for n in z.namelist() if n.startswith("ppt/theme/theme")]
            assert any(b"CA00CA" in z.read(n) for n in themes), "主题未继承"

        prs = Presentation(str(out))
        # 模板尺寸继承（4:3）
        assert abs(Emu(prs.slide_width).inches - 10.0) < 0.01
        assert abs(Emu(prs.slide_height).inches - 7.5) < 0.01
        # 模板旧内容页被清除
        assert not any(
            sh.has_text_frame and "OLD CONTENT" in sh.text_frame.text
            for slide in prs.slides for sh in slide.shapes
        )
        _assert_shapes_within(prs)


class TestFixOutlineCompleteness:
    """清单 684：fixable 问题的每种 fix_action 都必须有真实修复分支。

    历史条目（reduce_font/unify_font 仅 pass）已由 277 关闭；复核发现
    apply_template_size（模板尺寸不一致，fixable=True）在 fix_outline 中
    无分支被静默忽略，本批补齐。本类钉住"发出的每种 action 都有分支"。
    """

    def test_apply_template_size_action_fixes_outline_dimensions(self):
        from office_agent.ppt_agent.template_analyzer import TemplateConfig
        from office_agent.ppt_agent.quality_checker import PPTQualityReport, PPTQualityIssue

        checker = PPTQualityChecker()
        checker.template_config = TemplateConfig(slide_width=10.0, slide_height=7.5)
        outline = PPTOutline(
            title="t",
            slide_width=13.333, slide_height=7.5,
            slides=[SlideContent(layout="content", title="x")],
        )
        report = PPTQualityReport(file_path="x")
        report.issues.append(PPTQualityIssue(
            slide_index=-1, issue_type="template", severity="warning",
            message="页面尺寸与模板不一致",
            fixable=True, fix_action={"type": "apply_template_size"},
        ))
        fixed = checker.fix_outline(outline, report)
        assert fixed.slide_width == 10.0
        assert fixed.slide_height == 7.5
        # copy-on-write：原 outline 不受影响
        assert outline.slide_width == 13.333

    def test_every_emitted_action_type_has_fix_branch(self):
        """源码守卫：quality_checker 发出的每种 fix_action type 都必须在
        fix_outline 中有对应 elif 分支，防止再出现"标记 fixable 却无修复"。"""
        import ast
        import inspect
        import textwrap
        import office_agent.ppt_agent.quality_checker as qc_module

        src = inspect.getsource(qc_module)
        tree = ast.parse(src)
        emitted = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) \
                    and node.func.id == "PPTQualityIssue":
                for kw in node.keywords:
                    if kw.arg == "fix_action" and isinstance(kw.value, ast.Dict):
                        for k, v in zip(kw.value.keys, kw.value.values):
                            if isinstance(k, ast.Constant) and k.value == "type" \
                                    and isinstance(v, ast.Constant):
                                emitted.add(v.value)

        fixed_src = textwrap.dedent(inspect.getsource(
            qc_module.PPTQualityChecker.fix_outline))
        handled = set(re.findall(r'atype == "([a-z_]+)"', fixed_src))

        # add_slides 为刻意不自动修复（检查阶段标 fixable=False），除外
        assert emitted, "守卫失效：未收集到任何 fix_action"
        assert emitted <= handled | {"add_slides"}, (
            f"以下 fix_action 在 fix_outline 中无分支: {emitted - handled}"
        )

