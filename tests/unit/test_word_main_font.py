"""Word 主字体识别钉住测试（清单滞后核验）。

`word_scorer._analyze_paragraphs` 的主字体选择已修复为：
1. 优先从包含 CJK 字符的 run 中取主字体（段首英文/数字 run 不再把
   西文字体误当作中文主字体）；
2. 同类候选中按承载字符数（run 文本长度）取主导 run；
3. 无 CJK run 时回退到全部候选；
4. 无有效 run 时主字体为 None。

本文件把这些行为钉住，防止回退。
"""
from docx import Document
from docx.shared import Pt

from office_agent.quality_scoring.word_scorer import WordQualityScorer


def _paragraph_info(runs):
    """构造单段文档并返回其段落分析信息。runs: [(text, font_name)]"""
    doc = Document()
    para = doc.add_paragraph()
    for text, font_name in runs:
        run = para.add_run(text)
        run.font.name = font_name
        run.font.size = Pt(12)
    return WordQualityScorer()._analyze_paragraphs(doc)[0]


def test_cjk_run_wins_over_leading_western_run():
    info = _paragraph_info([
        ("2026 report: ", "Arial"),
        ("年度总结报告正文内容", "宋体"),
    ])
    assert info["main_font"] == "宋体"


def test_leading_digits_do_not_hijack_main_font():
    info = _paragraph_info([
        ("12345", "Times New Roman"),
        ("中文段落主体", "黑体"),
    ])
    assert info["main_font"] == "黑体"


def test_dominant_cjk_run_selected_by_char_count():
    info = _paragraph_info([
        ("短", "黑体"),
        ("这是一段更长的中文正文内容", "宋体"),
    ])
    assert info["main_font"] == "宋体"


def test_no_cjk_falls_back_to_longest_run():
    info = _paragraph_info([
        ("ab", "Courier New"),
        ("longer english text", "Arial"),
    ])
    assert info["main_font"] == "Arial"


def test_run_without_font_name_has_no_main_font():
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run("没有显式字体的文本")
    run.font.size = Pt(12)
    info = WordQualityScorer()._analyze_paragraphs(doc)[0]
    assert info["main_font"] is None
