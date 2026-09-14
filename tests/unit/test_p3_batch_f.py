# -*- coding: utf-8 -*-
"""P3-F 批（质量/Word）回归：P3-94/95/97/98/99/101。"""
import glob
import os
import tempfile
from types import SimpleNamespace

import pytest


# ---------- P3-94 公式类型精确匹配，SUM 不再误命中 SUMIF ----------
def test_p3_94_formula_exact_match():
    from office_agent.quality_scoring.excel_scorer import ExcelQualityScorer
    scorer = ExcelQualityScorer()
    analysis = {
        "total_formulas": 4, "formula_count": 4, "formula_error_rate": 0,
        "formula_types": {"SUMIF"},
    }
    base = scorer._score_formulas(analysis, None)
    hit_sumif = scorer._score_formulas(analysis, ["SUMIF"])
    miss_sum = scorer._score_formulas(analysis, ["SUM"])
    assert hit_sumif == pytest.approx(base + 10)
    assert miss_sum == pytest.approx(base)          # SUM 不得子串命中 SUMIF
    # 大小写/空白归一
    assert scorer._score_formulas(analysis, [" sumif "]) == pytest.approx(base + 10)


# ---------- P3-95 两个 workbook 确定性关闭（真实 xlsx） ----------
def test_p3_95_workbooks_closed(tmp_path, monkeypatch):
    import openpyxl
    import office_agent.quality_scoring.excel_scorer as es
    path = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=SUM(B1:B3)"
    wb.save(str(path))
    wb.close()

    closed = []
    orig = es.load_workbook

    class _Wrap:
        def __init__(self, inner):
            self._inner = inner

        def close(self):
            closed.append(1)
            self._inner.close()

        def __getattr__(self, name):
            return getattr(self._inner, name)

        def __getitem__(self, key):
            return self._inner[key]

    def fake(path, data_only):
        return _Wrap(orig(path, data_only=data_only))

    monkeypatch.setattr(es, "load_workbook", fake)
    result = es.ExcelQualityScorer().score(str(path))
    assert result.total_score >= 0
    assert len(closed) == 2                          # 公式视图 + 数据视图都关闭


# ---------- P3-97 缩进短段落无引用标点不算引用 ----------
def test_p3_97_indent_alone_not_quote():
    from office_agent.parsers.document_structure import DocumentStructureAnalyzer
    analyzer = DocumentStructureAnalyzer()
    para = SimpleNamespace(
        paragraph_format=SimpleNamespace(left_indent=SimpleNamespace(pt=72))
    )
    assert analyzer._is_quote("这是一段缩进的普通正文，没有任何引号标记。", para) is False
    assert analyzer._is_quote("“这是带引号的缩进引用。”", para) is True
    zero = SimpleNamespace(
        paragraph_format=SimpleNamespace(left_indent=None)
    )
    assert analyzer._is_quote("“无缩进但引号开头”", zero) is True  # QUOTE_INDICATORS


# ---------- P3-98 “行”结合字号换算 ----------
def test_p3_98_line_unit_uses_font_size():
    from office_agent.parsers.format_parser import FormatRuleParser
    p = FormatRuleParser()
    assert p._extract_space_before("段前1行") == pytest.approx(12.0)
    # 小四=12pt，单倍行距 1.3 倍
    assert p._extract_space_before("小四 段前1行") == pytest.approx(15.6)
    assert p._extract_space_after("三号 段后2行") == pytest.approx(16.0 * 1.3 * 2)


# ---------- P3-99 save 失败清理空临时文件 ----------
def test_p3_99_failed_save_removes_temp(tmp_path, monkeypatch):
    import docx
    import docx.document
    from office_agent.services import input_conversion
    txt = tmp_path / "a.txt"
    txt.write_text("第一段\n第二段", encoding="utf-8")

    before = set(glob.glob(os.path.join(tempfile.gettempdir(),
                                        input_conversion.OWNED_TEMP_PREFIX + "*.docx")))

    def boom(self, path):
        raise OSError("disk full")

    monkeypatch.setattr(docx.document.Document, "save", boom)
    with pytest.raises(OSError):
        input_conversion.ensure_docx_input(str(txt))

    after = set(glob.glob(os.path.join(tempfile.gettempdir(),
                                       input_conversion.OWNED_TEMP_PREFIX + "*.docx")))
    assert after == before                        # 没有遗留新的空临时文件
    # 非文本类输入原样返回
    assert input_conversion.ensure_docx_input(str(tmp_path / "x.xlsx")) == \
        str(tmp_path / "x.xlsx")


# ---------- P3-101 维度不匹配显式报错，空向量仍为 0 ----------
def test_p3_101_cosine_dimension_mismatch_loud():
    from office_agent.knowledge_base.embeddings import (
        cosine_similarity, InvalidEmbeddingVectorError,
    )
    assert cosine_similarity([1.0, 0.0], [1.0, 0.0]) == pytest.approx(1.0)
    assert cosine_similarity([], [1.0]) == 0.0
    assert cosine_similarity([0.0, 0.0], [1.0, 0.0]) == 0.0   # 零范数
    with pytest.raises(InvalidEmbeddingVectorError):
        cosine_similarity([1.0, 0.0, 0.0], [1.0, 0.0])        # 维度不符
