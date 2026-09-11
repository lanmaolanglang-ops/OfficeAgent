"""Excel 质量器"统一正文字体"不得抹掉用户强调格式（P1-12）。

断言对象是真实的 openpyxl Font 属性，而不是"没抛异常"。
"""
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from office_agent.excel_agent.quality_checker import (
    BODY_FONT_NAME,
    BODY_FONT_SIZE,
    ExcelQualityChecker,
    _build_body_font,
)
from office_agent.excel_agent.models import ExcelQualityIssue
from office_agent.quality.checker import IssueSeverity


def _font_issue():
    return ExcelQualityIssue(
        sheet_name="Sheet",
        issue_type="format",
        severity=IssueSeverity.INFO.value,
        message="正文字体种类过多（4种），建议统一",
        fixable=True,
    )


@pytest.fixture
def ws():
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Sheet"
    sheet["A1"], sheet["B1"], sheet["C1"] = "品名", "金额", "备注"
    sheet["A2"], sheet["B2"], sheet["C2"] = "甲", 100, "正常"
    sheet["A3"], sheet["B3"], sheet["C3"] = "乙", 200, "警告"
    sheet["A4"], sheet["B4"], sheet["C4"] = "丙", 300, "合计"
    return sheet


def _run_fix(ws):
    return ExcelQualityChecker()._fix_sheet(ws, [_font_issue()])


# ------------------------------------------------------------------
# 1. 强调格式必须保留
# ------------------------------------------------------------------

def test_bold_is_preserved(ws):
    ws["A2"].font = Font(name="宋体", size=12, bold=True)
    _run_fix(ws)

    assert ws["A2"].font.name == BODY_FONT_NAME
    assert ws["A2"].font.sz == BODY_FONT_SIZE
    assert ws["A2"].font.b is True, "用户加粗不得被统一字体抹掉"


def test_font_color_is_preserved(ws):
    ws["A2"].font = Font(name="宋体", size=12, color="FFFF0000")
    _run_fix(ws)

    assert ws["A2"].font.name == BODY_FONT_NAME
    assert ws["A2"].font.color is not None
    assert ws["A2"].font.color.rgb.endswith("FF0000"), "标红语义必须保留"


def test_italic_underline_strike_are_preserved(ws):
    ws["A2"].font = Font(name="宋体", size=12, italic=True, underline="single", strike=True)
    _run_fix(ws)

    assert ws["A2"].font.i is True
    assert ws["A2"].font.u == "single"
    assert ws["A2"].font.strike is True


def test_combined_emphasis_survives(ws):
    ws["A2"].font = Font(name="宋体", size=14, bold=True, italic=True, color="FFFFFF00")
    _run_fix(ws)

    f = ws["A2"].font
    assert (f.b, f.i) == (True, True)
    assert f.color.rgb.endswith("FFFF00")
    assert f.name == BODY_FONT_NAME and f.sz == BODY_FONT_SIZE


# ------------------------------------------------------------------
# 2. 统一字体本身仍然生效
# ------------------------------------------------------------------

def test_default_cell_gets_unified_font(ws):
    _run_fix(ws)
    for ref in ("A2", "B2", "C2", "A3", "B3"):
        assert ws[ref].font.name == BODY_FONT_NAME
        assert ws[ref].font.sz == BODY_FONT_SIZE


def test_multi_font_table_is_actually_unified(ws):
    ws["A2"].font = Font(name="宋体", size=9)
    ws["B2"].font = Font(name="Arial", size=11)
    ws["C2"].font = Font(name="Times New Roman", size=13)
    _run_fix(ws)

    assert {ws[c].font.name for c in ("A2", "B2", "C2")} == {BODY_FONT_NAME}
    assert {ws[c].font.sz for c in ("A2", "B2", "C2")} == {BODY_FONT_SIZE}


def test_fix_is_counted(ws):
    assert _run_fix(ws) >= 1


# ------------------------------------------------------------------
# 3. 不该被当成正文的单元格
# ------------------------------------------------------------------

def test_formula_cells_are_untouched(ws):
    ws["C2"] = "=SUM(A2:B2)"
    ws["C2"].font = Font(name="宋体", size=12, bold=True)
    _run_fix(ws)

    assert ws["C2"].font.name == "宋体", "公式格沿用旧行为：不参与正文统一"
    assert ws["C2"].font.sz == 12


def test_header_row_is_not_treated_as_body(ws):
    ws["A1"].font = Font(name="黑体", size=16, bold=True, color="FF1F4E79")
    _run_fix(ws)

    assert ws["A1"].font.name == "黑体", "表头不是正文，不能被统一掉"
    assert ws["A1"].font.sz == 16


def test_empty_cells_are_untouched(ws):
    ws["A3"] = None
    _run_fix(ws)
    assert ws["A3"].font.name != BODY_FONT_NAME or ws["A3"].font.sz != BODY_FONT_SIZE


# ------------------------------------------------------------------
# 4. _build_body_font 单元行为
# ------------------------------------------------------------------

def test_build_body_font_preserves_semantics():
    src = Font(name="宋体", size=12, bold=True, italic=True, color="FFFF0000")
    out = _build_body_font(src)

    assert out.name == BODY_FONT_NAME
    assert out.sz == BODY_FONT_SIZE
    assert out.b is True and out.i is True
    assert out.color.rgb.endswith("FF0000")


def test_build_body_font_handles_none_existing():
    out = _build_body_font(None)
    assert out.name == BODY_FONT_NAME
    assert out.sz == BODY_FONT_SIZE


def test_build_body_font_does_not_mutate_source():
    src = Font(name="宋体", size=12, bold=True)
    out = _build_body_font(src)

    assert src.name == "宋体" and src.sz == 12, "Font 不可原地修改"
    assert out is not src


def test_build_body_font_accepts_override():
    out = _build_body_font(Font(bold=True), name="Arial", size=8)
    assert out.name == "Arial" and out.sz == 8 and out.b is True


def test_false_flags_are_not_lost(ws):
    """显式设置为 False 的字段（如 b=False）也应按原值继承，不变成 None。"""
    ws["A2"].font = Font(name="宋体", size=12, bold=True, italic=False)
    _run_fix(ws)
    assert ws["A2"].font.i is False
    assert ws["A2"].font.b is True
