"""文本类输入 → docx 转换链专项测试（清单 273）。"""
import sys

import pytest
from docx import Document

from office_agent.services.input_conversion import ensure_docx_input, is_text_like


def _docx_text(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


class TestIsTextLike:
    def test_text_like_extensions(self):
        for path in ("a.txt", "b.md", "c.pdf", "D.TXT"):
            assert is_text_like(path)

    def test_office_files_not_text_like(self):
        for path in ("a.docx", "b.pptx", "c.xlsx", ""):
            assert not is_text_like(path)


class TestEnsureDocxInput:
    def test_docx_passthrough(self, tmp_path):
        path = tmp_path / "in.docx"
        Document().save(str(path))
        assert ensure_docx_input(str(path)) == str(path)

    def test_txt_converts_with_content(self, tmp_path):
        src = tmp_path / "报告.txt"
        src.write_text("第一季度总结\n营收增长 20%\n利润同步提升", encoding="utf-8")
        out = ensure_docx_input(str(src))
        assert out.endswith(".docx") and out != str(src)
        text = _docx_text(out)
        assert "营收增长 20%" in text
        assert "利润同步提升" in text

    def test_md_headings_become_word_headings(self, tmp_path):
        src = tmp_path / "notes.md"
        src.write_text("# 章节一\n正文内容\n## 小节\n更多内容", encoding="utf-8")
        out = ensure_docx_input(str(src))
        doc = Document(out)
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "章节一" in headings
        assert "小节" in headings
        assert "正文内容" in _docx_text(out)

    def test_gbk_encoded_txt(self, tmp_path):
        src = tmp_path / "legacy.txt"
        src.write_bytes("中文内容GBK编码".encode("gbk"))
        out = ensure_docx_input(str(src))
        assert "中文内容GBK编码" in _docx_text(out)

    def test_pdf_without_pymupdf_raises_clear_error(self, tmp_path, monkeypatch):
        monkeypatch.setitem(sys.modules, "pymupdf", None)
        src = tmp_path / "a.pdf"
        src.write_bytes(b"%PDF-1.4 fake")
        with pytest.raises(RuntimeError, match="pymupdf"):
            ensure_docx_input(str(src))


class TestMarkdownTitleAndLevels:
    """RC：存储落盘文件名是内部 ID，绝不能泄漏为成品标题；标题层级不得错位。"""

    @staticmethod
    def _heading_levels(path):
        doc = Document(ensure_docx_input(str(path)))
        return [(p.style.name, p.text) for p in doc.paragraphs
                if p.style.name.startswith("Heading")]

    def test_internal_storage_id_never_becomes_heading(self, tmp_path):
        # 存储层用 file_<hex> 落盘，模拟这一真实文件名
        src = tmp_path / "file_0c9309764523.md"
        src.write_text("# 真实标题\n正文\n## 次级\n内容", encoding="utf-8")
        headings = self._heading_levels(src)
        texts = [t for _, t in headings]
        assert "file_0c9309764523" not in texts
        assert headings[0] == ("Heading 1", "真实标题")

    def test_markdown_heading_levels_align_with_word(self, tmp_path):
        src = tmp_path / "file_aaaaaaaa1234.md"
        src.write_text("# H1\n\n## H2\n\n### H3\n", encoding="utf-8")
        headings = self._heading_levels(src)
        assert headings == [("Heading 1", "H1"), ("Heading 2", "H2"), ("Heading 3", "H3")]

    def test_txt_without_h1_has_no_internal_id_title(self, tmp_path):
        src = tmp_path / "out_deadbeef9999.txt"
        src.write_text("第一行\n第二行\n", encoding="utf-8")
        headings = self._heading_levels(src)
        assert headings == []
        assert "out_deadbeef9999" not in _docx_text(ensure_docx_input(str(src)))

    def test_markdown_pipe_table_becomes_real_table(self, tmp_path):
        src = tmp_path / "t.md"
        src.write_text(
            "# 标题\n\n| 季度 | 收入 |\n| --- | --- |\n| Q1 | 100 |\n| Q2 | 130 |\n\n表后正文\n",
            encoding="utf-8")
        doc = Document(ensure_docx_input(str(src)))
        assert len(doc.tables) == 1
        rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        assert rows == [["季度", "收入"], ["Q1", "100"], ["Q2", "130"]]
        text = _docx_text(ensure_docx_input(str(src)))
        assert "---" not in text  # 分隔行不得作为垃圾段落残留
        assert "表后正文" in text

    def test_lone_pipe_line_without_delimiter_is_paragraph(self, tmp_path):
        src = tmp_path / "t.md"
        src.write_text("# T\n\n备注 | 说明\n普通行\n", encoding="utf-8")
        doc = Document(ensure_docx_input(str(src)))
        assert len(doc.tables) == 0
        assert "备注 | 说明" in _docx_text(ensure_docx_input(str(src)))


class TestWordTaskWiring:
    def test_process_word_converts_before_engine(self):
        """转换必须发生在 WordService 调用之前（源码接线守卫）。"""
        import inspect

        from office_agent.task_queue.tasks import word_tasks
        src = inspect.getsource(word_tasks.process_word)
        assert src.index("ensure_docx_input") < src.index("WordService")
