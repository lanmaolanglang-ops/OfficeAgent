# -*- coding: utf-8 -*-
"""P3 batch B regression tests (Excel)."""
import json

import openpyxl
import pytest


def _col(name, idx=0, dt="number", sem="metric"):
    from office_agent.excel_agent.models import ColumnInfo
    return ColumnInfo(index=idx, name=name, data_type=dt, semantic_type=sem)


def _sheet(cols, rows=3):
    from office_agent.excel_agent.models import SheetInfo
    return SheetInfo(name="s", columns=cols, row_count=rows)


# ---------- P3-21: shared intent vocabulary, case-insensitive ----------
def test_p3_21_intent_keywords_shared_and_case_insensitive():
    from office_agent.excel_agent.excel_orchestrator import (
        _task_has_any, FORMULA_INTENT_KEYWORDS, CHART_INTENT_KEYWORDS,
    )
    assert _task_has_any("calculate the SUM", FORMULA_INTENT_KEYWORDS)
    assert _task_has_any("计算平均值", FORMULA_INTENT_KEYWORDS)
    assert _task_has_any("画个 CHART", CHART_INTENT_KEYWORDS)
    assert not _task_has_any("", FORMULA_INTENT_KEYWORDS)


# ---------- P3-23: short YoY series -> no orphan result column ----------
def test_p3_23_short_yoy_series_no_orphan_header():
    from office_agent.excel_agent.formula_generator import (
        FormulaGenerator, FormulaTemplate,
    )
    fg = FormulaGenerator()
    cols = [_col("销售额")]
    tpl = FormulaTemplate("growth_yoy", [r"同比"],
                          "=IFERROR(({cur}-{prev_year})/{prev_year},0)",
                          "math", "同比增长率")
    specs = fg._generate_row_wise(tpl, cols, _sheet(cols, 4),
                                  data_start_row=2, text="月度同比")
    assert specs == []


# ---------- P3-24: IF comparator negation/completeness ----------
@pytest.mark.parametrize("text,expected", [
    ("不大于80", "{cell}<=80"), ("小于等于60", "{cell}<=60"),
    ("不高于90", "{cell}<=90"), ("大于等于60", "{cell}>=60"),
    ("不低于70", "{cell}>=70"), ("不小于75", "{cell}>=75"),
    ("大于50", "{cell}>50"), ("小于30", "{cell}<30"),
    ("超过90", "{cell}>90"), ("低于5", "{cell}<5"),
])
def test_p3_24_if_comparators(text, expected):
    from office_agent.excel_agent.formula_generator import FormulaGenerator
    cond, _, _ = FormulaGenerator()._parse_if_condition(text, _col("x"))
    assert cond == expected


# ---------- P3-25: single-char header must not over-select ----------
def test_p3_25_fuzzy_column_selection_not_overbroad():
    from office_agent.excel_agent.formula_generator import FormulaGenerator
    cols = [_col("销售额", 0), _col("利润额", 1), _col("成本额", 2)]
    picked = FormulaGenerator()._detect_target_columns("计算销售额", _sheet(cols))
    assert [c.name for c in picked] == ["销售额"]


# ---------- P3-28: Timestamp preview values are JSON-safe ----------
def test_p3_28_timestamp_unique_values_json_safe():
    import pandas as pd
    from office_agent.excel_agent.data_analyzer import DataAnalyzer
    df = pd.DataFrame({"日期": pd.to_datetime(["2024-01-01", "2024-02-01"]),
                       "数值": [1, 2]})
    prof = DataAnalyzer().analyze_dataframe(df, "s")
    col = next(c for c in prof.sheets[0].columns if c.name == "日期")
    json.dumps(col.unique_values, ensure_ascii=False)  # must not raise


# ---------- P3-31: interior text column not mis-picked as time axis ----------
def test_p3_31_time_column_fallback_only_leading():
    from office_agent.excel_agent.analysis_engine import AnalysisEngine
    eng = AnalysisEngine()
    headers = ["数值", "产品名称", "备注"]
    # no date col, first text col is index 1 (interior) -> -1
    assert eng._find_time_column(headers, [], [1, 2]) == -1
    # leading text column is the conventional axis
    assert eng._find_time_column(["月份", "数值"], [], [0]) == 0


# ---------- P3-32/33: LOG10 not a reference; real out-of-range still flagged ----------
def test_p3_32_33_reference_word_boundary():
    from office_agent.excel_agent.quality_checker import ExcelQualityChecker
    qc = ExcelQualityChecker()
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 12):
        ws.cell(row=r, column=1, value=1)
    assert qc._check_formula_refs(ws, "=LOG10(A2)", "C1") == []
    assert qc._check_formula_refs(ws, "=POWER10(A2)", "C2") == []
    assert qc._check_formula_refs(ws, "=A999", "C3")


# ---------- P3-34: absolute-ref denominator div-zero recognized ----------
def test_p3_34_absolute_ref_div_zero():
    from office_agent.excel_agent.quality_checker import ExcelQualityChecker
    qc = ExcelQualityChecker()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = 0
    issues = qc._check_div_zero_risk(ws, "=A1/$B$2", "C2")
    assert issues and issues[0].severity == "error"


# ---------- P3-35: fill_template headers actually written (real XLSX) ----------
def test_p3_35_fill_template_applies_headers(tmp_path):
    from office_agent.excel_agent.template_analyzer import ExcelTemplateAnalyzer
    tpl = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "旧A"
    ws["B1"] = "旧B"
    ws["A2"] = 1
    wb.save(tpl)
    wb.close()
    out = tmp_path / "out.xlsx"
    ExcelTemplateAnalyzer().fill_template(
        str(tpl), [[10, 20]], str(out), sheet_name="Sheet1",
        headers=["新名称", "新数值"])
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2["Sheet1"]
    assert [ws2["A1"].value, ws2["B1"].value] == ["新名称", "新数值"]
    wb2.close()
