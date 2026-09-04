"""models.col_letter 统一列字母换算的针对性测试。

根因背景：chart_generator._col_letter 按 1-based 实现（调用方全部
`index + 1`），formula_generator._col_letter 按 0-based 实现，同名
方法跨类口径漂移。现统一为 models.col_letter（0-based，对齐
ColumnInfo.index 约定），两个生成器均委托之。
"""
import pytest
from openpyxl.utils import get_column_letter

from office_agent.excel_agent.models import col_letter


class TestColLetter:
    @pytest.mark.parametrize("index", [0, 1, 25, 26, 27, 51, 52, 701, 702, 16383])
    def test_matches_openpyxl(self, index):
        assert col_letter(index) == get_column_letter(index + 1)

    def test_boundary_letters(self):
        assert col_letter(0) == "A"
        assert col_letter(25) == "Z"
        assert col_letter(26) == "AA"
        assert col_letter(16383) == "XFD"

    @pytest.mark.parametrize("bad", [-1, "A", 1.5, None])
    def test_invalid_index_raises(self, bad):
        with pytest.raises(ValueError):
            col_letter(bad)


class TestGeneratorDelegation:
    """两个生成器的同名方法必须口径一致（0-based）。"""

    def test_chart_generator_delegates_zero_based(self):
        from office_agent.excel_agent.chart_generator import ChartGenerator
        assert ChartGenerator._col_letter(0) == "A"
        assert ChartGenerator._col_letter(2) == "C"

    def test_formula_generator_delegates_zero_based(self):
        from office_agent.excel_agent.formula_generator import FormulaGenerator
        assert FormulaGenerator._col_letter(0) == "A"
        assert FormulaGenerator._col_letter(2) == "C"

    def test_both_generators_agree(self):
        from office_agent.excel_agent.chart_generator import ChartGenerator
        from office_agent.excel_agent.formula_generator import FormulaGenerator
        for i in (0, 7, 25, 26, 100):
            assert ChartGenerator._col_letter(i) == FormulaGenerator._col_letter(i) == col_letter(i)

    def test_chart_generator_source_has_no_manual_plus_one(self):
        """chart_generator 调用方不得再手动 +1（口径在统一入口内）。"""
        import inspect
        from office_agent.excel_agent import chart_generator
        src = inspect.getsource(chart_generator)
        assert "_col_letter(cat_col.index + 1)" not in src
        assert "_col_letter(col.index + 1)" not in src


class TestSingleImplementation:
    """清单滞后核验钉住：列字母换算全包唯一实现，无隐藏兼容函数。

    2026-09-04 核验结论：models.col_letter 为唯一换算实现；
    chart_generator._col_letter / formula_generator._col_letter 均为
    纯委托；其余模块使用 openpyxl 官方 get_column_letter（worksheet
    层 1-based API，不同抽象层，非重新实现）。
    """

    def test_conversion_algorithm_lives_only_in_models(self):
        """换算算法（divmod 进制循环）只允许出现在 models.py。"""
        import office_agent.excel_agent as pkg
        from pathlib import Path
        pkg_dir = Path(pkg.__file__).parent
        offenders = []
        for py in pkg_dir.glob("*.py"):
            if "divmod" in py.read_text(encoding="utf-8"):
                offenders.append(py.name)
        assert offenders == ["models.py"]

    def test_delegates_contain_no_conversion_logic(self):
        """两个生成器的 _col_letter 必须是纯委托（无 divmod/chr 换算）。"""
        import inspect
        from office_agent.excel_agent.chart_generator import ChartGenerator
        from office_agent.excel_agent.formula_generator import FormulaGenerator
        for cls in (ChartGenerator, FormulaGenerator):
            src = inspect.getsource(cls._col_letter)
            assert "col_letter(index)" in src
            assert "divmod" not in src and "chr(" not in src

    def test_no_other_col_letter_definitions(self):
        """包内除 models.col_letter 与两个委托外无第三份定义。"""
        import office_agent.excel_agent as pkg
        from pathlib import Path
        pkg_dir = Path(pkg.__file__).parent
        definitions = []
        for py in pkg_dir.glob("*.py"):
            for line in py.read_text(encoding="utf-8").splitlines():
                stripped = line.strip()
                if stripped.startswith("def col_letter") or stripped.startswith("def _col_letter"):
                    definitions.append((py.name, stripped))
        definitions.sort()
        assert definitions == [
            ("chart_generator.py", "def _col_letter(index: int) -> str:"),
            ("formula_generator.py", "def _col_letter(index: int) -> str:"),
            ("models.py", "def col_letter(index: int) -> str:"),
        ]

    def test_column_info_index_stays_zero_based_at_boundary(self):
        """data_analyzer 产出的 ColumnInfo.index 保持 0-based 约定。"""
        import pandas as pd
        from office_agent.excel_agent.data_analyzer import DataAnalyzer
        analyzer = DataAnalyzer()
        df = pd.DataFrame({"地区": ["华东", "华北"], "销量": [10, 20]})
        profile = analyzer.analyze_dataframe(df)
        cols = {c.name: c for c in profile.sheets[0].columns}
        assert cols["地区"].index == 0
        assert cols["销量"].index == 1
