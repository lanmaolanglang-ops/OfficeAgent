"""PPT 渲染分页回归测试（P1-15）。

目录页 / 列表页原来的 items[:6] / items[:8] 是**单页布局容量**，
但超容量时直接把尾部条目丢掉。现在应分页续页，内容一条都不能少。
"""
import pytest
from pptx import Presentation

from office_agent.ppt_agent.ppt_service import (
    CONTENT_LIST_ITEMS_PER_PAGE,
    TOC_ITEMS_PER_PAGE,
    PPTService,
    _continuation_title,
    _paginate_items,
)
from office_agent.ppt_agent.models import SlideContent


@pytest.fixture
def svc():
    s = PPTService()
    # _base_deck 平时由 generate() 设置，这里直接驱动渲染器需自行初始化
    s.prs = Presentation()
    s._base_deck = False
    return s


def _slide_texts(slide):
    """取出一页里所有可见文本（含序号圆里的数字）。"""
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text
            if t:
                out.append(t)
    return out


def _all_texts(svc):
    return [_slide_texts(s) for s in svc.prs.slides]


# ------------------------------------------------------------------
# 目录页（单页 6 项）
# ------------------------------------------------------------------

@pytest.mark.parametrize("n,expected_pages", [
    (5, 1), (6, 1), (7, 2), (8, 2), (9, 2), (16, 3),
])
def test_toc_page_count(svc, n, expected_pages):
    items = [f"条目{i + 1}" for i in range(n)]
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=items, page_number=1))
    assert len(svc.prs.slides) == expected_pages


def test_toc_no_content_lost(svc):
    items = [f"条目{i + 1}" for i in range(16)]
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=items, page_number=1))

    flat = [t for page in _all_texts(svc) for t in page]
    for item in items:
        assert flat.count(item) == 1, f"{item} 必须且只能出现一次"


def test_toc_keeps_order(svc):
    items = [f"条目{i + 1}" for i in range(14)]
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=items, page_number=1))

    flat = [t for page in _all_texts(svc) for t in page]
    positions = [flat.index(i) for i in items]
    assert positions == sorted(positions), "条目顺序必须保持"


def test_toc_continuation_titles(svc):
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=[f"条目{i + 1}" for i in range(14)],
                                 page_number=1))
    texts = _all_texts(svc)
    assert texts[0][0] == "目录"
    assert texts[1][0] == "目录（续 2）"
    assert texts[2][0] == "目录（续 3）"


def test_toc_numbering_is_continuous_across_pages(svc):
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=[f"条目{i + 1}" for i in range(8)],
                                 page_number=1))
    texts = _all_texts(svc)
    assert "07" in texts[1], "续页序号应接着第一页编号，不能重新从 01 开始"


def test_toc_page_numbers_increment(svc):
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=[f"条目{i + 1}" for i in range(8)],
                                 page_number=3))
    texts = _all_texts(svc)
    assert "3" in texts[0]
    assert "4" in texts[1]


def test_toc_empty_does_not_produce_garbage_items(svc):
    svc._render_toc(SlideContent(layout="toc", title="目录", bullets=[], page_number=1))
    assert len(svc.prs.slides) == 1
    assert _all_texts(svc)[0] == ["目录", "1"]


def test_toc_single_page_behaviour_unchanged(svc):
    """5 条时仍是 1 页、仍从 01 开始编号（旧行为不回归）。"""
    svc._render_toc(SlideContent(layout="toc", title="目录",
                                 bullets=[f"条目{i + 1}" for i in range(5)],
                                 page_number=1))
    texts = _all_texts(svc)
    assert len(texts) == 1
    assert texts[0][0] == "目录"
    assert "01" in texts[0] and "05" in texts[0]


# ------------------------------------------------------------------
# 列表页（单页 8 项：2 列 × 4 行）
# ------------------------------------------------------------------

@pytest.mark.parametrize("n,expected_pages", [
    (5, 1), (8, 1), (9, 2), (16, 2), (17, 3),
])
def test_content_list_page_count(svc, n, expected_pages):
    svc._render_content_list(SlideContent(
        layout="content_list", title="要点",
        bullets=[f"要点{i + 1}" for i in range(n)], page_number=1))
    assert len(svc.prs.slides) == expected_pages


def test_content_list_no_content_lost(svc):
    items = [f"要点{i + 1}" for i in range(17)]
    svc._render_content_list(SlideContent(layout="content_list", title="要点",
                                          bullets=items, page_number=1))
    flat = [t for page in _all_texts(svc) for t in page]
    for item in items:
        assert flat.count(item) == 1


def test_content_list_numbering_continuous(svc):
    svc._render_content_list(SlideContent(
        layout="content_list", title="要点",
        bullets=[f"要点{i + 1}" for i in range(10)], page_number=1))
    texts = _all_texts(svc)
    assert "9" in texts[1] and "10" in texts[1], "续页序号应继续，不能回到 1"


def test_content_list_continuation_title(svc):
    svc._render_content_list(SlideContent(
        layout="content_list", title="要点",
        bullets=[f"要点{i + 1}" for i in range(10)], page_number=1))
    assert _all_texts(svc)[1][0] == "要点（续 2）"


def test_content_list_empty_is_single_page(svc):
    svc._render_content_list(SlideContent(layout="content_list", title="要点",
                                          bullets=[], page_number=1))
    assert len(svc.prs.slides) == 1
    assert _all_texts(svc)[0] == ["要点", "1"]


def test_content_list_handles_dict_items(svc):
    """dict 形态条目（LLM 常用）不能被当成空文本丢掉。"""
    items = [{"text": f"要点{i + 1}"} for i in range(9)]
    svc._render_content_list(SlideContent(layout="content_list", title="要点",
                                          bullets=items, page_number=1))
    flat = [t for page in _all_texts(svc) for t in page]
    for i in range(9):
        assert flat.count(f"要点{i + 1}") == 1


# ------------------------------------------------------------------
# 辅助函数
# ------------------------------------------------------------------

def test_paginate_items():
    assert _paginate_items(list(range(16)), 6) == [list(range(6)),
                                                   list(range(6, 12)),
                                                   list(range(12, 16))]
    assert _paginate_items([], 6) == []
    assert _paginate_items([1, 2], 0) == [[1, 2]], "容量非正时退化为单页，不产生空切片"


def test_continuation_title():
    assert _continuation_title("目录", 0) == "目录"
    assert _continuation_title("目录", 1) == "目录（续 2）"
    assert _continuation_title("目录", 2) == "目录（续 3）"


def test_capacity_constants_match_layout():
    assert TOC_ITEMS_PER_PAGE == 6
    assert CONTENT_LIST_ITEMS_PER_PAGE == 8
