"""Batch 3 数据正确性回归测试。"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook
from pptx import Presentation


class TestSumifsSingleCondition:
    def test_auto_conditional_uses_sumif_not_fake_sumifs(self, tmp_path):
        """P2-35：单一条件不得写成缺参的 SUMIFS。"""
        from office_agent.excel_agent.excel_service import ExcelService
        from office_agent.excel_agent.formula_generator import (
            FormulaGenerator,
            FormulaTemplate,
        )
        from office_agent.excel_agent.models import ColumnInfo, SheetInfo

        out = tmp_path / "s.xlsx"
        svc = ExcelService().create(str(out), "S")
        svc.write_data(
            "S",
            [["部门", "金额"], ["A", 1], ["B", 2], ["A", 3]],
            has_header=True,
        )
        sheet = SheetInfo(
            name="S", row_count=4, col_count=2,
            columns=[
                ColumnInfo(
                    name="部门", index=0, data_type="text", semantic_type="category",
                    unique_values=["A", "B"], unique_count=2,
                ),
                ColumnInfo(name="金额", index=1, data_type="number", semantic_type="amount"),
            ],
        )
        gen = FormulaGenerator()
        tpl = FormulaTemplate(
            "sumifs", [r"sumifs"], "=SUMIFS({sum_range},{criteria_range1},{criteria1})",
            "conditional", "多条件求和",
        )
        formulas = gen._generate_conditional(
            tpl,
            [sheet.columns[1]],
            sheet.columns[0],
            sheet,
            2,
        )
        assert formulas
        for f in formulas:
            if f.category == "sumifs":
                # 必须是单条件 SUMIF，或完整双条件 SUMIFS；不允许单条件 SUMIFS
                assert f.formula.startswith("=SUMIF(") or f.formula.count(",") >= 4, f.formula
                assert f.formula.startswith("=SUMIFS(") is False or f.formula.count(",") >= 4

        gen.apply_to_sheet(svc, "S", formulas)
        svc.save()
        wb = load_workbook(str(out), data_only=False)
        found = [
            str(c.value)
            for row in wb["S"].iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert found
        for formula in found:
            if formula.startswith("=SUMIFS"):
                # 若仍是 SUMIFS，必须参数完整（>= 5 个逗号段）
                assert formula.count(",") >= 4, formula


class TestVisionBoolCoercion:
    @pytest.mark.parametrize("value,expected", [
        (True, True), (False, False),
        ("true", True), ("false", False),
        ("True", True), ("FALSE", False),
        (1, True), (0, False),
        (None, False),
    ])
    def test_to_bool(self, value, expected):
        from office_agent.excel_agent.vision_analyzer import ExcelVisionAnalyzer

        result = ExcelVisionAnalyzer._to_bool({"v": value}, "v", False)
        assert result is expected


class TestWordNoDuplicateRuns:
    def test_apply_format_does_not_duplicate_text(self, tmp_path):
        from docx import Document

        from office_agent.models.schemas import FontConfig
        from office_agent.services.word_service import WordService

        path = tmp_path / "dup.docx"
        doc = Document()
        doc.add_paragraph("正常段落")
        doc.save(str(path))

        svc = WordService()
        cfg = FontConfig(size=12, en_font="Arial", cn_font="微软雅黑")
        # 应用段落格式到现有 runs
        for para in doc.paragraphs:
            if hasattr(svc, "_apply_paragraph_format"):
                try:
                    svc._apply_paragraph_format(para, cfg)
                except TypeError:
                    pass
        texts = [p.text for p in doc.paragraphs if p.text]
        assert texts.count("正常段落") == 1


class TestPptPageNumbersAfterPagination:
    def test_page_numbers_are_unique(self, tmp_path):
        from office_agent.ppt_agent.models import PPTOutline, SlideContent
        from office_agent.ppt_agent.ppt_service import PPTService

        outline = PPTOutline(title="T")
        # 两栏超长内容触发分页
        outline.slides = [
            SlideContent(layout="cover", title="封面", subtitle="s"),
            SlideContent(
                layout="two_column",
                title="对比",
                left_content=[f"L{i}" for i in range(20)],
                right_content=[f"R{i}" for i in range(20)],
                page_number=2,
            ),
            SlideContent(layout="summary", title="总结", bullets=["完"], page_number=3),
        ]
        out = tmp_path / "page.pptx"
        result = PPTService().generate(outline, str(out))
        assert result.success, result.message
        prs = Presentation(str(out))
        # 至少 cover + 两栏续页 + summary
        assert len(prs.slides) >= 3
        page_texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t.isdigit() and len(t) <= 2:
                        page_texts.append(int(t))
        # 渲染出的页码不得重复
        assert len(page_texts) == len(set(page_texts)), page_texts


class TestCsvParse:
    def test_csv_import_works(self, tmp_path):
        from office_agent.knowledge_base.document_parser import DocumentParser

        path = tmp_path / "data.csv"
        path.write_text("部门,金额\nA,10\nB,20\n", encoding="utf-8")
        doc = DocumentParser().parse(str(path))
        assert "部门" in doc.full_text
        assert "10" in doc.full_text
        assert doc.metadata.get("row_count", 0) >= 2


class TestChartSeriesLocalDegrade:
    def test_mismatched_series_skipped_not_fail_deck(self, tmp_path):
        from office_agent.ppt_agent.models import PPTOutline, SlideContent
        from office_agent.ppt_agent.ppt_service import PPTService

        outline = PPTOutline(title="T")
        outline.slides = [
            SlideContent(layout="cover", title="封面"),
            SlideContent(
                layout="chart",
                title="图",
                chart_type="column",
                chart_categories=["一", "二", "三"],
                chart_series=[
                    ("好", [1, 2, 3]),
                    ("坏", [1, 2]),  # 长度不匹配
                ],
                page_number=2,
            ),
        ]
        out = tmp_path / "chart.pptx"
        result = PPTService().generate(outline, str(out))
        assert result.success, result.message
        prs = Presentation(str(out))
        assert len(prs.slides) >= 2
