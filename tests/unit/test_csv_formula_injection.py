"""P4-3 回归测试：CSV 公式注入——前导空白不得绕过检测。

历史缺陷：`_sanitize_csv_cell` 只看 ``value[:1]``，因此

    " =SUM(A1:A2)"   "\\t=SUM(A1:A2)"   "\\r=..."   " +CMD"

都能绕过转义，被 openpyxl 当作公式写入输出 xlsx。

修复口径：
  * 检测前对前导空白（空格/tab/CR/LF/NBSP/BOM…）做规范化；
  * 输出保留用户原始文本语义，仅为转义追加前导单引号；
  * 负数、正数、普通文本不得被误杀。
"""
import pytest

from office_agent.task_queue.tasks import excel_tasks
from office_agent.task_queue.tasks.excel_tasks import (
    _csv_to_xlsx,
    _sanitize_csv_cell,
)


# ------------------------------------------------------------ 注入载荷


@pytest.mark.parametrize("payload", [
    "=SUM(A1:A2)",
    " =SUM(A1:A2)",          # 前导空格
    "  =SUM(A1:A2)",         # 多个空格
    "\t=SUM(A1:A2)",         # 前导 tab
    "\r=SUM(A1:A2)",         # 前导 CR
    "\n=SUM(A1:A2)",         # 前导 LF
    "\r\n=SUM(A1:A2)",
    "\x0b=SUM(A1:A2)",       # 垂直制表
    "\x0c=SUM(A1:A2)",       # 换页
    "\u00a0=SUM(A1:A2)",     # 不换行空格
    "\ufeff=SUM(A1:A2)",     # BOM
    " \t\r\n =SUM(A1:A2)",   # 混合空白
])
def test_leading_blank_does_not_bypass_equal(payload):
    result = _sanitize_csv_cell(payload)
    assert result.startswith("'"), f"未转义: {payload!r}"
    # 转义只追加单引号，原始文本（含前导空白）原样保留
    assert result[1:] == payload


@pytest.mark.parametrize("payload", [
    "+CMD",
    " +CMD",
    "\t+CMD",
    "\r+CMD",
    "-CMD",
    " -CMD",
    "@SUM(1)",
    " @SUM",
    "\t@SUM",
])
def test_leading_blank_does_not_bypass_other_prefixes(payload):
    result = _sanitize_csv_cell(payload)
    assert result.startswith("'"), f"未转义: {payload!r}"
    assert result[1:] == payload


def test_equal_always_escaped_even_if_numeric_looking():
    """= 开头即便后段可解析为数字也仍是公式，不得放行。"""
    for payload in ("=1", "=1e5", "=inf", "=nan", "=+3"):
        assert _sanitize_csv_cell(payload).startswith("'"), payload


# -------------------------------------------------------- 不得误杀


@pytest.mark.parametrize("value", [
    "普通文本",
    "张三",
    "hello world",
    "123",
    "-123",
    "-12.5",
    "+123",
    "+3.14",
    "  -123",        # 前导空格的负数仍是数值语义
    "\t-12.5",
    "1e5",
    "",
    "电话 13800138000",
    "-",             # 只有负号
])
def test_benign_values_are_untouched(value):
    assert _sanitize_csv_cell(value) == value


@pytest.mark.parametrize("value", [None, 123, -12.5, 0, True, [], {}])
def test_non_string_values_pass_through(value):
    assert _sanitize_csv_cell(value) == value


def test_negative_number_is_not_escaped_but_negative_text_is():
    assert _sanitize_csv_cell("-123") == "-123"
    assert _sanitize_csv_cell("-123abc") == "'-123abc"
    assert _sanitize_csv_cell(" -123") == " -123"
    assert _sanitize_csv_cell(" -123abc") == "' -123abc"


# ------------------------------------------------------------ 端到端


def _write_csv(tmp_path, rows):
    import csv

    path = tmp_path / "input.csv"
    with open(path, "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        for row in rows:
            writer.writerow(row)
    return str(path)


def test_csv_to_xlsx_writes_injection_as_text_not_formula(tmp_path, monkeypatch):
    """端到端：带前导空白的注入载荷写入 xlsx 后不得是公式单元格。"""
    import openpyxl

    monkeypatch.setattr(excel_tasks, "get_output_dir", lambda: str(tmp_path))
    csv_path = _write_csv(tmp_path, [
        ["名称", "备注"],
        ["a", " =SUM(A1:A2)"],
        ["b", "\t=SUM(A1:A2)"],
        ["c", "\r=cmd|'/c calc'!A1"],
        ["d", " @SUM(1)"],
        ["e", "-12.5"],
        ["f", "普通文本"],
    ])

    xlsx = _csv_to_xlsx(csv_path)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Sheet1"]

    def _cell(row):
        return ws.cell(row=row, column=2)

    # 注入载荷：不是公式单元格，且以单引号转义
    for row in (2, 3, 4, 5):
        cell = _cell(row)
        assert cell.data_type != "f", f"第 {row} 行被写成公式: {cell.value!r}"
        assert str(cell.value).lstrip(" \t\r\n").startswith("'"), cell.value

    # 正常数据未被误伤
    assert _cell(6).value in ("-12.5", -12.5)
    assert _cell(7).value == "普通文本"


def test_csv_to_xlsx_keeps_plain_payload_readable(tmp_path, monkeypatch):
    """无注入的普通 CSV 导入后内容不变（防过度转义）。"""
    import openpyxl

    monkeypatch.setattr(excel_tasks, "get_output_dir", lambda: str(tmp_path))
    csv_path = _write_csv(tmp_path, [
        ["产品", "价格"],
        ["A", "-12.5"],
        ["B", "100"],
    ])
    xlsx = _csv_to_xlsx(csv_path)
    ws = openpyxl.load_workbook(xlsx)["Sheet1"]
    assert ws.cell(row=2, column=1).value == "A"
    assert ws.cell(row=3, column=1).value == "B"
