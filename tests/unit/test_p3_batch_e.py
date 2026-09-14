# -*- coding: utf-8 -*-
"""P3-E 批（知识/嵌入/视觉）回归测试：P3-76~91 中的 REAL_OPEN。

覆盖：76/77/78/79/80/81/82/83/84/85/87/88/90/91。
Office 文档（docx/xlsx）使用真实产物读取验证，不只断言函数可调用。
"""
import json

import pytest


# ---------- P3-76 缺 doc_type 的 chunk 不得绕过类型过滤 ----------
def test_p3_76_missing_doc_type_filtered():
    from office_agent.knowledge_base.vector_store import VectorStore
    from office_agent.knowledge_base.models import KnowledgeChunk
    mf = VectorStore._matches_filters
    typed = KnowledgeChunk(id="a", content="a", metadata={"doc_type": "spec"})
    missing = KnowledgeChunk(id="b", content="b", metadata={})
    other = KnowledgeChunk(id="c", content="c", metadata={"doc_type": "rule"})
    assert mf(None, typed, ["spec"], None) is True
    assert mf(None, missing, ["spec"], None) is False
    assert mf(None, other, ["spec"], None) is False
    assert mf(None, missing, None, None) is True


# ---------- P3-77 save 不再物化全量向量副本且可往返 ----------
def test_p3_77_save_roundtrip(tmp_path):
    from office_agent.knowledge_base.knowledge_base import OfficeKnowledgeBase as KnowledgeBase
    kb = KnowledgeBase(storage_dir=str(tmp_path))
    kb.add_text("预算总额为一万元，用于差旅报销。", title="t", doc_type="spec")
    kb.save()
    chunks_file = tmp_path / "chunks.json"
    assert chunks_file.exists()
    reloaded = KnowledgeBase(storage_dir=str(tmp_path))
    assert len(reloaded.documents) == 1
    stored_embeddings = [s.embedding for s in reloaded.store._chunks]
    assert stored_embeddings and all(len(e) > 0 for e in stored_embeddings)


# ---------- P3-78 载入文档保留 updated_at ----------
def test_p3_78_updated_at_preserved_on_construct():
    from office_agent.knowledge_base.models import KnowledgeDocument
    doc = KnowledgeDocument(
        id="x", created_at="2020-01-01T00:00:00+00:00",
        updated_at="2021-05-05T00:00:00+00:00",
    )
    assert doc.updated_at == "2021-05-05T00:00:00+00:00"
    fresh = KnowledgeDocument(title="new")
    assert fresh.created_at and fresh.updated_at


# ---------- P3-79 Word 表格按 body 顺序归入正确章节（真实 docx） ----------
def test_p3_79_table_attached_to_preceding_section(tmp_path):
    import docx
    from office_agent.knowledge_base.document_parser import DocumentParser
    doc = docx.Document()
    doc.add_paragraph("第一章", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "键"
    table.rows[0].cells[1].text = "值"
    doc.add_paragraph("第一章正文")
    doc.add_paragraph("第二章", style="Heading 1")
    doc.add_paragraph("第二章正文")
    path = tmp_path / "s.docx"
    doc.save(str(path))

    parsed = DocumentParser().parse(str(path))
    sections = {s.title: s.content for s in parsed.sections}
    assert "[表格]" in sections.get("第一章", "")
    assert "[表格]" not in sections.get("第二章", "")
    # full_text 顺序：表格出现在第二章标题之前
    full = parsed.full_text
    assert full.index("键 | 值") < full.index("第二章")


# ---------- P3-80 Excel 中间空单元格不错位（真实 xlsx） ----------
def test_p3_80_excel_interior_empty_cell_kept(tmp_path):
    import openpyxl
    from office_agent.knowledge_base.document_parser import DocumentParser
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A", None, "C"])
    ws.append(["x", "y", "z"])
    path = tmp_path / "s.xlsx"
    wb.save(str(path))
    parsed = DocumentParser().parse(str(path))
    first_line = parsed.sections[0].content.splitlines()[0]
    assert first_line.startswith("A |  | C")


# ---------- P3-81 显式清空 API Key ----------
def test_p3_81_clear_api_key_distinct_from_blank_preserve(tmp_path):
    from office_agent.knowledge_base.embedding_config import EmbeddingConfigManager
    mgr = EmbeddingConfigManager(config_dir=str(tmp_path))
    mgr.save_config(provider="custom", api_key="sk-existing",
                    base_url="https://h/v1", model="m")
    # 留空仍保留旧值（既有契约）
    mgr.save_config(provider="custom", api_key="",
                    base_url="https://h2/v1", model="m2")
    assert mgr.get_config()["api_key"] == "sk-existing"
    # 显式清空
    mgr.clear_api_key()
    assert mgr.get_config()["api_key"] == ""
    assert mgr.is_configured() is False


# ---------- P3-82 embedding base_url 校验 ----------
def test_p3_82_embedder_base_url_validation():
    from office_agent.knowledge_base.embeddings import (
        APIEmbedder, EmbeddingBackendUnavailableError,
    )
    for bad in ("file:///etc/passwd", "api.example.com/v1",
                "https://u:p@host/v1", "ftp://host/v1"):
        with pytest.raises(EmbeddingBackendUnavailableError):
            APIEmbedder(api_key="k", base_url=bad)
    # 内网 http 合法，且尾斜杠被裁掉
    ok = APIEmbedder(api_key="k", base_url="http://10.0.0.5:9000/v1/")
    assert ok.base_url == "http://10.0.0.5:9000/v1"


# ---------- P3-83 “设计风格”不再误路由视觉 ----------
def test_p3_83_design_style_not_vision_keyword():
    from office_agent.model_gateway.model_router import VISION_KEYWORDS
    assert "设计风格" not in VISION_KEYWORDS


# ---------- P3-84 CJK 密度感知截断 ----------
def test_p3_84_cjk_aware_truncation():
    from office_agent.model_gateway.clients.base import BaseModelClient
    fn = BaseModelClient._truncate_to_token_budget
    cjk = fn("中" * 2000, 100)
    eng = fn("a" * 2000, 100)
    assert "截断" in cjk and len(cjk) < 300      # 中文 ~1 token/字
    assert 330 < len(eng) <= 355                 # 英文 ~4 字符/token + 标记
    assert fn("短文本", 100) == "短文本"
    assert fn("x", 0) == "x"


# ---------- P3-85 空 choices 明确错误而非 IndexError ----------
def test_p3_85_empty_choices_returns_clear_error(monkeypatch):
    from office_agent.model_gateway.clients import openai_client as mod
    from office_agent.model_gateway.clients.openai_client import OpenAIClient
    from office_agent.models.model_schemas import ModelConfig, ModelProvider

    class _Resp:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self):
            return json.dumps({"choices": []}).encode()

    monkeypatch.setattr(mod, "urlopen", lambda req, timeout: _Resp())
    client = OpenAIClient(ModelConfig(
        id="d", provider=ModelProvider.DEEPSEEK, display_name="d",
        api_key="x", base_url="https://api.deepseek.com/v1", model="m",
    ))
    result = client.chat([{"role": "user", "content": "hi"}])
    assert result.success is False
    assert "choices" in (result.error or "")


# ---------- P3-87 经复核改判 VERIFIED_NO_CHANGE ----------
# 既有 test_document_renderer_lifecycle 明确“构造即确定唯一 owned root、
# close/__exit__ 幂等清理”的契约；懒建会破坏单 root 与 gateway 复用语义，
# 而“构造却不 close”本就是调用方违约（context manager 已覆盖），故不改。


# ---------- P3-88 图片 URL scheme 校验 ----------
def test_p3_88_image_from_url_scheme():
    from office_agent.vision_gateway.vision_models import ImageInput
    for bad in ("file:///C:/x.png", "/tmp/x.png", "ftp://h/x.png", ""):
        with pytest.raises(ValueError):
            ImageInput.from_url(bad)
    img = ImageInput.from_url("https://host/a.png")
    assert img.url.endswith("a.png")


# ---------- P3-90 字符串内花括号不影响 JSON 配平 ----------
def test_p3_90_extract_json_braces_inside_string():
    from office_agent.quality.visual_checker import WordVisualChecker
    raw = '前缀 {"page_summary":"用法 {写} 法","page_score":88,"issues":[]} 后缀'
    extracted = WordVisualChecker._extract_json(raw)
    obj = json.loads(extracted)
    assert obj["page_score"] == 88
    assert "{写}" in obj["page_summary"]


# ---------- P3-91 单页非法分数不拖垮整报 ----------
def test_p3_91_safe_float_fallback():
    from office_agent.quality.visual_checker import WordVisualChecker as C
    assert C._safe_float("N/A", 70.0) == 70.0
    assert C._safe_float(None, 0.5) == 0.5
    assert C._safe_float("92", 70.0) == 92.0
