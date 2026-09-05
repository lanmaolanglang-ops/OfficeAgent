"""
降级路径结构化日志专项（清单 3.1「其他降级路径」收尾）

对每个治理过的降级腿钉住三件事：
1. fallback 仍然发生且业务结果与旧行为一致（默认结果/部分结果保留不变）；
2. 产生带上下文（操作、文件、原因）的结构化 warning/exception 日志；
3. 日志只含定位信息，不含文档正文、密钥等敏感内容。

按 skill 约定不依赖 caplog（套件级 logging 配置会改传播），
直接在各模块命名 logger 上挂接临时 handler。
"""
import logging

import pytest


class _ListHandler(logging.Handler):
    def __init__(self):
        super().__init__()
        self.records = []

    def emit(self, record):
        self.records.append(record)

    @property
    def messages(self):
        return [r.getMessage() for r in self.records]


@pytest.fixture
def capture_logs():
    attached = []

    def _cap(name):
        target = logging.getLogger(name)
        handler = _ListHandler()
        prev_level = target.level
        target.setLevel(logging.DEBUG)
        target.addHandler(handler)
        attached.append((target, handler, prev_level))
        return handler

    yield _cap

    for target, handler, prev_level in attached:
        target.removeHandler(handler)
        target.setLevel(prev_level)


def _make_xlsx(path, rows=(("月份", "销量"), ("一月", 10), ("二月", 20))):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return str(path)


# ---------------------------------------------------------------- analysis engine

class TestAnalysisEngineSchemaFallback:
    def test_schema_failure_degrades_to_none_and_logs(self, tmp_path, monkeypatch, capture_logs):
        from office_agent.excel_agent.analysis_engine import AnalysisEngine
        from office_agent.excel_agent.data_analyzer import DataAnalyzer

        handler = capture_logs("office_agent.excel_agent.analysis_engine")
        file_path = _make_xlsx(tmp_path / "data.xlsx")

        def _boom(self, *args, **kwargs):
            raise RuntimeError("profiling broken")
        monkeypatch.setattr(DataAnalyzer, "analyze_schema_frames", _boom)

        engine = AnalysisEngine()
        report = engine.analyze_file(file_path)

        # fallback 仍发生：schema 降级为 None，主分析报告照常产出
        assert engine.schema is None
        assert report is not None
        assert any("降级为无 schema" in m and "profiling broken" in m
                   for m in handler.messages)


# ---------------------------------------------------------------- chart generator

class TestChartGeneratorProfileFallback:
    def test_profile_failure_still_generates_file_and_logs(self, tmp_path, monkeypatch, capture_logs):
        from office_agent.excel_agent.chart_generator import ChartGenerator
        from office_agent.excel_agent.data_analyzer import DataAnalyzer

        handler = capture_logs("office_agent.excel_agent.chart_generator")
        file_path = _make_xlsx(tmp_path / "data.xlsx")
        output_path = str(tmp_path / "out.xlsx")

        def _boom(self, *args, **kwargs):
            raise RuntimeError("profile broken")
        monkeypatch.setattr(DataAnalyzer, "analyze", _boom)

        gen = ChartGenerator()
        saved, specs = gen.generate_to_file(file_path, text="", output_path=output_path)

        # fallback 仍发生：无画像继续，自动选图为空但文件照常写出
        assert gen.profile is None
        assert specs == []
        assert saved == output_path
        assert any("降级为无画像选图" in m and "profile broken" in m
                   for m in handler.messages)


# ---------------------------------------------------------------- excel orchestrator

class TestFragileDetectionFallback:
    def test_detection_failure_returns_empty_and_logs(self, tmp_path, capture_logs):
        from office_agent.excel_agent.excel_orchestrator import ExcelOrchestrator

        handler = capture_logs("office_agent.excel_agent.excel_orchestrator")
        bad = tmp_path / "broken.xlsx"
        bad.write_bytes(b"not a real xlsx")

        assert ExcelOrchestrator._detect_fragile_elements(str(bad)) == ""
        assert any("脆弱元素检测失败" in m for m in handler.messages)


# ---------------------------------------------------------------- excel cached scan

class TestCachedErrorScanFallback:
    def test_reopen_failure_skips_scan_and_logs(self, monkeypatch, capture_logs):
        import office_agent.excel_agent.quality_checker as qc
        from office_agent.excel_agent.quality_checker import ExcelQualityChecker

        handler = capture_logs("office_agent.excel_agent.quality_checker")

        def _boom(*args, **kwargs):
            raise RuntimeError("cannot reopen")
        monkeypatch.setattr(qc, "load_workbook", _boom)

        assert ExcelQualityChecker()._check_cached_errors("whatever.xlsx") == []
        assert any("缓存错误值扫描跳过" in m and "cannot reopen" in m
                   for m in handler.messages)

    def test_mid_scan_failure_keeps_partial_and_logs(self, monkeypatch, capture_logs):
        import office_agent.excel_agent.quality_checker as qc
        from office_agent.excel_agent.quality_checker import ExcelQualityChecker

        handler = capture_logs("office_agent.excel_agent.quality_checker")

        class _Cell:
            def __init__(self, value, coordinate):
                self.value = value
                self.coordinate = coordinate

        class _Ws:
            title = "Sheet1"

            def iter_rows(self, max_row=None):
                yield [_Cell("#DIV/0!", "A1")]
                raise RuntimeError("scan blew up")

        class _Wb:
            worksheets = [_Ws()]

            def close(self):
                pass

        monkeypatch.setattr(qc, "load_workbook", lambda *a, **k: _Wb())

        issues = ExcelQualityChecker()._check_cached_errors("whatever.xlsx")
        # 部分结果保留：中断前已收集的错误值不丢
        assert len(issues) == 1
        assert issues[0].cell_ref == "A1"
        assert any("缓存错误值扫描中断" in m and "scan blew up" in m
                   for m in handler.messages)


# ---------------------------------------------------------------- ppt theme fallback

class TestPptThemeFallback:
    def test_theme_parse_failure_falls_back_and_logs(self, monkeypatch, capture_logs):
        from pptx import Presentation
        from office_agent.ppt_agent.template_analyzer import (
            TemplateAnalyzer, TemplateConfig,
        )

        handler = capture_logs("office_agent.ppt_agent.template_analyzer")
        analyzer = TemplateAnalyzer()

        def _boom(self, *args, **kwargs):
            raise RuntimeError("theme xml broken")
        monkeypatch.setattr(TemplateAnalyzer, "_parse_color_scheme", _boom)

        called = {"colors": False}
        original = TemplateAnalyzer._fallback_extract_colors

        def _spy(self, prs, config):
            called["colors"] = True
            return original(self, prs, config)
        monkeypatch.setattr(TemplateAnalyzer, "_fallback_extract_colors", _spy)

        analyzer._extract_theme(Presentation(), TemplateConfig())

        assert called["colors"] is True
        assert any("降级为从幻灯片内容提取" in m and "theme xml broken" in m
                   for m in handler.messages)


# ---------------------------------------------------------------- pptx render fallback

class TestPptxRenderFallback:
    def test_render_falls_back_to_text_and_logs(self, tmp_path, monkeypatch, capture_logs):
        from pptx import Presentation
        from office_agent.vision_gateway.document_renderer import DocumentRenderer

        handler = capture_logs("office_agent.vision_gateway.document_renderer")
        pptx_path = tmp_path / "deck.pptx"
        Presentation().save(str(pptx_path))

        renderer = DocumentRenderer(output_dir=str(tmp_path / "out"))
        monkeypatch.setattr(renderer, "_convert_pptx_to_pdf", lambda *a: None)

        pages = renderer._render_pptx(str(pptx_path))
        # fallback 仍发生：走文本提取腿（本例空演示文稿，0 页）
        assert pages == []
        assert any("降级为文本提取" in m for m in handler.messages)

    def test_missing_libreoffice_logs_before_fallback(self, tmp_path, monkeypatch, capture_logs):
        import os
        import shutil
        from office_agent.vision_gateway.document_renderer import DocumentRenderer

        handler = capture_logs("office_agent.vision_gateway.document_renderer")
        monkeypatch.setattr(shutil, "which", lambda name: None)
        monkeypatch.setattr(os.path, "exists", lambda p: False)

        renderer = DocumentRenderer(output_dir=str(tmp_path / "out"))
        assert renderer._convert_pptx_to_pdf("deck.pptx") is None
        assert any("未找到 LibreOffice" in m for m in handler.messages)


# ---------------------------------------------------------------- upload metadata fallback

class TestUploadMetadataFallback:
    @pytest.mark.parametrize("suffix,meta_key", [
        (".docx", "paragraphs"),
        (".pptx", "slides"),
        (".pdf", "pages"),
    ])
    def test_corrupt_file_metadata_skipped_and_logged(self, tmp_path, capture_logs, suffix, meta_key):
        from office_agent.task_queue.tasks.file_tasks import process_upload

        handler = capture_logs("office_agent.tasks.file")
        bad = tmp_path / f"broken{suffix}"
        bad.write_bytes(b"garbage, not a real document")

        result = process_upload(str(bad), file_id="f1")

        # fallback 仍发生：任务成功，仅缺失该格式元数据
        assert result["status"] == "success"
        assert meta_key not in result["metadata"]
        assert any("元数据失败" in m for m in handler.messages)


# ---------------------------------------------------------------- knowledge base import

class TestKnowledgeBaseImportFallback:
    def test_bad_document_skipped_with_structured_log(self, tmp_path, capture_logs):
        from office_agent.knowledge_base.knowledge_base import OfficeKnowledgeBase

        handler = capture_logs("office_agent.knowledge_base")
        src = tmp_path / "src"
        src.mkdir()
        (src / "bad.docx").write_bytes(b"garbage, not a real docx")

        kb = OfficeKnowledgeBase(storage_dir=str(tmp_path / "kb"))
        docs = kb.import_directory(str(src))

        # 部分结果保留：坏文档跳过，其余继续（本例无其余文档）
        assert docs == []
        assert any("导入失败" in m and "bad.docx" in m for m in handler.messages)


# ---------------------------------------------------------------- ppt image availability fallback

class TestPptImageAvailabilityFallback:
    def test_availability_check_failure_degrades_and_logs(self, capture_logs):
        from types import SimpleNamespace
        from office_agent.ppt_agent.ppt_orchestrator import PPTOrchestrator

        handler = capture_logs("office_agent.ppt.orchestrator")

        class _BrokenGateway:
            def available(self):
                raise RuntimeError("config store down")

        orch = PPTOrchestrator(image_gateway=_BrokenGateway())
        outline = SimpleNamespace(slides=[])
        result = orch._generate_marked_images(outline)

        # fallback 仍发生：原大纲原样返回，错误进入 metadata 且可观测
        assert result is outline
        assert len(orch.image_generation["errors"]) == 1
        assert any("降级为纯文本大纲" in m for m in handler.messages)


# ---------------------------------------------------------------- task api db-session fallback

class TestTaskApiSessionFallback:
    def test_session_failure_returns_none_and_logs(self, monkeypatch, capture_logs):
        import office_agent.database.session as db_session
        from office_agent.api.router.task import _get_db_session

        handler = capture_logs("office_agent.api.task")

        def _boom():
            raise RuntimeError("database unavailable")
        monkeypatch.setattr(db_session, "SessionLocal", _boom)

        assert _get_db_session() is None
        errors = [r for r in handler.records if r.levelno >= logging.ERROR]
        assert any("降级到非数据库路径" in r.getMessage() for r in errors)
        # 日志不含连接串/密钥等敏感内容，仅为定位信息
        assert all("password" not in r.getMessage().lower()
                   and "://" not in r.getMessage() for r in handler.records)
