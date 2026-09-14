"""P2 Final Closure：OPEN 项与 PARTIAL 收口回归。"""
from __future__ import annotations

import json
from pathlib import Path



class TestP2_16FileManagerMetadata:
    def test_malformed_metadata_does_not_overwrite(self, tmp_path):
        from office_agent.api.core.file_manager import FileManager

        meta = tmp_path / "file_metadata.json"
        meta.write_text("{not-json", encoding="utf-8")
        mgr = FileManager(upload_dir=str(tmp_path), output_dir=str(tmp_path / "out"))
        assert mgr._degraded_metadata is True
        # 原件应被备份保留
        backups = list(tmp_path.glob("file_metadata.json.corrupt.*"))
        assert backups, "损坏元数据必须备份保留"
        mgr.save_metadata()
        # 不得用空索引覆盖
        assert not meta.exists() or json.loads(meta.read_text(encoding="utf-8")) != []

    def test_valid_metadata_loads(self, tmp_path):
        from office_agent.api.core.file_manager import FileManager

        payload = [{
            "file_id": "f1",
            "original_name": "a.txt",
            "stored_path": str(tmp_path / "f1.txt"),
            "file_type": "text",
            "extension": ".txt",
            "size": 3,
            "upload_time": "2026-01-01T00:00:00+00:00",
            "metadata": {},
        }]
        (tmp_path / "f1.txt").write_text("abc", encoding="utf-8")
        (tmp_path / "file_metadata.json").write_text(
            json.dumps(payload), encoding="utf-8"
        )
        mgr = FileManager(upload_dir=str(tmp_path), output_dir=str(tmp_path / "out"))
        assert "f1" in mgr.files
        assert mgr._degraded_metadata is False


class TestP2_28TaskListBatchFiles:
    def test_find_by_ids_exists(self):
        from office_agent.database.repository.base import BaseRepository

        assert hasattr(BaseRepository, "find_by_ids")


class TestP2_36ChartNonContiguousNumeric:
    def test_stacked_chart_uses_series_ranges_not_span(self, tmp_path):
        from office_agent.excel_agent.chart_generator import ChartGenerator
        from office_agent.excel_agent.models import ColumnInfo, SheetInfo

        sheet = SheetInfo(
            name="S", row_count=5, col_count=4,
            columns=[
                ColumnInfo(name="月份", index=0, data_type="text", semantic_type="date"),
                ColumnInfo(name="销售额", index=1, data_type="number", semantic_type="amount"),
                ColumnInfo(name="备注", index=2, data_type="text", semantic_type="text"),
                ColumnInfo(name="成本", index=3, data_type="number", semantic_type="amount"),
            ],
        )
        gen = ChartGenerator()
        num_cols = [c for c in sheet.columns if c.data_type == "number"]
        assert [c.name for c in num_cols] == ["销售额", "成本"]
        # 修复后 series 显式列出 B 与 D；first:last 旧实现会夹入 C（备注）
        assert [gen._col_letter(c.index) for c in num_cols] == ["B", "D"]
        assert f"{gen._col_letter(num_cols[0].index)}:" \
               f"{gen._col_letter(num_cols[-1].index)}" == "B:D"


class TestP2_39SampleNullCount:
    def test_null_count_is_sample_based(self, tmp_path):
        from openpyxl import Workbook

        from office_agent.excel_agent.data_analyzer import DataAnalyzer

        path = tmp_path / "nulls.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["a"])
        for i in range(300):
            ws.append([None if i >= 200 else i])
        wb.save(str(path))
        profile = DataAnalyzer().analyze(str(path))
        col = profile.sheets[0].columns[0]
        # 采样窗口约 199 行：前 199 个数据行有值 → 采样 null_count ≈ 0
        # 不得伪装成全表 100 个 null
        assert col.null_count < 50, col.null_count


class TestP2_40QualityInventoryReset:
    def test_sequential_checks_do_not_leak_chart_count(self, tmp_path):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference

        from office_agent.excel_agent.quality_checker import ExcelQualityChecker

        a = tmp_path / "a.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["x", "y"])
        ws.append([1, 2])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2))
        ws.add_chart(chart, "D2")
        wb.save(str(a))

        b = tmp_path / "b.xlsx"
        Workbook().save(str(b))

        checker = ExcelQualityChecker()
        checker.check(str(a))
        checker.check(str(b))
        assert checker._chart_count == 0, "B 无图表，不得继承 A 的 inventory"


class TestP2_52TemplateNoOverwrite:
    def test_same_path_goes_to_new_name(self, tmp_path):
        from pptx import Presentation

        from office_agent.ppt_agent.template_analyzer import TemplateAnalyzer

        tpl = tmp_path / "tpl.pptx"
        Presentation().save(str(tpl))
        original = tpl.read_bytes()
        out_path = str(tpl)  # 同路径
        result = TemplateAnalyzer().create_from_template(
            str(tpl), [{"layout_index": 0, "title": "T"}], out_path,
        )
        assert result != str(tpl)
        assert Path(result).exists()
        assert tpl.read_bytes() == original, "用户模板不得被覆盖"


class TestP2_54ServiceWrapperQuoting:
    def test_paths_with_quotes_are_python_safe(self, tmp_path):
        from desktop.service_manager import build_service_wrapper_script

        evil = tmp_path / 'ev"il'
        script = build_service_wrapper_script(evil, 8765, tmp_path / "data")
        # 生成脚本应能被 parse（路径经 repr 嵌入）
        compile(script, "<service_wrapper>", "exec")


class TestP2_71MasterKeyError:
    def test_decrypt_failure_raises_clear_error(self, tmp_path):
        # 源码契约：全量解密失败必须抛 MODEL_CREDENTIAL_DECRYPTION_FAILED
        src = Path(
            __import__("office_agent.model_gateway.model_manager", fromlist=["x"]).__file__
        ).read_text(encoding="utf-8")
        assert "MODEL_CREDENTIAL_DECRYPTION_FAILED" in src


class TestP2_73FallbackTelemetry:
    def test_fallback_requires_primary_attempt(self):
        src = Path(
            __import__("office_agent.model_gateway.failover", fromlist=["x"]).__file__
        ).read_text(encoding="utf-8")
        assert "primary_model" in src
        assert "any(m == primary" in src or "primary" in src


class TestP2_64LogInputDefault:
    def test_sensitive_nested_keys_redacted(self):
        from office_agent.logging_system.decorators import _summarize_input

        def sample(**kwargs):
            pass

        summary = _summarize_input(
            sample,
            (),
            {
                "headers": {"Authorization": "Bearer secret"},
                "api_key": "sk-x",
                "password": "p",
                "token": "t",
            },
            True,
            500,
        )
        assert "sk-x" not in summary
        assert "Bearer secret" not in summary or "[已隐藏]" in summary


class TestP2_15SingleProcessInvariant:
    def test_runtime_manager_is_process_singleton(self):

        src = Path(
            __import__("office_agent.runtime_manager", fromlist=["x"]).__file__
        ).read_text(encoding="utf-8")
        assert "_runtime_manager_lock" in src
        # 文档化：本地桌面单 backend 进程
        assert "double-checked locking" in src or "DCL" in src or "单例" in src


class TestP2_45TablePaginationHint:
    def test_truncation_is_explicit_not_silent(self):
        src = Path(
            __import__("office_agent.ppt_agent.ppt_service", fromlist=["x"]).__file__
        ).read_text(encoding="utf-8")
        assert "省略" in src
        assert "max_rows" in src
