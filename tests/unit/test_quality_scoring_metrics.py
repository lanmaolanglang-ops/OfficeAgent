"""P1-20：PPT / Excel 评分器口径修正。

修复前行为（旧源码副本实测）：
A. ``ppt_scorer._analyze_slides`` 的空页判定只看 ``text_length`` 与
   ``pictures``；"只有表格 / 只有图表"的页被误判为空白页，
   进而扣减 content 分并报"存在 N 页空白幻灯片"。
B. ``title_count`` 在 shape 循环内累加且不去重；同一页命中多个
   title-like 形状（标题占位符 + 顶部大字号文本框）会 +2，而下游
   ``title_ratio`` / ``title_coverage`` 的分母是**页数**，口径不一致。
C. ``excel_scorer._extract_function_name`` 用 ``func.isalpha()`` 判定，
   把 ``RANK.EQ`` / ``STDEV.P`` 这类带 ``.`` 的合法函数名整类丢弃，
   尽管 ``COMMON_FUNCTIONS`` 白名单里就写着 ``RANK.EQ``。
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt

from office_agent.quality_scoring.excel_scorer import ExcelQualityScorer
from office_agent.quality_scoring.ppt_scorer import PPTQualityScorer

# 1x1 透明 PNG，仅用于构造"只有图片"的幻灯片
_PNG_1X1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c6300010000050001"
    "0d0a2db40000000049454e44ae426082"
)


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_text_box(slide, text, *, top_inches=0.5, size=28):
    box = slide.shapes.add_textbox(
        Inches(1), Inches(top_inches), Inches(6), Inches(1),
    )
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    return box


def _score(path):
    return PPTQualityScorer().score(str(path))


def _content(path):
    return _score(path).content_details


class TestEmptySlideDetection:
    def test_text_only_is_not_empty(self, tmp_path):
        prs = Presentation()
        _add_text_box(_blank_slide(prs), "正文内容")
        path = tmp_path / "text.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 0

    def test_picture_only_is_not_empty(self, tmp_path):
        image = tmp_path / "p.png"
        image.write_bytes(_PNG_1X1)
        prs = Presentation()
        slide = _blank_slide(prs)
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
        path = tmp_path / "pic.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 0

    def test_table_only_is_not_empty(self, tmp_path):
        prs = Presentation()
        slide = _blank_slide(prs)
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1),
                                       Inches(4), Inches(2)).table
        table.cell(0, 0).text = "A"
        path = tmp_path / "table.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 0

    def test_chart_only_is_not_empty(self, tmp_path):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("收入", (1, 2))

        prs = Presentation()
        slide = _blank_slide(prs)
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
            Inches(6), Inches(3), chart_data,
        )
        path = tmp_path / "chart.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 0

    def test_genuinely_empty_slide_is_empty(self, tmp_path):
        prs = Presentation()
        _blank_slide(prs)
        path = tmp_path / "empty.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 1

    def test_placeholder_only_slide_is_empty(self, tmp_path):
        """按现有定义：没有任何文本/图片/表格/图表的占位符页仍属空白。"""
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])  # 标题 + 副标题占位符，均无文字
        path = tmp_path / "placeholder.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 1

    def test_mixed_content_slide_is_not_empty(self, tmp_path):
        prs = Presentation()
        slide = _blank_slide(prs)
        _add_text_box(slide, "正文")
        slide.shapes.add_table(2, 2, Inches(1), Inches(3),
                               Inches(4), Inches(2)).table
        path = tmp_path / "mixed.pptx"
        prs.save(path)
        assert _content(path)["empty_slides"] == 0

    def test_table_and_chart_slides_are_not_reported_as_blank_issues(self, tmp_path):
        prs = Presentation()
        slide = _blank_slide(prs)
        slide.shapes.add_table(2, 2, Inches(1), Inches(1),
                               Inches(4), Inches(2)).table
        path = tmp_path / "table2.pptx"
        prs.save(path)

        result = _score(path)
        assert not any("空白幻灯片" in issue for issue in result.issues)
        # 空页判定修正后，内容完整度不应被错误扣分
        assert result.content_completeness > 50


class TestTitleCountDedup:
    def _title_count(self, path):
        return _score(path).content_details["title_count"]

    def test_single_title_per_slide(self, tmp_path):
        prs = Presentation()
        _add_text_box(_blank_slide(prs), "唯一标题")
        path = tmp_path / "one.pptx"
        prs.save(path)
        assert self._title_count(path) == 1

    def test_two_title_like_shapes_on_one_slide_count_once(self, tmp_path):
        prs = Presentation()
        slide = _blank_slide(prs)
        _add_text_box(slide, "标题一", top_inches=0.3, size=28)
        _add_text_box(slide, "标题二", top_inches=0.8, size=32)
        path = tmp_path / "two.pptx"
        prs.save(path)
        assert self._title_count(path) == 1

    def test_same_title_text_on_two_slides_counts_twice(self, tmp_path):
        """按页去重，而不是按标题文本去重：两页都叫"目录"仍是两页有标题。"""
        prs = Presentation()
        _add_text_box(_blank_slide(prs), "目录")
        _add_text_box(_blank_slide(prs), "目录")
        path = tmp_path / "same.pptx"
        prs.save(path)
        assert self._title_count(path) == 2

    def test_slide_without_title_is_not_counted(self, tmp_path):
        prs = Presentation()
        slide = _blank_slide(prs)
        _add_text_box(slide, "正文小字", top_inches=3.0, size=18)
        path = tmp_path / "notitle.pptx"
        prs.save(path)
        assert self._title_count(path) == 0

    def test_title_placeholder_plus_textbox_counts_once(self, tmp_path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "占位符标题"
        _add_text_box(slide, "第二个标题形状", top_inches=0.4, size=30)
        path = tmp_path / "ph.pptx"
        prs.save(path)
        assert self._title_count(path) == 1

    def test_title_coverage_never_exceeds_one(self, tmp_path):
        prs = Presentation()
        slide = _blank_slide(prs)
        _add_text_box(slide, "标题一", top_inches=0.3, size=28)
        _add_text_box(slide, "标题二", top_inches=0.8, size=32)
        path = tmp_path / "cov.pptx"
        prs.save(path)
        assert _score(path).template_details["title_coverage"] <= 1.0


class TestExcelFunctionNameExtraction:
    @pytest.mark.parametrize("formula,expected", [
        ("=SUM(A1:A10)", "SUM"),
        ("=IF(A1>0,1,0)", "IF"),
        ("=VLOOKUP(A1,B:C,2,FALSE)", "VLOOKUP"),
        ("=RANK.EQ(B2,$B$2:$B$10)", "RANK.EQ"),
        ("=STDEV.P(A1:A10)", "STDEV.P"),
        ("=LOG10(A1)", "LOG10"),
        ("=ATAN2(A1,A2)", "ATAN2"),
        ("=_xlfn.CONCAT(A1,B1)", "_XLFN.CONCAT"),
        ("=Sheet1!SUM(A1:A2)", "SUM"),
    ])
    def test_valid_function_names_are_extracted(self, formula, expected):
        assert ExcelQualityScorer()._extract_function_name(formula) == expected

    @pytest.mark.parametrize("formula", [
        "=1+1",
        "=SUM 1(A1)",
        "=#REF!(A1)",
        "=(A1)",
        "=+",
        "=",
        "=A1+B1*(",
    ])
    def test_invalid_names_are_rejected(self, formula):
        assert ExcelQualityScorer()._extract_function_name(formula) is None

    def test_every_whitelisted_function_is_extractable(self):
        """白名单与提取器必须自洽——否则白名单形同虚设。"""
        scorer = ExcelQualityScorer()
        for name in sorted(ExcelQualityScorer.COMMON_FUNCTIONS):
            assert scorer._extract_function_name(f"={name}(A1)") == name

    def test_rank_eq_is_scored_end_to_end(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["姓名", "得分", "排名"])
        for row in range(2, 6):
            ws.append([f"人{row}", 100 - row, f"=RANK.EQ(B{row},$B$2:$B$5)"])
        ws.append(["合计", "=SUM(B2:B5)", ""])
        path = tmp_path / "rank.xlsx"
        wb.save(path)

        scorer = ExcelQualityScorer()
        result = scorer.score(str(path), expected_formulas=["RANK.EQ", "SUM"])
        types = result.formula_details["formula_types"]
        assert "RANK.EQ" in types
        assert "SUM" in types

    def test_rank_eq_matches_expected_formulas(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["得分", "排名"])
        ws.append([10, "=RANK.EQ(A2,$A$2:$A$3)"])
        ws.append([20, "=RANK.EQ(A3,$A$2:$A$3)"])
        path = tmp_path / "rank2.xlsx"
        wb.save(path)

        scorer = ExcelQualityScorer()
        hit = scorer.score(str(path), expected_formulas=["RANK.EQ"])
        miss = scorer.score(str(path), expected_formulas=["XIRR"])
        assert hit.formula_accuracy > miss.formula_accuracy

    def test_formula_count_increases_with_dotted_function(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["得分", "排名"])
        ws.append([10, "=RANK.EQ(A2,$A$2:$A$3)"])
        ws.append([20, "=RANK.EQ(A3,$A$2:$A$3)"])
        path = tmp_path / "rank3.xlsx"
        wb.save(path)

        details = ExcelQualityScorer().score(str(path)).formula_details
        assert details["formula_count"] == 1
        assert details["total_formulas"] == 2
