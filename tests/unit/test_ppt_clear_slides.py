"""PPT 清页（clear_slides）package 完整性专项测试。

实证背景：python-pptx 保存时按关系图遍历序列化部件，朴素 drop_rel 清页
在常规模板下不会产生孤儿部件；但模板含自定义放映（custShowLst）时，
slide rId 被二次引用触发 XmlPart.drop_rel 的引用计数守卫，导致清页
静默失败、留下孤儿 slide 部件。clear_slides 通过
"先移除 custShowLst → 再移除 sldId → 最后 drop_rel" 的顺序修复。
"""
import io
import posixpath
import re
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from office_agent.ppt_agent.template_analyzer import clear_slides

_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _png(color):
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (32, 32), color).save(buf, "PNG")
    return buf.getvalue()


SHARED_PNG = _png((200, 30, 30))
UNIQUE_PNG = _png((30, 30, 200))


def _package_violations(data: bytes):
    """返回 (dangling_rels, orphan_parts)：package 级关系完整性检查。"""
    z = zipfile.ZipFile(io.BytesIO(data))
    names = set(z.namelist())
    referenced = set()
    dangling = []
    for name in names:
        if not name.endswith(".rels"):
            continue
        base_dir = posixpath.dirname(posixpath.dirname(name))
        xml = z.read(name).decode("utf-8")
        for m in re.finditer(r'<Relationship\b[^>]*>', xml):
            tag = m.group(0)
            target = re.search(r'Target="([^"]+)"', tag).group(1)
            if 'TargetMode="External"' in tag:
                continue
            resolved = (target.lstrip("/") if target.startswith("/")
                        else posixpath.normpath(posixpath.join(base_dir, target)))
            if resolved not in names:
                dangling.append((name, target))
            else:
                referenced.add(resolved)
    orphans = sorted(
        n for n in names
        if not n.endswith(".rels") and n != "[Content_Types].xml" and n not in referenced
    )
    return dangling, orphans


def _build_template(with_custom_show: bool = False) -> bytes:
    """构造复杂模板：2 页共用一张图，第 1 页另有独占图 + chart(嵌入 workbook)，
    第 2 页含演讲备注。"""
    prs = Presentation()
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(io.BytesIO(SHARED_PNG), Inches(1), Inches(1))
    s1.shapes.add_picture(io.BytesIO(UNIQUE_PNG), Inches(3), Inches(1))
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("S1", (1, 2))
    s1.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), chart_data
    )

    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_picture(io.BytesIO(SHARED_PNG), Inches(1), Inches(1))
    s2.notes_slide.notes_text_frame.text = "演讲备注"

    if with_custom_show:
        rids = [e.get(qn("r:id")) for e in prs.slides._sldIdLst]
        cust = etree.SubElement(prs.part._element, f"{{{_P_NS}}}custShowLst")
        show = etree.SubElement(cust, f"{{{_P_NS}}}custShow")
        show.set("name", "cs")
        show.set("id", "0")
        lst = etree.SubElement(show, f"{{{_P_NS}}}sldLst")
        for rid in rids:
            e = etree.SubElement(lst, f"{{{_P_NS}}}sld")
            e.set(f"{{{_R_NS}}}id", rid)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class TestClearSlidesPackageIntegrity:
    def test_simple_template_no_orphans(self):
        prs = Presentation(io.BytesIO(_build_template()))
        clear_slides(prs)
        assert len(prs.slides) == 0

        buf = io.BytesIO()
        prs.save(buf)
        dangling, orphans = _package_violations(buf.getvalue())
        assert dangling == []
        assert orphans == []

    def test_slide_exclusive_parts_dropped(self):
        # slide 独占的图表 / 嵌入 workbook / 备注页 / 独占图片随清页剔除
        prs = Presentation(io.BytesIO(_build_template()))
        clear_slides(prs)
        buf = io.BytesIO()
        prs.save(buf)
        names = zipfile.ZipFile(io.BytesIO(buf.getvalue())).namelist()
        assert not [n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)]
        assert not [n for n in names if "charts/chart" in n]
        assert not [n for n in names if "embeddings/" in n]
        assert not [n for n in names if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n)]
        assert not [n for n in names if "media" in n]  # 两张图均不再被引用

    def test_master_layout_theme_survive(self):
        prs = Presentation(io.BytesIO(_build_template()))
        n_masters = len(prs.slide_masters)
        n_layouts = len(prs.slide_layouts)
        clear_slides(prs)
        buf = io.BytesIO()
        prs.save(buf)
        prs2 = Presentation(io.BytesIO(buf.getvalue()))
        assert len(prs2.slide_masters) == n_masters
        assert len(prs2.slide_layouts) == n_layouts

    def test_shared_media_survives_when_still_referenced(self):
        # 共享图片：旧页清除后由新页继续引用时，媒体部件必须保留
        prs = Presentation(io.BytesIO(_build_template()))
        clear_slides(prs)
        new_slide = prs.slides.add_slide(prs.slide_layouts[6])
        new_slide.shapes.add_picture(io.BytesIO(SHARED_PNG), Inches(1), Inches(1))
        buf = io.BytesIO()
        prs.save(buf)
        data = buf.getvalue()
        dangling, orphans = _package_violations(data)
        assert dangling == []
        assert orphans == []
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        assert [n for n in names if "media" in n]

    def test_custom_show_template_no_orphan_slides(self):
        # 回归钉住：custShowLst 二次引用 slide rId 曾使清页静默失败
        prs = Presentation(io.BytesIO(_build_template(with_custom_show=True)))
        clear_slides(prs)
        assert len(prs.slides) == 0

        buf = io.BytesIO()
        prs.save(buf)
        data = buf.getvalue()
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        # slide 部件必须真正从 package 消失，而不是留在包里成为孤儿
        assert not [n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)]
        dangling, orphans = _package_violations(data)
        assert dangling == []
        assert orphans == []
        # 自定义放映已随清页移除，不再引用已删除的 slide
        pres_xml = zipfile.ZipFile(io.BytesIO(data)).read("ppt/presentation.xml").decode("utf-8")
        assert "custShowLst" not in pres_xml

    def test_reopen_after_clear_and_regenerate(self):
        prs = Presentation(io.BytesIO(_build_template(with_custom_show=True)))
        clear_slides(prs)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "新页"
        buf = io.BytesIO()
        prs.save(buf)
        prs2 = Presentation(io.BytesIO(buf.getvalue()))
        assert len(prs2.slides) == 1
        assert len(prs2.slide_masters) == 1


class TestCallSiteConsistency:
    """两个产品调用点（ppt_service 模板基底 / TemplateAnalyzer.create_from_template）
    必须消费同一 clear_slides 实现，行为一致。"""

    def test_ppt_service_base_template_uses_clear_slides(self, tmp_path):
        from office_agent.ppt_agent.ppt_service import PPTService

        tpl = tmp_path / "base.pptx"
        tpl.write_bytes(_build_template(with_custom_show=True))

        outline = _make_outline(str(tpl))
        out = tmp_path / "out.pptx"
        result = PPTService().generate(outline, str(out))
        assert result.success if hasattr(result, "success") else True

        names = zipfile.ZipFile(str(out)).namelist()
        # 模板原有两页（含 chart/备注）不得残留为孤儿
        slide_texts = []
        for n in sorted(names):
            if re.match(r"ppt/slides/slide\d+\.xml$", n):
                slide_texts.append(
                    zipfile.ZipFile(str(out)).read(n).decode("utf-8"))
        joined = "".join(slide_texts)
        assert "演讲备注" not in joined
        dangling, orphans = _package_violations(out.read_bytes())
        assert dangling == []
        assert orphans == []

    def test_template_analyzer_static_delegate(self):
        from office_agent.ppt_agent.template_analyzer import TemplateAnalyzer

        prs = Presentation(io.BytesIO(_build_template(with_custom_show=True)))
        TemplateAnalyzer._clear_slides(prs)
        assert len(prs.slides) == 0
        buf = io.BytesIO()
        prs.save(buf)
        dangling, orphans = _package_violations(buf.getvalue())
        assert dangling == []
        assert orphans == []


def _make_outline(base_template_path: str):
    """构造最小 PPTOutline，附带模板基底路径。"""
    from office_agent.ppt_agent.models import PPTOutline

    outline = PPTOutline(title="测试汇报", slides=[])
    outline._base_template_path = base_template_path
    return outline
