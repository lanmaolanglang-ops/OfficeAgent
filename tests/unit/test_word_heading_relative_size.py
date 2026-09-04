"""Word 手动标题相对字号判断专项测试。

`_analyze_headings` 对无标题样式但加粗加大的段落，按文档正文字号
（出现次数最多的字号）的相对倍率判断：>= 1.25 倍视为手动标题，
>= 1.5 倍视为一级标题；不再使用硬编码 16/18pt 阈值。
"""
from docx import Document
from docx.shared import Pt

from office_agent.quality_scoring.word_scorer import WordQualityScorer


def _headings(body_size=12, body_count=4, extra=()):
    """构造 body_count 段正文 + extra 自定义段落，返回标题分析结果。

    extra: [(text, size_pt, bold)]
    """
    doc = Document()
    for i in range(body_count):
        para = doc.add_paragraph()
        run = para.add_run(f"正文段落内容第{i + 1}段，用于建立字号基准。")
        run.font.size = Pt(body_size)
        run.font.name = "宋体"
    for text, size, bold in extra:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "黑体"
    scorer = WordQualityScorer()
    return scorer._analyze_headings(doc, scorer._analyze_paragraphs(doc))


def test_bold_16pt_on_12pt_body_is_level2_heading():
    result = _headings(extra=[("手动二级标题", 16, True)])
    assert result["total_headings"] == 1
    assert result["headings"][0]["level"] == 2
    assert result["headings"][0]["is_manual"] is True


def test_bold_18pt_on_12pt_body_is_level1_heading():
    result = _headings(extra=[("手动一级标题", 18, True)])
    assert result["headings"][0]["level"] == 1


def test_relative_rule_catches_small_body_documents():
    # 正文 10pt 时 13pt 加粗即应视为标题（旧硬编码 >=16 会漏检）
    result = _headings(body_size=10, extra=[("小字文档的手动标题", 13, True)])
    assert result["total_headings"] == 1
    assert result["headings"][0]["is_manual"] is True


def test_bold_same_size_as_body_is_not_heading():
    result = _headings(extra=[("加粗但字号相同的强调句", 12, True)])
    assert result["total_headings"] == 0


def test_large_but_not_bold_is_not_heading():
    result = _headings(extra=[("只有字号没有加粗", 20, False)])
    assert result["total_headings"] == 0


def test_style_heading_unaffected_by_relative_rule():
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run("正文基准段落，用于提供正文字号。")
    run.font.size = Pt(12)
    doc.add_heading("样式标题", level=1)
    scorer = WordQualityScorer()
    result = scorer._analyze_headings(doc, scorer._analyze_paragraphs(doc))
    assert result["headings"][0]["level"] == 1
    assert result["headings"][0]["is_manual"] is False
