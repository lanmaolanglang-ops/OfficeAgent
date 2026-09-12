"""P1-17：PPT 规划器不得伪造占位正文，且标量字段必须净化。

修复前行为（独立旧源码副本实测，见最终报告 F 节）：
- ``plan_from_word`` 在章节/章没有可提取正文时写入 ``f"{X}相关内容"``，
  把"看起来像用户内容"的占位文本当成正式 PPT 正文输出。
- ``_parse_ai_outline`` 补封面时直接用 ``data.get("subtitle", "")``，
  LLM 返回 list/dict 时会被写成 Python repr（如 ``"['a', 'b']"``）进成品。

这些用例断言**最终产物内容**（大纲字段 / 生成出的 pptx 文本），
而不只是"函数没抛异常"。
"""

from __future__ import annotations

from types import SimpleNamespace

import pytest
from docx import Document
from pptx import Presentation

from office_agent.ppt_agent.content_planner import ContentPlanner
from office_agent.ppt_agent.ppt_orchestrator import PPTOrchestrator


def _gateway(content="", success=True):
    return SimpleNamespace(
        chat=lambda **kwargs: SimpleNamespace(
            success=success,
            content=content,
            error="" if success else "simulated failure",
            raw_response={},
        ),
    )


def _docx_with_empty_chapter(path) -> str:
    """构造一份"只有标题、没有正文"的 Word，用于触发空内容路径。"""
    doc = Document()
    doc.add_heading("第一章 空章节", level=1)
    doc.save(str(path))
    return str(path)


FABRICATION_MARKERS = ("相关内容", "暂无内容", "更多内容", "详细介绍", "待补充")


def _assert_not_python_repr(value: str):
    """项目既有 _scalar_text 用 JSON 序列化容器；Python repr 的特征是单引号包裹。"""
    assert isinstance(value, str)
    assert "['" not in value, f"出现 Python list repr: {value!r}"
    assert "{'" not in value, f"出现 Python dict repr: {value!r}"


def _assert_no_fabrication(texts: list[str]):
    joined = "\n".join(texts)
    for marker in FABRICATION_MARKERS:
        assert marker not in joined, f"出现伪造占位文本 {marker!r}: {joined!r}"


class TestNoFabricatedPlaceholderContent:
    def test_word_without_body_does_not_fabricate_bullets(self, temp_dir):
        path = _docx_with_empty_chapter(temp_dir / "empty.docx")
        outline = ContentPlanner().plan_from_word(path)

        content_slides = [s for s in outline.slides if s.layout == "content"]
        assert content_slides, "应保留内容页（受控缺失），而不是丢弃整页"

        for slide in content_slides:
            assert slide.bullets == [], f"空章节不得伪造要点: {slide.bullets!r}"

        _assert_no_fabrication([b for s in outline.slides for b in s.bullets])

    def test_empty_section_keeps_page_and_does_not_crash_renderer(self, temp_dir):
        """空内容页仍能完整生成 pptx：空要点是合法受控状态。"""
        path = _docx_with_empty_chapter(temp_dir / "empty2.docx")
        out = temp_dir / "empty2.pptx"

        result = PPTOrchestrator().generate_from_word(str(path), output_path=str(out))

        assert result.success is True, result.message
        prs = Presentation(str(out))
        texts = [
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        ]
        _assert_no_fabrication(texts)

    def test_planner_missing_bullets_key_yields_empty_list(self):
        outline = ContentPlanner().plan_from_outline_data("标题", [
            {"layout": "content", "title": "没有要点的页"},
        ])
        assert outline.slides[0].bullets == []

    @pytest.mark.parametrize("value", [[], None, "", [], {}])
    def test_bullets_empty_or_none_never_becomes_placeholder(self, value):
        outline = ContentPlanner().plan_from_outline_data("标题", [
            {"layout": "content", "title": "页", "bullets": value},
        ])
        assert outline.slides[0].bullets == []

    def test_llm_empty_content_falls_back_without_fabrication(self):
        planner = ContentPlanner(model_gateway=_gateway(content=""))
        outline = planner.plan_from_theme("项目汇报", slide_count=8)
        assert outline.used_template is True
        _assert_no_fabrication([b for s in outline.slides for b in s.bullets])

    def test_normal_content_is_preserved(self):
        outline = ContentPlanner().plan_from_outline_data("标题", [
            {"layout": "content", "title": "页", "bullets": ["真实要点一", "真实要点二"]},
        ])
        assert outline.slides[0].bullets == ["真实要点一", "真实要点二"]


class TestScalarSanitization:
    @pytest.mark.parametrize("value,expected", [
        ("副标题", "副标题"),
        (None, ""),
        (123, "123"),
        (True, "True"),
    ])
    def test_outline_data_title_and_subtitle_scalar(self, value, expected):
        outline = ContentPlanner().plan_from_outline_data("标题", [
            {"layout": "content", "title": value, "subtitle": value},
        ])
        slide = outline.slides[0]
        assert isinstance(slide.title, str)
        assert isinstance(slide.subtitle, str)
        assert slide.title == expected
        assert slide.subtitle == expected

    @pytest.mark.parametrize("value", [["a", "b"], {"x": 1}])
    def test_outline_data_container_becomes_string_not_python_repr(self, value):
        outline = ContentPlanner().plan_from_outline_data("标题", [
            {"layout": "content", "title": value, "subtitle": value},
        ])
        slide = outline.slides[0]
        _assert_not_python_repr(slide.title)
        _assert_not_python_repr(slide.subtitle)

    def test_plan_from_theme_sanitizes_theme_and_subtitle(self):
        planner = ContentPlanner()
        outline = planner.plan_from_theme(
            ["甲", "乙"], slide_count=6, subtitle={"k": "v"}, author=["作者"],
        )
        assert isinstance(outline.title, str)
        assert isinstance(outline.subtitle, str)
        assert isinstance(outline.author, str)
        _assert_not_python_repr(outline.subtitle)
        _assert_not_python_repr(outline.author)

    def test_cover_subtitle_from_llm_is_sanitized(self):
        """LLM 用 list 当 subtitle 时，补封面路径不得写入 Python repr。"""
        planner = ContentPlanner.__new__(ContentPlanner)
        data = {
            "subtitle": ["第一行", "第二行"],
            # 故意让首帧不是封面，触发补封面分支
            "slides": [{"layout": "content", "title": "内容"}],
        }
        outline = planner._parse_ai_outline(data, "主题", "professional")
        cover = outline.slides[0]
        assert cover.layout == "cover"
        assert isinstance(cover.subtitle, str)
        _assert_not_python_repr(cover.subtitle)

    def test_cover_subtitle_dict_is_sanitized(self):
        planner = ContentPlanner.__new__(ContentPlanner)
        data = {
            "subtitle": {"行1": "值"},
            "slides": [{"layout": "content", "title": "内容"}],
        }
        outline = planner._parse_ai_outline(data, "主题", "professional")
        cover = outline.slides[0]
        assert isinstance(cover.subtitle, str)
        _assert_not_python_repr(cover.subtitle)

    def test_cover_subtitle_none_and_number(self):
        for raw, expected in [(None, ""), (7, "7")]:
            planner = ContentPlanner.__new__(ContentPlanner)
            data = {"subtitle": raw, "slides": [{"layout": "content", "title": "内容"}]}
            outline = planner._parse_ai_outline(data, "主题", "professional")
            assert outline.slides[0].subtitle == expected

    def test_plan_from_text_sanitizes_title(self):
        outline = ContentPlanner().plan_from_text("### 页\n- 要点", title=["A", "B"])
        assert isinstance(outline.title, str)
        assert isinstance(outline.slides[0].subtitle, str)
        _assert_not_python_repr(outline.title)
