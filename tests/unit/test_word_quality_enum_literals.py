"""Word 质检 IssueSeverity/IssueType 枚举字面量收敛专项测试（清单 5.1 剩余部分）。

根因：quality/checker.py 定义了 IssueSeverity/IssueType 枚举，
但构造与比较点长期使用裸字符串字面量，枚举形同虚设。
收敛后：全部构造/比较经枚举 .value 单一来源，QualityIssue 对外仍是
普通字符串（API 不变），AST 守卫钉住不再有字面量回潮。
"""
import ast
import inspect

import pytest
from docx import Document
from docx.shared import Pt

from office_agent.quality.checker import (
    IssueSeverity, IssueType, QualityChecker, QualityIssue, QualityReport)


def _severity_values():
    return {s.value for s in IssueSeverity}


def _type_values():
    return {t.value for t in IssueType}


class TestEnumValuesPinned:
    def test_severity_values(self):
        assert _severity_values() == {"error", "warning", "info"}

    def test_type_values_cover_checker_vocabulary(self):
        assert _type_values() == {
            "font", "font_size", "line_spacing", "heading_level",
            "numbering", "table_format", "garbled", "empty_paragraph",
            "missing_text", "alignment", "indent", "page_setup",
        }


class TestReportSemanticsUnchanged:
    def test_score_weighting_unchanged(self):
        """error −10 / warning −3 / info −1 的计分口径与旧字面量版本一致。"""
        report = QualityReport()
        report.add_issue(QualityIssue(type="font", severity="error", message="e"))
        report.add_issue(QualityIssue(type="font", severity="warning", message="w"))
        report.add_issue(QualityIssue(type="font", severity="info", message="i"))
        report.compute_score()
        assert report.score == pytest.approx(86.0)
        assert report.passed is False
        assert len(report.errors()) == 1
        assert len(report.warnings()) == 1

    def test_error_marks_report_not_passed_on_add(self):
        report = QualityReport()
        report.add_issue(QualityIssue(
            type=IssueType.GARBLED.value,
            severity=IssueSeverity.ERROR.value, message="x"))
        assert report.passed is False

    def test_get_fix_config_matches_enum_typed_issues(self):
        report = QualityReport()
        report.add_issue(QualityIssue(
            type=IssueType.FONT.value, severity=IssueSeverity.WARNING.value,
            message="第1段正文字体为「黑体」，期望「宋体」", expected="宋体"))
        report.add_issue(QualityIssue(
            type=IssueType.LINE_SPACING.value,
            severity=IssueSeverity.WARNING.value,
            message="行距", expected="1.5 倍"))
        fix = report.get_fix_config()
        assert fix["font"] == "宋体"
        assert fix["line_spacing"] == pytest.approx(1.5)

    def test_to_text_groups_all_three_severities(self):
        report = QualityReport(file_path="x.docx")
        report.add_issue(QualityIssue(
            type=IssueType.FONT.value, severity=IssueSeverity.ERROR.value, message="e"))
        report.add_issue(QualityIssue(
            type=IssueType.FONT.value, severity=IssueSeverity.WARNING.value, message="w"))
        report.add_issue(QualityIssue(
            type=IssueType.FONT.value, severity=IssueSeverity.INFO.value, message="i"))
        text = report.to_text()
        assert "【错误】1项" in text
        assert "【警告】1项" in text
        assert "【提示】1项" in text


class TestCheckerEndToEnd:
    def test_issues_carry_enum_vocabulary_strings(self, tmp_path):
        """真实 docx 检查：issue 的 type/severity 仍是枚举词表内的普通字符串。"""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("这是一段足够长的正文内容，用于触发行距与字体检查。")
        run.font.name = "黑体"
        run.font.size = Pt(20)
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            "黑体")
        path = tmp_path / "sample.docx"
        doc.save(path)

        report = QualityChecker().check(str(path))
        assert report.issues, "构造的文档应至少产生一个质量问题"
        for issue in report.issues:
            assert isinstance(issue.severity, str)
            assert isinstance(issue.type, str)
            assert issue.severity in _severity_values()
            assert issue.type in _type_values()


class TestSourceGuard:
    def test_no_literal_type_or_severity_kwargs(self):
        """checker.py 中任何 type=/severity= 关键字实参不得回潮为字符串字面量。"""
        tree = ast.parse(inspect.getsource(
            __import__("office_agent.quality.checker", fromlist=["checker"])))
        offenders = []
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            for kw in node.keywords:
                if kw.arg in ("type", "severity") and isinstance(kw.value, ast.Constant):
                    offenders.append((node.lineno, kw.arg, kw.value.value))
        assert offenders == []

    def test_no_literal_severity_or_type_comparison(self):
        """checker.py 中不得再出现与 severity/type 词表字面量的 == 比较。"""
        vocabulary = _severity_values() | _type_values()
        tree = ast.parse(inspect.getsource(
            __import__("office_agent.quality.checker", fromlist=["checker"])))
        offenders = []
        for node in ast.walk(tree):
            if not isinstance(node, ast.Compare):
                continue
            for comp in node.comparators:
                if isinstance(comp, ast.Constant) and comp.value in vocabulary:
                    offenders.append((node.lineno, comp.value))
        assert offenders == []
