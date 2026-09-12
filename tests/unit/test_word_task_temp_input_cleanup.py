"""P2-9：Word 任务临时转换输入的生命周期回归测试。

覆盖重点：ensure_docx_input 对 .txt/.md/.pdf 会新建一个临时 .docx，其所有权
归调用方；process_word 此前在所有退出路径都没有清理它，任务跑一次留一个。
同时必须保证用户原始 docx（借用的路径）绝不被删。
"""
import os
import tempfile
from types import SimpleNamespace

import pytest
from docx import Document

from office_agent.services import input_conversion
from office_agent.services.input_conversion import (
    OWNED_TEMP_PREFIX, ensure_docx_input, is_owned_temp_input,
)
from office_agent.task_queue.tasks import word_tasks


class _FakeStorage:
    """替代全局 StorageService：只提供 process_word 需要的落点。"""

    def __init__(self, fail=False):
        self.fail = fail
        self.registered = []

    def save_new_output(self, **kwargs):
        if self.fail:
            raise RuntimeError("输出登记失败")
        self.registered.append(kwargs)
        return SimpleNamespace(file_id="file_out_1")


@pytest.fixture
def fake_storage(monkeypatch):
    from office_agent.storage import storage_service

    holder = {}

    def factory():
        return holder["service"]

    def install(service):
        holder["service"] = service
        return service

    monkeypatch.setattr(storage_service, "get_storage_service", factory)
    return install


@pytest.fixture
def no_output_dir_pollution(tmp_path, monkeypatch):
    monkeypatch.setenv("OFFICE_AGENT_OUTPUT_DIR", str(tmp_path / "outputs"))


@pytest.fixture
def captured_temps(monkeypatch):
    """记录 ensure_docx_input 的返回值，用于断言清理是否真的发生。"""
    seen = []
    real = ensure_docx_input

    def spy(path):
        resolved = real(path)
        seen.append(resolved)
        return resolved

    monkeypatch.setattr(input_conversion, "ensure_docx_input", spy)
    return seen


def _write_docx(path, text="正文内容"):
    doc = Document()
    doc.add_heading("标题", level=1)
    doc.add_paragraph(text)
    doc.save(str(path))
    return str(path)


def _write_txt(path, text="第一季度总结\n营收增长 20%"):
    path = path if isinstance(path, str) else str(path)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(text)
    return path


class TestOwnershipPredicate:
    def test_owned_temp_is_recognized(self, tmp_path):
        resolved = ensure_docx_input(_write_txt(tmp_path / "报告.txt"))
        try:
            assert is_owned_temp_input(resolved)
            assert os.path.basename(resolved).startswith(OWNED_TEMP_PREFIX)
        finally:
            os.remove(resolved)

    def test_original_docx_is_borrowed(self, tmp_path):
        original = _write_docx(tmp_path / "原始文档.docx")
        assert ensure_docx_input(original) == original
        assert not is_owned_temp_input(original)

    def test_user_file_named_like_temp_is_not_owned(self, tmp_path):
        """同名陷阱：用户文件恰好带模块前缀时不得被判成临时文件。"""
        trap = _write_docx(tmp_path / f"{OWNED_TEMP_PREFIX}trap.docx")
        assert not is_owned_temp_input(trap)
        assert os.path.exists(trap)

    def test_empty_and_odd_values(self):
        for value in ("", None, "relative/path.docx"):
            assert not is_owned_temp_input(value)

    def test_temp_file_sits_in_system_temp(self, tmp_path):
        resolved = ensure_docx_input(_write_txt(tmp_path / "a.txt"))
        try:
            assert os.path.dirname(os.path.realpath(resolved)) == \
                os.path.realpath(tempfile.gettempdir())
        finally:
            os.remove(resolved)


class TestProcessWordCleanup:
    def test_native_docx_is_never_deleted(self, tmp_path, fake_storage):
        fake_storage(_FakeStorage())
        original = _write_docx(tmp_path / "原始文档.docx")
        output = tmp_path / "out.docx"

        result = word_tasks.process_word(
            original, output_path=str(output), options={}, _task_id="t1",
        )

        assert result["status"] == "success"
        assert os.path.exists(original)  # 借用的用户文件必须保留

    def test_converted_temp_is_removed_on_success(self, tmp_path, fake_storage,
                                                  captured_temps):
        fake_storage(_FakeStorage())
        source = _write_txt(tmp_path / "报告.txt")
        output = tmp_path / "out.docx"

        result = word_tasks.process_word(
            source, output_path=str(output), options={}, _task_id="t2",
        )

        assert result["status"] == "success"
        assert len(captured_temps) == 1
        temp = captured_temps[0]
        assert temp != source
        assert is_owned_temp_input(temp)
        assert not os.path.exists(temp)      # 成功路径清理
        assert os.path.exists(source)        # 用户原始文件保留

    def test_converted_temp_is_removed_on_engine_failure(self, tmp_path,
                                                         fake_storage,
                                                         captured_temps,
                                                         monkeypatch):
        from office_agent.services.word_service import WordService

        fake_storage(_FakeStorage())
        monkeypatch.setattr(
            WordService, "process",
            lambda self, *a, **kw: (_ for _ in ()).throw(RuntimeError("引擎炸了")),
        )
        source = _write_txt(tmp_path / "报告.txt")

        result = word_tasks.process_word(
            source, output_path=str(tmp_path / "out.docx"), options={}, _task_id="t3",
        )

        assert result["status"] == "failed"
        assert "引擎炸了" in result["error"]
        assert not os.path.exists(captured_temps[0])

    def test_converted_temp_is_removed_on_output_registration_failure(
            self, tmp_path, fake_storage, captured_temps):
        fake_storage(_FakeStorage(fail=True))
        source = _write_txt(tmp_path / "报告.txt")

        result = word_tasks.process_word(
            source, output_path=str(tmp_path / "out.docx"), options={}, _task_id="t4",
        )

        assert result["status"] == "failed"
        assert "输出登记失败" in result["error"]
        assert not os.path.exists(captured_temps[0])

    def test_converted_temp_is_removed_when_input_missing_downstream(
            self, tmp_path, fake_storage, captured_temps, monkeypatch):
        """输出文件未生成（登记前就失败）也要清理临时输入。"""
        from office_agent.services.word_service import ProcessResult, WordService

        fake_storage(_FakeStorage())
        monkeypatch.setattr(
            WordService, "process",
            lambda self, *a, **kw: ProcessResult(
                success=True, message="ok", output_path=str(tmp_path / "ghost.docx"),
            ),
        )
        source = _write_txt(tmp_path / "报告.txt")

        result = word_tasks.process_word(
            source, output_path=str(tmp_path / "out.docx"), options={}, _task_id="t5",
        )

        assert result["status"] == "failed"
        assert "未生成输出文件" in result["error"]
        assert not os.path.exists(captured_temps[0])

    def test_cleanup_failure_preserves_original_error(self, tmp_path, fake_storage,
                                                      captured_temps, monkeypatch):
        """清理自身失败不得掩盖主错误信息（也不得把任务报成成功）。"""
        from office_agent.services.word_service import WordService

        fake_storage(_FakeStorage())
        monkeypatch.setattr(
            WordService, "process",
            lambda self, *a, **kw: (_ for _ in ()).throw(RuntimeError("真正的失败")),
        )
        source = _write_txt(tmp_path / "报告.txt")
        real_remove = os.remove

        def refuse(path, **kwargs):
            if path in captured_temps:
                raise OSError("临时文件被占用")
            return real_remove(path, **kwargs)

        monkeypatch.setattr(os, "remove", refuse)

        result = word_tasks.process_word(
            source, output_path=str(tmp_path / "out.docx"), options={}, _task_id="t6",
        )

        assert result["status"] == "failed"
        assert "真正的失败" in result["error"]
        assert "被占用" not in result["error"]

    def test_repeated_runs_leave_no_residue(self, tmp_path, fake_storage,
                                            captured_temps):
        fake_storage(_FakeStorage())
        source = _write_txt(tmp_path / "报告.txt")

        for index in range(3):
            word_tasks.process_word(
                source, output_path=str(tmp_path / f"out{index}.docx"),
                options={}, _task_id=f"t{index}",
            )

        assert len(captured_temps) == 3
        assert len(set(captured_temps)) == 3          # 每次都是独立临时文件
        assert not any(os.path.exists(p) for p in captured_temps)
        assert os.path.exists(source)

    def test_user_file_named_like_temp_survives_the_task(self, tmp_path,
                                                         fake_storage):
        """端到端：即便文件名与临时前缀同形，用户文件也不能被删。"""
        fake_storage(_FakeStorage())
        trap = _write_docx(tmp_path / f"{OWNED_TEMP_PREFIX}trap.docx")

        result = word_tasks.process_word(
            trap, output_path=str(tmp_path / "out.docx"), options={}, _task_id="t7",
        )

        assert result["status"] == "success"
        assert os.path.exists(trap)
