"""L557 异常边界专项测试。

按异常处理模式（而非逐 catch）验证本批收窄后的行为契约：
- 预期业务/数据异常仍被正确处理（fail-open / fail-closed 方向不变）
- 编程错误（AttributeError/KeyError 等）不再被静默吞掉
- fallback 行为与返回值契约不变
- 日志留痕且不包含敏感内容
"""
import logging
from contextlib import contextmanager

import pytest

from office_agent.security.auth.password import hash_password, verify_password
from office_agent.security.auth.jwt import JWTManager
from office_agent.config_system.loaders import YamlLoader
from office_agent.ppt_agent.quality_checker import PPTQualityChecker, PPTQualityReport
from office_agent.runtime_manager import AppConfig, AppStatus, ApplicationRuntimeManager


@contextmanager
def recorded_logs(logger_name: str):
    """直接录制指定模块 logger（套件级配置会改传播，不用 caplog）。"""
    target = logging.getLogger(logger_name)
    records = []

    class _ListHandler(logging.Handler):
        def emit(self, record):
            records.append(record)

    handler = _ListHandler()
    old_level = target.level
    target.addHandler(handler)
    target.setLevel(logging.DEBUG)
    try:
        yield records
    finally:
        target.removeHandler(handler)
        target.setLevel(old_level)


# ---------------------------------------------------------------- password

class TestVerifyPassword:
    def test_roundtrip(self):
        stored = hash_password("S3cret!")
        assert verify_password("S3cret!", stored) is True
        assert verify_password("wrong", stored) is False

    def test_malformed_hash_returns_false_and_logs(self):
        # 畸形存储哈希：fail-closed 返回 False，且留痕（此前静默 except Exception）
        with recorded_logs("office_agent.security.auth.password") as records:
            assert verify_password("pw", "pbkdf2_sha256$notanint$@@@$###") is False
        assert any(r.levelno >= logging.WARNING for r in records)
        # 日志不得包含密码或哈希内容
        assert all("pbkdf2_sha256$notanint" not in r.getMessage() for r in records)

    def test_non_string_inputs_return_false(self):
        assert verify_password(None, "whatever") is False
        assert verify_password("pw", None) is False

    def test_programming_error_propagates(self, monkeypatch):
        # AttributeError 等编程错误不得再被伪装成"密码错误"
        import office_agent.security.auth.password as password_module

        def _broken(*args, **kwargs):
            raise AttributeError("simulated programming bug")

        monkeypatch.setattr(password_module.hashlib, "pbkdf2_hmac", _broken)
        with pytest.raises(AttributeError):
            verify_password("pw", hash_password("pw"))


# ---------------------------------------------------------------- jwt

class TestJwtDecode:
    def test_roundtrip(self, tmp_path):
        manager = JWTManager(secret_key="test-secret", state_dir=tmp_path)
        token = manager.create_access_token(user_id="u1", username="u", role="admin")
        payload = manager.decode(token)
        assert payload.user_id == "u1"

    def test_tampered_token_raises_value_error(self, tmp_path):
        manager = JWTManager(secret_key="test-secret", state_dir=tmp_path)
        token = manager.create_access_token(user_id="u1", username="u", role="admin")
        with pytest.raises(ValueError):
            manager.decode(token[:-2] + "xx")

    def test_non_string_token_propagates_programming_error(self, tmp_path):
        # 此前 except Exception 会把 AttributeError 包装成"Token无效"
        manager = JWTManager(secret_key="test-secret", state_dir=tmp_path)
        with pytest.raises(AttributeError):
            manager.decode(123)


# ---------------------------------------------------------------- yaml loader

class TestYamlLoader:
    def test_load_failure_returns_empty_and_logs(self, tmp_path):
        import importlib.util

        (tmp_path / "config.yaml").write_text("a: [unclosed\n", encoding="utf-8")
        loader = YamlLoader(config_dir=str(tmp_path))
        with recorded_logs("office_agent.config.loader") as records:
            assert loader.load() == {}
        if importlib.util.find_spec("yaml") is None:
            # 当前环境无 PyYAML：ImportError 分支是活路径
            assert any(r.levelno >= logging.WARNING for r in records)
        else:
            assert any(r.levelno >= logging.ERROR for r in records)

    def test_valid_yaml_when_available(self, tmp_path):
        import importlib.util

        (tmp_path / "config.yaml").write_text("key: value\n", encoding="utf-8")
        loader = YamlLoader(config_dir=str(tmp_path))
        if importlib.util.find_spec("yaml") is None:
            assert loader.load() == {}
        else:
            assert loader.load() == {"key": "value"}

    def test_missing_file_returns_empty(self, tmp_path):
        assert YamlLoader(config_dir=str(tmp_path)).load() == {}


# ---------------------------------------------------------------- ppt quality checker

class _BrokenShape:
    """left/top/width/height 访问抛出指定异常的假元素。"""

    def __init__(self, exc):
        self._exc = exc

    def __getattr__(self, name):
        if name in ("left", "top", "width", "height", "text_frame"):
            raise self._exc
        raise AttributeError(name)


class _BrokenTextShape:
    has_text_frame = True

    @property
    def text_frame(self):
        raise TypeError("corrupt text frame")


class TestPPTQualityCheckerBoundaries:
    def setup_method(self):
        self.checker = PPTQualityChecker()
        self.report = PPTQualityReport()

    def test_unreadable_coordinates_skipped_and_logged(self):
        shape = _BrokenShape(AttributeError("legacy pptx quirk"))
        with recorded_logs("office_agent.ppt_agent.quality_checker") as records:
            self.checker._check_shape_bounds(shape, 10.0, 7.5, 0, self.report)
        assert self.report.issues == []
        assert any(r.levelno >= logging.DEBUG for r in records)

    def test_programming_error_propagates(self):
        shape = _BrokenShape(KeyError("simulated programming bug"))
        with pytest.raises(KeyError):
            self.checker._check_shape_bounds(shape, 10.0, 7.5, 0, self.report)

    def test_text_overflow_shape_error_skipped_and_logged(self):
        with recorded_logs("office_agent.ppt_agent.quality_checker") as records:
            self.checker._check_text_overflow(_BrokenTextShape(), 0, self.report)
        assert self.report.issues == []
        assert any(r.levelno >= logging.DEBUG for r in records)


# ---------------------------------------------------------------- runtime manager

class TestRuntimeManagerCallbacks:
    def test_failing_callback_logged_and_isolated(self):
        manager = ApplicationRuntimeManager(AppConfig())
        seen = []

        def bad_callback(status, state):
            raise RuntimeError("gui callback boom")

        manager.on_status_change(bad_callback)
        manager.on_status_change(lambda status, state: seen.append(status))
        with recorded_logs("office_agent.runtime_manager") as records:
            manager._set_status(AppStatus.STOPPED)
        # 失败回调不影响其他回调与主流程
        assert seen == [AppStatus.STOPPED]
        assert any(r.levelno >= logging.WARNING for r in records)


# ---------------------------------------------------------------- image config

class TestImageModelConfigLoad:
    def test_corrupt_config_falls_back_and_logs(self, tmp_path):
        # 惰性 import 且需先加载 api.main：直接导入本模块会经
        # model_gateway -> api.routing 形成部分初始化循环（与生产入口顺序一致即可）
        import office_agent.api.main  # noqa: F401
        from office_agent.image_generation.config import ImageModelConfigManager

        (tmp_path / "image_model.json").write_text("{corrupt", encoding="utf-8")
        with recorded_logs("office_agent.image_generation.config") as records:
            manager = ImageModelConfigManager(config_dir=str(tmp_path))
        config = manager.get_config()
        assert config["provider"] == "agnes"
        assert config["api_key"] == ""
        assert any(r.levelno >= logging.WARNING for r in records)

    def test_missing_config_uses_default(self, tmp_path):
        import office_agent.api.main  # noqa: F401
        from office_agent.image_generation.config import ImageModelConfigManager

        manager = ImageModelConfigManager(config_dir=str(tmp_path))
        assert manager.get_config()["provider"] == "agnes"


# ---------------------------------------------------------------- sqlite pragma

class TestSqlitePragmaBoundary:
    def test_memory_engine_usable_despite_pragma_tolerance(self):
        from sqlalchemy import text

        from office_agent.database.connection import get_engine

        engine = get_engine("sqlite:///:memory:")
        try:
            with engine.connect() as conn:
                assert conn.execute(text("SELECT 1")).scalar() == 1
        finally:
            engine.dispose()


# ---------------------------------------------------------------- formula generator

class TestFormulaGeneratorProfileFallback:
    def test_analyze_failure_continues_and_logs(self, tmp_path, monkeypatch):
        from openpyxl import Workbook

        from office_agent.excel_agent.data_analyzer import DataAnalyzer
        from office_agent.excel_agent.formula_generator import FormulaGenerator

        workbook = Workbook()
        workbook.active["A1"] = 1
        source = tmp_path / "data.xlsx"
        workbook.save(source)

        def _broken_analyze(self, file_path):
            raise RuntimeError("pandas exploded")

        monkeypatch.setattr(DataAnalyzer, "analyze", _broken_analyze)
        generator = FormulaGenerator()
        output = tmp_path / "out.xlsx"
        with recorded_logs("office_agent.excel_agent.formula_generator") as records:
            result_path, formulas = generator.apply_to_file(
                str(source), "在B1写入1", output_path=str(output)
            )
        # fallback 行为不变：无画像也能完成公式生成
        assert generator.profile is None
        assert result_path == str(output)
        assert output.exists()
        assert any(r.levelno >= logging.WARNING for r in records)
