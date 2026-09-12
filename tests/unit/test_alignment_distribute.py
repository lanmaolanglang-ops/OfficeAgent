"""P1-26：「分散对齐」的 canonical 值必须被消费者接收，不得静默回退成两端对齐。

修复前行为（旧源码副本实测）：
``parsers/format_parser.py`` 把「分散对齐」归一为 ``"distribute"``，
``parsers/document_structure.py`` 把 ``WD_ALIGN_PARAGRAPH.DISTRIBUTE`` 也归一为
``"distribute"``，但消费端 ``word_service.ALIGNMENT_STR_MAP`` 没有这个键，
``Alignment`` 枚举也没有对应成员 → ``.get(..., Alignment.JUSTIFY)`` 静默回退，
用户选「分散对齐」最终得到「两端对齐」。
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from office_agent.models.schemas import Alignment, FormatConfig
from office_agent.parsers.document_structure import DocumentStructureAnalyzer
from office_agent.parsers.format_parser import ALIGNMENT_MAP as PARSER_ALIGNMENT_MAP
from office_agent.parsers.format_parser import parse_format_rule
from office_agent.services.word_service import (
    ALIGNMENT_MAP,
    ALIGNMENT_STR_MAP,
    WordService,
)

BODY_TEXT = "这是一段足够长的正文段落，用于核对对齐方式在最终文档里的真实取值。"

#: 自然语言 → canonical → python-docx 枚举
ALIGNMENT_TABLE = [
    ("左对齐", "left", WD_ALIGN_PARAGRAPH.LEFT),
    ("居中", "center", WD_ALIGN_PARAGRAPH.CENTER),
    ("居中对齐", "center", WD_ALIGN_PARAGRAPH.CENTER),
    ("右对齐", "right", WD_ALIGN_PARAGRAPH.RIGHT),
    ("两端对齐", "justify", WD_ALIGN_PARAGRAPH.JUSTIFY),
    ("分散对齐", "distribute", WD_ALIGN_PARAGRAPH.DISTRIBUTE),
]


def _written_alignment(tmp_path, config_dict: dict):
    """真正跑一次 WordService.process，读产物 docx 里正文段落的 alignment。"""
    source = tmp_path / "in.txt"
    source.write_text(f"# 标题\n{BODY_TEXT}\n", encoding="utf-8")
    output = tmp_path / "out.docx"

    result = WordService().process(str(source), config_dict=config_dict, output_path=str(output))
    assert result.success, result.message

    document = Document(str(output))
    target = next(p for p in document.paragraphs if BODY_TEXT[:8] in p.text)
    return target.alignment


class TestCanonicalValueContract:
    def test_distribute_is_a_real_alignment_member(self):
        assert Alignment.DISTRIBUTE.value == "distribute"

    def test_consumer_map_accepts_the_parser_canonical_value(self):
        assert ALIGNMENT_STR_MAP["distribute"] is Alignment.DISTRIBUTE

    @pytest.mark.parametrize("alias", ["distribute", "distributed", "分散对齐"])
    def test_aliases_resolve_to_the_same_enum(self, alias):
        assert ALIGNMENT_STR_MAP[alias] is Alignment.DISTRIBUTE

    def test_enum_has_a_python_docx_mapping(self):
        assert ALIGNMENT_MAP[Alignment.DISTRIBUTE] is WD_ALIGN_PARAGRAPH.DISTRIBUTE

    def test_parser_output_is_always_consumable(self):
        """通用契约：parser 能产出的每个 canonical 值，consumer 都必须认识。"""
        unfriendly = [
            value for value in set(PARSER_ALIGNMENT_MAP.values())
            if value not in ALIGNMENT_STR_MAP
        ]
        assert unfriendly == [], f"parser 会产出 consumer 不认识的值: {unfriendly}"

    @pytest.mark.parametrize("text,canonical,_expected", ALIGNMENT_TABLE)
    def test_parser_canonical_matches_contract(self, text, canonical, _expected):
        parsed = parse_format_rule(f"正文宋体小四，{text}")
        assert parsed["alignment"] == canonical
        assert ALIGNMENT_STR_MAP[parsed["alignment"]] in tuple(Alignment)


class TestConfigResolution:
    @pytest.mark.parametrize("text,canonical,expected", ALIGNMENT_TABLE)
    def test_config_from_dict_resolves_natural_language(self, text, canonical, expected):
        parsed = parse_format_rule(f"正文{text}")
        config = WordService().config_from_dict(parsed)
        assert config.body_paragraph.alignment is ALIGNMENT_STR_MAP[canonical]

    @pytest.mark.parametrize("alias", ["distribute", "distributed", "分散对齐"])
    def test_config_from_dict_accepts_distribute_aliases(self, alias):
        config = WordService().config_from_dict({"alignment": alias})
        assert config.body_paragraph.alignment is Alignment.DISTRIBUTE
        assert config.body_paragraph.alignment is not Alignment.JUSTIFY

    def test_unknown_value_keeps_the_existing_fallback_policy(self):
        """真正未知的值仍按项目既有策略回退 JUSTIFY（本批不改全局策略）。"""
        config = WordService().config_from_dict({"alignment": "对角线对齐"})
        assert config.body_paragraph.alignment is Alignment.JUSTIFY

    def test_default_alignment_unchanged(self):
        assert WordService().config_from_dict({}).body_paragraph.alignment is Alignment.JUSTIFY
        assert FormatConfig().body_paragraph.alignment is Alignment.JUSTIFY


class TestWrittenDocumentAlignment:
    @pytest.mark.parametrize("text,canonical,expected", ALIGNMENT_TABLE)
    def test_natural_language_reaches_the_document(self, tmp_path, text, canonical, expected):
        parsed = parse_format_rule(f"正文宋体小四，{text}")
        assert _written_alignment(tmp_path, parsed) == expected

    def test_distribute_is_not_silently_downgraded_to_justify(self, tmp_path):
        """核心回归：分散对齐必须真的写成 DISTRIBUTE，而不是 JUSTIFY。"""
        parsed = parse_format_rule("正文宋体小四，分散对齐")
        assert parsed["alignment"] == "distribute"

        written = _written_alignment(tmp_path, parsed)
        assert written is WD_ALIGN_PARAGRAPH.DISTRIBUTE
        assert written is not WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_justify_still_writes_justify(self, tmp_path):
        parsed = parse_format_rule("正文宋体小四，两端对齐")
        written = _written_alignment(tmp_path, parsed)
        assert written is WD_ALIGN_PARAGRAPH.JUSTIFY
        assert written is not WD_ALIGN_PARAGRAPH.DISTRIBUTE


class TestProducerConsumerClosure:
    def test_document_structure_producer_output_is_consumable(self, tmp_path):
        """document_structure 是另一个 "distribute" 生产者，契约同样必须闭合。"""
        path = tmp_path / "aligned.docx"
        document = Document()
        paragraph = document.add_paragraph(BODY_TEXT)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
        document.save(str(path))

        tree = DocumentStructureAnalyzer().analyze(Document(str(path)))
        produced = {node.alignment for node in tree.nodes if node.alignment}
        assert "distribute" in produced

        for value in produced:
            if value not in ALIGNMENT_STR_MAP:
                continue
            config = WordService().config_from_dict({"alignment": value})
            assert config.body_paragraph.alignment in tuple(Alignment)

        config = WordService().config_from_dict({"alignment": "distribute"})
        assert config.body_paragraph.alignment is Alignment.DISTRIBUTE

    def test_heading_alignment_path_also_accepts_distribute(self):
        config = WordService().config_from_dict({
            "headings": {"1": {"alignment": "分散对齐"}},
        })
        assert config.headings[1].paragraph.alignment is Alignment.DISTRIBUTE
