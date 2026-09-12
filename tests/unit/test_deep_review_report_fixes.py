"""Regression coverage for verified defects from the 2026-08-31 deep review."""
import asyncio
import json
from types import SimpleNamespace

import pytest


def test_gemini_vision_uses_camel_case_image_fields(monkeypatch):
    from office_agent.vision_gateway.clients.gemini_client import GeminiVisionClient
    from office_agent.vision_gateway.vision_models import ImageInput, VisionRequest

    captured = {}
    client = GeminiVisionClient("secret")

    def fake_post(_url, _headers, data):
        captured.update(json.loads(data))
        return {"candidates": [{"content": {"parts": [{"text": "{}"}]}}]}

    monkeypatch.setattr(client, "_http_post", fake_post)
    response = client.analyze(VisionRequest(images=[ImageInput.from_buffer(b"image")]))
    assert response.success
    image_part = captured["contents"][0]["parts"][1]
    assert image_part["inlineData"]["mimeType"] == "image/png"
    assert "inline_data" not in image_part


def test_percent_strings_are_scaled_to_ratios():
    from office_agent.excel_agent.analysis_engine import AnalysisEngine
    from office_agent.excel_agent.data_analyzer import DataAnalyzer
    import pandas as pd

    engine = AnalysisEngine()
    assert engine._to_float("50%") == pytest.approx(0.5)
    assert engine._get_numeric_values([["50%"], [0.25]], 0) == pytest.approx([0.5, 0.25])
    numeric = DataAnalyzer._to_numeric_series(pd.Series(["50%", "12.5%", 0.2]))
    assert numeric.tolist() == pytest.approx([0.5, 0.125, 0.2])


def test_upload_processing_assigns_non_office_file_type(tmp_path):
    from office_agent.task_queue.tasks.file_tasks import process_upload

    source = tmp_path / "notes.txt"
    source.write_text("hello", encoding="utf-8")
    result = process_upload(str(source), file_id="f1")
    assert result["status"] == "success"
    assert result["file_type"] == "text"


def test_unknown_general_task_is_not_reported_as_success():
    from office_agent.task_queue.tasks import process_general

    result = process_general("do something completely unsupported")
    assert result["status"] == "failed"
    assert result["error"]


def test_sandbox_injects_data_and_malformed_output_is_error(tmp_path, monkeypatch):
    from office_agent.security.sandbox.sandbox import Sandbox, SandboxStatus
    import office_agent.security.sandbox.sandbox as sandbox_module

    sandbox = Sandbox(work_dir=tmp_path, allow_unsafe_subprocess=True)
    result = sandbox.execute_data_analysis("result = data['value']", {"value": 7})
    assert result.status == SandboxStatus.SUCCESS
    assert result.result == "7"

    monkeypatch.setattr(
        sandbox_module.subprocess,
        "run",
        lambda *_args, **_kwargs: SimpleNamespace(
            stdout="{" * 100, stderr="", returncode=0
        ),
    )
    malformed = sandbox.execute("result = 1")
    assert malformed.status == SandboxStatus.ERROR
    assert "输出格式无效" in malformed.error


def test_vision_json_parser_handles_braces_strings_and_empty_table_cells():
    from office_agent.vision_gateway.clients.gemini_client import GeminiVisionClient

    client = GeminiVisionClient("test")
    text = 'prefix ```JSON\n{"summary":"contains {x} and [y]","elements":[{"confidence":null}]}\n``` suffix'
    extracted = client._extract_json(text)
    assert json.loads(extracted)["summary"] == "contains {x} and [y]"

    tables = client._extract_markdown_tables("|A|B|C|\n|---|---|---|\n|1||3|")
    assert tables[0].data == [["1", "", "3"]]


def test_model_gateway_accepts_chat_message_objects():
    from office_agent.model_gateway.gateway import ModelGateway
    from office_agent.models.model_schemas import ChatMessage, ModelResponse, AITaskType

    gateway = object.__new__(ModelGateway)
    observed = {}
    gateway.manager = SimpleNamespace(get_default_model_id=lambda: None)
    gateway.router = SimpleNamespace(
        infer_task_type=lambda text: observed.setdefault("text", text) or AITaskType.SIMPLE_TEXT,
        select_model=lambda *_args, **_kwargs: [],
    )
    gateway.failover = SimpleNamespace(
        execute_with_failover=lambda **_kwargs: ModelResponse(success=True, content="ok")
    )
    response = gateway.chat(messages=[ChatMessage(role="user", content="hello")])
    assert response.success
    assert observed["text"] == "hello"


def test_format_parser_prefers_compound_color_and_preserves_three_line():
    from office_agent.parsers.format_parser import FormatRuleParser

    config = FormatRuleParser().parse("正文深蓝色，表格使用三线表").to_config_dict()
    assert config["color"] == "1F4E79"
    assert config["table_three_line"] is True


def test_old_versions_keeps_latest_per_file(tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from office_agent.database.base import Base
    from office_agent.database.models.file import File, FileVersion
    from office_agent.database.repository.file_repo import FileRepository

    engine = create_engine(f"sqlite:///{tmp_path / 'versions.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add_all([
            File(id="f1", filename="a", original_name="a", file_type="text",
                 extension=".txt", storage_path="a"),
            File(id="f2", filename="b", original_name="b", file_type="text",
                 extension=".txt", storage_path="b"),
        ])
        for file_id in ("f1", "f2"):
            for version in range(1, 5):
                session.add(FileVersion(
                    parent_file_id=file_id, version_number=version,
                    storage_path=f"{file_id}/{version}",
                ))
        session.commit()
        old = FileRepository(session).get_old_versions(keep=2)
        assert {(item.parent_file_id, item.version_number) for item in old} == {
            ("f1", 1), ("f1", 2), ("f2", 1), ("f2", 2)
        }


def test_ppt_scorer_checks_required_content_and_title_identity(tmp_path):
    from pptx import Presentation
    from office_agent.quality_scoring.ppt_scorer import PPTQualityScorer

    path = tmp_path / "required.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Verified title"
    presentation.save(path)

    result = PPTQualityScorer().score(str(path), required_content=["must appear"])
    assert any("must appear" in issue for issue in result.issues)
    assert result.template_details["title_coverage"] == 1.0


def test_excel_chart_chinese_matching_and_static_formula_errors():
    from office_agent.quality_scoring.excel_scorer import ExcelQualityScorer

    scorer = ExcelQualityScorer()
    base = {"chart_count": 1, "chart_types": ["BarChart"], "total_data_cells": 20}
    assert scorer._score_charts(base, ["柱状图"]) > scorer._score_charts(base, ["饼图"])
    assert scorer._has_static_formula_error("=#REF!+1")
    assert scorer._has_static_formula_error("=10/0")


def test_fixed_excel_errors_do_not_fail_report():
    from office_agent.excel_agent.models import ExcelQualityIssue
    from office_agent.excel_agent.quality_checker import ExcelQualityChecker, QualityReport

    issue = ExcelQualityIssue(issue_type="formula", severity="error", message="fixed")
    issue.fixed = True
    report = QualityReport(issues=[issue], fixed_count=1)
    ExcelQualityChecker()._calc_score(report)
    assert report.error_count == 0
    assert report.passed is True


def test_feedback_missing_task_and_database_unavailable_are_not_success(monkeypatch):
    from office_agent.api.core.exceptions import APIError
    from office_agent.api.router import task as task_router
    from office_agent.api.schemas.request import FeedbackRequest

    monkeypatch.setattr(task_router, "_get_db_session", lambda: None)
    monkeypatch.setattr(task_router.task_manager, "get_task", lambda _task_id: None)
    with pytest.raises(APIError) as caught:
        asyncio.run(task_router.task_feedback("missing", FeedbackRequest(rating=5)))
    assert caught.value.status_code == 503


def test_ppt_missing_data_does_not_render_fabricated_values(tmp_path):
    from pptx import Presentation
    from office_agent.ppt_agent.models import PPTOutline, SlideContent
    from office_agent.ppt_agent.ppt_service import PPTService

    output = tmp_path / "no-fakes.pptx"
    outline = PPTOutline(title="Real topic")
    outline.add_slide(SlideContent(layout="toc", title="目录"))
    outline.add_slide(SlideContent(layout="data_cards", title="数据"))
    outline.add_slide(SlideContent(layout="timeline", title="计划"))
    generated = PPTService().generate(outline, str(output))
    assert generated.success
    presentation = Presentation(output)
    text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
    for forbidden in ("目录项 1", "指标一", "2024 Q1", "增长率", "50 万"):
        assert forbidden not in text


def test_document_vision_result_exposes_partial_state():
    from office_agent.vision_gateway.vision_models import DocumentVisionResult

    partial = DocumentVisionResult(successful_pages=2, failed_pages=1, partial=True)
    assert partial.success is True
    assert partial.partial is True
    failed = DocumentVisionResult(error="全部失败", failed_pages=3)
    assert failed.success is False


def test_vision_gateway_reports_failed_pages_as_partial():
    from office_agent.vision_gateway.gateway import VisionGateway
    from office_agent.vision_gateway.vision_models import (
        DocumentPage, ImageInput, VisionResponse,
    )

    gateway = object.__new__(VisionGateway)
    ok_image = ImageInput.from_buffer(b"one")
    ok_image.label = "ok"
    failed_image = ImageInput.from_buffer(b"two")
    failed_image.label = "fail"
    gateway.renderer = SimpleNamespace(
        dpi=150,
        max_pages=30,
        render=lambda _path: [
            DocumentPage(page_number=1, image=ok_image),
            DocumentPage(page_number=2, image=failed_image),
        ],
    )
    gateway.analyze = lambda request, _model_key: (
        VisionResponse(success=True, content="page one")
        if request.images[0].label == "ok"
        else VisionResponse(success=False, error="provider error")
    )
    result = gateway.analyze_document("example.pdf", max_concurrency=2)
    assert result.success is True
    assert result.partial is True
    assert result.successful_pages == 1
    assert result.failed_page_numbers == [2]


def test_config_update_does_not_mutate_cache_when_persistence_fails():
    import threading

    from office_agent.config_system.config_manager import ConfigManager

    manager = object.__new__(ConfigManager)
    manager._lock = threading.RLock()
    manager._model_store = None
    manager._models = {"m1": {"model_id": "m1", "model_name": "before"}}
    manager._agents = {}
    manager._prompts = {}
    manager._skills = {}
    manager._workflows = {}
    manager._change_listeners = []
    manager._session_factory = lambda: (_ for _ in ()).throw(RuntimeError("disk full"))

    with pytest.raises(RuntimeError):
        manager.update_model("m1", {"model_name": "after"})
    assert manager._models["m1"]["model_name"] == "before"

    with pytest.raises(RuntimeError):
        manager.update_prompt("p1", "new prompt")
    assert "p1" not in manager._prompts


def test_model_manager_skips_one_bad_entry_without_losing_all(tmp_path, monkeypatch):
    from office_agent.model_gateway.model_manager import ModelManager

    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    (tmp_path / "models.json").write_text(json.dumps({
        "models": [
            {"provider": "openai"},
            {"id": "valid", "provider": "openai", "display_name": "Valid", "model": "gpt-test"},
        ],
        "routing": {},
    }), encoding="utf-8")
    manager = ModelManager(str(tmp_path))
    assert manager.get_model("valid") is not None
    assert len(manager.list_models()) == 1


def test_worker_cancellation_stops_at_progress_checkpoint(monkeypatch):
    import threading
    import time
    from office_agent.task_queue.worker import LocalWorker

    worker = LocalWorker()
    worker._session_factory = None
    started = threading.Event()
    stopped = threading.Event()

    def cancellable(progress=None, **_kwargs):
        started.set()
        try:
            while True:
                time.sleep(0.01)
                progress.update(10, "checkpoint")
        finally:
            stopped.set()

    worker.register("test.cancellable", cancellable)
    task_id = worker.submit("test.cancellable")
    assert started.wait(1)
    worker.revoke(task_id)
    assert stopped.wait(1)
    worker.shutdown(wait=True)


def test_terminal_task_state_cannot_be_overwritten(tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from office_agent.database.base import Base
    from office_agent.database.models.task import Task
    from office_agent.database.repository.task_repo import TaskRepository

    engine = create_engine(f"sqlite:///{tmp_path / 'tasks.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(Task(id="t1", task_type="general", instruction="x", status="running"))
        session.commit()
        repo = TaskRepository(session)
        assert repo.complete_task("t1") is True
        session.commit()
        assert repo.fail_task("t1", "late timeout") is False
        session.commit()
        session.expire_all()
        assert repo.get_by_id("t1").status == "success"
