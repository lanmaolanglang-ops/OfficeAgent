"""P1-24：Word 首行缩进必须是单一来源，不得出现"24 磅 / 2 字符"双轨。

修复前行为（旧源码副本实测）：
- ``FormatConfig.body_paragraph`` 默认 ``first_line_indent=0 磅 + chars=2``；
- ``WordService.config_from_dict({})`` 兜底 ``first_line_indent=24 磅 + chars=0``。
12pt 下视觉等价（24 磅 = 2 字符），但正文字号≠12pt 时同一份配置在两条入口下
缩进不同（14pt：一条 28 磅、一条 24 磅）。

canonical 规则（项目内已存在两处独立实现，本批把入口对齐到它）：
1. ``_apply_paragraph_format``：``first_line_indent_chars`` 非 0 时按字符换算
   ``chars × 当前字号``；否则用绝对磅值。
2. ``quality/checker._check_indent``：同一口径用于判定文档是否缩进正确。

断言落到**最终 docx 段落的真实缩进值**，而不是只看 config 对象。
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.shared import Pt

from office_agent.models.schemas import (
    DEFAULT_BODY_FIRST_LINE_INDENT_CHARS,
    FontConfig,
    FormatConfig,
    ParagraphConfig,
)
from office_agent.parsers.format_parser import parse_format_rule
from office_agent.services.word_service import WordService

BODY_TEXT = "这是一段足够长的正文段落，用于核对首行缩进在最终文档里的真实取值。"


def _config(**overrides):
    return WordService().config_from_dict(dict(overrides))


def _indent_points(config: FormatConfig, font_size: float | None = None) -> float:
    """按生产代码同一条路径算出缩进磅值（不落盘，用于口径对齐断言）。"""
    paragraph = Document().add_paragraph("正文")
    font = config.body_font
    if font_size is not None:
        font = FontConfig(**{**font.__dict__, "size": font_size})
    WordService()._apply_paragraph_format(paragraph, font, config.body_paragraph)
    indent = paragraph.paragraph_format.first_line_indent
    return 0.0 if indent is None else indent.pt


def _written_indent(tmp_path, config_dict: dict, *, font_size: float | None = None) -> float:
    """真正跑一次 WordService.process，读出产物 docx 里的首行缩进磅值。"""
    source = tmp_path / "in.txt"
    source.write_text(f"# 标题\n{BODY_TEXT}\n", encoding="utf-8")
    output = tmp_path / "out.docx"
    payload = dict(config_dict)
    if font_size is not None:
        payload["size"] = font_size

    result = WordService().process(str(source), config_dict=payload, output_path=str(output))
    assert result.success, result.message

    document = Document(str(output))
    target = next(p for p in document.paragraphs if BODY_TEXT[:8] in p.text)
    indent = target.paragraph_format.first_line_indent
    return 0.0 if indent is None else indent.pt


class TestSingleSourceDefault:
    def test_default_field_value_matches_canonical_constant(self):
        assert ParagraphConfig().first_line_indent == 0.0
        assert DEFAULT_BODY_FIRST_LINE_INDENT_CHARS == 2.0

    def test_format_config_default_uses_canonical_constant(self):
        config = FormatConfig()
        assert config.body_paragraph.first_line_indent_chars == \
            DEFAULT_BODY_FIRST_LINE_INDENT_CHARS
        assert config.body_paragraph.first_line_indent == 0.0

    def test_config_from_dict_default_matches_format_config_default(self):
        """核心回归：两条入口在"用户未给缩进"时必须得到同一个缩进口径。"""
        from_dict = _config()
        default = FormatConfig()
        assert from_dict.body_paragraph.first_line_indent_chars == \
            default.body_paragraph.first_line_indent_chars
        assert from_dict.body_paragraph.first_line_indent == \
            default.body_paragraph.first_line_indent
        # 修复前这里会是 (chars=0, points=24)
        assert from_dict.body_paragraph.first_line_indent == 0.0

    @pytest.mark.parametrize("font_size", [10.5, 12.0, 14.0, 16.0])
    def test_two_entries_agree_at_every_font_size(self, font_size):
        a = _indent_points(_config(), font_size)
        b = _indent_points(FormatConfig(), font_size)
        assert a == pytest.approx(b)
        assert a == pytest.approx(DEFAULT_BODY_FIRST_LINE_INDENT_CHARS * font_size)


class TestIndentUnitResolution:
    def test_chars_only(self):
        config = _config(first_line_indent_chars=3)
        assert config.body_paragraph.first_line_indent_chars == 3
        assert config.body_paragraph.first_line_indent == 0
        assert _indent_points(config, 14.0) == pytest.approx(42.0)

    def test_points_only_is_honored_not_overridden_by_default_chars(self):
        config = _config(first_line_indent=30)
        assert config.body_paragraph.first_line_indent == 30
        # 显式磅值时字符分量必须为 0，否则默认的 2 字符会在 chars 优先规则里抢走它
        assert config.body_paragraph.first_line_indent_chars == 0
        assert _indent_points(config, 14.0) == pytest.approx(30.0)

    def test_both_given_chars_wins(self):
        config = _config(first_line_indent_chars=2, first_line_indent=30)
        assert _indent_points(config, 14.0) == pytest.approx(28.0)

    def test_neither_given_uses_canonical_two_chars(self):
        config = _config()
        assert _indent_points(config, 12.0) == pytest.approx(24.0)
        assert _indent_points(config, 16.0) == pytest.approx(32.0)

    def test_explicit_zero_chars_disables_indent(self):
        config = _config(first_line_indent_chars=0)
        assert config.body_paragraph.first_line_indent_chars == 0
        assert _indent_points(config, 12.0) == 0.0

    def test_direct_construction_still_supported(self):
        paragraph = Document().add_paragraph("正文")
        WordService()._apply_paragraph_format(
            paragraph, FontConfig(size=15), ParagraphConfig(first_line_indent_chars=2),
        )
        assert paragraph.paragraph_format.first_line_indent.pt == pytest.approx(30.0)

    def test_negative_value_has_no_extra_validation_on_either_entry(self):
        """项目当前对缩进无范围校验：负数在两条入口下都原样透传（不新增策略）。"""
        from_dict = _config(first_line_indent_chars=-2)
        direct = ParagraphConfig(first_line_indent_chars=-2)
        assert from_dict.body_paragraph.first_line_indent_chars == -2
        assert _indent_points(from_dict, 12.0) == pytest.approx(
            _indent_points(
                FormatConfig(body_paragraph=direct), 12.0,
            )
        )

    def test_string_and_none_values_are_coerced(self):
        assert _config(first_line_indent_chars="2.5").body_paragraph.first_line_indent_chars == 2.5
        assert _config(first_line_indent_chars=None).body_paragraph.first_line_indent_chars == 0
        assert _config(first_line_indent="24pt").body_paragraph.first_line_indent == 24.0


class TestWrittenDocumentIndent:
    def test_default_path_writes_canonical_indent(self, tmp_path):
        assert _written_indent(tmp_path, {}, font_size=12.0) == pytest.approx(24.0)

    def test_non_default_font_size_scales_the_indent(self, tmp_path):
        """14pt 正文 + 未指定缩进 → 28 磅（2 字符）；旧实现恒为 24 磅。"""
        assert _written_indent(tmp_path, {}, font_size=14.0) == pytest.approx(28.0)

    def test_explicit_points_survive_to_the_document(self, tmp_path):
        assert _written_indent(
            tmp_path, {"first_line_indent": 30}, font_size=14.0,
        ) == pytest.approx(30.0)

    def test_explicit_chars_survive_to_the_document(self, tmp_path):
        assert _written_indent(
            tmp_path, {"first_line_indent_chars": 3}, font_size=14.0,
        ) == pytest.approx(42.0)

    def test_parser_to_document_indent_is_single_sourced(self, tmp_path):
        """自然语言"首行缩进2字符" → parser → config → 最终 docx 缩进。"""
        parsed = parse_format_rule("正文宋体小四，首行缩进2字符")
        assert parsed["first_line_indent_chars"] == 2
        assert "first_line_indent" not in parsed
        assert _written_indent(tmp_path, parsed) == pytest.approx(24.0)

    def test_parser_point_indent_still_uses_points(self, tmp_path):
        parsed = parse_format_rule("正文宋体小四，首行缩进18磅")
        assert parsed["first_line_indent"] == 18
        assert "first_line_indent_chars" not in parsed
        assert _written_indent(tmp_path, parsed) == pytest.approx(18.0)


class TestQualityCheckerSharesTheRule:
    def test_checker_uses_the_same_chars_first_conversion(self):
        """质量检查器的"期望缩进"必须与渲染口径一致（防止再次漂移）。"""
        config = FormatConfig()
        expected = (
            config.body_paragraph.first_line_indent_chars * config.body_font.size
            if config.body_paragraph.first_line_indent_chars
            else config.body_paragraph.first_line_indent
        )
        assert expected == pytest.approx(
            DEFAULT_BODY_FIRST_LINE_INDENT_CHARS * config.body_font.size,
        )
        assert _indent_points(config) == pytest.approx(expected)

    def test_checker_accepts_written_document(self, tmp_path):
        from office_agent.quality.checker import QualityChecker

        source = tmp_path / "in.txt"
        source.write_text(f"# 标题\n{BODY_TEXT}\n", encoding="utf-8")
        output = tmp_path / "out.docx"
        service = WordService()
        config = service.config_from_dict({"size": 14})
        result = service.process(str(source), config, str(output))
        assert result.success, result.message

        report = QualityChecker().check(str(output), config, str(source))
        indent_issues = [i for i in report.issues if i.type == "indent"]
        assert indent_issues == [], [i.message for i in indent_issues]
