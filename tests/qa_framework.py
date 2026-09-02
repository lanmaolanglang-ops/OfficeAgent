"""
OfficeAgent QA Test Framework
本地质量保证测试框架
"""
import os
import sys
import time
import json
import tempfile
import shutil
from pathlib import Path
from dataclasses import dataclass, field
from enum import Enum
from typing import Optional, Callable, Any

# 确保项目路径
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


class TestCategory(str, Enum):
    __test__ = False
    UNIT = "unit"
    INTEGRATION = "integration"
    AGENT = "agent"
    OFFICE_COMPAT = "office_compatibility"
    RUNTIME = "runtime"
    PERFORMANCE = "performance"
    RECOVERY = "recovery"


class TestSeverity(str, Enum):
    __test__ = False
    CRITICAL = "critical"  # 阻塞发布
    HIGH = "high"          # 必须修复
    MEDIUM = "medium"      # 建议修复
    LOW = "low"            # 可延后


@dataclass
class TestResult:
    __test__ = False
    """测试结果"""
    test_name: str
    category: TestCategory
    severity: TestSeverity
    passed: bool
    duration_ms: int = 0
    score: float = 0.0  # 0-100
    message: str = ""
    details: dict = field(default_factory=dict)
    error: str = ""
    timestamp: float = field(default_factory=time.time)

    def to_dict(self) -> dict:
        return {
            "test_name": self.test_name,
            "category": self.category.value,
            "severity": self.severity.value,
            "passed": self.passed,
            "duration_ms": self.duration_ms,
            "score": self.score,
            "message": self.message,
            "details": self.details,
            "error": self.error,
            "timestamp": self.timestamp,
        }


@dataclass
class PerformanceMetrics:
    """性能指标"""
    test_name: str
    duration_ms: int = 0
    peak_memory_mb: float = 0
    avg_cpu_percent: float = 0
    file_size_kb: float = 0
    operations_per_second: float = 0

    def to_dict(self) -> dict:
        return {
            "test_name": self.test_name,
            "duration_ms": self.duration_ms,
            "peak_memory_mb": round(self.peak_memory_mb, 2),
            "avg_cpu_percent": round(self.avg_cpu_percent, 2),
            "file_size_kb": round(self.file_size_kb, 2),
            "operations_per_second": round(self.operations_per_second, 2),
        }


class QATestRunner:
    """QA测试运行器"""

    def __init__(self):
        self.results: list[TestResult] = []
        self.metrics: list[PerformanceMetrics] = []
        self._start_time = 0
        self._end_time = 0

    def run_test(self, name: str, category: TestCategory, severity: TestSeverity,
                 func: Callable, *args, **kwargs) -> TestResult:
        start = time.time()
        result = TestResult(
            test_name=name,
            category=category,
            severity=severity,
            passed=False,
        )
        try:
            ret = func(*args, **kwargs)
            if isinstance(ret, tuple):
                result.passed = ret[0]
                result.message = str(ret[1]) if len(ret) > 1 else ""
                if len(ret) > 2 and isinstance(ret[2], dict):
                    result.details = ret[2]
                if len(ret) > 3:
                    result.score = float(ret[3])
            elif isinstance(ret, bool):
                result.passed = ret
            elif isinstance(ret, dict):
                result.passed = ret.get("passed", True)
                result.message = ret.get("message", "")
                result.details = ret.get("details", {})
                result.score = ret.get("score", 100.0 if result.passed else 0.0)
            else:
                result.passed = True
        except Exception as e:
            result.passed = False
            result.error = f"{type(e).__name__}: {e}"
            import traceback
            result.details["traceback"] = traceback.format_exc()
        result.duration_ms = int((time.time() - start) * 1000)
        if not result.score:
            result.score = 100.0 if result.passed else 0.0
        self.results.append(result)
        return result

    def add_metric(self, metric: PerformanceMetrics):
        self.metrics.append(metric)

    def get_summary(self) -> dict:
        total = len(self.results)
        passed = sum(1 for r in self.results if r.passed)
        failed = total - passed
        by_category = {}
        by_severity = {}
        for r in self.results:
            cat = r.category.value
            sev = r.severity.value
            by_category.setdefault(cat, {"total": 0, "passed": 0, "failed": 0})
            by_category[cat]["total"] += 1
            if r.passed:
                by_category[cat]["passed"] += 1
            else:
                by_category[cat]["failed"] += 1
            by_severity.setdefault(sev, {"total": 0, "passed": 0, "failed": 0})
            by_severity[sev]["total"] += 1
            if r.passed:
                by_severity[sev]["passed"] += 1
            else:
                by_severity[sev]["failed"] += 1
        critical_failed = by_severity.get("critical", {}).get("failed", 0)
        avg_score = sum(r.score for r in self.results) / total if total else 0
        return {
            "total": total,
            "passed": passed,
            "failed": failed,
            "pass_rate": round(passed / total * 100, 1) if total else 0,
            "avg_score": round(avg_score, 1),
            "critical_failed": critical_failed,
            "by_category": by_category,
            "by_severity": by_severity,
            "total_duration_ms": int((self._end_time - self._start_time) * 1000) if self._end_time else 0,
            "release_blocked": critical_failed > 0,
        }

    def start(self):
        self._start_time = time.time()

    def end(self):
        self._end_time = time.time()

    def get_failed_tests(self) -> list[TestResult]:
        return [r for r in self.results if not r.passed]


def get_test_dataset_dir() -> Path:
    return PROJECT_ROOT / "test_dataset"


def get_temp_output_dir() -> Path:
    d = Path(tempfile.mkdtemp(prefix="officeagent_qa_"))
    return d


def create_test_word(path: Path, content_type: str = "simple") -> Path:
    """创建测试Word文件"""
    from docx import Document
    doc = Document()
    if content_type == "simple":
        doc.add_heading("测试文档", 0)
        doc.add_paragraph("这是一个简单的测试文档。")
    elif content_type == "academic":
        doc.add_heading("学术论文测试", 0)
        for i in range(1, 4):
            doc.add_heading(f"第{i}章 标题", level=1)
            doc.add_paragraph(f"这是第{i}章的内容。" * 20)
    elif content_type == "business":
        doc.add_heading("企业报告", 0)
        doc.add_paragraph("公司年度报告内容。")
        table = doc.add_table(rows=5, cols=4)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = f"R{i}C{j}"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def create_test_excel(path: Path, rows: int = 100, cols: int = 10) -> Path:
    """创建测试Excel文件"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "测试数据"
    headers = [f"列{i+1}" for i in range(cols)]
    ws.append(headers)
    for r in range(rows):
        row_data = [r * cols + c + 1 for c in range(cols)]
        ws.append(row_data)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def create_test_ppt(path: Path, slides: int = 10) -> Path:
    """创建测试PPT文件"""
    from pptx import Presentation
    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"幻灯片 {i+1}"
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = f"内容 {i+1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def measure_performance(func: Callable, *args, **kwargs) -> tuple[Any, PerformanceMetrics]:
    """测量性能"""
    import tracemalloc
    tracemalloc.start()
    start = time.time()
    result = func(*args, **kwargs)
    duration = time.time() - start
    current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    metrics = PerformanceMetrics(
        test_name=getattr(func, '__name__', 'unknown'),
        duration_ms=int(duration * 1000),
        peak_memory_mb=peak / (1024 * 1024),
    )
    return result, metrics
