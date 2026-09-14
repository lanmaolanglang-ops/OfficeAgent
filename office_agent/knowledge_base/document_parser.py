"""
文档解析器 - 解析各种格式的规范文档

支持：docx, pptx, xlsx, txt, md, pdf
输出：纯文本 + 章节结构
"""
import os
import re
from typing import List, Dict, Tuple, Any
from dataclasses import dataclass, field

from ..text_encoding import read_text_file


@dataclass
class ParsedSection:
    """解析出的章节"""
    title: str = ""
    level: int = 0           # 标题级别（1=一级标题）
    content: str = ""
    page_number: int = 0


@dataclass
class ParsedDocument:
    """解析后的文档"""
    title: str = ""
    doc_type: str = "general"
    full_text: str = ""
    sections: List[ParsedSection] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    raw_blocks: List[Dict[str, Any]] = field(default_factory=list)

    def get_section_texts(self) -> List[Tuple[str, str]]:
        """返回 (标题, 内容) 列表"""
        return [(s.title, s.content) for s in self.sections if s.content.strip()]


class DocumentParser:
    """文档解析器"""

    # 文件扩展名到类型的映射
    EXT_MAP = {
        ".docx": "word",
        ".pptx": "ppt",
        ".xlsx": "excel",
        ".xls": "excel",
        ".csv": "csv",
        ".txt": "text",
        ".md": "markdown",
        ".pdf": "pdf",
    }

    def parse(self, file_path: str, doc_type: str = "") -> ParsedDocument:
        """解析文档"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        file_type = self.EXT_MAP.get(ext, "text")

        if file_type == "word":
            return self._parse_docx(file_path, doc_type)
        elif file_type == "ppt":
            return self._parse_pptx(file_path, doc_type)
        elif file_type == "excel":
            return self._parse_excel(file_path, doc_type)
        elif file_type == "csv":
            return self._parse_csv(file_path, doc_type)
        elif file_type == "pdf":
            return self._parse_pdf(file_path, doc_type)
        else:
            return self._parse_text(file_path, doc_type)

    def _parse_docx(self, file_path: str, doc_type: str) -> ParsedDocument:
        """解析 Word 文档"""
        from docx import Document

        doc = Document(file_path)
        result = ParsedDocument(
            title=os.path.splitext(os.path.basename(file_path))[0],
            doc_type=doc_type or "word_spec",
        )

        current_section = ParsedSection(title="正文", level=0)
        sections = [current_section]
        all_text_parts = []

        # P3-79: 按 body 真实顺序交错处理段落与表格。原先先遍历全部段落、
        # 再遍历 doc.tables，会把所有表格都塞进最后一个 section，丢失归属。
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        def _iter_body_blocks(document):
            for child in document.element.body.iterchildren():
                if child.tag == qn("w:p"):
                    yield "p", Paragraph(child, document)
                elif child.tag == qn("w:tbl"):
                    yield "t", Table(child, document)

        for kind, block in _iter_body_blocks(doc):
            if kind == "t":
                table_text_parts = []
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    table_text_parts.append(" | ".join(cells))
                table_text = "\n".join(table_text_parts)
                current_section.content += "\n[表格]\n" + table_text + "\n"
                all_text_parts.append(table_text)
                continue

            para = block
            text = para.text.strip()
            if not text:
                continue

            all_text_parts.append(text)

            # 判断是否为标题
            style = para.style
            style_name = (style.name or "").lower() if style is not None else ""
            is_heading = False
            heading_level = 0

            if style_name.startswith("heading") or style_name.startswith("标题"):
                is_heading = True
                # 提取级别数字
                level_match = re.search(r'(\d+)', style_name)
                heading_level = int(level_match.group(1)) if level_match else 1
            elif style_name == "title" or style_name == "标题":
                is_heading = True
                heading_level = 1
            elif re.match(r'^[一二三四五六七八九十]+[、.．]', text):
                is_heading = True
                heading_level = 1
            elif (re.match(r'^\d+[、.．]\s*\S', text) and len(text) < 50
                  and not text.rstrip().endswith(
                      ('。', '！', '？', '；', '，', '.', '!', '?', ';', ','))):
                # 以句读结尾的短编号文本是完整句子而非标题，避免误判（P5-4）
                is_heading = True
                heading_level = 2
            elif re.match(r'^\d+\.\d+[、.．\s]', text) and len(text) < 60:
                is_heading = True
                heading_level = 3

            if is_heading:
                current_section = ParsedSection(
                    title=text, level=heading_level, content=""
                )
                sections.append(current_section)
            else:
                current_section.content += text + "\n"

        result.sections = sections
        result.full_text = "\n".join(all_text_parts)
        result.metadata["paragraph_count"] = len(doc.paragraphs)
        result.metadata["table_count"] = len(doc.tables)

        return result

    def _parse_pptx(self, file_path: str, doc_type: str) -> ParsedDocument:
        """解析 PPT 文档"""
        from pptx import Presentation

        prs = Presentation(file_path)
        result = ParsedDocument(
            title=os.path.splitext(os.path.basename(file_path))[0],
            doc_type=doc_type or "ppt_spec",
        )

        sections = []
        all_text_parts = []

        for slide_idx, slide in enumerate(prs.slides, 1):
            slide_texts = []
            slide_title = ""

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
                            if not slide_title and len(text) < 50:
                                slide_title = text

                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        slide_texts.append(" | ".join(cells))

            if slide_texts:
                section = ParsedSection(
                    title=slide_title or f"第{slide_idx}页",
                    level=1,
                    content="\n".join(slide_texts),
                    page_number=slide_idx,
                )
                sections.append(section)
                all_text_parts.extend(slide_texts)

        result.sections = sections
        result.full_text = "\n".join(all_text_parts)
        result.metadata["slide_count"] = len(prs.slides)

        return result

    def _parse_csv(self, file_path: str, doc_type: str) -> ParsedDocument:
        """CSV 不能走 openpyxl（那是 OOXML），用标准库解析。"""
        import csv

        result = ParsedDocument(
            title=os.path.splitext(os.path.basename(file_path))[0],
            doc_type=doc_type or "csv_data",
        )
        lines: list[str] = []
        with open(file_path, "r", encoding="utf-8-sig", newline="", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                cells = [c.strip() for c in row]
                if any(cells):
                    lines.append(" | ".join(cells))
        result.sections = [
            ParsedSection(
                title=os.path.basename(file_path),
                level=1,
                content="\n".join(lines),
            )
        ] if lines else []
        result.full_text = "\n".join(lines)
        result.metadata["row_count"] = len(lines)
        return result

    def _parse_excel(self, file_path: str, doc_type: str) -> ParsedDocument:
        """解析 Excel 文档"""
        from openpyxl import load_workbook

        wb = load_workbook(file_path, data_only=True, read_only=True)
        result = ParsedDocument(
            title=os.path.splitext(os.path.basename(file_path))[0],
            doc_type=doc_type or "excel_rule",
        )

        sections = []
        all_text_parts = []
        sheet_names = list(wb.sheetnames)  # close 前固化，避免关闭后再读（P5-37）

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            sheet_texts = []

            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                # P3-80: 保留中间空单元格以维持列对齐（与 CSV 路径一致），
                # 只裁掉尾部连续空单元格；过滤非空会让后面的列错位左移。
                while cells and not cells[-1].strip():
                    cells.pop()
                if any(c.strip() for c in cells):
                    line = " | ".join(cells)
                    sheet_texts.append(line)

            if sheet_texts:
                section = ParsedSection(
                    title=sheet_name,
                    level=1,
                    content="\n".join(sheet_texts),
                )
                sections.append(section)
                all_text_parts.append(f"## {sheet_name}")
                all_text_parts.extend(sheet_texts)

        wb.close()

        result.sections = sections
        result.full_text = "\n".join(all_text_parts)
        result.metadata["sheet_count"] = len(sheet_names)

        return result

    def _parse_pdf(self, file_path: str, doc_type: str) -> ParsedDocument:
        """解析 PDF 文档。

        PDF 解析只有 PyMuPDF 一个 backend，二进制 PDF 绝不能当文本文件
        "降级"读取——那会把乱码当成解析成功。缺库时抛出带安装指引的
        明确错误；损坏/加密等打开失败由 PyMuPDF 自己的异常如实上报。
        """
        try:
            import pymupdf
        except ImportError as exc:
            raise RuntimeError(
                "PDF 解析需要 PyMuPDF 库，请安装依赖：pip install pymupdf"
            ) from exc
        doc = pymupdf.open(file_path)

        result = ParsedDocument(
            title=os.path.splitext(os.path.basename(file_path))[0],
            doc_type=doc_type or "general",
        )

        sections = []
        all_text_parts = []

        for page_num, page in enumerate(doc.pages(), 1):
            text = page.get_text().strip()
            if text:
                # 尝试识别标题（第一行较短的文本）
                lines = text.split("\n")
                title = lines[0].strip() if lines and len(lines[0].strip()) < 80 else f"第{page_num}页"

                section = ParsedSection(
                    title=title,
                    level=1,
                    content=text,
                    page_number=page_num,
                )
                sections.append(section)
                all_text_parts.append(text)

        doc.close()

        result.sections = sections
        result.full_text = "\n".join(all_text_parts)
        result.metadata["page_count"] = len(sections)

        return result

    def _parse_text(self, file_path: str, doc_type: str) -> ParsedDocument:
        """解析纯文本/Markdown（编码探测统一走 text_encoding 共享入口）"""
        content = read_text_file(file_path)

        result = ParsedDocument(
            title=os.path.splitext(os.path.basename(file_path))[0],
            doc_type=doc_type or "general",
            full_text=content,
        )

        # Markdown 标题识别
        sections = []
        current = ParsedSection(title="正文", level=0)
        sections.append(current)

        for line in content.split("\n"):
            line_stripped = line.strip()

            # Markdown 标题
            md_match = re.match(r'^(#{1,6})\s+(.+)$', line_stripped)
            if md_match:
                level = len(md_match.group(1))
                title = md_match.group(2).strip()
                current = ParsedSection(title=title, level=level, content="")
                sections.append(current)
            else:
                current.content += line + "\n"

        result.sections = sections

        return result

    def parse_text(self, text: str, title: str = "直接输入",
                   doc_type: str = "general") -> ParsedDocument:
        """直接解析文本内容"""
        result = ParsedDocument(title=title, doc_type=doc_type, full_text=text)

        sections = []
        current = ParsedSection(title="正文", level=0)
        sections.append(current)

        for line in text.split("\n"):
            line_stripped = line.strip()
            md_match = re.match(r'^(#{1,6})\s+(.+)$', line_stripped)
            if md_match:
                level = len(md_match.group(1))
                title_text = md_match.group(2).strip()
                current = ParsedSection(title=title_text, level=level, content="")
                sections.append(current)
            else:
                current.content += line + "\n"

        result.sections = sections
        return result
