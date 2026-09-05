"""
Word Service - Word 文档自动排版引擎
使用 python-docx + lxml 实现完整的 Word 文档处理

功能：
1. 文件读取（docx / txt）
2. 文档结构分析（标题层级、正文、表格）
3. 格式处理（字体、字号、加粗、斜体、对齐、行距、段间距、缩进）
4. 标题自动编号（1, 1.1, 1.1.1, 1.1.1.1）
5. 三线表处理（1.5/0.75/1.5磅）+ 自动表编号
6. 自定义规则接口 FormatConfig
"""
import re
import logging
from pathlib import Path
from typing import Optional

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

from ..models.schemas import (
    FormatConfig, FontConfig, ParagraphConfig, HeadingConfig,
    TableConfig, Alignment, ProcessResult, PageSetupConfig,
)
from ..parsers.document_structure import DocumentStructureAnalyzer, DocumentTree
from ..text_encoding import read_text_file

logger = logging.getLogger("office_agent.services.word")


# ============== 常量映射 ==============

ALIGNMENT_MAP = {
    Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}

ALIGNMENT_STR_MAP = {
    "left": Alignment.LEFT,
    "right": Alignment.RIGHT,
    "center": Alignment.CENTER,
    "justify": Alignment.JUSTIFY,
    "左对齐": Alignment.LEFT,
    "右对齐": Alignment.RIGHT,
    "居中": Alignment.CENTER,
    "居中对齐": Alignment.CENTER,
    "两端对齐": Alignment.JUSTIFY,
}

# 中文字号 -> 磅值
CHINESE_SIZE_MAP = {
    "初号": 42, "小初": 36,
    "一号": 26, "小一": 24,
    "二号": 22, "小二": 18,
    "三号": 16, "小三": 15,
    "四号": 14, "小四": 12,
    "五号": 10.5, "小五": 9,
    "六号": 7.5, "小六": 6.5,
    "七号": 5.5, "八号": 5,
}


# ============== 文档结构分析结果 ==============

class DocumentStructure:
    """文档结构分析结果"""
    def __init__(self):
        self.paragraphs_count = 0
        self.headings: dict[int, list] = {1: [], 2: [], 3: [], 4: []}
        self.tables_count = 0
        self.body_paragraphs = 0
        self.detected_levels: dict[int, int] = {}  # 段落索引 -> 检测到的级别

    def summary(self) -> str:
        parts = []
        for level in range(1, 5):
            count = len(self.headings[level])
            if count > 0:
                parts.append(f"{level}级标题{count}个")
        parts.append(f"正文段落{self.body_paragraphs}个")
        parts.append(f"表格{self.tables_count}个")
        return "，".join(parts)


# ============== WordService ==============

class WordService:
    """Word 文档自动排版引擎"""

    def __init__(self):
        self.changes: list[str] = []
        self.structure = DocumentStructure()
        self.document_tree: Optional[DocumentTree] = None
        self._table_counter = 0

    # ==========================================
    # 1. 文件读取
    # ==========================================

    def read_file(self, file_path: str) -> DocumentType:
        """
        读取文件，支持 .docx 和 .txt

        Args:
            file_path: 文件路径

        Returns:
            python-docx Document 对象（txt 会转换为 docx）
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()

        if ext == ".docx":
            return Document(str(path))

        elif ext == ".txt":
            return self._txt_to_docx(str(path))

        else:
            raise ValueError(f"不支持的文件格式: {ext}，仅支持 .docx 和 .txt")

    def _txt_to_docx(self, txt_path: str) -> DocumentType:
        """将 txt 文件转换为 Document 对象，支持 Markdown 标题和表格"""
        # 编码探测统一走 text_encoding 共享入口（失败抛 TextDecodeError，属 ValueError）
        content = read_text_file(txt_path)

        doc = Document()
        lines = content.split("\n")

        i = 0
        while i < len(lines):
            line = lines[i].rstrip("\r")
            stripped = line.strip()

            if not stripped:
                doc.add_paragraph("")
                i += 1
                continue

            # Markdown 表格识别：连续的 | 开头行
            if stripped.startswith("|") and stripped.endswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                # 解析表格（跳过分隔行 |---|---|）
                rows_data = []
                for tl in table_lines:
                    if re.match(r"^\|[\s\-:|]+\|$", tl):
                        continue  # 跳过分隔行
                    cells = [c.strip() for c in tl.strip("|").split("|")]
                    rows_data.append(cells)
                if rows_data:
                    ncols = max(len(r) for r in rows_data)
                    table = doc.add_table(rows=len(rows_data), cols=ncols)
                    table.style = "Table Grid"
                    for r_idx, row_data in enumerate(rows_data):
                        for c_idx, cell_text in enumerate(row_data):
                            if c_idx < ncols:
                                table.rows[r_idx].cells[c_idx].text = cell_text
                continue

            # Markdown 标题
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith("#### "):
                doc.add_heading(line[5:].strip(), level=4)
            else:
                doc.add_paragraph(line)

            i += 1

        return doc

    # ==========================================
    # 2. 文档结构分析
    # ==========================================

    def analyze_structure(self, doc: DocumentType) -> DocumentStructure:
        """
        分析文档结构

        使用 DocumentStructureAnalyzer 进行增强识别：
        - 中文/数字/英文标题
        - 列表项
        - 引用块
        - 语义关键词（摘要/参考文献等）

        同时填充旧的 DocumentStructure 以保持向后兼容。
        """
        self.structure = DocumentStructure()
        self.structure.paragraphs_count = len(doc.paragraphs)
        self.structure.tables_count = len(doc.tables)

        # 使用新的增强分析器
        analyzer = DocumentStructureAnalyzer()
        self.document_tree = analyzer.analyze(doc)

        # 将新分析结果转换为旧格式（向后兼容）
        for node in self.document_tree.nodes:
            if node.type in (
                "chapter", "section", "subsection", "subsubsection",
                "abstract", "reference", "appendix", "keywords",
            ) and 1 <= node.level <= 4:
                level = node.level
                self.structure.headings[level].append((node.index, node.text))
                self.structure.detected_levels[node.index] = level
            elif node.type == "title":
                # 文档标题：跳过正文格式，但不参与标题编号（level 0）
                self.structure.detected_levels[node.index] = 0
            elif node.type == "paragraph":
                self.structure.body_paragraphs += 1

        return self.structure

    # ==========================================
    # 3. 格式处理引擎
    # ==========================================

    def process(self, input_path: str,
                format_config: Optional[FormatConfig] = None,
                output_path: Optional[str] = None,
                config_dict: Optional[dict] = None) -> ProcessResult:
        """
        处理 Word 文档（主入口）

        Args:
            input_path: 输入文件路径
            format_config: FormatConfig 对象
            output_path: 输出路径
            config_dict: 字典格式配置（方便 Agent 直接传 dict）
        """
        self.changes = []
        self._table_counter = 0

        # 支持字典配置
        if config_dict and not format_config:
            format_config = self.config_from_dict(config_dict)
        if not format_config:
            format_config = FormatConfig()

        try:
            # 读取文件
            doc = self.read_file(input_path)
            is_txt = Path(input_path).suffix.lower() == ".txt"

            # 分析结构
            self.analyze_structure(doc)
            self.changes.append(f"文档结构: {self.structure.summary()}")

            # 应用正文格式
            self._apply_body_format(doc, format_config)

            # 应用标题格式 + 自动编号
            self._apply_heading_format(doc, format_config)

            # 处理表格（三线表 + 编号）
            if format_config.table_config:
                self._process_tables(doc, format_config.table_config)

            # 页面设置
            self._apply_page_setup(doc, format_config)

            # 生成输出路径
            if not output_path:
                p = Path(input_path)
                suffix = "_formatted.docx"
                output_path = str(p.with_name(p.stem + suffix))

            doc.save(output_path)

            # 质量检查 + 自动修复
            quality_report = None
            try:
                from ..quality.checker import QualityChecker
                checker = QualityChecker()
                quality_report = checker.check(output_path, format_config, input_path)

                # 如果有可修复的问题，自动修复一次
                if not quality_report.passed and quality_report.fixable_issues():
                    fix_config = quality_report.get_fix_config()
                    if fix_config:
                        self.changes.append(
                            f"质量检查发现{len(quality_report.errors())}个错误，"
                            f"{len(quality_report.warnings())}个警告，自动修复中..."
                        )
                        # 合并修复配置
                        merged = self._merge_config(format_config, fix_config)
                        doc2 = self.read_file(input_path)
                        self.analyze_structure(doc2)
                        self._apply_body_format(doc2, merged)
                        self._apply_heading_format(doc2, merged)
                        self._process_tables(doc2, merged.table_config)
                        self._apply_page_setup(doc2, merged)
                        doc2.save(output_path)
                        # 重新检查
                        quality_report = checker.check(output_path, merged, input_path)
            except Exception as exc:
                # 质量检查失败不影响主流程，但必须留痕便于排查
                self.changes.append(f"质量检查跳过: {type(exc).__name__}")
                logger.debug("Word quality check skipped: %s", exc)

            suggestions = [
                "建议检查标题层级是否正确",
                "建议核对表格编号连续性",
                "可根据需要调整首行缩进",
            ]
            if quality_report and not quality_report.passed:
                suggestions.append(
                    f"质量检查发现{len(quality_report.errors())}个错误，"
                    f"{len(quality_report.warnings())}个警告"
                )

            return ProcessResult(
                success=True,
                message=f"文档排版完成: {Path(output_path).name}",
                output_path=output_path,
                changes=self.changes,
                suggestions=suggestions,
                metadata={
                    "input_format": "txt" if is_txt else "docx",
                    "structure": self.structure.summary(),
                    "quality_score": quality_report.score if quality_report else None,
                    "quality_passed": quality_report.passed if quality_report else None,
                    "quality_issues": len(quality_report.issues) if quality_report else 0,
                }
            )

        except Exception as e:
            return ProcessResult(
                success=False,
                message=f"排版失败: {str(e)}",
                changes=self.changes,
            )

    def _apply_body_format(self, doc: DocumentType, config: FormatConfig):
        """应用正文格式"""
        font = config.body_font
        para_cfg = config.body_paragraph
        count = 0

        for idx, para in enumerate(doc.paragraphs):
            # 跳过标题
            if idx in self.structure.detected_levels:
                continue
            # 跳过空段落
            if not para.text.strip():
                continue

            self._apply_paragraph_format(para, font, para_cfg)
            count += 1

        if count > 0:
            self.changes.append(
                f"正文格式: {font.cn_font}/{font.size}pt/"
                f"{para_cfg.line_spacing}倍行距，共{count}段"
            )

    def _apply_heading_format(self, doc: DocumentType, config: FormatConfig):
        """应用标题格式（含自动编号）

        按文档顺序遍历段落，遇到标题时：
        1. 更新对应级别的计数器
        2. 重置所有下级计数器
        3. 应用格式和编号
        """
        # 编号计数器 [c1, c2, c3, c4] 对应 1-4 级
        counters = [0, 0, 0, 0]
        level_counts = {1: 0, 2: 0, 3: 0, 4: 0}

        # 按文档顺序遍历所有段落
        for idx, para in enumerate(doc.paragraphs):
            level = self.structure.detected_levels.get(idx)
            if not level or level > 4:
                continue

            heading_cfg = config.headings.get(level)
            if not heading_cfg:
                continue

            # 更新编号计数器
            counters[level - 1] += 1
            for i in range(level, 4):
                counters[i] = 0

            level_counts[level] += 1

            # 应用格式
            self._apply_paragraph_format(
                para, heading_cfg.font, heading_cfg.paragraph
            )

            # 自动编号
            if heading_cfg.numbering:
                number_str = self._build_number(counters, level)
                self._prepend_number(para, number_str)

        # 记录变更
        for level in range(1, 5):
            count = level_counts[level]
            heading_cfg = config.headings.get(level)
            if count > 0 and heading_cfg:
                self.changes.append(
                    f"{level}级标题: {heading_cfg.font.cn_font}/"
                    f"{heading_cfg.font.size}pt/加粗，共{count}个"
                )

    def _build_number(self, counters: list[int], level: int) -> str:
        """
        构建编号字符串
        level=1: "1"
        level=2: "1.1"
        level=3: "1.1.1"
        level=4: "1.1.1.1"
        """
        parts = [str(counters[i]) for i in range(level)]
        return ".".join(parts)

    def _prepend_number(self, para, number_str: str):
        """在标题前添加编号，同时保留原有 run 级字体、语言和强调格式。"""
        raw_text = "".join(run.text or "" for run in para.runs)
        text = raw_text.lstrip()
        leading_chars = len(raw_text) - len(text)

        # 匹配已有编号的模式：
        # 1. 多级数字编号: 1, 1.1, 1.1.1, 1.1.1.1（后面可跟.、、空格或直接跟中文）
        # 2. 单级数字编号: 1. 1、 1 （后面跟空格或标点）
        # 3. 中文编号: 第一章、一、、（一）
        existing_num_patterns = [
            r"^\d+\.\d+\.\d+\.\d+[\.、\s]*",   # 1.1.1.1
            r"^\d+\.\d+\.\d+[\.、\s]*",         # 1.1.1
            r"^\d+\.\d+[\.、\s]*",              # 1.1（注意：必须先匹配多级，避免"1.1"被拆成"1."）
            r"^第[一二三四五六七八九十百千\d]+[章节部分篇][\s]*",  # 第一章
            r"^[一二三四五六七八九十]+[、\.][\s]*",               # 一、
            r"^（[一二三四五六七八九十]+）[\s]*",                  # （一）
            r"^\d+[、\.][\s]+",                 # 1. （必须跟空格，避免匹配"1.1"）
        ]

        remove_chars = leading_chars
        for pattern in existing_num_patterns:
            m = re.match(pattern, text)
            if m:
                remove_chars += m.end()
                break

        # 仅从 run 序列头部删掉原编号；其余 run 不重建，局部加粗、
        # 中英文字体和超链接等格式均保持原位。
        remaining = remove_chars
        for run in para.runs:
            if remaining <= 0:
                break
            text_len = len(run.text or "")
            if text_len <= remaining:
                run.text = ""
                remaining -= text_len
            else:
                run.text = run.text[remaining:].lstrip()
                remaining = 0

        target_run = next((run for run in para.runs if run.text), None)
        if target_run is None:
            target_run = para.runs[0] if para.runs else para.add_run()
        separator = " " if any(run.text for run in para.runs) else ""
        target_run.text = f"{number_str}{separator}{target_run.text}"

    def _apply_paragraph_format(self, para, font_cfg: FontConfig,
                                para_cfg: ParagraphConfig):
        """应用单个段落格式"""
        # 对齐方式
        if para_cfg.alignment:
            para.alignment = ALIGNMENT_MAP.get(
                para_cfg.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY
            )

        # 段落格式
        pf = para.paragraph_format

        if para_cfg.line_spacing:
            if para_cfg.line_spacing_rule == "exactly":
                pf.line_spacing = Pt(para_cfg.line_spacing)
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            else:
                pf.line_spacing = para_cfg.line_spacing
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

        if para_cfg.space_before is not None:
            pf.space_before = Pt(para_cfg.space_before)
        if para_cfg.space_after is not None:
            pf.space_after = Pt(para_cfg.space_after)

        if para_cfg.first_line_indent_chars:
            pf.first_line_indent = Pt(para_cfg.first_line_indent_chars * font_cfg.size)
        elif para_cfg.first_line_indent:
            pf.first_line_indent = Pt(para_cfg.first_line_indent)
        if para_cfg.hanging_indent:
            pf.hanging_indent = Pt(para_cfg.hanging_indent)

        # 字体格式（应用到所有 run）
        for run in para.runs:
            self._apply_run_font(run, font_cfg)

        # 如果段落有文字但没有 run（特殊情况）
        if para.text and not para.runs:
            run = para.add_run(para.text)
            self._apply_run_font(run, font_cfg)

    def _apply_run_font(self, run, font_cfg: FontConfig):
        """应用 run 级别的字体格式"""
        # 英文字体
        run.font.name = font_cfg.en_font
        run.font.size = Pt(font_cfg.size)

        if font_cfg.bold is not None:
            run.font.bold = font_cfg.bold
        if font_cfg.italic is not None:
            run.font.italic = font_cfg.italic
        if font_cfg.underline is not None:
            run.font.underline = font_cfg.underline

        if font_cfg.color:
            try:
                run.font.color.rgb = RGBColor.from_string(font_cfg.color)
            except ValueError:
                pass

        # 中文字体（通过 lxml 设置 eastAsia）
        self._set_chinese_font(run, font_cfg.cn_font)

    def _set_chinese_font(self, run, font_name: str):
        """设置中文字体（lxml 操作）"""
        rPr = run._element.get_or_add_rPr()

        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)

        rFonts.set(qn("w:eastAsia"), font_name)
        rFonts.set(qn("w:ascii"), run.font.name or "Times New Roman")
        rFonts.set(qn("w:hAnsi"), run.font.name or "Times New Roman")

    # ==========================================
    # 5. 表格处理（三线表）
    # ==========================================

    def _process_tables(self, doc: DocumentType, table_config: TableConfig):
        """处理所有表格：三线表 + 自动编号（已有题注则不重复添加）"""
        for table in doc.tables:
            self._table_counter += 1
            if table_config.three_line:
                self._apply_three_line_table(table, table_config)

            if table_config.auto_number:
                # 检查表格上方是否已有题注（"表N"开头的段落）
                if not self._has_existing_caption(table):
                    self._add_table_caption(table, self._table_counter)

        if self._table_counter > 0:
            self.changes.append(
                f"表格处理: {self._table_counter}个三线表，自动编号"
            )

    def _has_existing_caption(self, table) -> bool:
        """检查表格上方是否已有'表N'题注"""
        tbl_element = table._tbl
        parent = tbl_element.getparent()
        tbl_index = list(parent).index(tbl_element)

        # 查看表格前1-2个段落
        for i in range(max(0, tbl_index - 2), tbl_index):
            prev_elem = parent[i]
            if prev_elem.tag.endswith('}p'):
                # 获取段落文本
                texts = prev_elem.findall('.//' + qn('w:t'))
                para_text = ''.join(t.text or '' for t in texts).strip()
                if re.match(r'^表\d+', para_text):
                    return True
        return False

    def _apply_three_line_table(self, table, config: TableConfig):
        """
        应用三线表格式
        - 顶线: 1.5磅
        - 表头底线: 0.75磅
        - 底线: 1.5磅
        - 无竖线
        """
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 先清除所有单元格边框
        self._clear_table_borders(table)

        # 设置三线表边框
        tbl = table._tbl
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr)

        # 顶线和底线（表格级）
        top_sz = str(int(config.top_border * 8))    # 1.5磅 = 12 (八分之一磅)
        bot_sz = str(int(config.bottom_border * 8)) # 1.5磅 = 12

        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{top_sz}" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="{bot_sz}" w:space="0" w:color="000000"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
        '''
        tblBorders = parse_xml(borders_xml)
        self._set_tbl_borders(tblPr, tblBorders)

        # 表头行底线（0.75磅）
        if len(table.rows) > 0:
            header_row = table.rows[0]
            for cell in header_row.cells:
                self._set_cell_bottom_border(cell, config.middle_border)
                # 表头加粗
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True

    # tblPr 子元素按 OOXML schema 的顺序，tblBorders 之后允许出现的元素
    _TBLPR_AFTER_BORDERS = (
        "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription",
    )

    def _set_tbl_borders(self, tblPr, tblBorders):
        """替换 tblBorders：先移除已有的（重复元素会让 Word/WPS 判定文档损坏），
        再按 schema 顺序插入到 tblLayout/tblCellMar/tblLook 等后继元素之前。"""
        for existing in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(existing)
        for tag in self._TBLPR_AFTER_BORDERS:
            successor = tblPr.find(qn(tag))
            if successor is not None:
                successor.addprevious(tblBorders)
                return
        tblPr.append(tblBorders)

    def _clear_table_borders(self, table):
        """清除表格所有边框"""
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = tcPr.find(qn("w:tcBorders"))
                if tcBorders is not None:
                    tcPr.remove(tcBorders)
                # 添加无边框设置
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "left", "bottom", "right"]:
                    b = OxmlElement(f"w:{edge}")
                    b.set(qn("w:val"), "none")
                    b.set(qn("w:sz"), "0")
                    borders.append(b)
                tcPr.append(borders)

    def _set_cell_bottom_border(self, cell, size_pt: float):
        """设置单元格底边框"""
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            tcPr.append(tcBorders)

        sz = str(int(size_pt * 8))
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), sz)
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        tcBorders.append(bottom)

    def _add_table_caption(self, table, table_num: int):
        """在表格上方添加"表N"题注"""
        tbl_element = table._tbl
        parent = tbl_element.getparent()
        tbl_index = list(parent).index(tbl_element)

        caption_xml = f'''
        <w:p {nsdecls("w")}>
            <w:pPr>
                <w:jc w:val="center"/>
                <w:spacing w:before="60" w:after="60"/>
            </w:pPr>
            <w:r>
                <w:rPr>
                    <w:rFonts w:eastAsia="宋体" w:ascii="Times New Roman"/>
                    <w:sz w:val="21"/>
                    <w:b w:val="false"/>
                </w:rPr>
                <w:t>表{table_num}</w:t>
            </w:r>
        </w:p>
        '''
        caption_p = parse_xml(caption_xml)
        parent.insert(tbl_index, caption_p)

    # ==========================================
    # 页面设置
    # ==========================================

    def _apply_page_setup(self, doc: DocumentType, config: FormatConfig):
        """应用页面设置（页边距、纸张大小等）"""
        ps = config.page_setup
        for section in doc.sections:
            section.top_margin = Cm(ps.margin_top)
            section.bottom_margin = Cm(ps.margin_bottom)
            section.left_margin = Cm(ps.margin_left)
            section.right_margin = Cm(ps.margin_right)
            if ps.page_width > 0 and ps.page_height > 0:
                section.page_width = Cm(ps.page_width)
                section.page_height = Cm(ps.page_height)

    # ==========================================
    # 6. 自定义规则接口
    # ==========================================

    @staticmethod
    def _coerce_float(value, default: float) -> float:
        """宽容的数值解析：LLM 可能返回 "12pt"、None、bool 等非法字号。"""
        if isinstance(value, bool) or value is None:
            return default
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            m = re.search(r"\d+(?:\.\d+)?", value)
            if m:
                return float(m.group())
        return default

    def config_from_dict(self, d: dict) -> FormatConfig:
        """
        从字典创建 FormatConfig

        示例:
        {
            "font": "宋体",
            "size": "小四",  # 或 12
            "line_spacing": 1.25,
            "alignment": "justify",
            "first_line_indent": 24,
            "bold": False,
            "headings": {
                "1": {"font": "黑体", "size": 15, "bold": True},
                "2": {"font": "黑体", "size": 14, "bold": True}
            }
        }
        """
        # 解析正文字体
        cn_font = d.get("font", d.get("cn_font", "宋体"))
        en_font = d.get("en_font", "Times New Roman")

        # 解析字号
        size_raw = d.get("size", 12)
        if isinstance(size_raw, str):
            size = CHINESE_SIZE_MAP.get(size_raw, None)
            if size is None:
                size = self._coerce_float(size_raw, 12.0)
        else:
            size = self._coerce_float(size_raw, 12.0)

        # 解析对齐
        align_str = d.get("alignment", "justify")
        alignment = ALIGNMENT_STR_MAP.get(align_str, Alignment.JUSTIFY)

        body_font = FontConfig(
            cn_font=cn_font,
            en_font=en_font,
            size=size,
            bold=d.get("bold", False),
            italic=d.get("italic", False),
        )

        body_para = ParagraphConfig(
            alignment=alignment,
            line_spacing=self._coerce_float(d.get("line_spacing"), 1.25),
            line_spacing_rule=(
                "exactly" if d.get("line_spacing_rule") == "exactly" else "multiple"
            ),
            space_before=self._coerce_float(d.get("space_before"), 0),
            space_after=self._coerce_float(d.get("space_after"), 0),
            first_line_indent=self._coerce_float(d.get("first_line_indent"), 24),
            first_line_indent_chars=self._coerce_float(
                d.get("first_line_indent_chars"), 0
            ),
        )

        # 解析标题配置
        headings = {}
        headings_dict = d.get("headings", {})
        if not isinstance(headings_dict, dict):
            headings_dict = {}
        defaults = FormatConfig._default_headings()
        for level in range(1, 5):
            level_key = str(level)
            h_cfg = headings_dict.get(level_key, {})
            if not isinstance(h_cfg, dict):
                h_cfg = {}
            default_heading = defaults[level]
            h_size_raw = h_cfg.get("size", default_heading.font.size)
            if isinstance(h_size_raw, str):
                h_size = CHINESE_SIZE_MAP.get(h_size_raw, None)
                if h_size is None:
                    h_size = self._coerce_float(h_size_raw, default_heading.font.size)
            else:
                h_size = self._coerce_float(h_size_raw, default_heading.font.size)

            h_align = ALIGNMENT_STR_MAP.get(
                h_cfg.get("alignment", "left"), Alignment.LEFT
            )

            headings[level] = HeadingConfig(
                level=level,
                font=FontConfig(
                    cn_font=h_cfg.get("font", default_heading.font.cn_font),
                    en_font=h_cfg.get("en_font", default_heading.font.en_font),
                    size=float(h_size),
                    bold=h_cfg.get("bold", default_heading.font.bold),
                    italic=h_cfg.get("italic", default_heading.font.italic),
                    color=h_cfg.get("color", default_heading.font.color),
                ),
                paragraph=ParagraphConfig(
                    alignment=h_align,
                    line_spacing=self._coerce_float(h_cfg.get("line_spacing"), 1.0),
                    line_spacing_rule=(
                        "exactly" if h_cfg.get("line_spacing_rule") == "exactly"
                        else "multiple"
                    ),
                    space_before=self._coerce_float(h_cfg.get("space_before"), 6),
                    space_after=self._coerce_float(h_cfg.get("space_after"), 6),
                    first_line_indent=self._coerce_float(h_cfg.get("first_line_indent"), 0),
                    first_line_indent_chars=self._coerce_float(
                        h_cfg.get("first_line_indent_chars"), 0
                    ),
                ),
                numbering=h_cfg.get("numbering", True),
            )

        # 表格配置
        table_cfg = TableConfig(
            three_line=bool(d.get("table_three_line", True)),
            top_border=self._coerce_float(d.get("table_top"), 1.5),
            middle_border=self._coerce_float(d.get("table_middle"), 0.75),
            bottom_border=self._coerce_float(d.get("table_bottom"), 1.5),
            auto_number=d.get("table_numbering", True),
        )

        # 支持 table 字典格式（来自模板分析器）
        table_dict = d.get("table", {})
        if isinstance(table_dict, dict):
            if "three_line" in table_dict:
                table_cfg.three_line = bool(table_dict["three_line"])
            if "top_border" in table_dict:
                table_cfg.top_border = self._coerce_float(table_dict["top_border"], 1.5)
            if "middle_border" in table_dict:
                table_cfg.middle_border = self._coerce_float(table_dict["middle_border"], 0.75)
            if "bottom_border" in table_dict:
                table_cfg.bottom_border = self._coerce_float(table_dict["bottom_border"], 1.5)

        # 页面设置
        page_setup = PageSetupConfig()
        page_dict = d.get("page", {})
        if isinstance(page_dict, dict):
            page_setup.margin_top = self._coerce_float(page_dict.get("margin_top"), 2.54)
            page_setup.margin_bottom = self._coerce_float(page_dict.get("margin_bottom"), 2.54)
            page_setup.margin_left = self._coerce_float(page_dict.get("margin_left"), 3.17)
            page_setup.margin_right = self._coerce_float(page_dict.get("margin_right"), 3.17)
            if "page_width" in page_dict:
                page_setup.page_width = self._coerce_float(page_dict["page_width"], 21.0)
            if "page_height" in page_dict:
                page_setup.page_height = self._coerce_float(page_dict["page_height"], 29.7)

        return FormatConfig(
            body_font=body_font,
            body_paragraph=body_para,
            headings=headings,
            table_config=table_cfg,
            page_setup=page_setup,
        )

    def _merge_config(self, base: FormatConfig, override: dict) -> FormatConfig:
        """合并配置：base + override dict"""
        if not override:
            return base
        # 用 config_from_dict 解析 override，然后合并到 base
        override_config = self.config_from_dict(override)

        if override.get("font"):
            base.body_font.cn_font = override_config.body_font.cn_font
        if override.get("size"):
            base.body_font.size = override_config.body_font.size
        if override.get("line_spacing"):
            base.body_paragraph.line_spacing = override_config.body_paragraph.line_spacing
        if override.get("alignment"):
            base.body_paragraph.alignment = override_config.body_paragraph.alignment
        if override.get("first_line_indent"):
            base.body_paragraph.first_line_indent = override_config.body_paragraph.first_line_indent
            base.body_paragraph.first_line_indent_chars = 0
        if override.get("first_line_indent_chars"):
            base.body_paragraph.first_line_indent_chars = (
                override_config.body_paragraph.first_line_indent_chars
            )

        return base

    # ==========================================
    # 便捷方法
    # ==========================================

    def set_font(self, input_path: str, font_name: str,
                 output_path: Optional[str] = None) -> ProcessResult:
        """仅设置字体"""
        config = FormatConfig()
        config.body_font.cn_font = font_name
        for h in config.headings.values():
            h.font.cn_font = font_name
        return self.process(input_path, config, output_path)

    def set_spacing(self, input_path: str, line_spacing: float,
                    output_path: Optional[str] = None) -> ProcessResult:
        """仅设置行距"""
        config = FormatConfig()
        config.body_paragraph.line_spacing = line_spacing
        return self.process(input_path, config, output_path)

    def create_document(self, content: str,
                       format_config: Optional[FormatConfig] = None,
                       output_path: str = "output.docx",
                       config_dict: Optional[dict] = None) -> ProcessResult:
        """从文本内容创建新文档"""
        try:
            if config_dict:
                format_config = self.config_from_dict(config_dict)
            if not format_config:
                format_config = FormatConfig()

            doc = Document()

            # 按行解析（支持 Markdown 标题与表格，与 txt 入口能力一致）
            lines = content.split("\n")
            i = 0
            while i < len(lines):
                line = lines[i]
                line = line.rstrip("\r")
                if not line.strip():
                    doc.add_paragraph("")
                    i += 1
                    continue

                stripped = line.strip()
                if stripped.startswith("|") and stripped.endswith("|"):
                    table_lines = []
                    while i < len(lines) and lines[i].strip().startswith("|"):
                        table_lines.append(lines[i].strip())
                        i += 1
                    rows_data = [
                        [cell.strip() for cell in table_line.strip("|").split("|")]
                        for table_line in table_lines
                        if not re.match(r"^\|[\s\-:|]+\|$", table_line)
                    ]
                    if rows_data:
                        ncols = max(len(row) for row in rows_data)
                        table = doc.add_table(rows=len(rows_data), cols=ncols)
                        table.style = "Table Grid"
                        for row_idx, row in enumerate(rows_data):
                            for col_idx, cell_text in enumerate(row):
                                table.rows[row_idx].cells[col_idx].text = cell_text
                    continue

                if line.startswith("# "):
                    doc.add_heading(line[2:].strip(), level=1)
                elif line.startswith("## "):
                    doc.add_heading(line[3:].strip(), level=2)
                elif line.startswith("### "):
                    doc.add_heading(line[4:].strip(), level=3)
                elif line.startswith("#### "):
                    doc.add_heading(line[5:].strip(), level=4)
                else:
                    doc.add_paragraph(line)
                i += 1

            # 分析并应用格式
            self.analyze_structure(doc)
            self._apply_body_format(doc, format_config)
            self._apply_heading_format(doc, format_config)
            self._apply_page_setup(doc, format_config)

            doc.save(output_path)

            return ProcessResult(
                success=True,
                message=f"文档创建成功: {Path(output_path).name}",
                output_path=output_path,
                changes=[f"创建文档，共{len(lines)}行"],
            )
        except Exception as e:
            return ProcessResult(
                success=False,
                message=f"创建失败: {str(e)}",
            )
