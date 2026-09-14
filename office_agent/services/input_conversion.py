"""文本类输入（.txt/.md/.pdf）到 .docx 的统一转换。

WordService 基于 python-docx，只能消费 .docx。本模块在 Word 任务入口
把文本类输入一次性转换为临时 .docx，API 路由因此可以把这些扩展名
直接派发给 word_agent（清单 273 的转换链）。转换是全仓唯一接缝：
不要在 Agent 内部再各自实现文本读取。
"""
from __future__ import annotations

import os
import re
import tempfile

TEXT_LIKE_EXTENSIONS = (".txt", ".md", ".pdf")

# 本模块生成的临时转换副本统一前缀。
# 调用方据此区分"借用的用户原文件"与"本模块创建、需由调用方清理的临时副本"。
OWNED_TEMP_PREFIX = "office_agent_input_"

# 存储层落盘文件名是内部 ID（file_/out_/tmp_ + 十六进制）。文本类输入的
# ParsedDocument.title 会回退为文件名 stem，若不拦截，这个内部 ID 会被当成
# 文档一级标题写进用户拿到的 Word（RC：内部标识泄漏 + 标题层级错位）。
_INTERNAL_STORAGE_ID_RE = re.compile(r"^(?:file|out|tmp)_[0-9a-f]{6,}$", re.IGNORECASE)


def _safe_document_title(parsed) -> str:
    """返回可安全写入正文的文档标题；内部存储 ID 一律不渲染为标题。"""
    title = (getattr(parsed, "title", "") or "").strip()
    if title and not _INTERNAL_STORAGE_ID_RE.match(title):
        return title
    return ""


def _split_pipe_cells(line: str):
    """拆分一行 GitHub 风格管道表；不像表格（无竖线）返回 None。"""
    s = (line or "").strip()
    if "|" not in s:
        return None
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [cell.strip() for cell in s.split("|")]


def _is_table_delimiter(cells) -> bool:
    """判断是否为 | --- | :--: | 形式的分隔行。"""
    if not cells:
        return False
    seen = False
    for cell in cells:
        token = cell.replace(" ", "")
        if token == "":
            continue
        if not re.fullmatch(r":?-{1,}:?", token):
            return False
        seen = True
    return seen


def _emit_content_lines(document, lines) -> bool:
    """把一段正文行写入 docx；连续的 Markdown 管道表渲染为真实表格。

    缺了它，``| --- |`` 分隔行会作为可见垃圾段落残留，且表格数据不会成为
    可被 Word/WPS 识别的表格（RC W1：markdown 输入必须真实落表）。
    """
    wrote = False
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        header = _split_pipe_cells(line)
        if header and i + 1 < n:
            delim = _split_pipe_cells(lines[i + 1].strip())
            if delim and len(delim) == len(header) and _is_table_delimiter(delim):
                data_rows = []
                j = i + 2
                while j < n:
                    row = _split_pipe_cells(lines[j].strip())
                    if not row:
                        break
                    data_rows.append(row)
                    j += 1
                table = document.add_table(rows=1 + len(data_rows), cols=len(header))
                try:
                    table.style = "Table Grid"
                except KeyError:
                    pass
                for col, text in enumerate(header):
                    table.rows[0].cells[col].text = text
                for ri, row in enumerate(data_rows, start=1):
                    for col in range(len(header)):
                        table.rows[ri].cells[col].text = (
                            row[col] if col < len(row) else "")
                wrote = True
                i = j
                continue
        document.add_paragraph(line)
        wrote = True
        i += 1
    return wrote


def is_text_like(path: str) -> bool:
    return os.path.splitext(path or "")[1].lower() in TEXT_LIKE_EXTENSIONS


def is_owned_temp_input(path: str) -> bool:
    """判断路径是否为本模块生成的临时转换副本（调用方负责清理）。

    必须同时满足"文件名带本模块前缀"与"位于系统临时目录"两个条件：
    只看前缀的话，一个恰好命名为 ``office_agent_input_xxx.docx`` 的用户上传
    文件会被误判成临时文件，进而在任务结束时被删除——那是不可逆的用户数据
    丢失。对非文本类输入，``ensure_docx_input`` 原样返回用户路径，这里必须
    给出明确的 False。
    """
    if not path:
        return False
    name = os.path.basename(path)
    if not name.startswith(OWNED_TEMP_PREFIX) or not name.endswith(".docx"):
        return False
    try:
        parent = os.path.realpath(os.path.dirname(os.path.abspath(path)))
        temp_root = os.path.realpath(tempfile.gettempdir())
    except OSError:
        return False
    return parent == temp_root


def ensure_docx_input(path: str) -> str:
    """返回可供 python-docx 消费的输入路径。

    非文本类输入原样返回（借用，调用方**不得**删除）；文本类输入解析章节
    结构后写成临时 .docx 并返回新路径——该临时文件由**调用方持有所有权**，
    任务结束后必须清理（用 ``is_owned_temp_input`` 判定所有权）。
    PDF 依赖可选的 pymupdf，缺失时给出明确错误而不是退化成乱码文本。
    """
    ext = os.path.splitext(path or "")[1].lower()
    if ext not in TEXT_LIKE_EXTENSIONS:
        return path
    if ext == ".pdf":
        try:
            import pymupdf  # noqa: F401
        except ImportError as exc:
            raise RuntimeError(
                "处理 PDF 输入需要安装 pymupdf（pip install pymupdf）"
            ) from exc

    from ..knowledge_base.document_parser import DocumentParser
    parsed = DocumentParser().parse(path)

    from docx import Document
    document = Document()
    # Markdown 的首个一级标题（#）已在解析期提升为文档标题；其余 ##/### 直接
    # 对应 Word Heading 2/3。其它来源（docx/pdf…）维持历史的 level+1 映射。
    is_markdown = ext == ".md"
    title = _safe_document_title(parsed)
    if title:
        document.add_heading(title, level=1)
    wrote_content = False
    for section in parsed.sections:
        if section.title and section.level >= 1 and section.title != title:
            heading_level = section.level if is_markdown else min(section.level + 1, 9)
            document.add_heading(section.title, level=heading_level)
        section_lines = (section.content or "").split("\n")
        wrote_content = _emit_content_lines(document, section_lines) or wrote_content
    if not wrote_content and parsed.full_text.strip():
        wrote_content = _emit_content_lines(document, parsed.full_text.split("\n"))

    fd, out_path = tempfile.mkstemp(
        suffix=".docx", prefix=OWNED_TEMP_PREFIX
    )
    os.close(fd)
    # P3-99: mkstemp 已创建空文件；若 document.save 失败必须清掉这个空壳，
    # 不能让异常路径在系统临时目录里遗留零字节临时文件。
    try:
        document.save(out_path)
    except BaseException:
        try:
            os.remove(out_path)
        except OSError:
            pass
        raise
    return out_path
