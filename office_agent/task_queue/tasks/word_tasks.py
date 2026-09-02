"""
Word 相关后台任务

任务流程：
    0%   初始化
    20%  文件读取
    40%  结构分析
    60%  格式处理
    80%  生成文件
    100% 完成
"""
import os
import time
import uuid
import logging
import json
from ...security.error_sanitizer import sanitize_error
from ...runtime_config import get_output_dir

logger = logging.getLogger("office_agent.tasks.word")

def _safe_output(ext: str = ".docx") -> str:
    output_dir = str(get_output_dir())
    os.makedirs(output_dir, exist_ok=True)
    # 时间戳只精确到秒，normal 队列并发下同秒完成的任务会写同一路径互相覆盖
    unique = f"{time.strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}"
    return os.path.join(output_dir, f"word_{unique}{ext}")


def _display_stem(input_path: str, options: dict) -> str:
    """输出文件名应基于用户上传时的原始文件名，而非内部 file_id 路径。"""
    stem = os.path.splitext(os.path.basename(input_path or ""))[0]
    input_ids = (options or {}).get("input_file_ids") or []
    if input_ids:
        try:
            from ...storage.storage_service import get_storage_service
            original = get_storage_service().get_info(input_ids[0]).original_name or ""
            original_stem = os.path.splitext(original)[0]
            if original_stem and not original_stem.startswith("file_"):
                stem = original_stem
        except Exception:
            pass
    return stem or "output"


def _llm_format_config(instruction: str, input_path: str, options: dict,
                       progress=None):
    """Translate conversational Word feedback into the engine's config schema."""
    try:
        from ...model_gateway.gateway import ModelGateway
        from ...models.model_schemas import AITaskType
        from docx import Document
        doc = Document(input_path)
        structure = {
            "paragraphs": len(doc.paragraphs),
            "headings": [p.text[:120] for p in doc.paragraphs
                         if p.style and p.style.name.startswith("Heading")][:30],
            "tables": len(doc.tables),
        }
        history = options.get("history", []) if isinstance(options, dict) else []
        prompt = json.dumps({"instruction": instruction,
                             "previous_instruction": options.get("previous_instruction", ""),
                             "is_follow_up": bool(options.get("is_follow_up")),
                             "history": history,
                             "document": structure}, ensure_ascii=False)
        system = ("你是Word排版参数解析器。理解用户当前要求和历史反馈，输出严格JSON，不能输出解释。"
                  "只允许字段：font,en_font,size,line_spacing,alignment,first_line_indent,bold,italic,"
                  "space_before,space_after,headings。headings的键只能是1到4，值可含font,size,bold,italic,"
                  "alignment,line_spacing,numbering。只输出用户明确要求或合理修正所需字段；若无法确定输出{}。")
        response = ModelGateway(
            cancel_event=getattr(progress, "cancel_event", None)
        ).chat(user_message=prompt, system_prompt=system,
               task_type=AITaskType.DOCUMENT_UNDERSTANDING,
               temperature=0.1, max_tokens=1200)
        if isinstance(options, dict):
            options["model_call"] = {
                "called": True, "success": bool(response.success),
                "model": getattr(response, "model_used", None),
                "provider": getattr(response, "provider", None),
                "fallback_used": not bool(response.success),
                "error": getattr(response, "error", ""),
                "attempts": (response.raw_response or {}).get("attempts", 1) if isinstance(getattr(response, "raw_response", None), dict) else 1,
            }
        if not response.success or not response.content:
            return None
        text = response.content.strip()
        if "```" in text:
            text = text.replace("```json", "").replace("```", "").strip()
        config = json.loads(text)
        if not isinstance(config, dict):
            return None
        allowed = {"font", "en_font", "size", "line_spacing", "alignment",
                   "first_line_indent", "bold", "italic", "space_before", "space_after", "headings"}
        config = {k: v for k, v in config.items() if k in allowed}
        if isinstance(config.get("headings"), dict):
            config["headings"] = {str(k): v for k, v in config["headings"].items()
                                   if str(k) in {"1", "2", "3", "4"} and isinstance(v, dict)}
        return config or None
    except Exception as exc:
        if isinstance(options, dict):
            options["model_call"] = {"called": True, "success": False,
                                      "fallback_used": True, "error": sanitize_error(exc, "模型解析不可用"), "attempts": 0}
        logger.warning("Word LLM解析不可用，回退规则解析: %s", sanitize_error(exc, "模型解析不可用"))
        return None


def process_word(input_path: str, output_path: str = None,
                 instruction: str = "", options: dict = None,
                 progress=None, _task_id: str = None, **kwargs) -> dict:
    """
    Word 文档处理主任务

    Args:
        input_path: 输入文件路径
        output_path: 输出文件路径
        instruction: 处理指令
        options: 选项
        progress: 进度回调对象
        _task_id: 任务ID
    """
    options = options or {}
    result = {"status": "success", "output_files": [], "steps": []}

    try:
        if progress:
            progress.update(5, "初始化Word处理")

        # 1. 检查文件
        if not os.path.exists(input_path):
            raise FileNotFoundError(f"文件不存在: {input_path}")

        if not output_path:
            output_path = _safe_output()

        if progress:
            progress.update(20, "读取文档内容")

                # 2. 调用真实存在的 Word 排版引擎（python-docx + 结构分析 + 格式应用）
        from ...services.word_service import WordService

        service = WordService()

        if progress:
            progress.update(40, "分析文档结构")

        # 3. 处理文档：统一走 WordService.process 主入口。
        if progress:
            progress.update(60, "执行格式处理")

        config_dict = options.get("config") if isinstance(options, dict) else None
        # Convert the user's natural-language request into the format schema. The
        # old implementation passed an opaque `instruction` key to WordService,
        # which is ignored by config_from_dict and made chat requests no-ops.
        if config_dict is None and instruction:
            config_dict = _llm_format_config(
                instruction, input_path, options, progress
            )
            if progress:
                progress.check_cancelled()
        if config_dict is None and instruction:
            from ...parsers.format_parser import FormatRuleParser
            parsed = FormatRuleParser().parse(instruction)
            config_dict = parsed.to_config_dict()
            if not config_dict:
                config_dict = None
        process_result = service.process(
            input_path,
            output_path=output_path,
            config_dict=config_dict,
        )

        if not process_result or not process_result.success:
            message = (getattr(process_result, "message", None) or "Word处理失败")
            raise RuntimeError(f"Word引擎处理失败: {message}")

        output = process_result.output_path
        if not output or not os.path.exists(output):
            raise RuntimeError(f"Word引擎未生成输出文件: {output}")

        if progress:
            progress.update(80, "登记输出文件到 Storage")

        # 4. 输出文件以独立新文件登记（outputs 桶，新 file_id，不触发版本覆盖）。
        from ...storage.storage_service import get_storage_service

        storage = get_storage_service()
        original_name = options.get("output_filename") or (
            f"{_display_stem(input_path, options)}_formatted.docx"
        )
        file_info = storage.save_new_output(
            source_path=output,
            original_name=original_name,
            owner_id=options.get("_owner_id"),
            change_description=instruction or "Word 文档处理",
        )

        if not file_info or not file_info.file_id:
            raise RuntimeError("输出文件登记失败: 未获得 file_id")

        result["output_files"].append(file_info.file_id)
        result["output_path"] = output
        result["output_file_id"] = file_info.file_id
        result["message"] = getattr(process_result, "message", "Word处理完成")

        if progress:
            progress.update(100, "Word处理完成")

        result["steps"].append("文档处理完成")
        logger.info(f"Word任务 {_task_id} 完成: file_id={file_info.file_id}")

    except Exception as e:
        logger.error("Word任务 %s 失败: %s", _task_id, sanitize_error(e))
        result["status"] = "failed"
        result["error"] = sanitize_error(e)

    if options.get("model_call"):
        result["model_call"] = options["model_call"]
    return result


def format_document(input_path: str, output_path: str = None,
                    template_type: str = None, instruction: str = "",
                    options: dict = None,
                    progress=None, _task_id: str = None, **kwargs) -> dict:
    """文档格式排版"""
    return process_word(
        input_path=input_path,
        output_path=output_path,
        instruction=instruction or f"按照{template_type or '标准'}格式排版",
        options=options or {},
        progress=progress,
        _task_id=_task_id,
    )
