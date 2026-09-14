"""
文件处理后台任务
"""
import os
import logging
from pathlib import Path
from datetime import datetime, timedelta, timezone
from typing import Any, cast

from sqlalchemy.engine import CursorResult

from ...security.error_sanitizer import sanitize_error
from ...runtime_config import get_data_root

logger = logging.getLogger("office_agent.tasks.file")


def process_upload(file_path: str, file_id: str | None = None,
                   file_type: str | None = None, options: dict | None = None,
                   progress=None, _task_id: str | None = None, **kwargs) -> dict:
    """
    上传后处理任务（验证、提取元数据、生成预览等）
    """
    options = options or {}
    metadata: dict = {}
    result = {"status": "success", "file_id": file_id, "metadata": metadata}

    try:
        if progress:
            progress.update(10, "验证文件")

        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        file_size = os.path.getsize(file_path)
        metadata["size"] = file_size

        if progress:
            progress.update(30, "提取文件元数据")

        # 根据文件类型处理
        ext = Path(file_path).suffix.lower()
        type_by_extension = {
            ".docx": "word", ".doc": "word",
            ".pptx": "ppt", ".ppt": "ppt",
            ".xlsx": "excel", ".xls": "excel", ".csv": "excel",
            ".pdf": "pdf",
            ".txt": "text", ".md": "text", ".rtf": "text",
            ".png": "image", ".jpg": "image", ".jpeg": "image",
            ".gif": "image", ".bmp": "image", ".webp": "image",
            ".tif": "image", ".tiff": "image",
        }
        result["file_type"] = file_type or type_by_extension.get(
            ext, ext.lstrip(".") or "unknown"
        )

        if ext == ".docx":
            result["file_type"] = "word"
            try:
                from docx import Document
                doc = Document(file_path)
                metadata["paragraphs"] = len(doc.paragraphs)
                metadata["tables"] = len(doc.tables)
            except Exception as exc:
                logger.warning("提取 Word 元数据失败 %s: %s", file_path, exc)

        elif ext == ".pptx":
            result["file_type"] = "ppt"
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                metadata["slides"] = len(prs.slides)
            except Exception as exc:
                logger.warning("提取 PPT 元数据失败 %s: %s", file_path, exc)

        elif ext == ".xlsx":
            result["file_type"] = "excel"
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, read_only=True)
                try:
                    metadata["sheets"] = wb.sheetnames
                finally:
                    # read_only 工作簿持有文件句柄，不 close 会泄漏到 GC
                    wb.close()
            except Exception as exc:
                logger.warning("提取 Excel 元数据失败 %s: %s", file_path, exc)

        elif ext == ".pdf":
            result["file_type"] = "pdf"
            try:
                import pymupdf
                pdf_doc = pymupdf.open(file_path)
                metadata["pages"] = len(pdf_doc)
                pdf_doc.close()
            except Exception as exc:
                logger.warning("提取 PDF 元数据失败 %s: %s", file_path, exc)

        if progress:
            progress.update(70, "生成文件信息")

        if progress:
            progress.update(100, "文件处理完成")

        logger.info(f"文件处理 {file_id} 完成: {result['file_type']}")

    except Exception as e:
        logger.error("文件处理 %s 失败: %s", file_id, sanitize_error(e))
        result["status"] = "failed"
        result["error"] = sanitize_error(e)
        raise

    return result


def convert_format(input_path: str, output_path: str,
                   target_format: str, options: dict | None = None,
                   progress=None, _task_id: str | None = None, **kwargs) -> dict:
    """
    文件格式转换
    """
    result = {"status": "success", "output_files": []}

    try:
        if progress:
            progress.update(10, "准备转换")

        if not os.path.exists(input_path):
            raise FileNotFoundError(input_path)

        if progress:
            progress.update(30, f"转换为{target_format}")

        # 转换引擎尚未实现：与其假装成功并返回空产物，不如明确失败。
        # （实现真实转换引擎时在此接入处理链，并在成功路径上报
        # progress.update(100, "转换完成")——此前的不可达代码已删除。）
        raise RuntimeError(
            f"文件格式转换（→{target_format}）暂未实现，请使用 Word/PPT/Excel Agent 处理"
        )

    except Exception as e:
        result["status"] = "failed"
        result["error"] = sanitize_error(e)
        raise

    return result


def cleanup_old_logs(days: int | None = None, progress=None, _task_id: str | None = None, **kwargs) -> dict:
    """清理超过保留期的执行/模型调用/错误日志（防 SQLite 无限膨胀）"""
    import os
    if not days or days <= 0:
        days = int(os.environ.get("LOG_RETENTION_DAYS", "30"))
    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    deleted: dict[str, int | str] = {}
    try:
        from sqlalchemy import delete
        from ...database.session import SessionLocal
        from ...database.models.execution import ExecutionLog, ModelCallLog, ErrorLog
        session = SessionLocal()
        try:
            for model in (ExecutionLog, ModelCallLog, ErrorLog):
                try:
                    res = cast("CursorResult[Any]", session.execute(delete(model).where(model.created_at < cutoff)))
                    deleted[model.__tablename__] = res.rowcount or 0
                except Exception as exc:
                    deleted[getattr(model, "__tablename__", model.__name__)] = f"error: {sanitize_error(exc, '清理失败')}"
            session.commit()
        finally:
            session.close()
        logger.info("日志保留清理（>%d天）: %s", days, deleted)
        return {"days": days, "deleted": deleted}
    except Exception as exc:
        logger.warning("日志清理失败: %s", sanitize_error(exc, "清理失败"))
        return {"days": days, "error": sanitize_error(exc, "清理失败")}


def cleanup_temp_files(progress=None, _task_id: str | None = None, **kwargs) -> dict:
    """
    清理临时文件（定时任务，每天执行）

    扫描真实数据目录（OFFICE_AGENT_DATA_DIR 优先）：
    - 中间产物目录 outputs/（成品已复制进 storage 桶，这里的是处理过程文件）
    - 分片上传残留 multipart/
    - 存储桶内原子写残留的 *.tmp-* 文件
    - 兼容旧版本的 temp/uploads/cache 目录
    存储桶内的正式文件（DB 有记录）一律不按 mtime 删除，避免破坏下载。
    """
    result = {"status": "success", "cleaned": 0, "freed_bytes": 0}

    try:
        if progress:
            progress.update(10, "扫描临时文件")

        data_root = get_data_root()
        dirs_to_clean = [
            data_root / "outputs",
            data_root / "temp",
            data_root / "uploads",
            data_root / "cache",
        ]
        multipart_dir = data_root / "storage" / "multipart"
        storage_root = data_root / "storage"

        cutoff = datetime.now(timezone.utc) - timedelta(days=7)
        cleaned = 0
        freed = 0

        def _unlink_expired(path: Path):
            nonlocal cleaned, freed
            try:
                if not path.is_file():
                    return
                mtime = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc)
                if mtime < cutoff:
                    size = path.stat().st_size
                    path.unlink()
                    cleaned += 1
                    freed += size
            except Exception as exc:
                logger.debug("跳过 %s: %s", path, exc)

        for d in dirs_to_clean:
            if not d.exists():
                continue
            for f in d.iterdir():
                _unlink_expired(f)

        # 分片上传残留：递归清理过期 part 文件
        if multipart_dir.exists():
            for f in multipart_dir.rglob("*"):
                _unlink_expired(f)

        # 原子写残留（LocalStorage 的临时文件名含 .tmp-），DB 从不引用
        if storage_root.exists():
            for f in storage_root.rglob("*.tmp-*"):
                _unlink_expired(f)

        result["cleaned"] = cleaned
        result["freed_bytes"] = freed

        if progress:
            progress.update(100, f"清理完成，释放 {freed / 1024 / 1024:.1f}MB")

        logger.info(f"临时文件清理: 删除 {cleaned} 个文件，释放 {freed} 字节")

    except Exception as e:
        result["status"] = "failed"
        from ...security.error_sanitizer import sanitize_error
        result["error"] = sanitize_error(e)

    return result


def system_health_check(progress=None, _task_id: str | None = None, **kwargs) -> dict:
    """
    系统健康检查（定时任务，每小时执行）
    """
    result: dict = {"status": "success", "checks": {}}

    try:
        if progress:
            progress.update(20, "检查存储")

        data_dir = get_data_root()
        if data_dir.exists():
            total_size = sum(f.stat().st_size for f in data_dir.rglob("*") if f.is_file())
            result["checks"]["storage_mb"] = round(total_size / 1024 / 1024, 2)

        if progress:
            progress.update(50, "检查数据库")

        try:
            from ...database.session import session_scope
            from ...database.repository import TaskRepository
            with session_scope() as session:
                repo = TaskRepository(session)
                result["checks"]["total_tasks"] = repo.count()
                result["checks"]["pending_tasks"] = len(repo.get_pending_tasks(limit=100))
        except Exception as e:
            # P3-74: 内层 DB 异常可能含连接串/文件路径，统一脱敏后再入结果
            result["checks"]["database"] = f"error: {sanitize_error(e, '数据库检查失败')}"

        if progress:
            progress.update(80, "检查Worker")

        result["checks"]["timestamp"] = datetime.now(timezone.utc).isoformat()

        if progress:
            progress.update(100, "健康检查完成")

    except Exception as e:
        result["status"] = "failed"
        # sanitize_error 已在模块顶部导入；此处不得再局部 import，否则该名
        # 会成为整个函数的局部变量，让内层 P3-74 脱敏点 UnboundLocalError。
        result["error"] = sanitize_error(e)

    return result
