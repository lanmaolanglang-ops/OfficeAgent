"""
Excel Service - Excel 底层操作服务

封装 openpyxl + pandas，提供：
- 读写 .xlsx 文件
- 数据增删改查
- 公式写入
- 格式化（字体、颜色、边框、对齐、数字格式）
- 图表生成
- 数据筛选/排序
"""
from pathlib import Path
from typing import Optional, List, Any, Tuple
from copy import copy

import math
import re

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule

import pandas as pd

from .models import (
    ChartSpec, FormatSpec, FormulaSpec, ExcelResult,
)


# ==========================================
# 样式预设
# ==========================================

HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT = Font(name="微软雅黑", size=10)
BODY_ALIGN = Alignment(horizontal="left", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

HEADER_BORDER = Border(
    left=Side(style="thin", color="1F4E79"),
    right=Side(style="thin", color="1F4E79"),
    top=Side(style="thin", color="1F4E79"),
    bottom=Side(style="medium", color="1F4E79"),
)

# 斑马纹
ZEBRA_FILL = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")


def _derive_output_path(file_path: str, suffix: str) -> str:
        """默认输出路径：大小写不敏感替换 .xlsx；非 xlsx 扩展名直接追加后缀，
        避免替换失败时输出路径与输入相同而覆写原文件。"""
        p = Path(file_path)
        if p.suffix.lower() == ".xlsx":
            return str(p.with_name(p.stem + suffix + ".xlsx"))
        return str(p) + suffix + ".xlsx"


def _sanitize_cell_value(value):
    """NaN/Infinity/NaT 会写出非法 OOXML（空 <v> 触发 Excel 修复提示），统一转 None"""
    if value is None:
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    try:
        # pandas NaT / numpy nan
        if value != value:  # NaN 自反性
            return None
    except Exception:
        pass
    if value is pd.NaT:
        return None
    return value


def hex_to_color(hex_str: str):
    """十六进制颜色转 openpyxl Color"""
    from openpyxl.styles import Color
    if not isinstance(hex_str, str):
        raise ValueError("颜色值必须是字符串")
    h = hex_str.strip().lstrip("#")
    if len(h) == 3:
        h = "".join(ch * 2 for ch in h)
    if not re.fullmatch(r"[0-9A-Fa-f]{6}([0-9A-Fa-f]{2})?", h):
        raise ValueError(f"无效的十六进制颜色: {hex_str!r}")
    # openpyxl 使用 ARGB；六位 RGB 显式补不透明 alpha，避免被解释为透明色。
    if len(h) == 6:
        h = f"FF{h}"
    return Color(rgb=h.upper())


class ExcelService:
    """
    Excel 底层操作服务

    用法:
        svc = ExcelService()
        svc.open("data.xlsx") 或 svc.create("output.xlsx")
        svc.write_data("Sheet1", data)
        svc.add_formula(FormulaSpec(formula="=SUM(A1:A10)", target_cell="A11"))
        svc.apply_format(FormatSpec(range_str="A1:D1", bold=True, bg_color="1F4E79"))
        svc.add_chart(ChartSpec(chart_type="column", data_range="A1:D10"))
        svc.save()
    """

    def __init__(self):
        self.wb: Optional[Workbook] = None
        self.file_path: str = ""
        self._opened_from: Optional[str] = None
        self.changes: List[str] = []

    def create(self, file_path: str, sheet_name: str = "Sheet1") -> 'ExcelService':
        """创建新工作簿"""
        self.wb = Workbook()
        self.wb.active.title = sheet_name
        self.file_path = file_path
        # 新建工作簿没有"磁盘上的原始数据"，保存到同一路径是安全的
        self._opened_from: Optional[str] = None
        self.changes.append(f"创建新工作簿: {file_path}")
        return self

    def open(self, file_path: str) -> 'ExcelService':
        """打开已有工作簿"""
        if not Path(file_path).exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        # 宏工作簿必须 keep_vba，否则另存时 VBA/宏被静默剥离
        keep_vba = Path(file_path).suffix.lower() in (".xlsm", ".xltm")
        self.wb = load_workbook(file_path, data_only=False, keep_vba=keep_vba)
        self.file_path = file_path
        self._opened_from = str(file_path)
        self.changes.append(f"打开文件: {Path(file_path).name}")
        return self

    def save(self, output_path: str = "") -> ExcelResult:
        """保存文件"""
        if not self.wb:
            return ExcelResult(success=False, message="没有打开的工作簿")

        path = output_path or self.file_path
        if not path:
            return ExcelResult(success=False, message="未指定输出路径")

        # 禁止把输出写回"从磁盘打开的输入文件"（防止磁盘满/中途失败损坏
        # 用户原始数据）；create() 新建的工作簿不受限制
        if self._opened_from and Path(path).resolve() == Path(self._opened_from).resolve():
            return ExcelResult(
                success=False,
                message=f"拒绝保存：输出路径与输入文件相同（{Path(path).name}），请另存为新文件",
            )

        # 确保目录存在
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)

        return ExcelResult(
            success=True,
            message=f"保存成功: {Path(path).name}",
            output_path=path,
            sheet_count=len(self.wb.sheetnames),
            changes=self.changes.copy(),
        )

    # ==========================================
    # 工作表操作
    # ==========================================

    def get_sheet(self, name: str = None):
        """获取工作表"""
        if name is None:
            return self.wb.active
        if name in self.wb.sheetnames:
            return self.wb[name]
        return None

    def create_sheet(self, name: str, index: int = None):
        """创建新工作表"""
        if name in self.wb.sheetnames:
            return self.wb[name]
        ws = self.wb.create_sheet(name, index)
        self.changes.append(f"创建工作表: {name}")
        return ws

    def rename_sheet(self, old_name: str, new_name: str):
        """重命名工作表"""
        if old_name in self.wb.sheetnames:
            self.wb[old_name].title = new_name
            self.changes.append(f"重命名工作表: {old_name} → {new_name}")

    def delete_sheet(self, name: str):
        """删除工作表"""
        if name in self.wb.sheetnames:
            del self.wb[name]
            self.changes.append(f"删除工作表: {name}")

    def list_sheets(self) -> List[str]:
        """列出所有工作表"""
        return list(self.wb.sheetnames)

    # ==========================================
    # 数据读写
    # ==========================================

    def write_data(self, sheet_name: str, data: list,
                   start_row: int = 1, start_col: int = 1,
                   has_header: bool = True):
        """
        写入二维数据

        Args:
            sheet_name: 工作表名
            data: 二维列表 [[行1], [行2], ...]
            start_row/start_col: 起始位置（1-based）
            has_header: 第一行是否为表头（自动应用表头样式）
        """
        ws = self.get_sheet(sheet_name)
        if ws is None:
            ws = self.create_sheet(sheet_name)

        for r_idx, row_data in enumerate(data):
            for c_idx, value in enumerate(row_data):
                value = _sanitize_cell_value(value)
                cell = ws.cell(row=start_row + r_idx, column=start_col + c_idx)
                if isinstance(cell, MergedCell):
                    continue
                cell.value = value
                # 表头样式
                if has_header and r_idx == 0:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = HEADER_ALIGN
                    cell.border = HEADER_BORDER
                else:
                    cell.font = BODY_FONT
                    cell.alignment = BODY_ALIGN
                    cell.border = THIN_BORDER
                    # 斑马纹
                    if (r_idx % 2 == 0) and has_header:
                        cell.fill = ZEBRA_FILL

        # 自动调整列宽
        if data:
            for c_idx in range(len(data[0])):
                max_len = max(
                    len(str(row[c_idx])) if c_idx < len(row) else 0
                    for row in data
                )
                col_letter = get_column_letter(start_col + c_idx)
                ws.column_dimensions[col_letter].width = min(max(max_len * 1.5 + 2, 8), 40)

        self.changes.append(f"写入数据: {sheet_name}, {len(data)}行 x {len(data[0]) if data else 0}列")

    def write_dataframe(self, sheet_name: str, df: pd.DataFrame,
                        start_row: int = 1, start_col: int = 1):
        """写入 pandas DataFrame"""
        data = [list(df.columns)] + df.values.tolist()
        self.write_data(sheet_name, data, start_row, start_col, has_header=True)

    def read_data(self, sheet_name: str = None,
                  start_row: int = 1, start_col: int = 1,
                  max_rows: int = None, max_cols: int = None) -> list:
        """读取数据为二维列表"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return []

        rows = []
        for r_idx, row in enumerate(ws.iter_rows(min_row=start_row, min_col=start_col,
                                                  values_only=True)):
            if max_rows and r_idx >= max_rows:
                break
            row_data = list(row)
            if max_cols:
                row_data = row_data[:max_cols]
            rows.append(row_data)
        return rows

    def read_dataframe(self, sheet_name: str = None) -> pd.DataFrame:
        """读取为 pandas DataFrame"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return pd.DataFrame()
        data = self.read_data(sheet_name)
        if not data:
            return pd.DataFrame()
        return pd.DataFrame(data[1:], columns=data[0])

    def get_cell(self, sheet_name: str, cell_ref: str):
        """获取单元格"""
        ws = self.get_sheet(sheet_name)
        return ws[cell_ref] if ws else None

    def set_cell(self, sheet_name: str, cell_ref: str, value: Any):
        """设置单元格值"""
        ws = self.get_sheet(sheet_name)
        if ws:
            cell = ws[cell_ref]
            if not isinstance(cell, MergedCell):
                cell.value = value

    # ==========================================
    # 公式
    # ==========================================

    def add_formula(self, spec: FormulaSpec, sheet_name: str = None):
        """写入公式"""
        ws = self.get_sheet(sheet_name)
        if ws is None or not spec.target_cell:
            return

        formula = spec.formula or ""
        category = getattr(spec, "category", "") or ""
        # header/label 规格是纯文本（如 "环比增长率"、"华东"），写成
        # "=中文" 会变成 #NAME? 错误公式；仅真正的公式才加 "=" 前缀
        is_plain_text = category in ("header", "label") or (
            formula.startswith("=")
            and not re.search(r"[A-Za-z0-9$(]", formula)
        )
        if is_plain_text:
            cell = ws[spec.target_cell]
            if isinstance(cell, MergedCell):
                self.changes.append(f"跳过合并从属单元格: {spec.target_cell}")
                return
            cell.value = formula.lstrip("=")
            cell.font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
            self.changes.append(f"写入文本 {spec.target_cell}: {formula.lstrip('=')}")
            return
        if not formula.startswith("="):
            formula = "=" + formula

        cell = ws[spec.target_cell]
        if isinstance(cell, MergedCell):
            self.changes.append(f"跳过合并从属单元格: {spec.target_cell}")
            return
        cell.value = formula
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
        self.changes.append(f"添加公式 {spec.target_cell}: {formula}")

    def add_formulas(self, specs: List[FormulaSpec], sheet_name: str = None):
        """批量添加公式"""
        for spec in specs:
            self.add_formula(spec, sheet_name)

    # ==========================================
    # 格式化
    # ==========================================

    def apply_format(self, spec: FormatSpec, sheet_name: str = None):
        """应用格式"""
        ws = self.get_sheet(sheet_name)
        if ws is None or not spec.range_str:
            return

        cell_range = ws[spec.range_str]
        # 统一为二维结构：((cell, cell, ...), (cell, cell, ...))
        if isinstance(cell_range, tuple):
            # 检查第一个元素是否是 Cell（单行情况）
            from openpyxl.cell.cell import Cell
            if cell_range and isinstance(cell_range[0], Cell):
                cell_range = (cell_range,)
        else:
            cell_range = ((cell_range,),)

        for row in cell_range:
            for cell in row:
                if spec.font_name or spec.font_size:
                    cell.font = Font(
                        name=spec.font_name or "微软雅黑",
                        size=spec.font_size or 10,
                        bold=spec.bold,
                        italic=spec.italic,
                        color=spec.font_color.lstrip("#") if spec.font_color else "000000",
                    )
                elif spec.bold or spec.italic or spec.font_color:
                    f = cell.font
                    cell.font = Font(
                        name=f.name, size=f.size,
                        bold=spec.bold if spec.bold else f.bold,
                        italic=spec.italic if spec.italic else f.italic,
                        color=spec.font_color.lstrip("#") if spec.font_color else f.color,
                    )

                if spec.bg_color:
                    cell.fill = PatternFill(
                        start_color=spec.bg_color.lstrip("#"),
                        end_color=spec.bg_color.lstrip("#"),
                        fill_type="solid"
                    )

                if spec.alignment:
                    align_map = {
                        "left": "left", "center": "center", "right": "right",
                    }
                    cell.alignment = Alignment(
                        horizontal=align_map.get(spec.alignment, "left"),
                        vertical="center",
                        wrap_text=True,
                    )

                if spec.number_format:
                    cell.number_format = spec.number_format

                if spec.border:
                    cell.border = THIN_BORDER

        self.changes.append(f"应用格式: {spec.range_str}")

    def apply_header_style(self, sheet_name: str, range_str: str = "1:1"):
        """应用表头样式"""
        self.apply_format(FormatSpec(
            range_str=range_str,
            font_name="微软雅黑", font_size=11, bold=True,
            font_color="FFFFFF", bg_color="1F4E79",
            alignment="center", border=True,
        ), sheet_name)

    def auto_width(self, sheet_name: str = None):
        """自动调整列宽"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        cell_len = len(str(cell.value))
                        # 中文字符算2
                        cell_len = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                        max_len = max(max_len, cell_len)
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_len * 0.8 + 2, 8), 50)

    def set_column_width(self, col_letter: str, width: float,
                         sheet_name: str = None):
        """设置单列列宽"""
        ws = self.get_sheet(sheet_name)
        if ws is None or not col_letter:
            return
        try:
            ws.column_dimensions[col_letter].width = float(width)
        except (TypeError, ValueError):
            pass

    def freeze_header(self, sheet_name: str = None):
        """冻结首行"""
        ws = self.get_sheet(sheet_name)
        if ws:
            ws.freeze_panes = "A2"
            self.changes.append(f"冻结首行: {ws.title}")

    def add_filter(self, sheet_name: str = None, range_str: str = ""):
        """添加自动筛选"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return
        if not range_str:
            if ws.max_row > 1:
                range_str = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        if not range_str:
            return
        ws.auto_filter.ref = range_str
        self.changes.append(f"添加筛选: {ws.title}")

    # ==========================================
    # 图表
    # ==========================================

    def add_chart(self, spec: ChartSpec, sheet_name: str = None):
        """添加图表，并完整应用 ChartSpec 的类型、堆积、图例和组合图字段。"""
        from .chart_generator import ChartGenerator
        ChartGenerator()._render_chart(self, spec, sheet_name)

    def add_charts(self, specs: List[ChartSpec], sheet_name: str = None):
        """批量添加图表"""
        for spec in specs:
            self.add_chart(spec, sheet_name)

    # ==========================================
    # 数据操作
    # ==========================================

    def sort_data(self, sheet_name: str, sort_col: int,
                  ascending: bool = True, has_header: bool = True):
        """排序数据，同时保留公式、单元格样式、批注和超链接。"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return

        # 数据区域存在合并单元格时跳过重写（写入 MergedCell 会抛异常，
        # 且整表重写会破坏合并结构）
        ws_target = ws
        if getattr(ws_target, "merged_cells", None) and ws_target.merged_cells.ranges:
            self.changes.append(f"排序跳过: {sheet_name} 含合并单元格")
            return

        first_data_row = 2 if has_header else 1
        if ws.max_row < first_data_row:
            return

        sort_column = sort_col + 1
        if sort_column < 1 or sort_column > ws.max_column:
            raise ValueError(f"排序列超出范围: {sort_col}")

        # 不能经 values_only/整表 write_data 往返，否则公式会固化、样式会被重建。
        # 这里为每一行保留完整单元格状态；公式在换行时按 Excel 的相对引用规则平移。
        row_snapshots = []
        for source_row in range(first_data_row, ws.max_row + 1):
            cells = []
            for cell in ws[source_row]:
                cells.append({
                    "coordinate": cell.coordinate,
                    "value": cell.value,
                    "style": copy(cell._style),
                    "comment": copy(cell.comment),
                    "hyperlink": copy(cell.hyperlink),
                })
            row_snapshots.append((source_row, cells))

        def sort_key(snapshot):
            _source_row, cells = snapshot
            value = cells[sort_column - 1]["value"]
            if value is None:
                return (2, "")
            if isinstance(value, bool):
                return (0, int(value))
            if isinstance(value, (int, float)):
                return (0, float(value))
            if hasattr(value, "timestamp"):
                try:
                    return (0, value.timestamp())
                except (OSError, ValueError):
                    pass
            return (1, str(value).casefold())

        non_blank = [row for row in row_snapshots
                     if row[1][sort_column - 1]["value"] is not None]
        blank = [row for row in row_snapshots
                 if row[1][sort_column - 1]["value"] is None]
        non_blank.sort(key=sort_key, reverse=not ascending)
        sorted_rows = non_blank + blank

        from openpyxl.formula.translate import Translator
        for target_row, (_source_row, cells) in enumerate(sorted_rows, start=first_data_row):
            for col_idx, snapshot in enumerate(cells, start=1):
                target = ws.cell(row=target_row, column=col_idx)
                value = snapshot["value"]
                if isinstance(value, str) and value.startswith("="):
                    try:
                        value = Translator(
                            value, origin=snapshot["coordinate"]
                        ).translate_formula(target.coordinate)
                    except (TypeError, ValueError):
                        # 外部链接、结构化引用等 Translator 不认识时保留原公式，
                        # 仍优于把公式固化为缓存值。
                        pass
                target.value = value
                target._style = copy(snapshot["style"])
                target.comment = copy(snapshot["comment"])
                target._hyperlink = copy(snapshot["hyperlink"])

        self.changes.append(f"排序: {sheet_name} 第{sort_col+1}列 {'升序' if ascending else '降序'}")

    def add_summary_row(self, sheet_name: str, label: str = "合计",
                        sum_cols: List[int] = None, data_start_row: int | None = None):
        """添加汇总行"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return

        normalized_label = str(label).strip()
        summary_labels = {"合计", "总计", "汇总", normalized_label}
        for row in range(max(1, ws.max_row - 4), ws.max_row + 1):
            existing = ws.cell(row=row, column=1).value
            if isinstance(existing, str) and existing.strip() in summary_labels:
                self.changes.append(f"汇总行已存在，跳过: {existing.strip()}")
                return

        if data_start_row is None:
            data_start_row = 2
            # 表格对象给出了比“固定第 2 行”更可靠的数据起点。
            if ws.tables:
                from openpyxl.utils.cell import range_boundaries
                table = next(iter(ws.tables.values()))
                _min_col, min_row, _max_col, _max_row = range_boundaries(table.ref)
                data_start_row = min_row + 1
        if data_start_row < 1 or data_start_row > ws.max_row:
            raise ValueError(f"汇总数据起始行无效: {data_start_row}")

        last_row = ws.max_row + 1
        ws.cell(row=last_row, column=1, value=label).font = Font(
            name="微软雅黑", size=10, bold=True
        )

        if sum_cols:
            for col in sum_cols:
                col_letter = get_column_letter(col + 1)
                formula = f"=SUM({col_letter}{data_start_row}:{col_letter}{last_row - 1})"
                cell = ws.cell(row=last_row, column=col + 1, value=formula)
                cell.font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")

        self.changes.append(f"添加汇总行: {label}")

    def add_conditional_format(self, sheet_name: str, range_str: str,
                               rule_type: str = "color_scale"):
        """添加条件格式"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return

        if rule_type == "color_scale":
            rule = ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B"
            )
            ws.conditional_formatting.add(range_str, rule)
        elif rule_type == "data_bar":
            rule = DataBarRule(
                start_type="min", end_type="max",
                color="638EC6", showValue=True
            )
            ws.conditional_formatting.add(range_str, rule)

        self.changes.append(f"条件格式: {range_str} ({rule_type})")

    # ==========================================
    # 工具方法
    # ==========================================

    def get_dimensions(self, sheet_name: str = None) -> Tuple[int, int]:
        """获取数据维度 (行数, 列数)"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return (0, 0)
        return (ws.max_row, ws.max_column)

    def get_preview(self, sheet_name: str = None, rows: int = 10) -> dict:
        """获取数据预览"""
        ws = self.get_sheet(sheet_name)
        if ws is None:
            return {}
        data = self.read_data(sheet_name, max_rows=rows)
        return {
            "sheet": ws.title,
            "total_rows": ws.max_row,
            "total_cols": ws.max_column,
            "preview": data,
        }
