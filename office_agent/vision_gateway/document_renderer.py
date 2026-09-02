"""
文档渲染器 - 将 PDF/PPT/图片 转为可分析的图片
支持：
- PDF → 每页图片（PyMuPDF）
- PPTX → 每页图片（通过提取形状渲染或转PDF再截图）
- 图片 → 直接使用
- 扫描文档（图片PDF）→ 直接截图
"""
import os
import logging
import shutil
import tempfile
import uuid
from pathlib import Path
from typing import List, Optional

from .vision_models import ImageInput, DocumentPage


logger = logging.getLogger(__name__)


class DocumentRenderer:
    """文档渲染器"""

    # 支持的图片格式
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tiff", ".webp"}
    PDF_EXTENSIONS = {".pdf"}
    PPT_EXTENSIONS = {".pptx"}

    def __init__(self, dpi: int = 200, max_pages: int = 50,
                 output_dir: Optional[str] = None):
        """
        Args:
            dpi: 渲染 DPI（越高越清晰，但图片越大）
            max_pages: 最大渲染页数
            output_dir: 输出目录（默认临时目录）
        """
        self.dpi = dpi
        self.max_pages = max_pages
        self._owns_output_dir = output_dir is None
        self.output_dir = output_dir or tempfile.mkdtemp(prefix="vision_")
        os.makedirs(self.output_dir, exist_ok=True)

    def close(self):
        """清理由渲染器自行创建的临时目录。"""
        if self._owns_output_dir and self.output_dir:
            shutil.rmtree(self.output_dir, ignore_errors=True)
            self.output_dir = ""

    def __enter__(self):
        return self

    def __exit__(self, _exc_type, _exc, _tb):
        self.close()

    def render(self, file_path: str) -> List[DocumentPage]:
        """
        渲染文档为图片页面列表

        Args:
            file_path: 文件路径

        Returns:
            DocumentPage 列表
        """
        ext = Path(file_path).suffix.lower()

        if ext in self.IMAGE_EXTENSIONS:
            return self._render_image(file_path)
        elif ext in self.PDF_EXTENSIONS:
            return self._render_pdf(file_path)
        elif ext in self.PPT_EXTENSIONS:
            return self._render_pptx(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    def render_to_images(self, file_path: str) -> List[ImageInput]:
        """渲染并返回 ImageInput 列表"""
        pages = self.render(file_path)
        return [p.image for p in pages if p.image]

    def _render_image(self, file_path: str) -> List[DocumentPage]:
        """单张图片"""
        img = ImageInput.from_file(file_path, page_number=1, label=Path(file_path).name)
        return [DocumentPage(
            page_number=1,
            image=img,
            text_hint=Path(file_path).stem,
        )]

    def _render_pdf(self, file_path: str) -> List[DocumentPage]:
        """PDF 渲染为图片"""
        import pymupdf

        pages = []
        doc = pymupdf.open(file_path)
        try:
            zoom = self.dpi / 72  # 72 是 PDF 默认 DPI
            matrix = pymupdf.Matrix(zoom, zoom)
            total = min(len(doc), self.max_pages)

            for i in range(total):
                page = doc[i]
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                img_path = os.path.join(self.output_dir, f"page_{i+1:04d}.png")
                pix.save(img_path)
                text = page.get_text().strip()[:500]
                img = ImageInput.from_file(img_path, page_number=i+1)
                img.width = pix.width
                img.height = pix.height
                pages.append(DocumentPage(
                    page_number=i + 1,
                    image=img,
                    width=pix.width,
                    height=pix.height,
                    text_hint=text,
                ))
        finally:
            doc.close()
        return pages

    def _render_pptx(self, file_path: str) -> List[DocumentPage]:
        """
        PPTX 渲染为图片

        策略：
        1. 优先尝试 LibreOffice 转 PDF 再渲染
        2. 如果没有 LibreOffice，使用 python-pptx 提取文本+形状信息作为提示
        """
        # 尝试 LibreOffice 转 PDF
        pdf_path = self._convert_pptx_to_pdf(file_path)
        if pdf_path and os.path.exists(pdf_path):
            try:
                return self._render_pdf(pdf_path)
            finally:
                try:
                    os.unlink(pdf_path)
                except OSError:
                    pass

        # 降级：提取每页文本信息
        return self._extract_pptx_text_pages(file_path)

    def _convert_pptx_to_pdf(self, file_path: str) -> Optional[str]:
        """使用 LibreOffice 将 PPTX 转为 PDF"""
        import subprocess

        # 查找 LibreOffice
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            # Windows 常见路径
            win_paths = [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            ]
            for p in win_paths:
                if os.path.exists(p):
                    soffice = p
                    break

        if not soffice:
            return None

        conversion_dir = tempfile.mkdtemp(prefix="lo-convert-", dir=self.output_dir)
        profile_dir = tempfile.mkdtemp(prefix="lo-profile-")
        try:
            converted_pdf = os.path.join(conversion_dir, Path(file_path).stem + ".pdf")
            out_pdf = os.path.join(
                self.output_dir, f"{Path(file_path).stem}-{uuid.uuid4().hex[:8]}.pdf"
            )
            cmd = [
                soffice, "--headless",
                f"-env:UserInstallation={Path(profile_dir).resolve().as_uri()}",
                "--convert-to", "pdf",
                "--outdir", conversion_dir,
                file_path
            ]
            result = subprocess.run(
                cmd, capture_output=True, timeout=60,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            )
            if result.returncode == 0 and os.path.exists(converted_pdf):
                shutil.move(converted_pdf, out_pdf)
                return out_pdf
            output = result.stderr or result.stdout or b""
            logger.warning("LibreOffice 转换失败: %s", output.decode(errors="replace"))
        except Exception as exc:
            logger.warning("LibreOffice 转换异常: %s", exc, exc_info=True)
        finally:
            shutil.rmtree(conversion_dir, ignore_errors=True)
            shutil.rmtree(profile_dir, ignore_errors=True)

        return None

    def _extract_pptx_text_pages(self, file_path: str) -> List[DocumentPage]:
        """降级方案：从 PPTX 提取文本（无图片渲染）"""
        from pptx import Presentation

        prs = Presentation(file_path)
        pages = []

        for i, slide in enumerate(prs.slides):
            if i >= self.max_pages:
                break

            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            texts.append(text)
                elif shape.has_table:
                    texts.append("[表格]")
                elif shape.shape_type == 13:  # Picture
                    texts.append("[图片]")

            # 生成一个占位图片（纯色背景+文字）
            img_path = self._create_placeholder_image(
                i + 1, texts,
                width=int(prs.slide_width / 914400 * 96),
                height=int(prs.slide_height / 914400 * 96),
            )

            img = ImageInput.from_file(img_path, page_number=i+1)
            pages.append(DocumentPage(
                page_number=i + 1,
                image=img,
                text_hint="\n".join(texts),
            ))

        return pages

    def _create_placeholder_image(self, page_num: int, texts: List[str],
                                   width: int = 960, height: int = 540) -> str:
        """创建占位图片（用于无法渲染的PPT）"""
        from PIL import Image, ImageDraw, ImageFont

        img = Image.new("RGB", (width, height), "#FFFFFF")
        draw = ImageDraw.Draw(img)

        # 标题
        try:
            font_title = ImageFont.truetype("msyh.ttc", 24)
            font_body = ImageFont.truetype("msyh.ttc", 14)
        except Exception:
            font_title = ImageFont.load_default()
            font_body = ImageFont.load_default()

        draw.text((40, 30), f"第 {page_num} 页", fill="#1F4E79", font=font_title)
        draw.line([(40, 65), (width - 40, 65)], fill="#1F4E79", width=2)

        y = 80
        for text in texts[:20]:
            draw.text((50, y), f"• {text[:60]}", fill="#333333", font=font_body)
            y += 24
            if y > height - 40:
                break

        img_path = os.path.join(self.output_dir, f"slide_{page_num:04d}.png")
        img.save(img_path)
        return img_path

    def render_images_from_bytes(self, data: bytes, filename: str = "") -> List[DocumentPage]:
        """从字节数据渲染"""
        ext = Path(filename).suffix.lower() if filename else ".png"
        tmp_path = os.path.join(self.output_dir, f"input-{uuid.uuid4().hex}{ext}")
        with open(tmp_path, "wb") as f:
            f.write(data)
        return self.render(tmp_path)
