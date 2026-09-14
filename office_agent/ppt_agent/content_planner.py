"""
Content Planner - PPT 内容规划器

负责：
1. 从主题生成 PPT 大纲（封面、目录、章节、内容、总结）
2. 从文本/Word 文档提取内容并分配到各页
3. 可选 LLM 生成更丰富的内容
4. 每页结构：标题、要点、备注
"""
import re
import json
import logging
from pathlib import Path
from typing import Optional

from .models import (
    PPTOutline, SlideContent, SlideLayout, items_per_page,
)

logger = logging.getLogger("office_agent.ppt.content_planner")


# ==========================================
# 主题 → 大纲模板
# ==========================================

THEME_TEMPLATES = {
    "项目汇报": [
        ("cover", "项目汇报", "{subtitle}"),
        ("toc", "目录", ""),
        ("section", "项目背景", ""),
        ("content", "背景与目标", ["项目立项背景", "核心目标与KPI", "预期成果"]),
        ("content", "现状分析", ["当前进展", "存在问题", "面临挑战"]),
        ("section", "实施方案", ""),
        ("content", "技术方案", ["技术架构", "核心模块", "关键技术选型"]),
        ("content", "实施计划", ["里程碑节点", "资源分配", "时间安排"]),
        ("data_cards", "关键数据", []),
        ("section", "成果与展望", ""),
        ("content", "阶段成果", ["已完成工作", "核心产出", "价值体现"]),
        ("summary", "总结与展望", ["感谢支持", "Q&A"]),
    ],
    "产品介绍": [
        ("cover", "产品介绍", "{subtitle}"),
        ("toc", "目录", ""),
        ("content", "产品概述", ["产品定位", "目标用户", "核心价值"]),
        ("content", "市场背景", ["市场规模", "用户痛点", "竞品分析"]),
        ("content_image", "核心功能", []),  # 占位文案不得写入用户成品
        ("two_column", "产品优势", []),
        ("data_cards", "核心数据", []),
        ("content", "应用场景", []),  # 占位文案不得写入用户成品
        ("timeline", "发展规划", []),
        ("summary", "感谢聆听", []),
    ],
    "工作总结": [
        ("cover", "工作总结", "{subtitle}"),
        ("toc", "目录", ""),
        ("content", "工作概述", ["工作背景", "职责范围", "总体评价"]),
        ("content", "重点工作", []),  # 占位文案不得写入用户成品
        ("data_cards", "工作成果", []),
        ("content", "经验总结", ["成功经验", "不足之处", "改进方向"]),
        ("content", "下一步计划", ["工作计划", "个人成长", "资源需求"]),
        ("summary", "感谢", []),
    ],
    "培训课件": [
        ("cover", "培训课件", "{subtitle}"),
        ("toc", "课程大纲", ""),
        ("section", "基础知识", ""),
        ("content", "概念介绍", ["核心定义", "基本原理", "关键术语"]),
        ("content", "理论框架", ["理论体系", "核心模型", "发展历程"]),
        ("section", "实践应用", ""),
        ("content", "方法步骤", ["操作流程", "注意事项", "常见问题"]),
        ("content_list", "案例分析", []),
        ("two_column", "要点对比", []),
        ("content", "总结回顾", ["核心要点", "学习建议", "延伸阅读"]),
        ("summary", "Q&A", []),
    ],
    "通用": [
        ("cover", "{title}", "{subtitle}"),
        ("toc", "目录", ""),
        ("content", "背景介绍", ["背景概述", "现状分析", "核心问题"]),
        ("content", "核心内容", []),  # 占位文案不得写入用户成品
        ("content", "详细说明", ["具体内容", "关键细节", "实施路径"]),
        ("data_cards", "数据展示", []),
        ("content", "总结", ["主要结论", "后续计划", "行动建议"]),
        ("summary", "感谢聆听", []),
    ],
}

# 缺少真实输入时保持为空。示例数字/时间线绝不能进入用户成品。
DEFAULT_DATA: list = []
DEFAULT_TIMELINE: list = []


class ContentPlanner:
    """
    PPT 内容规划器

    用法:
        planner = ContentPlanner()
        outline = planner.plan_from_theme("AI项目汇报", slide_count=10)
        outline = planner.plan_from_text(text_content)
        outline = planner.plan_from_word("document.docx")
    """

    def __init__(self, model_gateway=None):
        self.model_gateway = model_gateway

    @staticmethod
    def _enforce_slide_budget(outline: PPTOutline, budget: int) -> PPTOutline:
        """把大纲裁剪到预算页数：保留首页（封面）与末页（总结/致谢），
        从中间超量部分截断；页数不足不伪造内容。"""
        try:
            budget = max(2, int(budget))
        except (TypeError, ValueError):
            budget = 10
        slides = outline.slides
        if len(slides) <= budget:
            # 页数仍受绝对上限保护
            if len(slides) > 50:
                outline.slides = slides[:49] + [slides[-1]]
                for i, s in enumerate(outline.slides):
                    s.page_number = i + 1
            return outline
        keep_head = budget - 1  # 留 1 页给总结
        head = slides[:keep_head]
        tail = slides[-1:]
        outline.slides = head + tail
        for i, s in enumerate(outline.slides):
            s.page_number = i + 1
        return outline

    def plan_from_theme(self, theme: str,
                        slide_count: int = 10,
                        style: str = "professional",
                        subtitle: str = "",
                        author: str = "") -> PPTOutline:
        """
        从主题生成 PPT 大纲

        优先使用 LLM 生成真实内容；LLM 不可用时回退到模板。
        """
        # 0. 标量净化：上游（Agent 指令解析 / API 请求）可能传来 list/dict/None，
        #    统一走既有 _scalar_text，避免非字符串进入 PPTOutline 后被渲染成
        #    Python repr（如 "['a', 'b']"）写进成品。
        theme = self._scalar_text(theme)
        subtitle = self._scalar_text(subtitle)
        author = self._scalar_text(author)

        # 1. 清理标题：从用户指令中提取真正的主题
        clean_title = self._extract_title(theme)

        outline = PPTOutline(
            title=clean_title,
            subtitle=subtitle,
            author=author,
            theme=style,
        )

        # 2. 尝试用 LLM 生成完整大纲
        ai_outline = None
        self._ai_failure_reason = ""
        if self.model_gateway:
            ai_outline = self._generate_with_ai(clean_title, slide_count, style)

        if ai_outline and ai_outline.slides:
            # 页数预算强制：LLM 无视 slide_count 要求时以用户要求为准
            # （保留封面与总结页，裁掉中间超量页），杜绝"要 10 页给 30 页"
            return self._enforce_slide_budget(ai_outline, slide_count)

        # 3. LLM 失败/未配置时回退模板，保证有产出，同时用 used_template 标记避免冒充 AI
        if self.model_gateway:
            logger.warning("LLM生成PPT大纲失败，回退模板生成: %s", clean_title)
        else:
            logger.warning(f"未配置LLM，使用模板生成: {clean_title}")
        outline = self._generate_from_template(clean_title, slide_count, style, subtitle, author)
        outline.used_template = True
        # P3-41: 回退原因不再与“LLM 空响应”混同，显式记录到 changes
        if self.model_gateway:
            reason_text = {
                "length_truncated": "LLM 输出被 max_tokens 截断、JSON 不完整",
                "empty": "LLM 返回为空",
                "invalid_json": "LLM 返回的内容不是合法 JSON",
                "error": "LLM 生成过程出错",
            }.get(self._ai_failure_reason, "LLM 生成失败")
            outline.changes.append(f"已回退模板生成：{reason_text}")
        return outline

    def plan_from_text(self, text: str,
                       style: str = "professional",
                       title: str = "",
                       slide_count: int | None = None) -> PPTOutline:
        """
        从文本内容生成 PPT 大纲

        支持格式：
        - # 标题 → 封面
        - ## 章节名 → 章节页
        - ### 页标题 → 内容页
        - 普通段落 → 要点
        - - 或 • → 列表项
        """
        # 标量净化：文本与标题可能来自不可信输入，非字符串不进入大纲
        text = self._scalar_text(text)
        title = self._scalar_text(title)
        outline = PPTOutline(
            title=title or "演示文稿",
            theme=style,
        )

        lines = text.strip().split("\n")
        current_slide = None
        current_bullets: list[str] = []

        for line in lines:
            line = line.rstrip()
            if not line.strip():
                continue

            # Markdown 标题
            if line.startswith("# ") and not line.startswith("## "):
                # 封面
                if current_slide:
                    current_slide.bullets = current_bullets
                    outline.add_slide(current_slide)
                cover = SlideContent(
                    layout="cover",
                    title=line[2:].strip(),
                    subtitle=title,
                )
                outline.add_slide(cover)
                current_slide = None
                current_bullets = []
                continue

            if line.startswith("## "):
                # 章节页
                if current_slide:
                    current_slide.bullets = current_bullets
                    outline.add_slide(current_slide)
                section = SlideContent(
                    layout="section",
                    title=line[3:].strip(),
                )
                outline.add_slide(section)
                current_slide = None
                current_bullets = []
                continue

            if line.startswith("### "):
                # 内容页
                if current_slide:
                    current_slide.bullets = current_bullets
                    outline.add_slide(current_slide)
                current_slide = SlideContent(
                    layout="content",
                    title=line[4:].strip(),
                )
                current_bullets = []
                continue

            # 列表项
            if re.match(r"^[-*•·]\s+", line):
                bullet = re.sub(r"^[-*•·]\s+", "", line).strip()
                current_bullets.append(bullet)
                continue

            # 数字列表
            if re.match(r"^\d+[.、)]\s*", line):
                bullet = re.sub(r"^\d+[.、)]\s*", "", line).strip()
                current_bullets.append(bullet)
                continue

            # 普通文本 → 如果有当前页，作为要点；否则新建内容页
            if current_slide:
                if line.strip():
                    current_bullets.append(line.strip())
            else:
                current_slide = SlideContent(
                    layout="content",
                    title=line.strip()[:30],
                )

        # 最后一页
        if current_slide:
            current_slide.bullets = current_bullets
            outline.add_slide(current_slide)

        # 确保有封面
        if not outline.slides or outline.slides[0].layout != "cover":
            cover = SlideContent(layout="cover", title=outline.title)
            outline.slides.insert(0, cover)

        # 添加总结页
        has_summary = any(s.layout == "summary" for s in outline.slides)
        if not has_summary:
            outline.add_slide(SlideContent(layout="summary", title="感谢聆听"))

        # 重新编号
        for i, s in enumerate(outline.slides):
            s.page_number = i + 1

        if slide_count:
            outline = self._enforce_slide_budget(outline, slide_count)
        return outline

    def plan_from_word(self, docx_path: str,
                       style: str = "professional",
                       slide_count: int | None = None) -> PPTOutline:
        """
        从 Word 文档生成 PPT 大纲

        利用已有的 DocumentStructureAnalyzer 分析文档结构，
        将章节标题转为 PPT 页面。
        """
        from docx import Document
        from ..parsers.document_structure import DocumentStructureAnalyzer

        doc = Document(docx_path)
        analyzer = DocumentStructureAnalyzer()
        tree = analyzer.analyze(doc)

        outline = PPTOutline(
            title=tree.title or Path(docx_path).stem,
            theme=style,
        )

        # 封面
        outline.add_slide(SlideContent(
            layout="cover",
            title=tree.title or Path(docx_path).stem,
        ))

        # 目录
        chapters = tree.get_chapters()
        if chapters:
            toc_items = [c.text for c in chapters]
            outline.add_slide(SlideContent(
                layout="toc",
                title="目录",
                bullets=toc_items,
            ))

        # 按章节组织内容
        for chapter in chapters:
            # 章节过渡页
            outline.add_slide(SlideContent(
                layout="section",
                title=chapter.text,
            ))

            # 收集章节下的内容
            sections = chapter.find_children("section")
            if sections:
                for section in sections:
                    bullets = []
                    # 收集子节点文本
                    for child in section.children:
                        if child.type == "paragraph" and child.text:
                            text = child.text.strip()
                            if len(text) > 5:
                                # 不再 [:80] 硬截断正文，长句由渲染/自适应字号处理（P3-40）
                                bullets.append(text)
                        elif child.type == "list_item" and child.text:
                            bullets.append(child.text.strip())

                    if not bullets:
                        # 文档里确实没有可提取的正文时，不伪造「XX相关内容」
                        # 这类占位文本冒充用户内容；保留空要点页，由预检
                        # 以结构化提示（内容较少）呈现，属受控缺失状态。
                        logger.warning(
                            "章节 %r 未提取到正文，保留空内容页（不生成占位文本）",
                            section.text,
                        )

                    outline.add_slide(SlideContent(
                        layout="content",
                        title=section.text,
                        # 容量口径与渲染器一致（content=8），超出由 fix_outline 续页（P3-40/47）
                        bullets=bullets[:items_per_page("content")],
                    ))
            else:
                # 没有子节，直接用章节内容
                paragraphs = chapter.find_children("paragraph")
                bullets = [p.text.strip() for p in paragraphs
                           if p.text and len(p.text.strip()) > 10
                           ][:items_per_page("content")]
                if not bullets:
                    logger.warning(
                        "章节 %r 未提取到正文，保留空内容页（不生成占位文本）",
                        chapter.text,
                    )
                outline.add_slide(SlideContent(
                    layout="content",
                    title=chapter.text,
                    bullets=bullets,
                ))

        # 总结
        outline.add_slide(SlideContent(layout="summary", title="感谢聆听"))

        if slide_count:
            outline = self._enforce_slide_budget(outline, slide_count)
        return outline

    def plan_from_outline_data(self, title: str, slides_data: list,
                               style: str = "professional") -> PPTOutline:
        """
        从结构化数据生成大纲

        slides_data: [{"layout": "content", "title": "...", "bullets": [...]}]
        """
        title = self._scalar_text(title)
        outline = PPTOutline(title=title, theme=style)

        for data in slides_data:
            if not isinstance(data, dict):
                logger.warning("跳过非字典的大纲项: %r", type(data).__name__)
                continue
            slide = SlideContent(
                layout=data.get("layout", "content"),
                title=self._scalar_text(data.get("title")),
                subtitle=self._scalar_text(data.get("subtitle")),
                bullets=self._normalize_str_list(data.get("bullets")),
                body_text=self._scalar_text(data.get("body_text")),
                left_content=self._normalize_str_list(data.get("left_content")),
                right_content=self._normalize_str_list(data.get("right_content")),
                image_path=self._scalar_text(data.get("image_path")),
                image_alt=self._scalar_text(data.get("image_alt")),
                image_prompt=self._scalar_text(data.get("image_prompt")),
                data=data.get("data", []),
                table_data=self._normalize_table(data.get("table_data")),
                table_header=data.get("table_header", True),
                chart_type=self._scalar_text(data.get("chart_type")) or "bar",
                chart_title=self._scalar_text(data.get("chart_title")),
                chart_categories=self._normalize_str_list(data.get("chart_categories")),
                chart_series=data.get("chart_series", []),
                quote_text=self._scalar_text(data.get("quote_text")),
                quote_source=self._scalar_text(data.get("quote_source")),
                timeline_items=data.get("timeline_items", []),
                notes=self._scalar_text(data.get("notes")),
            )
            outline.add_slide(slide)

        return outline

    # ==========================================
    # 辅助方法
    # ==========================================

    def _match_template(self, theme: str) -> list:
        """匹配最接近的模板"""
        theme_lower = theme.lower()

        # 关键词匹配
        keywords = {
            "项目汇报": ["项目", "汇报", "进展", "里程碑"],
            "产品介绍": ["产品", "介绍", "发布", "功能"],
            "工作总结": ["总结", "工作", "年度", "季度", "复盘"],
            "培训课件": ["培训", "课件", "教程", "课程", "学习"],
        }

        for template_name, kws in keywords.items():
            for kw in kws:
                if kw in theme_lower:
                    return THEME_TEMPLATES[template_name]

        return THEME_TEMPLATES["通用"]

    def _adjust_slide_count(self, outline: PPTOutline, target: int) -> PPTOutline:
        """调整页数到目标数量"""
        current = len(outline.slides)

        if current == target:
            return outline

        if current < target:
            # 内容不足时宁可少页，也不为满足页数伪造“补充内容/要点”。
            return outline
        else:
            essential_indices = [
                i for i, slide in enumerate(outline.slides)
                if slide.layout in ("cover", "toc", "section", "summary")
            ]
            required_indices = set(essential_indices)
            # 每个章节页至少保留其后的一个内容页，防止出现“只有章节封面”。
            for section_index in [
                i for i, slide in enumerate(outline.slides) if slide.layout == "section"
            ]:
                for follower in range(section_index + 1, current):
                    if outline.slides[follower].layout in ("section", "summary"):
                        break
                    if outline.slides[follower].layout not in ("cover", "toc"):
                        required_indices.add(follower)
                        break
            # 目标小于结构完整所需页数时宁可超出，也不截掉总结或孤立章节。
            if target < len(required_indices):
                selected = required_indices
            else:
                removable_indices = [
                    i for i in range(current) if i not in required_indices
                ]
                keep_count = min(target - len(required_indices), len(removable_indices))
                chosen = {
                    removable_indices[(i * len(removable_indices)) // keep_count]
                    for i in range(keep_count)
                } if keep_count else set()
                selected = required_indices | chosen
            # 只按索引筛选，保持章节页与内容页的原始相对顺序。
            outline.slides = [
                slide for i, slide in enumerate(outline.slides) if i in selected
            ]

        # 重新编号
        for i, s in enumerate(outline.slides):
            s.page_number = i + 1

        return outline

    # ==========================================
    # 新的智能内容生成方法
    # ==========================================

    def _extract_title(self, raw: str) -> str:
        """从用户指令中提取干净的PPT标题"""
        if not raw:
            return "演示文稿"
        title = raw.strip()
        # 去掉常见前缀
        for prefix in ["生成一份", "生成一个", "帮我生成", "做一个", "做一份", "创建一个",
                        "帮我做", "请生成", "请帮我", "生成", "做", "创建", "写一份", "写一个"]:
            if title.startswith(prefix):
                title = title[len(prefix):].strip()
                break
        # 去掉常见后缀
        for suffix in ["的PPT", "的ppt", "PPT", "ppt", "演示文稿", "幻灯片",
                        "相关内容联网搜索", "联网搜索", "相关内容"]:
            if title.endswith(suffix):
                title = title[:-len(suffix)].strip()
        # 去掉"关于"前缀
        if title.startswith("关于"):
            title = title[2:].strip()
        # 长指令在首个冒号/破折号之后是分页与内容说明，标题只取其前的短主题
        for sep in ("：", ":", "——"):
            if sep in title:
                head = title.split(sep, 1)[0].strip()
                if 2 <= len(head) <= 40:
                    title = head
                    break
        # 去掉开头页数限定，如 "8 页的 AI 行业介绍" -> "AI 行业介绍"
        page_match = re.match(r"^\d+\s*页(?:的)?\s*(.+)$", title)
        if page_match and page_match.group(1).strip():
            title = page_match.group(1).strip()
        # 截断后可能仍残留尾部 PPT/演示文稿等词，再剥离一次
        for tail in ("的PPT", "的ppt", "PPT", "ppt", "演示文稿", "幻灯片"):
            if title.endswith(tail) and len(title) > len(tail):
                title = title[: -len(tail)].strip()
        return title if title else raw.strip()

    def _generate_with_ai(self, title: str, slide_count: int, style: str) -> Optional[PPTOutline]:
        """使用LLM生成完整PPT大纲"""
        try:
            prompt = f"""请为主题"{title}"生成一份{slide_count}页的PPT大纲，风格为{style}。

要求：
1. 每页包含真实、具体、有信息量的内容，不要使用占位符
2. 内容要紧扣主题，有实际价值
3. 每页的要点要具体、详实，不要写"要点一"这种空话

请直接输出 JSON 结果，不要输出任何思考过程、推理或解释。严格按以下JSON格式返回，不要返回其他内容：
{{
  "subtitle": "副标题（一句话概括）",
  "slides": [
    {{
      "layout": "cover",
      "title": "{title}",
      "subtitle": "副标题"
    }},
    {{
      "layout": "toc",
      "title": "目录",
      "bullets": ["目录项1", "目录项2", "目录项3", "目录项4", "目录项5"]
    }},
    {{
      "layout": "content",
      "title": "具体的页面标题",
      "bullets": ["具体要点1（有实际内容）", "具体要点2（有实际内容）", "具体要点3（有实际内容）"]
    }},
    {{
      "layout": "data_cards",
      "title": "关键数据",
      "data": [["指标名称1", "数值", "单位"], ["指标名称2", "数值", "单位"], ["指标名称3", "数值", "单位"], ["指标名称4", "数值", "单位"]]
    }},
    {{
      "layout": "summary",
      "title": "总结",
      "bullets": ["总结要点1", "总结要点2", "展望"]
    }}
  ]
}}

可用layout类型: cover(封面), toc(目录), section(章节页), content(内容页，可用body_text放整段正文或bullets放要点), content_image(图文页，需要配图，附image_prompt图片描述词), two_column(两栏), content_list(要点列表), data_cards(数据卡片，附data数组), table(表格页，附table_data二维数组，第一行为表头), timeline(时间线), chart(图表页，附chart_type、chart_categories数组、chart_series系列数组), quote(金句页，附quote_text引用文字与quote_source来源), summary(总结)

当某一页内容更适合用图片表达时，使用 content_image 布局，并提供一句具体的 image_prompt 描述画面（横向构图、商务风格、避免文字）。当需要展示结构化数据时，使用 table 布局并提供 table_data。请确保生成{slide_count}页左右，内容充实专业。"""

            result = self.model_gateway.chat(
                user_message=prompt,
                task_type_str="ppt_content",
                temperature=0.7,
                max_tokens=8192,
            )

            if not result or not result.success or not result.content:
                err = (result.error if result and getattr(result, "error", None) else "") or "unknown"
                logger.warning("LLM返回为空: %s", err)
                self._ai_failure_reason = "empty"
                return None

            # 解析JSON
            content = result.content.strip()
            # 去掉可能的markdown代码块标记
            fenced = re.search(
                r"```(?:json)?\s*(.*?)```", content,
                flags=re.IGNORECASE | re.DOTALL,
            )
            if fenced:
                content = fenced.group(1).strip()

            data = json.loads(content)
            return self._parse_ai_outline(data, title, style)

        except json.JSONDecodeError as e:
            # P3-41: 被 max_tokens 截断的 JSON 通常缺少闭合花括号，
            # 与“返回了乱码/非 JSON”区分开，便于上层显式标注回退原因。
            tail = content.rstrip()[-1:]
            self._ai_failure_reason = (
                "length_truncated" if tail != "}" else "invalid_json"
            )
            logger.warning("AI返回的JSON解析失败(%s): %s", self._ai_failure_reason, e)
            return None
        except Exception as e:
            self._ai_failure_reason = "error"
            logger.warning(f"AI生成大纲失败: {e}")
            return None

    @staticmethod
    def _scalar_text(value) -> str:
        """LLM 可能对任意字段返回 null/数字/嵌套结构，统一转为可渲染文本"""
        if value is None:
            return ""
        if isinstance(value, str):
            return value
        if isinstance(value, (dict, list)):
            import json as _json
            try:
                return _json.dumps(value, ensure_ascii=False)
            except (TypeError, ValueError):
                return str(value)
        return str(value)

    @classmethod
    def _normalize_str_list(cls, value) -> list:
        """bullets/left/right 列表规范化：全部元素转为字符串"""
        if not isinstance(value, list):
            return []
        out = []
        for item in value:
            if isinstance(item, str):
                out.append(item)
            elif isinstance(item, dict) and isinstance(item.get("text"), str):
                out.append(item["text"])
            elif item is None:
                continue
            else:
                out.append(cls._scalar_text(item))
        return out

    @classmethod
    def _normalize_table(cls, value) -> list:
        """table_data 规范化：二维字符串列表"""
        if not isinstance(value, list):
            return []
        rows = []
        for row in value:
            if isinstance(row, (list, tuple)) and len(row) > 0:
                rows.append([cls._scalar_text(c) for c in row])
            elif isinstance(row, str):
                rows.append([row])
        return rows

    def _parse_ai_outline(self, data: dict, title: str, style: str) -> PPTOutline:
        """解析AI返回的JSON为PPTOutline（字段类型不可信，全部规范化）"""
        outline = PPTOutline(
            title=title,
            subtitle=self._scalar_text(data.get("subtitle")),
            theme=style,
        )

        slides_raw = data.get("slides")
        if not isinstance(slides_raw, list):
            slides_raw = []
        for slide_data in slides_raw:
            if not isinstance(slide_data, dict):
                continue
            layout = slide_data.get("layout", "content")
            if not isinstance(layout, str):
                layout = "content"
            elif not SlideLayout.is_valid(layout):
                # LLM 输出的不可信版式：显式记录并规范化为标准内容页，
                # 而不是静默按 content 渲染（构造点 SlideContent 是严格校验的）
                logger.warning("LLM 返回未知版式 %r，已规范化为 content", layout)
                layout = SlideLayout.CONTENT.value
            chart_series_raw = slide_data.get("chart_series")
            slide = SlideContent(
                layout=layout,
                title=self._scalar_text(slide_data.get("title")),
                subtitle=self._scalar_text(slide_data.get("subtitle")),
                bullets=self._normalize_str_list(slide_data.get("bullets")),
                body_text=self._scalar_text(slide_data.get("body_text")),
                image_prompt=self._scalar_text(slide_data.get("image_prompt")),
                table_data=self._normalize_table(slide_data.get("table_data")),
                chart_type=self._scalar_text(slide_data.get("chart_type")) or "bar",
                chart_title=self._scalar_text(slide_data.get("chart_title")),
                chart_categories=self._normalize_str_list(slide_data.get("chart_categories")),
                chart_series=chart_series_raw if isinstance(chart_series_raw, list) else [],
                quote_text=self._scalar_text(slide_data.get("quote_text")),
                quote_source=self._scalar_text(slide_data.get("quote_source")),
                notes=self._scalar_text(slide_data.get("notes")),
            )

            # 数据卡片：[(label, value, unit), ...]
            if layout == "data_cards" and isinstance(slide_data.get("data"), list):
                cards = []
                for d in slide_data["data"]:
                    if isinstance(d, (list, tuple)) and len(d) >= 2:
                        cards.append(tuple(str(x) if x is not None else "" for x in d[:3]))
                    elif isinstance(d, dict):
                        cards.append((self._scalar_text(d.get("label")),
                                      self._scalar_text(d.get("value")),
                                      self._scalar_text(d.get("unit"))))
                slide.data = cards
            elif layout == "data_cards":
                slide.data = []

            # 时间线
            if layout == "timeline" and isinstance(slide_data.get("timeline"), list):
                items = []
                for t in slide_data["timeline"]:
                    if isinstance(t, (list, tuple)) and len(t) >= 2:
                        items.append(tuple(str(x) if x is not None else "" for x in t[:3]))
                    elif isinstance(t, dict):
                        items.append((self._scalar_text(t.get("time")),
                                      self._scalar_text(t.get("title")),
                                      self._scalar_text(t.get("desc"))))
                slide.timeline_items = items
            elif layout == "timeline":
                slide.timeline_items = []

            # 两栏内容
            left = self._normalize_str_list(slide_data.get("left_content"))
            right = self._normalize_str_list(slide_data.get("right_content"))
            if left:
                slide.left_content = left
            if right:
                slide.right_content = right

            outline.add_slide(slide)

        # 确保至少有封面和总结
        if not outline.slides:
            return self._generate_from_template(title, 10, style, "", "")

        if outline.slides[0].layout != "cover":
            cover = SlideContent(layout="cover", title=title,
                                 subtitle=self._scalar_text(data.get("subtitle")))
            outline.slides.insert(0, cover)

        has_summary = any(s.layout == "summary" for s in outline.slides)
        if not has_summary:
            outline.add_slide(SlideContent(layout="summary", title="感谢聆听",
                                           bullets=["谢谢观看", "Q&A"]))

        # 重新编号
        for i, s in enumerate(outline.slides):
            s.page_number = i + 1

        return outline

    def _generate_from_template(self, theme: str, slide_count: int,
                                 style: str, subtitle: str, author: str) -> PPTOutline:
        """模板回退方案（原逻辑）"""
        theme = self._scalar_text(theme)
        subtitle = self._scalar_text(subtitle)
        author = self._scalar_text(author)
        template = self._match_template(theme)

        outline = PPTOutline(
            title=theme,
            subtitle=subtitle,
            author=author,
            theme=style,
        )

        for layout, title, bullets in template:
            if layout == "cover" and theme:
                # 封面永远展示用户主题：命名模板的封面标题是硬编码的"项目汇报"等，
                # 只替换 {title} 占位会让用户的真实题目在成品封面丢失（RC）。
                actual_title = theme
            else:
                actual_title = title.replace("{title}", theme).replace("{subtitle}", subtitle or theme)
            actual_bullets = []

            if isinstance(bullets, list):
                actual_bullets = bullets
            elif isinstance(bullets, str) and bullets:
                actual_bullets = [bullets]

            slide = SlideContent(
                layout=layout,
                title=actual_title,
                bullets=actual_bullets,
            )

            if layout == "data_cards" and not actual_bullets:
                slide.data = []
            elif layout == "timeline" and not actual_bullets:
                slide.timeline_items = []
            elif layout == "cover":
                slide.subtitle = subtitle or ""
                slide.notes = author or ""

            outline.add_slide(slide)

        if slide_count and len(outline.slides) != slide_count:
            outline = self._adjust_slide_count(outline, slide_count)

        return outline
