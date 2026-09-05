"""Deprecated compatibility adapter for the legacy ``SlidePlanner`` API.

Authoritative PPT planning lives in :class:`ContentPlanner` and is consumed by
:class:`PPTOrchestrator`.  This module only adapts the production
``PPTOutline`` contract to the legacy ``SlidePlan`` shape for callers that
still import the old names.  It contains no independent planning logic.
"""
from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import List

from ..text_encoding import TextDecodeError, read_text_file
from .content_planner import ContentPlanner
from .models import SlideContent, SlideLayout


@dataclass
class SlidePlanItem:
    """Legacy single-slide plan contract."""

    index: int
    layout: str
    title: str
    suggested_bullets: List[str] = field(default_factory=list)
    notes: str = ""
    importance: str = "normal"
    content_source: str = ""
    estimated_time: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class SlidePlan:
    """Legacy slide-plan envelope, mapped from ``PPTOutline``."""

    topic: str
    scene: str
    scene_name: str
    audience: str
    audience_name: str
    target_slides: int
    actual_slides: int
    style: str
    slides: List[SlidePlanItem] = field(default_factory=list)
    document_summary: str = ""
    key_points: List[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "meta": {
                "topic": self.topic,
                "scene": self.scene,
                "scene_name": self.scene_name,
                "audience": self.audience,
                "audience_name": self.audience_name,
                "target_slides": self.target_slides,
                "actual_slides": self.actual_slides,
                "style": self.style,
                "document_summary": self.document_summary,
                "key_points": self.key_points,
            },
            "slides": [slide.to_dict() for slide in self.slides],
        }

    def to_json(self, indent=2, ensure_ascii=False) -> str:
        return json.dumps(self.to_dict(), indent=indent, ensure_ascii=ensure_ascii)

    def to_outline_data(self) -> list:
        """Convert the legacy shape back to ``PPTOrchestrator`` input data."""
        result = []
        for slide in self.slides:
            item = {"layout": slide.layout, "title": slide.title}
            if slide.suggested_bullets:
                if slide.layout == SlideLayout.CONTENT_TWO_COL.value:
                    midpoint = len(slide.suggested_bullets) // 2
                    item["left_content"] = slide.suggested_bullets[:midpoint]
                    item["right_content"] = slide.suggested_bullets[midpoint:]
                elif slide.layout == SlideLayout.QUOTE.value:
                    item["quote_text"] = slide.suggested_bullets[0]
                elif slide.layout in {
                        SlideLayout.DATA_CARDS.value,
                        SlideLayout.TIMELINE.value,
                }:
                    item["layout"] = SlideLayout.CONTENT.value
                    item["bullets"] = slide.suggested_bullets[:4]
                else:
                    item["bullets"] = slide.suggested_bullets
            if slide.notes:
                item["notes"] = slide.notes
            result.append(item)
        return result


def _slide_bullets(slide: SlideContent) -> List[str]:
    if slide.bullets:
        return list(slide.bullets)
    if slide.layout == SlideLayout.CONTENT_TWO_COL.value:
        return [*slide.left_content, *slide.right_content]
    if slide.layout == SlideLayout.QUOTE.value:
        return [slide.quote_text] if slide.quote_text else []
    if slide.layout == SlideLayout.DATA_CARDS.value:
        return [
            " ".join(str(part) for part in card if part).strip()
            for card in slide.data
        ]
    if slide.layout == SlideLayout.TIMELINE.value:
        return [str(item[1]) for item in slide.timeline_items if len(item) > 1]
    if slide.layout == SlideLayout.TABLE.value and slide.table_data:
        return [str(value) for value in slide.table_data[0]]
    return []


def _outline_to_plan(outline, topic: str, scene: str, scene_name: str,
                     audience: str, audience_name: str) -> SlidePlan:
    slides = [
        SlidePlanItem(
            index=index,
            layout=slide.layout,
            title=slide.title,
            suggested_bullets=_slide_bullets(slide),
            notes=slide.notes,
        )
        for index, slide in enumerate(outline.slides, 1)
    ]
    return SlidePlan(
        topic=topic,
        scene=scene or "general",
        scene_name=scene_name or "通用",
        audience=audience or "general",
        audience_name=audience_name or "普通听众",
        target_slides=len(slides),
        actual_slides=len(slides),
        style=outline.theme or "professional",
        slides=slides,
    )


class SlidePlanner:
    """Deprecated adapter delegating all planning to ``ContentPlanner``."""

    def __init__(self):
        self._planner = ContentPlanner()

    def plan(self, topic: str,
             scene: str = "general",
             audience: str = "general",
             slide_count: int = 0,
             style: str = "",
             user_requirements: str = "") -> SlidePlan:
        outline = self._planner.plan_from_theme(
            theme=topic,
            slide_count=slide_count or 10,
            style=style or "professional",
        )
        return _outline_to_plan(
            outline, topic, scene, "", audience, "",
        )

    def plan_from_text(self, topic: str, text: str,
                       scene: str = "general",
                       audience: str = "general",
                       slide_count: int = 0,
                       style: str = "") -> SlidePlan:
        outline = self._planner.plan_from_text(
            text=text,
            style=style or "professional",
            title=topic,
        )
        return _outline_to_plan(
            outline, topic, scene, "", audience, "",
        )

    def plan_from_document(self, topic: str,
                           document_path: str,
                           scene: str = "general",
                           audience: str = "general",
                           slide_count: int = 0,
                           style: str = "") -> SlidePlan:
        path = Path(document_path)
        if not path.exists():
            raise FileNotFoundError(f"文档不存在: {path}")

        suffix = path.suffix.lower()
        if suffix == ".docx":
            outline = self._planner.plan_from_word(
                docx_path=str(path),
                style=style or "professional",
            )
        else:
            content = self._read_document(str(path))
            outline = self._planner.plan_from_text(
                text=content,
                style=style or "professional",
                title=topic,
            )
        return _outline_to_plan(
            outline, topic, scene, "", audience, "",
        )

    def _read_document(self, path: str) -> str:
        """Read legacy planner inputs; kept only for compatibility callers."""
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"文档不存在: {path}")

        suffix = path.suffix.lower()
        if suffix == ".docx":
            try:
                from docx import Document
                document = Document(str(path))
                return "\n".join(
                    paragraph.text.strip()
                    for paragraph in document.paragraphs
                    if paragraph.text.strip()
                )
            except Exception as exc:
                raise RuntimeError(f"Word 文档读取失败: {exc}") from exc
        if suffix in (".txt", ".md"):
            try:
                return read_text_file(path)
            except TextDecodeError as exc:
                raise RuntimeError(f"文本编码无法识别: {path.name}") from exc
        raise ValueError(f"不支持的文档格式: {suffix}")

    def generate_ppt(self, plan: SlidePlan, output_path: str = ""):
        from .ppt_orchestrator import PPTOrchestrator

        return PPTOrchestrator().generate_from_outline(
            title=plan.topic,
            slides_data=plan.to_outline_data(),
            style=plan.style,
            output_path=output_path or f"{plan.topic[:20]}.pptx",
        )
