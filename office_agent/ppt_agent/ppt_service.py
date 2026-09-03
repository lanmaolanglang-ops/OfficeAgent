"""
PPT Service - PPT 底层生成服务
封装 python-pptx，提供幻灯片创建、布局、样式等原子操作
"""
import logging
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn

from .models import (
    PPTOutline, SlideContent, ColorScheme, FontScheme,
    PPTGenerationResult,
)


logger = logging.getLogger(__name__)


# ==========================================
# 内置主题配色
# ==========================================

THEME_COLORS = {
    "professional": ColorScheme(
        primary="#1F4E79", secondary="#2E75B6", accent="#FFC000",
        bg="#FFFFFF", bg_dark="#1F4E79",
        text="#333333", text_light="#FFFFFF", text_muted="#666666",
        line="#D9D9D9",
    ),
    "minimal": ColorScheme(
        primary="#000000", secondary="#595959", accent="#C00000",
        bg="#FFFFFF", bg_dark="#000000",
        text="#000000", text_light="#FFFFFF", text_muted="#808080",
        line="#E0E0E0",
    ),
    "creative": ColorScheme(
        primary="#ED7D31", secondary="#FFC000", accent="#70AD47",
        bg="#FFFFFF", bg_dark="#ED7D31",
        text="#333333", text_light="#FFFFFF", text_muted="#666666",
        line="#F0D5B8",
    ),
    "tech": ColorScheme(
        primary="#0066CC", secondary="#00B0F0", accent="#00FF88",
        bg="#0A1628", bg_dark="#0066CC",
        text="#E0E0E0", text_light="#FFFFFF", text_muted="#999999",
        line="#1A3050",
    ),
    "academic": ColorScheme(
        primary="#2E4057", secondary="#048A81", accent="#D1495B",
        bg="#FFFFFF", bg_dark="#2E4057",
        text="#2E4057", text_light="#FFFFFF", text_muted="#6B7B8C",
        line="#D0D5DD",
    ),
    "nature": ColorScheme(
        primary="#2D6A4F", secondary="#40916C", accent="#D4A373",
        bg="#FFFFFF", bg_dark="#2D6A4F",
        text="#1B4332", text_light="#FFFFFF", text_muted="#52796F",
        line="#B7E4C7",
    ),
}

DEFAULT_FONTS = FontScheme()


def hex_to_rgb(hex_color: str) -> RGBColor:
    """十六进制颜色转 RGBColor（解析口径统一在 office_agent.colors）"""
    from ..colors import parse_hex_color

    r, g, b, _a = parse_hex_color(hex_color)
    return RGBColor(r, g, b)


class PPTService:
    """
    PPT 底层生成服务

    负责将 PPTOutline 渲染为 .pptx 文件。
    每种版式对应一个 _render_* 方法。
    """

    # 16:9 设计稿坐标 → 实际页面坐标的缩放因子（4:3 模板等场景）
    _sx: float = 1.0
    _sy: float = 1.0

    def __init__(self):
        self.prs: Optional[Presentation] = None
        self.colors: ColorScheme = THEME_COLORS["professional"]
        self.fonts: FontScheme = DEFAULT_FONTS
        self.changes: list = []
        self.template_config = None  # TemplateConfig，设置后遵循模板布局

    def _x(self, v: float):
        """设计稿 x/宽度 → 实际 Emu"""
        return Inches(v * self._sx)

    def _y(self, v: float):
        """设计稿 y/高度 → 实际 Emu"""
        return Inches(v * self._sy)

    def _surface_color(self, alternate: bool = False) -> str:
        """Return a readable card/table surface for both light and dark themes."""
        bg = self.colors.bg or "#FFFFFF"
        raw = bg.lstrip("#")
        try:
            red, green, blue = (int(raw[i:i + 2], 16) for i in (0, 2, 4))
            is_dark = (0.2126 * red + 0.7152 * green + 0.0722 * blue) < 128
        except (TypeError, ValueError):
            is_dark = False
        if is_dark:
            return self.colors.line if alternate else bg
        return "#F5F7FA" if alternate else "#FFFFFF"

    def generate(self, outline: PPTOutline, output_path: str) -> PPTGenerationResult:
        """
        根据大纲生成 PPTX 文件

        Args:
            outline: PPT 大纲
            output_path: 输出路径

        Returns:
            PPTGenerationResult
        """
        try:
            self.changes = []

            # 基底模板：以用户模板文件为容器（继承其母版/主题/背景/页面尺寸），
            # 只清空内容页。任何一步失败都回退到全新空白演示文稿。
            base_path = getattr(outline, '_base_template_path', None)
            self.prs = None
            self._base_deck = False
            if (base_path and Path(base_path).exists()
                    and Path(base_path).resolve() != Path(output_path).resolve()):
                try:
                    from pptx.oxml.ns import qn as _qn
                    base_prs = Presentation(base_path)
                    slide_ids = base_prs.slides._sldIdLst
                    for sld_id in list(slide_ids):
                        rel_id = sld_id.get(_qn('r:id'))
                        base_prs.part.drop_rel(rel_id)
                        slide_ids.remove(sld_id)
                    self.prs = base_prs
                    self._base_deck = True
                    # 页面尺寸以模板实际值为准（渲染缩放因子基于它计算）
                    outline.slide_width = self.prs.slide_width / 914400
                    outline.slide_height = self.prs.slide_height / 914400
                    self.changes.append(f"使用模板: {Path(base_path).name}（保留母版与主题）")
                except Exception as exc:
                    logger.warning("模板基底加载失败，回退空白演示文稿: %s", exc)
                    self.prs = None
                    self._base_deck = False
            if self.prs is None:
                self.prs = Presentation()

            # 模板配置（如果有）
            self.template_config = getattr(outline, '_template_config', None)

            # 页面尺寸（无基底模板时按 outline 设置）
            if not self._base_deck:
                self.prs.slide_width = Inches(outline.slide_width)
                self.prs.slide_height = Inches(outline.slide_height)
            # 渲染器均按 16:9 (13.333 x 7.5 英寸) 设计稿取坐标，
            # 模板尺寸不同（如 4:3）时等比缩放，避免元素越界出页
            self._sx = outline.slide_width / 13.333
            self._sy = outline.slide_height / 7.5

            # 配色（模板配置优先）
            if self.template_config:
                self.colors = self.template_config.get_color_scheme()
            elif outline.color_scheme:
                self.colors = outline.color_scheme
            elif outline.theme in THEME_COLORS:
                self.colors = THEME_COLORS[outline.theme]

            # 字体（模板配置优先）
            if self.template_config:
                self.fonts = self.template_config.get_font_scheme()
            else:
                self.fonts = outline.font_scheme or DEFAULT_FONTS

            # 渲染每一页
            for slide_content in outline.slides:
                self._render_slide(slide_content)

            self.prs.save(output_path)

            return PPTGenerationResult(
                success=True,
                message=f"PPT 生成成功，共 {len(outline.slides)} 页",
                output_path=output_path,
                slide_count=len(outline.slides),
                changes=self.changes,
            )

        except Exception as e:
            return PPTGenerationResult(
                success=False,
                message=f"PPT 生成失败: {str(e)}",
            )

    # ==========================================
    # 模板位置支持
    # ==========================================

    def _tpos(self, layout_type: str, region: str = "title"):
        """
        从模板配置获取位置（模板优先，无模板则用默认值）

        Args:
            layout_type: 版式类型 (cover/content/section/...)
            region: 区域 (title/content/picture)

        Returns:
            (left, top, width, height) 元组（英寸）
        """
        if not self.template_config:
            return None

        # 查找匹配的版式
        layout_info = self.template_config.find_layout(layout_type)
        if not layout_info:
            # 尝试模糊匹配
            for li in self.template_config.layouts:
                if layout_type in li.name.lower() or li.layout_type == layout_type:
                    layout_info = li
                    break

        if not layout_info:
            return None

        if region == "title" and layout_info.has_title:
            return (layout_info.title_left, layout_info.title_top,
                    layout_info.title_width, layout_info.title_height)
        elif region == "content" and layout_info.has_content:
            return (layout_info.content_left, layout_info.content_top,
                    layout_info.content_width, layout_info.content_height)
        elif region == "picture" and layout_info.has_picture:
            for ph in layout_info.placeholders:
                if ph.ph_type == "pic":
                    return (ph.left, ph.top, ph.width, ph.height)

        return None

    def _render_slide(self, content: SlideContent):
        """根据版式渲染单页"""
        layout = content.layout

        renderers = {
            "cover": self._render_cover,
            "toc": self._render_toc,
            "section": self._render_section,
            "content": self._render_content,
            "content_image": self._render_content_image,
            "two_column": self._render_two_column,
            "content_list": self._render_content_list,
            "data_cards": self._render_data_cards,
            "timeline": self._render_timeline,
            "quote": self._render_quote,
            "table": self._render_table,
            "chart": self._render_chart,
            "summary": self._render_summary,
            "blank": self._render_blank,
        }

        renderer = renderers.get(layout, self._render_content)
        renderer(content)
        self.changes.append(f"第{content.page_number}页: {content.title or layout}")

    # ==========================================
    # 版式渲染
    # ==========================================

    def _add_blank_slide(self) -> object:
        """添加空白幻灯片"""
        # 任意模板的版式集合不保证 [6] 是空白布局：优先选无占位符的版式
        layout = None
        for cand in self.prs.slide_layouts:
            try:
                if len(cand.placeholders) == 0:
                    layout = cand
                    break
            except Exception:
                continue
        if layout is None:
            layouts = list(self.prs.slide_layouts)
            layout = layouts[6] if len(layouts) > 6 else layouts[0]
        slide = self.prs.slides.add_slide(layout)
        # 基底模板时保留母版背景（覆盖填充会抹掉模板的背景/装饰）
        if not self._base_deck:
            bg = slide.background
            fill = bg.fill
            fill.solid()
            fill.fore_color.rgb = hex_to_rgb(self.colors.bg)
        return slide

    def _set_shape_bg(self, shape, color: str):
        """设置形状填充色"""
        shape.fill.solid()
        shape.fill.fore_color.rgb = hex_to_rgb(color)
        shape.line.fill.background()

    def _add_text_box(self, slide, left: float, top: float, width: float,
                      height: float, text: str, font_size: int = 18,
                      bold: bool = False, color: str = "",
                      alignment: str = "left", font_name: str = "",
                      anchor: str = "top") -> object:
        """添加文本框"""
        box = slide.shapes.add_textbox(
            self._x(left), self._y(top), self._x(width), self._y(height)
        )
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

        # 垂直对齐
        anchor_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        tf.vertical_anchor = anchor_map.get(anchor, MSO_ANCHOR.TOP)

        # 叶子层防御：None/数字等一律转字符串（p.text = None 会抛 TypeError）
        if not isinstance(text, str):
            text = "" if text is None else str(text)

        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = hex_to_rgb(color or self.colors.text)
        p.font.name = font_name or self.fonts.body_cn
        # 设置中文字体
        self._set_cn_font(p, font_name or self.fonts.body_cn)

        align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                     "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
        p.alignment = align_map.get(alignment, PP_ALIGN.LEFT)

        return box

    def _add_bullet_list(self, slide, left: float, top: float, width: float,
                         height: float, bullets: list, font_size: int = 18,
                         color: str = "", bullet_char: str = "•",
                         line_spacing: float = 1.5) -> object:
        """添加要点列表"""
        box = slide.shapes.add_textbox(
            self._x(left), self._y(top), self._x(width), self._y(height)
        )
        tf = box.text_frame
        tf.word_wrap = True

        for i, bullet in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()

            if isinstance(bullet, dict):
                raw_text = bullet.get("text", "")
                text = "" if raw_text is None else str(raw_text)
                try:
                    level = max(0, min(4, int(bullet.get("level", 0))))
                except (TypeError, ValueError):
                    level = 0
                sub_bullets = bullet.get("children", []) or []
            else:
                text = "" if bullet is None else str(bullet)
                level = 0
                sub_bullets = []

            p.text = f"{'  ' * level}{bullet_char} {text}"
            p.font.size = Pt(max(8, font_size - level * 2))
            p.font.color.rgb = hex_to_rgb(color or self.colors.text)
            p.font.name = self.fonts.body_cn
            self._set_cn_font(p, self.fonts.body_cn)
            p.space_after = Pt(8)
            p.line_spacing = line_spacing

            # 子要点
            for sub in sub_bullets:
                sp = tf.add_paragraph()
                sp.text = f"    ◦ {'' if sub is None else str(sub)}"
                sp.font.size = Pt(font_size - 2)
                sp.font.color.rgb = hex_to_rgb(self.colors.text_muted)
                sp.font.name = self.fonts.body_cn
                self._set_cn_font(sp, self.fonts.body_cn)
                sp.space_after = Pt(4)

        return box

    def _add_decorative_line(self, slide, left: float, top: float,
                             width: float, height: float = 0.04,
                             color: str = "") -> object:
        """添加装饰线"""
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            self._x(left), self._y(top), self._x(width), self._y(height)
        )
        self._set_shape_bg(line, color or self.colors.secondary)
        return line

    def _add_rectangle(self, slide, left: float, top: float, width: float,
                       height: float, color: str = "") -> object:
        """添加矩形色块"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            self._x(left), self._y(top), self._x(width), self._y(height)
        )
        self._set_shape_bg(shape, color or self.colors.primary)
        return shape

    def _add_rounded_rect(self, slide, left: float, top: float, width: float,
                          height: float, color: str = "") -> object:
        """添加圆角矩形"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            self._x(left), self._y(top), self._x(width), self._y(height)
        )
        self._set_shape_bg(shape, color or self.colors.primary)
        return shape

    def _add_page_number(self, slide, num: int):
        """添加页码"""
        self._add_text_box(
            slide, 12.0, 7.0, 1.0, 0.4,
            str(num), font_size=self.fonts.caption_size,
            color=self.colors.text_muted, alignment="right"
        )

    @staticmethod
    def _set_cn_font(run_or_para, font_name: str):
        """设置中文字体（通过 lxml，兼容 PPT drawingml）"""
        # 处理段落中的所有 run
        runs = run_or_para.runs if hasattr(run_or_para, 'runs') else [run_or_para]
        for run in runs:
            run.font.name = font_name
            rPr = run._r.get_or_add_rPr()
            # 设置东亚字体 a:ea
            ea = rPr.find(qn('a:ea'))
            if ea is None:
                ea = rPr.makeelement(qn('a:ea'), {})
                rPr.append(ea)
            ea.set('typeface', font_name)
            # 设置拉丁字体 a:latin
            latin = rPr.find(qn('a:latin'))
            if latin is None:
                latin = rPr.makeelement(qn('a:latin'), {})
                rPr.append(latin)
            latin.set('typeface', font_name)

    # ==========================================
    # 各版式具体实现
    # ==========================================

    def _render_cover(self, content: SlideContent):
        """封面页"""
        slide = self._add_blank_slide()

        # 深色背景
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(self.colors.bg_dark)

        # 装饰色块
        self._add_rectangle(slide, 0, 0, 0.15, 7.5, self.colors.accent)
        self._add_rectangle(slide, 0, 5.8, 13.333, 0.08, self.colors.accent)

        # 主标题
        self._add_text_box(
            slide, 1.5, 2.2, 10.333, 1.5,
            content.title, font_size=44, bold=True,
            color=self.colors.text_light, alignment="center",
            font_name=self.fonts.title_cn, anchor="middle"
        )

        # 副标题
        if content.subtitle:
            self._add_text_box(
                slide, 1.5, 3.8, 10.333, 0.8,
                content.subtitle, font_size=22,
                color=self.colors.text_light, alignment="center",
                anchor="middle"
            )

        # 作者/日期
        if content.notes:
            self._add_text_box(
                slide, 1.5, 5.0, 10.333, 0.5,
                content.notes, font_size=16,
                color=self.colors.text_muted, alignment="center"
            )

    def _render_toc(self, content: SlideContent):
        """目录页"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.8, 0.5, 5, 0.8,
            content.title or "目录", font_size=self.fonts.title_size,
            bold=True, color=self.colors.primary,
            font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.8, 1.3, 2, 0.05, self.colors.accent)

        # 目录项
        items = content.bullets or content.left_content or []

        for i, item in enumerate(items[:6]):
            y = 2.0 + i * 0.85
            # 序号
            self._add_text_box(
                slide, 1.2, y, 0.8, 0.6,
                f"{i+1:02d}", font_size=24, bold=True,
                color=self.colors.secondary, alignment="center"
            )
            # 文本
            if isinstance(item, str):
                text = item
            elif isinstance(item, dict):
                text = str(item.get("text", ""))
            else:
                text = str(item)
            self._add_text_box(
                slide, 2.2, y, 8, 0.6,
                text, font_size=20, color=self.colors.text,
                anchor="middle"
            )
            # 底线
            self._add_decorative_line(slide, 2.2, y + 0.65, 8.5, 0.01, self.colors.line)

        self._add_page_number(slide, content.page_number)

    def _render_section(self, content: SlideContent):
        """章节过渡页"""
        slide = self._add_blank_slide()

        # 半屏深色背景
        self._add_rectangle(slide, 0, 0, 5.5, 7.5, self.colors.primary)

        # 章节号
        self._add_text_box(
            slide, 0.8, 2.5, 4, 1.5,
            f"{content.page_number:02d}", font_size=72, bold=True,
            color=self.colors.accent, font_name=self.fonts.title_en
        )

        # 章节标题
        self._add_text_box(
            slide, 6, 2.8, 6.5, 1.2,
            content.title, font_size=36, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )

        if content.subtitle:
            self._add_text_box(
                slide, 6, 4.2, 6.5, 0.8,
                content.subtitle, font_size=18,
                color=self.colors.text_muted
            )

        self._add_decorative_line(slide, 6, 4.0, 2, 0.05, self.colors.accent)

    def _render_content(self, content: SlideContent):
        """标题+要点页"""
        slide = self._add_blank_slide()

        # 标题栏
        self._add_rectangle(slide, 0, 0, 13.333, 1.2, self.colors.primary)
        self._add_text_box(
            slide, 0.8, 0.2, 11, 0.8,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.text_light, font_name=self.fonts.title_cn,
            anchor="middle"
        )

        # 要点
        bullets = content.bullets or []
        body_font_size = content.body_font_size or self.fonts.body_size
        if bullets:
            self._add_bullet_list(
                slide, 1.0, 1.8, 11.333, 5.0,
                bullets, font_size=body_font_size,
                color=self.colors.text
            )
        elif content.body_text:
            self._add_text_box(
                slide, 1.0, 1.8, 11.333, 5.0,
                content.body_text, font_size=body_font_size,
                color=self.colors.text
            )

        self._add_page_number(slide, content.page_number)

    def _render_content_image(self, content: SlideContent):
        """左文右图页"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.8, 0.5, 11, 0.8,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.8, 1.3, 11.7, 0.04, self.colors.secondary)

        # 左侧要点
        bullets = content.bullets or content.left_content or []
        body_font_size = content.body_font_size or self.fonts.body_size
        if bullets:
            self._add_bullet_list(
                slide, 0.8, 1.8, 6.5, 5.0,
                bullets, font_size=body_font_size,
                color=self.colors.text
            )

        # 右侧图片占位
        if content.image_path and Path(content.image_path).exists():
            try:
                slide.shapes.add_picture(
                    content.image_path,
                    self._x(7.8), self._y(1.8),
                    width=self._x(4.8)
                )
            except Exception:
                self._add_image_placeholder(slide, 7.8, 1.8, 4.8, 4.5)
        else:
            self._add_image_placeholder(slide, 7.8, 1.8, 4.8, 4.5)

        self._add_page_number(slide, content.page_number)

    def _add_image_placeholder(self, slide, left, top, width, height):
        """添加图片占位符"""
        shape = self._add_rounded_rect(slide, left, top, width, height, self.colors.line)
        shape.line.color.rgb = hex_to_rgb(self.colors.secondary)
        shape.line.width = Pt(1)
        # 占位文字
        tf = shape.text_frame
        tf.text = "图片"
        p = tf.paragraphs[0]
        p.font.size = Pt(16)
        p.font.color.rgb = hex_to_rgb(self.colors.text_muted)
        p.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    def _render_two_column(self, content: SlideContent):
        """两栏对比页"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.8, 0.5, 11, 0.8,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.8, 1.3, 11.7, 0.04, self.colors.secondary)

        # 左栏
        left_items = content.left_content or []
        right_items = content.right_content or content.bullets or []

        # 左栏背景
        self._add_rounded_rect(slide, 0.6, 1.8, 5.8, 5.0, self._surface_color(True))
        if left_items:
            self._add_bullet_list(
                slide, 1.0, 2.0, 5.0, 4.6,
                left_items, font_size=16, color=self.colors.text
            )

        # 右栏背景
        self._add_rounded_rect(slide, 6.9, 1.8, 5.8, 5.0, self._surface_color(True))
        if right_items:
            self._add_bullet_list(
                slide, 7.3, 2.0, 5.0, 4.6,
                right_items, font_size=16, color=self.colors.text
            )

        # 中间分隔
        self._add_rectangle(slide, 6.6, 2.0, 0.04, 4.6, self.colors.line)

        self._add_page_number(slide, content.page_number)

    def _render_content_list(self, content: SlideContent):
        """列表页（带序号卡片）"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.8, 0.5, 11, 0.8,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.8, 1.3, 11.7, 0.04, self.colors.secondary)

        items = content.bullets or []
        cols = 2

        for i, item in enumerate(items[:8]):  # 最多8项
            col = i % cols
            row = i // cols
            x = 0.8 + col * 6.2
            y = 1.8 + row * 1.3

            if isinstance(item, str):
                text = item
            elif isinstance(item, dict):
                raw = item.get("text", "")
                text = "" if raw is None else str(raw)
            else:
                text = "" if item is None else str(item)

            # 序号圆
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                self._x(x), self._y(y), self._x(0.5), self._y(0.5)
            )
            self._set_shape_bg(circle, self.colors.secondary)
            tf = circle.text_frame
            tf.text = str(i + 1)
            p = tf.paragraphs[0]
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb(self.colors.text_light)
            p.alignment = PP_ALIGN.CENTER
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE

            # 文本
            self._add_text_box(
                slide, x + 0.7, y + 0.05, 5.2, 0.6,
                text, font_size=16, color=self.colors.text,
                anchor="middle"
            )

        self._add_page_number(slide, content.page_number)

    def _render_data_cards(self, content: SlideContent):
        """数据卡片页"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.8, 0.5, 11, 0.8,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.8, 1.3, 11.7, 0.04, self.colors.secondary)

        data = content.data or []
        data = list(data[:4])

        card_width = 2.6
        gap = 0.35
        total_width = len(data) * card_width + (len(data) - 1) * gap
        start_x = (13.333 - total_width) / 2

        colors = [self.colors.primary, self.colors.secondary,
                  self.colors.accent, self.colors.primary]

        for i, item in enumerate(data[:4]):
            if isinstance(item, (list, tuple)) and len(item) >= 2:
                label = item[0]
                value = item[1]
                unit = item[2] if len(item) > 2 else ""
            elif isinstance(item, dict):
                label = item.get("label", "")
                value = item.get("value", "")
                unit = item.get("unit", "")
            else:
                label, value, unit = str(item), "", ""

            x = start_x + i * (card_width + gap)
            card_color = colors[i % len(colors)]

            # 卡片背景
            self._add_rounded_rect(slide, x, 2.2, card_width, 3.0, card_color)

            # 数值
            self._add_text_box(
                slide, x, 2.8, card_width, 1.2,
                str(value), font_size=40, bold=True,
                color=self.colors.text_light, alignment="center",
                font_name=self.fonts.title_en, anchor="middle"
            )
            # 单位
            if unit:
                self._add_text_box(
                    slide, x, 4.0, card_width, 0.5,
                    unit, font_size=16,
                    color=self.colors.text_light, alignment="center"
                )
            # 标签
            self._add_text_box(
                slide, x, 4.6, card_width, 0.5,
                label, font_size=14,
                color=self.colors.text_light, alignment="center"
            )

        self._add_page_number(slide, content.page_number)

    def _render_timeline(self, content: SlideContent):
        """时间线页"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.8, 0.5, 11, 0.8,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.8, 1.3, 11.7, 0.04, self.colors.secondary)

        items = content.timeline_items or []
        n = len(items)
        if n == 0:
            return

        # 时间线主轴（节点文本框左右各占约 1.2 英寸，
        # 首末节点内收，避免描述文本越出页面右缘）
        line_y = 3.8
        self._add_decorative_line(slide, 1.0, line_y, 9.6, 0.04, self.colors.secondary)

        spacing = 9.6 / max(n - 1, 1) if n > 1 else 0

        for i, item in enumerate(items):
            if isinstance(item, (list, tuple)):
                time_str = item[0] if len(item) > 0 else ""
                title = item[1] if len(item) > 1 else ""
                desc = item[2] if len(item) > 2 else ""
            elif isinstance(item, dict):
                time_str = item.get("time", "")
                title = item.get("title", "")
                desc = item.get("desc", "")
            else:
                time_str, title, desc = str(item), "", ""

            x = 1.3 + i * spacing

            # 节点圆
            node = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                self._x(x - 0.15), self._y(line_y - 0.15),
                self._x(0.3), self._y(0.3)
            )
            self._set_shape_bg(node, self.colors.accent)

            # 时间（上方）
            self._add_text_box(
                slide, x - 1.0, line_y - 1.0, 2.0, 0.5,
                time_str, font_size=14, bold=True,
                color=self.colors.secondary, alignment="center"
            )
            # 标题（下方）
            self._add_text_box(
                slide, x - 1.0, line_y + 0.4, 2.0, 0.5,
                title, font_size=16, bold=True,
                color=self.colors.text, alignment="center"
            )
            # 描述
            if desc:
                self._add_text_box(
                    slide, x - 1.2, line_y + 0.9, 2.4, 0.8,
                    desc, font_size=12,
                    color=self.colors.text_muted, alignment="center"
                )

        self._add_page_number(slide, content.page_number)

    def _render_quote(self, content: SlideContent):
        """引用/金句页"""
        slide = self._add_blank_slide()

        # 深色背景
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(self.colors.bg_dark)

        # 大引号
        self._add_text_box(
            slide, 1.5, 1.5, 2, 2,
            "\u201C", font_size=120, bold=True,
            color=self.colors.accent, font_name="Georgia"
        )

        # 引用文字
        quote = content.quote_text or content.title or ""
        self._add_text_box(
            slide, 2.5, 2.5, 8.5, 2.5,
            quote, font_size=28,
            color=self.colors.text_light, alignment="center",
            anchor="middle"
        )

        # 来源
        source = content.quote_source or content.subtitle or ""
        if source:
            self._add_text_box(
                slide, 2.5, 5.2, 8.5, 0.5,
                f"— {source}", font_size=16,
                color=self.colors.text_muted, alignment="center"
            )

    def _render_summary(self, content: SlideContent):
        """总结/致谢页"""
        slide = self._add_blank_slide()

        # 深色背景
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(self.colors.bg_dark)

        # 装饰
        self._add_rectangle(slide, 0, 0, 13.333, 0.1, self.colors.accent)
        self._add_rectangle(slide, 0, 7.4, 13.333, 0.1, self.colors.accent)

        # 主文字
        main_text = content.title or "感谢聆听"
        self._add_text_box(
            slide, 1, 2.5, 11.333, 1.5,
            main_text, font_size=48, bold=True,
            color=self.colors.text_light, alignment="center",
            font_name=self.fonts.title_cn, anchor="middle"
        )

        if content.subtitle:
            self._add_text_box(
                slide, 1, 4.2, 11.333, 0.8,
                content.subtitle, font_size=20,
                color=self.colors.text_muted, alignment="center"
            )

        # 要点
        if content.bullets:
            self._add_bullet_list(
                slide, 3.5, 5.2, 6.3, 1.5,
                content.bullets[:3], font_size=14,
                color=self.colors.text_muted, bullet_char="✓"
            )

    def _render_blank(self, content: SlideContent):
        """空白页（仅标题）"""
        slide = self._add_blank_slide()
        if content.title:
            self._add_text_box(
                slide, 0.8, 0.5, 11, 0.8,
                content.title, font_size=self.fonts.title_size, bold=True,
                color=self.colors.primary
            )
        self._add_page_number(slide, content.page_number)

    # ==========================================
    # 表格页
    # ==========================================

    def _render_table(self, content: SlideContent):
        """表格页：标题 + 表格"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.5, 0.3, 12.333, 0.9,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        # 装饰线
        self._add_decorative_line(slide, 0.5, 1.2, 12.333, 0.03, self.colors.secondary)

        # 表格
        table_data = content.table_data
        if not table_data and content.bullets:
            # 从 bullets 转换（简单格式）
            table_data = [["项目", "内容"]] + [[b, ""] for b in content.bullets]

        if table_data:
            # 规范化：丢掉非列表行/空行，全部转字符串
            norm_rows = []
            for row in table_data:
                if isinstance(row, (list, tuple)) and len(row) > 0:
                    norm_rows.append([str(c) if c is not None else "" for c in row])
                elif isinstance(row, str):
                    norm_rows.append([row])
            table_data = norm_rows
            if table_data:
                max_cols = max(len(r) for r in table_data)
                if max_cols == 0 or len(table_data) == 0:
                    table_data = []
            if table_data:
                if len(table_data) > 40:
                    overflow = len(table_data) - 40
                    table_data = table_data[:40]
                    if not content.body_text:
                        content.body_text = f"（内容较长，仅显示前 40 行，省略 {overflow} 行）"
                rows = len(table_data)
                # 表格位置
                t_left, t_top = 0.8, 1.6
                t_width, t_height = 11.733, min(5.2, 0.5 * rows + 0.3)
                self._add_table(
                    slide, t_left, t_top, t_width, t_height,
                    table_data, has_header=content.table_header
                )

        # 备注/说明
        if content.body_text:
            self._add_text_box(
                slide, 0.8, 6.5, 11.733, 0.5,
                content.body_text, font_size=self.fonts.caption_size,
                color=self.colors.text_muted, font_name=self.fonts.body_cn
            )

        self._add_page_number(slide, content.page_number)

    def _add_table(self, slide, left: float, top: float, width: float,
                   height: float, data: list, has_header: bool = True):
        """
        添加表格

        Args:
            slide: 幻灯片
            left/top/width/height: 位置和大小（英寸）
            data: 二维列表，第一行为表头
            has_header: 第一行是否为表头
        """
        rows = len(data)
        cols = max((len(r) for r in data), default=0)
        if rows == 0 or cols == 0:
            return None

        table_shape = slide.shapes.add_table(
            rows, cols,
            self._x(left), self._y(top),
            self._x(width), self._y(height)
        )
        table = table_shape.table

        # 填充数据
        for r_idx, row_data in enumerate(data):
            for c_idx in range(cols):
                cell = table.cell(r_idx, c_idx)
                text = str(row_data[c_idx]) if c_idx < len(row_data) else ""
                cell.text = text

                # 单元格样式
                for para in cell.text_frame.paragraphs:
                    para.alignment = PP_ALIGN.CENTER
                    for run in para.runs:
                        run.font.size = Pt(self.fonts.body_size - 2)
                        run.font.name = self.fonts.body_cn
                        self._set_cn_font(run, self.fonts.body_cn)

                        if has_header and r_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = hex_to_rgb(self.colors.text_light)
                        else:
                            run.font.color.rgb = hex_to_rgb(self.colors.text)

                # 单元格填充
                if has_header and r_idx == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = hex_to_rgb(self.colors.primary)
                elif r_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = hex_to_rgb(self._surface_color(True))
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = hex_to_rgb(self._surface_color(False))

                # 垂直居中
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

        return table_shape

    # ==========================================
    # 图表页
    # ==========================================

    def _render_chart(self, content: SlideContent):
        """图表页：标题 + 图表"""
        slide = self._add_blank_slide()

        # 标题
        self._add_text_box(
            slide, 0.5, 0.3, 12.333, 0.9,
            content.title, font_size=self.fonts.title_size, bold=True,
            color=self.colors.primary, font_name=self.fonts.title_cn
        )
        self._add_decorative_line(slide, 0.5, 1.2, 12.333, 0.03, self.colors.secondary)

        # 图表数据
        categories = content.chart_categories
        series = content.chart_series

        # 如果没有显式图表数据，尝试从 data 字段转换
        if not categories and content.data:
            pairs = []
            for item in content.data:
                if isinstance(item, dict):
                    label = item.get("label", item.get("name", item.get("category", "")))
                    value = item.get("value", item.get("amount", 0))
                    pairs.append((label, value))
                elif isinstance(item, (list, tuple)) and len(item) >= 2:
                    pairs.append((item[0], item[1]))
            categories = [pair[0] for pair in pairs]
            series = [("数值", [pair[1] for pair in pairs])] if pairs else []

        if categories and series:
            chart_type = content.chart_type
            c_left, c_top = 0.8, 1.5
            c_width, c_height = 11.733, 5.2
            self._add_chart(
                slide, c_left, c_top, c_width, c_height,
                chart_type=chart_type,
                categories=categories,
                series=series,
                chart_title=content.chart_title or content.title
            )

        # 补充要点
        if content.bullets:
            self._add_bullet_list(
                slide, 0.8, 6.2, 11.733, 0.8,
                content.bullets[:3], font_size=self.fonts.caption_size,
                color=self.colors.text_muted
            )

        self._add_page_number(slide, content.page_number)

    def _add_chart(self, slide, left: float, top: float, width: float,
                   height: float, chart_type: str, categories: list,
                   series: list, chart_title: str = ""):
        """
        添加图表

        Args:
            slide: 幻灯片
            left/top/width/height: 位置和大小（英寸）
            chart_type: bar/column/line/pie
            categories: X轴类别列表
            series: [(系列名, [数值列表]), ...]
            chart_title: 图表标题
        """
        # 图表类型映射
        chart_type_map = {
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
            "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
            "area": XL_CHART_TYPE.AREA,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }
        xl_type = chart_type_map.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)

        # 构建图表数据（类别与数值强类型化：非法数值会写出损坏的
        # c:numCache XML，导致 PowerPoint/WPS 打开时提示修复）
        def _chart_num(v):
            if isinstance(v, bool) or v is None:
                return 0.0
            try:
                return float(v)
            except (TypeError, ValueError):
                return 0.0

        normalized_series = []
        for item in series or []:
            if isinstance(item, dict):
                name = item.get("name", item.get("label", "系列"))
                values = item.get("values", item.get("data", []))
            elif isinstance(item, (list, tuple)) and len(item) >= 2:
                name, values = item[0], item[1]
            else:
                continue
            if not isinstance(values, (list, tuple)):
                values = [values]
            normalized_series.append((name, values))

        if not normalized_series:
            return None

        normalized_categories = [str(c) for c in (categories or [])]
        if not normalized_categories:
            raise ValueError("图表类别不能为空")
        expected_len = len(normalized_categories)
        for name, values in normalized_series:
            if len(values) != expected_len:
                raise ValueError(
                    f"图表系列 {name!r} 有 {len(values)} 个值，但类别有 {expected_len} 个"
                )

        chart_data = CategoryChartData()
        chart_data.categories = normalized_categories
        for name, values in normalized_series:
            chart_data.add_series(str(name), [_chart_num(v) for v in (values or [])])

        # 添加图表
        chart_frame = slide.shapes.add_chart(
            xl_type,
            self._x(left), self._y(top),
            self._x(width), self._y(height),
            chart_data
        )
        chart = chart_frame.chart

        # 图表标题
        if chart_title:
            chart.has_title = True
            chart.chart_title.text_frame.text = chart_title
            for para in chart.chart_title.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(self.fonts.body_size)
                    run.font.bold = True
                    run.font.color.rgb = hex_to_rgb(self.colors.primary)
                    run.font.name = self.fonts.title_cn
        else:
            chart.has_title = False

        # 图例
        if len(normalized_series) > 1 and chart_type != "pie":
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(self.fonts.caption_size)
        else:
            chart.has_legend = (chart_type == "pie")
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.RIGHT
                chart.legend.font.size = Pt(self.fonts.caption_size)

        # 系列颜色
        series_colors = [
            self.colors.primary, self.colors.secondary, self.colors.accent,
            "#70AD47", "#5B9BD5", "#C0504D", "#8064A2"
        ]
        for i, s in enumerate(chart.series):
            color = series_colors[i % len(series_colors)]
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = hex_to_rgb(color)

        # 饼图数据标签
        if chart_type == "pie":
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_category_name = True
            plot.data_labels.position = XL_LABEL_POSITION.BEST_FIT
            plot.data_labels.font.size = Pt(self.fonts.caption_size)

        return chart_frame
