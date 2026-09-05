"""
Document Quality Checker - 文档质量检查器

Word 文档生成后自动检查：
1. 格式检查：字体、字号、行距、标题层级、编号连续性、表格格式
2. 内容检查：乱码检测、空段落、文字遗漏

输出检查报告，支持自动修复建议。
"""
import logging
import re
from difflib import SequenceMatcher
from dataclasses import dataclass, field, asdict
from enum import Enum
from typing import Optional, Dict
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn

from ..models.schemas import FormatConfig
from ..parsers.document_structure import DocumentStructureAnalyzer, DocumentTree

logger = logging.getLogger(__name__)


class IssueSeverity(Enum):
    """问题严重程度（Word/Excel/PPT 质量问题共用的唯一权威词汇表）"""
    ERROR = "error"
    WARNING = "warning"
    INFO = "info"

    @classmethod
    def values(cls) -> list:
        """全部合法 severity 字符串值。"""
        return [member.value for member in cls]

    @classmethod
    def is_valid(cls, value) -> bool:
        """value 是否为合法 severity（先经 normalize 同款规范化再判定）。"""
        if not isinstance(value, str):
            return False
        return value.strip().lower() in cls.values()

    @classmethod
    def normalize(cls, value, *, fallback=None) -> str:
        """信任边界规范化：大小写/空白不敏感；未知值不静默漂移。

        未知值给定 fallback 时返回 fallback，否则抛 ValueError。
        """
        if isinstance(value, str):
            normalized = value.strip().lower()
            if normalized in cls.values():
                return normalized
        if fallback is not None:
            return fallback
        raise ValueError(f"未知 IssueSeverity: {value!r}")


class IssueType(Enum):
    """问题类型"""
    FONT = "font"
    FONT_SIZE = "font_size"
    LINE_SPACING = "line_spacing"
    HEADING_LEVEL = "heading_level"
    NUMBERING = "numbering"
    TABLE_FORMAT = "table_format"
    GARBLED = "garbled"
    EMPTY_PARAGRAPH = "empty_paragraph"
    MISSING_TEXT = "missing_text"
    ALIGNMENT = "alignment"
    INDENT = "indent"
    PAGE_SETUP = "page_setup"


@dataclass
class QualityIssue:
    """单个质量问题"""
    type: str
    severity: str
    message: str
    paragraph_index: int = -1
    paragraph_text: str = ""
    expected: str = ""
    actual: str = ""
    fixable: bool = True
    fix_suggestion: str = ""

    def __post_init__(self):
        # 构造边界统一校验：枚举是唯一权威，未知 severity 不得静默漂移
        self.severity = IssueSeverity.normalize(self.severity)

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class QualityReport:
    """质量检查报告"""
    file_path: str = ""
    total_paragraphs: int = 0
    total_headings: int = 0
    total_tables: int = 0
    issues: list = field(default_factory=list)
    passed: bool = True
    score: float = 100.0

    def add_issue(self, issue: QualityIssue):
        self.issues.append(issue)
        if issue.severity == IssueSeverity.ERROR.value:
            self.passed = False

    def compute_score(self):
        error_count = sum(1 for i in self.issues if i.severity == IssueSeverity.ERROR.value)
        warning_count = sum(1 for i in self.issues if i.severity == IssueSeverity.WARNING.value)
        info_count = sum(1 for i in self.issues if i.severity == IssueSeverity.INFO.value)
        score = 100.0 - error_count * 10 - warning_count * 3 - info_count * 1
        self.score = max(0.0, min(100.0, score))
        self.passed = error_count == 0

    def errors(self) -> list:
        return [i for i in self.issues if i.severity == IssueSeverity.ERROR.value]

    def warnings(self) -> list:
        return [i for i in self.issues if i.severity == IssueSeverity.WARNING.value]

    def fixable_issues(self) -> list:
        return [i for i in self.issues if i.fixable]

    def to_dict(self) -> dict:
        self.compute_score()
        return {
            "file_path": self.file_path,
            "passed": self.passed,
            "score": self.score,
            "stats": {
                "total_paragraphs": self.total_paragraphs,
                "total_headings": self.total_headings,
                "total_tables": self.total_tables,
                "total_issues": len(self.issues),
                "errors": len(self.errors()),
                "warnings": len(self.warnings()),
            },
            "issues": [i.to_dict() for i in self.issues],
        }

    def to_text(self) -> str:
        self.compute_score()
        lines = []
        lines.append("=" * 50)
        lines.append("  文档质量检查报告")
        lines.append("=" * 50)
        lines.append(f"文件: {Path(self.file_path).name}")
        lines.append(f"段落: {self.total_paragraphs}  标题: {self.total_headings}  表格: {self.total_tables}")
        lines.append(f"质量分数: {self.score:.0f}/100  {'✓ 通过' if self.passed else '✗ 未通过'}")
        lines.append("")

        if not self.issues:
            lines.append("✓ 未发现问题，文档质量良好！")
        else:
            for severity_name, severity_label in [
                (IssueSeverity.ERROR.value, "错误"),
                (IssueSeverity.WARNING.value, "警告"),
                (IssueSeverity.INFO.value, "提示"),
            ]:
                items = [i for i in self.issues if i.severity == severity_name]
                if items:
                    lines.append(f"【{severity_label}】{len(items)}项")
                    for i, issue in enumerate(items, 1):
                        lines.append(f"  {i}. {issue.message}")
                        if issue.expected or issue.actual:
                            lines.append(f"     期望: {issue.expected}  实际: {issue.actual}")
                        if issue.fix_suggestion:
                            lines.append(f"     建议: {issue.fix_suggestion}")
                    lines.append("")

        return "\n".join(lines)

    def get_fix_config(self) -> dict:
        fix = {}
        for issue in self.issues:
            if not issue.fixable:
                continue
            if issue.type == IssueType.FONT.value and "正文" in issue.message:
                fix["font"] = issue.expected
            elif issue.type == IssueType.FONT_SIZE.value and "正文" in issue.message:
                fix["size"] = issue.expected
            elif issue.type == IssueType.LINE_SPACING.value:
                match = re.search(r"\d+(?:\.\d+)?", issue.expected or "")
                if match:
                    fix["line_spacing"] = float(match.group())
        return fix


# 标题类型集合
HEADING_TYPES = {"chapter", "section", "subsection", "subsubsection", "title"}


class QualityChecker:
    """
    文档质量检查器

    用法:
        checker = QualityChecker()
        report = checker.check("output.docx", format_config=config)
        print(report.to_text())
    """

    GARBLED_PATTERNS = [
        re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]"),
        re.compile(r"[\ufffd\uFFFE\uFFFF]"),
        re.compile(r"[\u200b\u200c\u200d\u2060]"),
        re.compile(r"[\ue000-\uf8ff]"),
    ]

    def __init__(self):
        self.structure_analyzer = DocumentStructureAnalyzer()

    def check(self, file_path: str,
              format_config: Optional[FormatConfig] = None,
              original_path: Optional[str] = None) -> QualityReport:
        doc = Document(file_path)
        report = QualityReport(file_path=file_path)

        # 分析文档结构
        tree = self.structure_analyzer.analyze(doc)
        report.total_paragraphs = len([p for p in doc.paragraphs if p.text.strip()])
        report.total_headings = len([n for n in tree.nodes if n.type in HEADING_TYPES])
        report.total_tables = len(doc.tables)

        if not format_config:
            format_config = FormatConfig()

        # 建立 paragraph_index → node_type 映射
        self.node_map: Dict[int, str] = {}
        for node in tree.nodes:
            if node.index >= 0:
                self.node_map[node.index] = node.type

        # 1. 格式检查
        self._check_fonts(doc, format_config, report)
        self._check_font_sizes(doc, format_config, report)
        self._check_line_spacing(doc, format_config, report)
        self._check_heading_levels(tree, report)
        self._check_numbering(tree, report)
        self._check_tables(doc, format_config, report)
        self._check_indent(doc, format_config, report)

        # 2. 内容检查
        self._check_garbled_text(doc, report)
        self._check_empty_paragraphs(doc, report)

        # 3. 对比原文
        if original_path:
            self._check_missing_text(doc, original_path, report)

        report.compute_score()
        return report

    def _is_heading(self, idx: int) -> bool:
        return self.node_map.get(idx, "paragraph") in HEADING_TYPES

    # ==========================================
    # 格式检查
    # ==========================================

    def _check_fonts(self, doc: DocumentType, config: FormatConfig, report: QualityReport):
        expected_cn = config.body_font.cn_font
        font_errors = 0

        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
            if self._is_heading(idx):
                continue
            if self.node_map.get(idx) in ("list_item", "quote"):
                continue

            for run in para.runs:
                if not run.text.strip():
                    continue
                cn_font = self._get_cn_font(run)
                if cn_font and expected_cn and cn_font != expected_cn:
                    if not self._is_likely_english(run.text):
                        font_errors += 1
                        if font_errors <= 3:  # 只报告前3个
                            report.add_issue(QualityIssue(
                                type=IssueType.FONT.value,
                                severity=IssueSeverity.WARNING.value,
                                message=f"第{idx+1}段正文字体为「{cn_font}」，期望「{expected_cn}」",
                                paragraph_index=idx,
                                paragraph_text=text[:40],
                                expected=expected_cn,
                                actual=cn_font,
                                fixable=True,
                                fix_suggestion=f"将正文字体统一为{expected_cn}",
                            ))
                break

        if font_errors > 3:
            report.add_issue(QualityIssue(
                type=IssueType.FONT.value,
                severity=IssueSeverity.WARNING.value,
                message=f"另有{font_errors - 3}段正文字体不符",
                expected=expected_cn,
                fixable=True,
                fix_suggestion=f"将正文字体统一为{expected_cn}",
            ))

    def _check_font_sizes(self, doc: DocumentType, config: FormatConfig, report: QualityReport):
        expected_size = config.body_font.size
        size_errors = 0

        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
            if self._is_heading(idx):
                continue
            if self.node_map.get(idx) in ("list_item", "quote"):
                continue

            for run in para.runs:
                if not run.text.strip():
                    continue
                if run.font.size:
                    actual_size = run.font.size.pt
                    if abs(actual_size - expected_size) > 0.5:
                        size_errors += 1
                        if size_errors <= 3:
                            report.add_issue(QualityIssue(
                                type=IssueType.FONT_SIZE.value,
                                severity=IssueSeverity.WARNING.value,
                                message=f"第{idx+1}段正文字号为{actual_size}pt，期望{expected_size}pt",
                                paragraph_index=idx,
                                paragraph_text=text[:40],
                                expected=f"{expected_size}pt",
                                actual=f"{actual_size}pt",
                                fixable=True,
                                fix_suggestion=f"将正文字号统一为{expected_size}pt",
                            ))
                break

        if size_errors > 3:
            report.add_issue(QualityIssue(
                type=IssueType.FONT_SIZE.value,
                severity=IssueSeverity.WARNING.value,
                message=f"另有{size_errors - 3}段正文字号不符",
                expected=f"{expected_size}pt",
                fixable=True,
            ))

    def _check_line_spacing(self, doc: DocumentType, config: FormatConfig, report: QualityReport):
        expected_spacing = config.body_paragraph.line_spacing
        spacing_errors = 0

        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text or len(text) < 10:
                continue
            if self._is_heading(idx):
                continue

            pf = para.paragraph_format
            if pf.line_spacing:
                ls = pf.line_spacing
                actual = ls if isinstance(ls, (int, float)) else (ls.pt if hasattr(ls, 'pt') else None)
                if actual and isinstance(actual, (int, float)) and actual < 5:
                    if abs(actual - expected_spacing) > 0.1:
                        spacing_errors += 1

        if spacing_errors > 0:
            report.add_issue(QualityIssue(
                type=IssueType.LINE_SPACING.value,
                severity=IssueSeverity.WARNING.value if spacing_errors <= 3 else IssueSeverity.INFO.value,
                message=f"有{spacing_errors}段正文行距为非{expected_spacing}倍",
                expected=f"{expected_spacing}倍",
                actual=f"{spacing_errors}段不符",
                fixable=True,
                fix_suggestion=f"将正文行距统一为{expected_spacing}倍",
            ))

    def _check_heading_levels(self, tree: DocumentTree, report: QualityReport):
        headings = [(n.level, n.text) for n in tree.nodes if n.type in HEADING_TYPES]

        prev_level = 0
        for level, text in headings:
            if prev_level > 0 and level > prev_level + 1:
                report.add_issue(QualityIssue(
                    type=IssueType.HEADING_LEVEL.value,
                    severity=IssueSeverity.WARNING.value,
                    message=f"标题层级跳跃：「{text[:30]}」从{prev_level}级跳到{level}级",
                    paragraph_text=text[:50],
                    expected=f"不超过{prev_level + 1}级",
                    actual=f"{level}级",
                    fixable=False,
                    fix_suggestion="检查标题层级是否正确，避免跳级",
                ))
            prev_level = level

    def _check_numbering(self, tree: DocumentTree, report: QualityReport):
        """检查标题编号是否连续"""
        counters = {1: 0, 2: 0, 3: 0, 4: 0}

        for node in tree.nodes:
            if node.type not in HEADING_TYPES:
                continue
            level = node.level
            if level < 1 or level > 4:
                continue

            text = node.text.strip()

            # 无编号标题不参与编号序列；否则它会把后续正常编号整体推后一位。
            num = self._extract_heading_number(text, level)
            if num is None:
                continue

            # 更新计数器
            if level == 1:
                counters[1] += 1
                counters[2] = counters[3] = counters[4] = 0
            elif level == 2:
                counters[2] += 1
                counters[3] = counters[4] = 0
            elif level == 3:
                counters[3] += 1
                counters[4] = 0
            elif level == 4:
                counters[4] += 1

            # 获取实际编号（最后一级）
            if isinstance(num, tuple):
                actual_last = num[-1]
                expected_last = counters[level]
            else:
                actual_last = num
                expected_last = counters[level]

            if actual_last != expected_last:
                report.add_issue(QualityIssue(
                    type=IssueType.NUMBERING.value,
                    severity=IssueSeverity.ERROR.value,
                    message=f"标题编号不连续：「{text[:30]}」编号为{actual_last}，期望{expected_last}",
                    paragraph_text=text[:50],
                    expected=str(expected_last),
                    actual=str(actual_last),
                    fixable=True,
                    fix_suggestion="重新编排标题编号",
                ))

        # 表编号
        table_nums = []
        for node in tree.nodes:
            m = re.match(r"^表(\d+)", node.text.strip())
            if m:
                table_nums.append((int(m.group(1)), node.text))

        for i, (num, text) in enumerate(table_nums, 1):
            if num != i:
                report.add_issue(QualityIssue(
                    type=IssueType.NUMBERING.value,
                    severity=IssueSeverity.ERROR.value,
                    message=f"表编号不连续：「{text[:30]}」编号为{num}，期望{i}",
                    paragraph_text=text[:50],
                    expected=str(i),
                    actual=str(num),
                    fixable=True,
                    fix_suggestion="重新编排表编号",
                ))

    def _check_tables(self, doc: DocumentType, config: FormatConfig, report: QualityReport):
        tc = config.table_config

        for t_idx, table in enumerate(doc.tables):
            tbl = table._tbl
            tblPr = tbl.find(qn("w:tblPr"))
            if tblPr is not None:
                tblBorders = tblPr.find(qn("w:tblBorders"))
                if tblBorders is not None:
                    borders = self._read_table_borders(tblBorders)

                    if not tc.show_left_right:
                        if borders.get("left", 0) > 0 or borders.get("right", 0) > 0:
                            report.add_issue(QualityIssue(
                                type=IssueType.TABLE_FORMAT.value,
                                severity=IssueSeverity.WARNING.value,
                                message=f"表{t_idx+1}有左右边框，三线表应无左右边框",
                                expected="无左右边框",
                                actual=f"左{borders.get('left',0)}磅，右{borders.get('right',0)}磅",
                                fixable=True,
                                fix_suggestion="去除表格左右边框",
                            ))

    def _check_indent(self, doc: DocumentType, config: FormatConfig, report: QualityReport):
        expected_indent = (
            config.body_paragraph.first_line_indent_chars * config.body_font.size
            if config.body_paragraph.first_line_indent_chars
            else config.body_paragraph.first_line_indent
        )
        if expected_indent <= 0:
            return

        mismatch = 0
        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text or len(text) < 15:
                continue
            if self._is_heading(idx):
                continue
            if self.node_map.get(idx) in ("list_item", "quote"):
                continue

            pf = para.paragraph_format
            if pf.first_line_indent:
                if abs(pf.first_line_indent.pt - expected_indent) > 2:
                    mismatch += 1

        if mismatch > 3:
            report.add_issue(QualityIssue(
                type=IssueType.INDENT.value,
                severity=IssueSeverity.INFO.value,
                message=f"有{mismatch}段正文首行缩进不符（期望{expected_indent}pt）",
                expected=f"{expected_indent}pt",
                actual=f"{mismatch}段不符",
                fixable=True,
                fix_suggestion=f"统一首行缩进为{expected_indent}pt",
            ))

    # ==========================================
    # 内容检查
    # ==========================================

    def _check_garbled_text(self, doc: DocumentType, report: QualityReport):
        for idx, para in enumerate(doc.paragraphs):
            text = para.text
            if not text.strip():
                continue

            for pattern in self.GARBLED_PATTERNS:
                matches = pattern.findall(text)
                if matches:
                    report.add_issue(QualityIssue(
                        type=IssueType.GARBLED.value,
                        severity=IssueSeverity.ERROR.value,
                        message=f"第{idx+1}段发现乱码字符（{len(matches)}个）",
                        paragraph_index=idx,
                        paragraph_text=text[:50],
                        expected="正常文字",
                        actual=f"含{len(matches)}个异常字符",
                        fixable=False,
                        fix_suggestion="检查原文编码，删除或替换乱码字符",
                    ))
                    break

            if re.search(r"(.)\1{6,}", text):
                report.add_issue(QualityIssue(
                    type=IssueType.GARBLED.value,
                    severity=IssueSeverity.WARNING.value,
                    message=f"第{idx+1}段存在连续重复字符",
                    paragraph_index=idx,
                    paragraph_text=text[:50],
                    fixable=False,
                ))

    def _check_empty_paragraphs(self, doc: DocumentType, report: QualityReport):
        empty_count = sum(1 for p in doc.paragraphs if not p.text.strip())
        total = len(doc.paragraphs)
        if total > 0 and empty_count > total * 0.3:
            report.add_issue(QualityIssue(
                type=IssueType.EMPTY_PARAGRAPH.value,
                severity=IssueSeverity.INFO.value,
                message=f"空段落过多（{empty_count}/{total}），建议清理",
                expected="少量空段落",
                actual=f"{empty_count}个空段落",
                fixable=True,
                fix_suggestion="删除多余空段落",
            ))

    def _check_missing_text(self, doc: DocumentType, original_path: str, report: QualityReport):
        try:
            orig = Document(original_path)
            orig_text = "\n".join(p.text for p in orig.paragraphs if p.text.strip())
            curr_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

            orig_sentences = re.split(r"[。！？\n]", orig_text)
            current_sentences = [
                sentence.strip() for sentence in re.split(r"[。！？\n]", curr_text)
                if sentence.strip()
            ]
            missing = []
            for sent in orig_sentences:
                sent = sent.strip()
                if (len(sent) > 15 and sent not in curr_text
                        and not any(self._text_similarity(sent, current) >= 0.72
                                    for current in current_sentences)):
                    missing.append(sent)

            if missing:
                report.add_issue(QualityIssue(
                    type=IssueType.MISSING_TEXT.value,
                    severity=IssueSeverity.ERROR.value,
                    message=f"可能遗漏{len(missing)}处原文内容",
                    expected="完整保留原文",
                    actual=f"缺失{len(missing)}处",
                    fixable=False,
                    fix_suggestion="检查并补全：" + "；".join(m[:30] for m in missing[:3]),
                ))
        except Exception as exc:
            logger.exception("原文保留检查失败: %s", original_path)
            report.add_issue(QualityIssue(
                type=IssueType.MISSING_TEXT.value, severity=IssueSeverity.INFO.value,
                message=f"原文保留检查未完成：{exc}", fixable=False,
                fix_suggestion="确认原文文件可读取后重新检查",
            ))

    @staticmethod
    def _text_similarity(left: str, right: str) -> float:
        def normalize(value: str) -> str:
            return re.sub(r"[\W_]+", "", value, flags=re.UNICODE).lower()

        return SequenceMatcher(None, normalize(left), normalize(right)).ratio()

    # ==========================================
    # 辅助方法
    # ==========================================

    def _get_cn_font(self, run) -> str:
        rPr = run._element.find(qn("w:rPr"))
        if rPr is not None:
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is not None:
                return rFonts.get(qn("w:eastAsia"), "")
        return ""

    @staticmethod
    def _is_likely_english(text: str) -> bool:
        en_chars = sum(1 for c in text if c.isascii() and (c.isalpha() or c.isdigit() or c in " .,-"))
        return en_chars > len(text) * 0.7

    def _extract_heading_number(self, text: str, level: int):
        """从标题文本中提取编号"""
        # "1 第一章 绪论" → 提取开头的数字
        m = re.match(r"^(\d+)\s+第", text)
        if m:
            return int(m.group(1))

        # "第一章 绪论"
        m = re.match(r"^第([一二三四五六七八九十百千\d]+)[章节]", text)
        if m:
            return self._cn_to_int(m.group(1))

        # "1.1 标题" 或 "1.1标题"
        m = re.match(r"^(\d+)\.(\d+)\.(\d+)\.(\d+)", text)
        if m:
            return tuple(int(g) for g in m.groups())

        m = re.match(r"^(\d+)\.(\d+)\.(\d+)", text)
        if m:
            return tuple(int(g) for g in m.groups())

        m = re.match(r"^(\d+)\.(\d+)", text)
        if m:
            return tuple(int(g) for g in m.groups())

        # "一、标题"
        m = re.match(r"^([一二三四五六七八九十]+)[、.]", text)
        if m:
            return self._cn_to_int(m.group(1))

        return None

    @staticmethod
    def _cn_to_int(cn: str) -> int:
        if cn.isdigit():
            return int(cn)
        cn_map = {"一":1, "二":2, "三":3, "四":4, "五":5,
                  "六":6, "七":7, "八":8, "九":9, "十":10}
        if cn in cn_map:
            return cn_map[cn]
        if "十" in cn:
            parts = cn.split("十")
            tens = cn_map.get(parts[0], 1) if parts[0] else 1
            ones = cn_map.get(parts[1], 0) if len(parts) > 1 and parts[1] else 0
            return tens * 10 + ones
        return 0

    def _read_table_borders(self, tblBorders) -> dict:
        borders = {}
        for tag, name in [("top", "top"), ("bottom", "bottom"), ("left", "left"),
                          ("right", "right"), ("insideH", "inside_h")]:
            elem = tblBorders.find(qn(f"w:{tag}"))
            if elem is not None:
                val = elem.get(qn("w:val"), "none")
                if val not in ("none", "nil"):
                    sz = elem.get(qn("w:sz"), "4")
                    try:
                        borders[name] = int(sz) / 8.0
                    except ValueError:
                        borders[name] = 0
        return borders


def check_document(file_path: str,
                   format_config: Optional[FormatConfig] = None,
                   original_path: Optional[str] = None) -> QualityReport:
    checker = QualityChecker()
    return checker.check(file_path, format_config, original_path)
